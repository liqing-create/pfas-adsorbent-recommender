"""Consolidate standardized adsorbent properties across studies.

This helper is called by ``normalize.py``. It owns identity mapping,
categorical harmonization, cross-study aggregation, and traceable workbook
layout. It deliberately does not own unit conversion and has no CLI.

* ``Adsorbent_Database``: one row per mapped/reported commercial adsorbent.
* ``Consolidation_Audit``: every raw source row plus identity, normalization,
  parsing, conversion, and inclusion decisions.

The explicit adsorbent map is authoritative.  Rows that do not match the map
are consolidated only when ``name_commercial`` is reported; generic study
labels such as "GAC" or "PAC" are never merged across studies implicitly.
"""

from __future__ import annotations

import csv
import difflib
import math
import os
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from normalization_rules import (
    normalize_bounded_fraction_values,
    normalize_elemental_fraction_value_lists,
)


DATABASE_SHEET = "Adsorbent_Database"
AUDIT_SHEET = "Consolidation_Audit"
AUTHORITATIVE_SHEET = "Adsorbent_Properties"
AUTHORITATIVE_AUDIT_SHEET = "Merge_Audit"

IDENTITY_SOURCE_COLUMNS = (
    "name_commercial",
    "name_abbreviation",
    "name_full",
)

NULL_LABELS = {"", "nan", "none", "null", "n/a", "na", "not reported", "unclassified"}

ADSORBENT_IDENTITY_STATUS_COLUMN = "adsorbent_identity_status"
ADSORBENT_IDENTITY_BASIS_COLUMN = "adsorbent_identity_basis"
ADSORBENT_IDENTITY_KEY_COLUMN = "adsorbent_identity_key"
ADSORBENT_STUDY_INSTANCE_KEY_COLUMN = "adsorbent_study_instance_key"
FUNCTIONAL_GROUP_IDENTITY_CONFLICT_COLUMN = "functional_group_identity_conflict"


@dataclass(frozen=True)
class NumericProperty:
    output_col: str
    source_col: str
    target_unit: str
    converter: str = "identity"
    unit_col: str | None = None
    specific_exchange_capacity: bool = False


@dataclass(frozen=True)
class AdsorbentProcessingConfig:
    raw_input_path: Path
    database_output_path: Path
    raw_input_sheet: str
    mapping_path: Path
    online_properties_path: Path
    online_properties_sheet: str
    database_properties_path: Path
    database_properties_sheet: str
    enable_online_injection: bool = True
    enable_database_injection: bool = True
    authoritative_output_path: Path | None = None
    authoritative_properties_sheet: str = AUTHORITATIVE_SHEET


ADSORBENT_PROPERTY_COLUMN_MAP = {
    "Adsorbent_category": "adsorbent_category",
    "Adsorbent_subcategory": "adsorbent_subcategory",
    "Vendor": "vendor",
    "particle_size_value_mm": "Nominal_grain_size_value_mm",
    "Nominal_grain_size_value_mm": "Nominal_grain_size_value_mm",
    "AC_raw_material": "AC_raw_material",
    "AC_activation_method": "AC_activation_method",
    "Polymer_matrix": "polymer_matrix",
    "Functional_group": "functional_group",
    "Elemental_composition_method": "elemental_composition_method",
    "ion_exchange_capacity_value_meq/g": "ion_exchange_capacity_value_meq/g",
    "ion_exchange_capacity_value_meq/L": "ion_exchange_capacity_value_meq/L",
    "ion_exchange_capacity_value_mmol/g": "ion_exchange_capacity_value_meq/g",
    "element_C_value": "element_C_value",
    "element_N_value": "element_N_value",
    "element_O_value": "element_O_value",
    "Pore_class": "pore_class",
    "Porosity": "porosity_total_value",
    "Pore_volume_(cm3/g)": "pore_volume_total_value",
    "Total_pore_percentage": "porosity_total_value",
    "Micro_pore_percentage": "porosity_micro_value",
    "Meso_pore_percentage": "porosity_meso_value",
    "Macro_pore_percentage": "porosity_macro_value",
    "Total_pore_volume_(cm3/g)": "pore_volume_total_value",
    "Micro_pore_volume_(cm3/g)": "pore_volume_micro_value",
    "Meso_pore_volume_(cm3/g)": "pore_volume_meso_value",
    "Macro_pore_volume_(cm3/g)": "pore_volume_macro_value",
    "Average_pore_diameter_(Angstrom)": "average_pore_diameter_(angstrom)",
    "SSA_(m2/g)": "ssa_(m2/g)_avg",
    "pHPZC": "phpzc",
    "Zeta_potential_(mV)": "zeta_potential_value",
    "Zeta_potential_pH": "zeta_potential_pH",
    "iodine_number_mg/g": "iodine_number_mg/g",
}


NUMERIC_PROPERTIES = (
    NumericProperty(
        "Nominal_grain_size_value_mm",
        "Nominal_grain_size_value",
        "mm",
        "grain_size",
        "Nominal_grain_size_unit",
    ),
    NumericProperty(
        "Average_pore_diameter_(Angstrom)",
        "average_pore_diameter_(angstrom)",
        "Angstrom",
    ),
    NumericProperty("SSA_(m2/g)", "ssa_(m2/g)", "m2/g"),
    NumericProperty("pHPZC", "phpzc", "pH"),
    NumericProperty("Total_pore_percentage", "porosity_total_value", "reported"),
    NumericProperty("Micro_pore_percentage", "porosity_micro_value", "reported"),
    NumericProperty("Meso_pore_percentage", "porosity_meso_value", "reported"),
    NumericProperty("Macro_pore_percentage", "porosity_macro_value", "reported"),
    NumericProperty(
        "Total_pore_volume_(cm3/g)",
        "pore_volume_total_value",
        "cm3/g",
        "pore_volume",
        "pore_volume_unit",
    ),
    NumericProperty(
        "Micro_pore_volume_(cm3/g)",
        "pore_volume_micro_value",
        "cm3/g",
        "pore_volume",
        "pore_volume_unit",
    ),
    NumericProperty(
        "Meso_pore_volume_(cm3/g)",
        "pore_volume_meso_value",
        "cm3/g",
        "pore_volume",
        "pore_volume_unit",
    ),
    NumericProperty(
        "Macro_pore_volume_(cm3/g)",
        "pore_volume_macro_value",
        "cm3/g",
        "pore_volume",
        "pore_volume_unit",
    ),
    NumericProperty("element_C_value", "element_C_value", "reported %"),
    NumericProperty("element_N_value", "element_N_value", "reported %"),
    NumericProperty("element_O_value", "element_O_value", "reported %"),
    NumericProperty("Zeta_potential_(mV)", "zeta_potential_value", "mV"),
    NumericProperty("Zeta_potential_pH", "zeta_potential_pH", "pH"),
    NumericProperty(
        "ion_exchange_capacity_value_meq/g",
        "ion_exchange_capacity_value",
        "meq/g",
        "exchange_mass",
        "ion_exchange_capacity_unit",
    ),
    NumericProperty(
        "ion_exchange_capacity_value_meq/L",
        "ion_exchange_capacity_value",
        "meq/L",
        "exchange_volume",
        "ion_exchange_capacity_unit",
    ),
)

CATEGORICAL_PROPERTIES = (
    ("Adsorbent_category", "adsorbent_category"),
    ("Adsorbent_subcategory", "adsorbent_subcategory"),
    ("AC_raw_material", "AC_raw_material"),
    ("AC_activation_method", "AC_activation_method"),
    ("Polymer_matrix", "polymer_matrix_normalized"),
    ("Elemental_composition_method", "elemental_composition_method"),
    ("Pore_class", "pore_class_normalized"),
)

AUTHORITATIVE_PROPERTY_COLUMNS = [
    "Product_key",
    "Name_Commercial",
    "Adsorbent_category",
    "Adsorbent_subcategory",
    "AC_raw_material",
    "AC_activation_method",
    "Polymer_matrix",
    "Functional_group",
    "Pore_class",
    "Nominal_grain_size_value_mm",
    "SSA_(m2/g)",
    "Total_pore_volume_(cm3/g)",
    "Micro_pore_volume_(cm3/g)",
    "Meso_pore_volume_(cm3/g)",
    "Macro_pore_volume_(cm3/g)",
    "Average_pore_diameter_(Angstrom)",
    "ion_exchange_capacity_value_meq/g",
    "ion_exchange_capacity_value_meq/L",
    "element_C_value",
    "element_N_value",
    "element_O_value",
    "pHPZC",
    "Zeta_potential_(mV)",
    "Zeta_potential_pH",
    "iodine_number_mg/g",
]

AUTHORITATIVE_AUDIT_COLUMNS = [
    "online_row_index",
    "online_name",
    "resolved_Product_key",
    "match_status",
    "match_method",
    "literature_name",
    "action",
    "fields_filled",
    "fields_overridden",
    "fields_unchanged",
    "fields_conflicted",
    "unmatched_candidates",
    "vendor",
    "data_sources",
    "vendor_verification_notes",
    "online_row_count_for_key",
    "notes",
]

ONLINE_TO_AUTHORITATIVE_COLUMN_MAP = {
    "Adsorbent_category": "Adsorbent_category",
    "Adsorbent_subcategory": "Adsorbent_subcategory",
    "Nominal_grain_size_value_mm": "Nominal_grain_size_value_mm",
    "Raw_Material": "AC_raw_material",
    "Activation_method": "AC_activation_method",
    "Polymer_matrix": "Polymer_matrix",
    "Functional_group": "Functional_group",
    "ion_exchange_capacity_value_meq/g": "ion_exchange_capacity_value_meq/g",
    "ion_exchange_capacity_value_mmol/g": "ion_exchange_capacity_value_meq/g",
    "ion_exchange_capacity_value_meq/L": "ion_exchange_capacity_value_meq/L",
    "Pore_class": "Pore_class",
    "SSA_(m2/g)": "SSA_(m2/g)",
    "Total_pore_volume_(cm3/g)": "Total_pore_volume_(cm3/g)",
    "Micro_pore_volume_(cm3/g)": "Micro_pore_volume_(cm3/g)",
    "Meso_pore_volume_(cm3/g)": "Meso_pore_volume_(cm3/g)",
    "Macro_pore_volume_(cm3/g)": "Macro_pore_volume_(cm3/g)",
    "Average_pore_diameter_(Angstrom)": "Average_pore_diameter_(Angstrom)",
    "element_C_value": "element_C_value",
    "element_N_value": "element_N_value",
    "element_O_value": "element_O_value",
    "pHPZC": "pHPZC",
    "Zeta_potential_(mV)": "Zeta_potential_(mV)",
    "Zeta_potential_pH": "Zeta_potential_pH",
    "iodine_number_mg/g": "iodine_number_mg/g",
}


def _text(value: Any) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    text = str(value).strip()
    return "" if text.casefold() in NULL_LABELS else text


def _collapse_spaces(text: str) -> str:
    return " ".join(text.replace("\u00a0", " ").split())


def _clean_product_text(value: Any) -> str:
    text = _text(value)
    text = (
        text.replace("\u00ae", "")
        .replace("\u2122", "")
        .replace("\u00a9", "")
    )
    text = unicodedata.normalize("NFKC", text)
    text = (
        text.replace("\u00ae", "")
        .replace("\u2122", "")
        .replace("\u00a9", "")
        .replace("\u00d7", "x")
        .replace("Ã‚Â®", "")
        .replace("Ã¢â€žÂ¢", "")
        .replace("Ã‚Â©", "")
        .replace("Ãƒâ€šÃ‚Â®", "")
        .replace("ÃƒÂ¢Ã¢â‚¬Å¾Ã‚Â¢", "")
        .replace("Ãƒâ€šÃ‚Â©", "")
    )
    return text


def normalize_product_key(value: Any) -> str:
    """Formatting-insensitive product key; does not collapse real variants."""
    text = _clean_product_text(value)
    text = (
        text.replace("Â®", "")
        .replace("â„¢", "")
        .replace("Â©", "")
        .replace("Ã‚Â®", "")
        .replace("Ã¢â€žÂ¢", "")
        .replace("Ã‚Â©", "")
    )
    text = re.sub(r"[\s\-_]+", "", text)
    return text.upper()


def clean_product_display(value: Any) -> str:
    text = _clean_product_text(value)
    text = (
        text.replace("Â®", "")
        .replace("â„¢", "")
        .replace("Â©", "")
        .replace("Ã‚Â®", "")
        .replace("Ã¢â€žÂ¢", "")
        .replace("Ã‚Â©", "")
    )
    return _collapse_spaces(text).strip(" -_,")

def canonical_product_key(name: Any) -> str:
    """Normalize a commercial adsorbent/product label for property matching."""
    if not isinstance(name, str):
        return ""

    text = _clean_product_text(name)
    text = (
        text.replace("Â®", "")
        .replace("â„¢", "")
        .replace("Â©", "")
        .replace("Ã‚Â®", "")
        .replace("Ã¢â€žÂ¢", "")
        .replace("Ã‚Â©", "")
    )

    # Vendor strings sometimes wrap the actual product in parentheses.
    match = re.search(r"\(([^)]*)\)", text)
    if match:
        text = match.group(1)

    text = re.sub(r"[\s\-_]+", "", text)
    return text.upper()


def _read_csv_rows(path: Path) -> list[dict[str, str]]:
    last_error: UnicodeDecodeError | None = None
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            with path.open("r", newline="", encoding=encoding) as handle:
                return list(csv.DictReader(handle))
        except UnicodeDecodeError as exc:
            last_error = exc
    raise UnicodeError(f"Could not decode {path}: {last_error}")


def load_adsorbent_map(path: Path) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for row in _read_csv_rows(path):
        raw_key = _text(row.get("key"))
        canonical = clean_product_display(row.get("canonical"))
        key = normalize_product_key(raw_key)
        if not key or not canonical:
            continue
        previous = mapping.get(key)
        if previous and previous != canonical:
            raise ValueError(
                f"Conflicting canonical values for normalized map key {key!r}: "
                f"{previous!r} and {canonical!r}"
            )
        mapping[key] = canonical
    return mapping


def _resolve_mapped_product(
    value: Any,
    mapping: dict[str, str],
    source: str,
) -> tuple[str, str, str]:
    raw_key = normalize_product_key(value)
    if not raw_key:
        return "", "", "unresolved"
    canonical = mapping.get(raw_key)
    if canonical:
        return normalize_product_key(canonical), clean_product_display(canonical), f"map:{source}"
    return raw_key, clean_product_display(value), f"direct:{source}"


def _is_extraction_missing_placeholder(value: Any) -> bool:
    if not isinstance(value, str):
        return False
    text = value.strip().lower()
    text = text.removeprefix("[").removesuffix("]").strip()
    return text in {
        "value not provided in chunk",
        "value not available in chunk",
        "not provided in chunk",
    }


def _blank(value: Any) -> bool:
    return (
        value is None
        or (isinstance(value, float) and math.isnan(value))
        or (isinstance(value, str) and not value.strip())
        or _is_extraction_missing_placeholder(value)
    )


_UNMATCHED_ADSORBENT_WARNING_PREFIX = (
    "No matching adsorbent property record for adsorbent_id="
)
_COMMERCIAL_MATCH_TRACE = "Commercial adsorbent property match resolved prior merge warning"


def _clear_resolved_commercial_property_warning(df_in: pd.DataFrame, idx: object) -> None:
    """Remove only the merge warning made obsolete by an injected property record."""
    if "warning_message" not in df_in.columns:
        return
    adsorbent_id = _text(df_in.at[idx, "adsorbent_id"]) if "adsorbent_id" in df_in.columns else ""
    if not adsorbent_id:
        return
    resolved_warning = f"{_UNMATCHED_ADSORBENT_WARNING_PREFIX}{adsorbent_id}"
    warning = _text(df_in.at[idx, "warning_message"])
    retained = [
        part.strip()
        for part in re.split(r";|\n", warning)
        if part.strip() and part.strip() != resolved_warning
    ]
    if len(retained) == len([part for part in re.split(r";|\n", warning) if part.strip()]):
        return

    df_in.at[idx, "warning_message"] = "; ".join(retained)
    if "normalization_trace" not in df_in.columns:
        df_in["normalization_trace"] = ""
    existing_trace = _text(df_in.at[idx, "normalization_trace"])
    if _COMMERCIAL_MATCH_TRACE not in existing_trace:
        df_in.at[idx, "normalization_trace"] = (
            f"{existing_trace}; {_COMMERCIAL_MATCH_TRACE}"
            if existing_trace
            else _COMMERCIAL_MATCH_TRACE
        )


ADSORBENT_CATEGORICAL_RAW_TRACE_FIELDS = (
    "adsorbent_subcategory",
    "AC_raw_material",
    "functional_group",
    "polymer_matrix",
)


def preserve_raw_adsorbent_categorical_fields(df_in: pd.DataFrame) -> None:
    """Retain reported categorical values before adsorbent enrichment changes them.

    The ``*_raw`` columns belong to ``Standardized_Full`` only. They let the
    normalized model fields be audited without leaking raw variants into
    ``Equilibrium_Data``.
    """
    for column in ADSORBENT_CATEGORICAL_RAW_TRACE_FIELDS:
        if column not in df_in.columns:
            continue
        raw_column = f"{column}_raw"
        if raw_column not in df_in.columns:
            df_in[raw_column] = df_in[column].astype("object").copy()


def _missing_or_unclassified(value: Any) -> bool:
    if _blank(value):
        return True
    return isinstance(value, str) and value.strip().lower() == "unclassified"


def _ensure_columns(df: pd.DataFrame, cols: Iterable[str], default: Any = "") -> None:
    missing: list[str] = []
    for col in cols:
        if col not in df.columns and col not in missing:
            missing.append(col)
    if missing:
        additions = pd.DataFrame(
            {col: pd.Series(default, index=df.index, dtype="object") for col in missing}
        )
        df._update_inplace(pd.concat([df, additions], axis=1))
    for col in cols:
        if col in df.columns:
            df[col] = df[col].astype("object")


def _norm_adsorbent_key(value: Any) -> str:
    """Name-map key normalization used for commercial adsorbent aliases."""
    if not isinstance(value, str):
        return ""
    text = _clean_product_text(value)
    text = text.replace("Ã‚Â®", "").replace("Ã¢â€žÂ¢", "").replace("Ã‚Â©", "")
    match = re.search(r"\(([^)]*)\)", text)
    if match:
        text = match.group(1)
    text = re.sub(r"[\s\-_]+", "", text)
    return text.upper()


def _load_adsorbent_name_map(path: Path) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for row in _read_csv_rows(path):
        key = _norm_adsorbent_key(row.get("key", ""))
        canonical = row.get("canonical", "")
        if canonical is None:
            canonical = ""
        canonical = str(canonical)
        if key:
            mapping[key] = canonical
    return mapping


def standardize_commercial_adsorbent_names(
    df: pd.DataFrame,
    mapping_path: Path,
) -> None:
    """Canonicalize explicit commercial aliases in ``name_commercial``.

    ``adsorbent_id`` is considered only through the curated alias map.  This
    lets a performance-only identifier such as ``PFA694E`` resolve to its
    commercial product without treating arbitrary study-local IDs as brands.
    """
    source_cols = [
        "name_commercial",
        "name_abbreviation",
        "name_full",
        "adsorbent_id",
    ]
    if "name_commercial" not in df.columns:
        df["name_commercial"] = ""
    if "name_commercial_raw" not in df.columns:
        df["name_commercial_raw"] = df["name_commercial"]

    mapping = _load_adsorbent_name_map(mapping_path)
    for idx, row in df.iterrows():
        canonical = None
        for col in source_cols:
            value = row.get(col, "")
            if not isinstance(value, str) or not value.strip():
                continue
            canonical = mapping.get(_norm_adsorbent_key(value))
            if canonical:
                break
        if canonical:
            df.at[idx, "name_commercial"] = canonical


def cleanup_adsorbent_subcategory(df_in: pd.DataFrame) -> None:
    """
    Clean up activated-carbon subcategories and move pore/raw-material hints
    into their dedicated fields.
    """
    _ensure_columns(
        df_in,
        ("adsorbent_category", "adsorbent_subcategory", "pore_class", "AC_raw_material"),
        "",
    )

    def _norm(value: Any) -> str:
        if not isinstance(value, str):
            return ""
        text = unicodedata.normalize("NFKC", value)
        text = text.replace("\u2013", "-").replace("\u2014", "-")
        text = re.sub(r"\s+", " ", text)
        return text.strip().lower()

    mask_ac = (
        df_in["adsorbent_category"]
        .astype(str)
        .str.strip()
        .str.lower()
        .eq("activated carbon")
    )
    sub_norm = df_in["adsorbent_subcategory"].apply(_norm)

    blank_mask = mask_ac & sub_norm.eq("")
    df_in.loc[blank_mask, "adsorbent_subcategory"] = "unclassified"

    both_str = _norm("powdered activated carbon (PAC); granular activated carbon (GAC)")
    both_mask = mask_ac & sub_norm.eq(both_str)
    df_in.loc[both_mask, "adsorbent_subcategory"] = "unclassified"

    pore_raw_re = re.compile(
        r"^(?P<pore>(?:micro|meso|macro)porous)\s*;\s*(?P<raw>[a-z]+)-based\.?$"
    )
    for idx in df_in.index[mask_ac]:
        text = _norm(df_in.at[idx, "adsorbent_subcategory"])
        match = pore_raw_re.match(text)
        if not match:
            continue
        pore = match.group("pore")
        raw = match.group("raw")
        if not str(df_in.at[idx, "pore_class"]).strip():
            df_in.at[idx, "pore_class"] = pore
        if not str(df_in.at[idx, "AC_raw_material"]).strip():
            df_in.at[idx, "AC_raw_material"] = raw
        df_in.at[idx, "adsorbent_subcategory"] = "unclassified"


def infer_ac_raw_material_from_name(df_in: pd.DataFrame) -> None:
    """Infer raw material for activated carbons from ``name_full``."""
    _ensure_columns(df_in, ("adsorbent_category", "AC_raw_material", "name_full"), "")

    def _norm_text(value: Any) -> str:
        if not isinstance(value, str):
            return ""
        text = unicodedata.normalize("NFKC", value)
        text = text.replace("Ã‚Â®", "").replace("Ã¢â€žÂ¢", "").replace("Ã‚Â©", "")
        text = text.replace("\u2013", "-").replace("\u2014", "-")
        text = re.sub(r"[-_/]+", " ", text)
        text = re.sub(r"\s+", " ", text)
        return text.strip().lower()

    raw_material_keywords = [
        ("bituminous coal", [r"\bbituminous\s+coal\b", r"\bbituminous\b"]),
        ("bamboo", [r"\bbamboo\b"]),
        ("wood", [r"\bwood\b"]),
        ("nutshell", [r"\bnutshell\b", r"\bnut\s+shells?\b"]),
        ("charcoal", [r"\bcharcoal\b"]),
        ("coconut", [r"\bcoconut\b"]),
        ("coal", [r"\bcoal\b"]),
    ]
    compiled_patterns = [
        (canonical, [re.compile(pattern) for pattern in patterns])
        for canonical, patterns in raw_material_keywords
    ]

    for idx, row in df_in.iterrows():
        category = str(row.get("adsorbent_category", "")).strip().lower()
        if category != "activated carbon":
            continue
        if not _blank(row.get("AC_raw_material")):
            continue
        name_full = row.get("name_full", "")
        if not isinstance(name_full, str) or not name_full.strip():
            continue
        text = _norm_text(name_full)
        for canonical, pattern_list in compiled_patterns:
            if any(pattern.search(text) for pattern in pattern_list):
                df_in.at[idx, "AC_raw_material"] = canonical
                break


def _categorical_key(value: Any) -> str:
    """Return a formatting-insensitive key for a reported categorical value."""
    text = unicodedata.normalize("NFKC", _text(value))
    text = (
        text.replace("Ã¢â‚¬â€œ", "-")
        .replace("Ã¢â‚¬â€", "-")
        .replace("\u2013", "-")
        .replace("\u2014", "-")
    )
    text = re.sub(r"[-_/]+", " ", text)
    return _collapse_spaces(text).casefold()


RAW_MATERIAL_ALIASES = {
    "agglomerated bituminous coal": "bituminous coal",
    "anthracite coal": "anthracite coal",
    "bamboo": "bamboo",
    "bituminous coal": "bituminous coal",
    "bituminous coal (agglomerated)": "bituminous coal",
    "charcoal": "charcoal",
    "coal": "coal",
    "coconut": "coconut",
    "coconut husk": "coconut",
    "coconut shell": "coconut",
    "lignite": "lignite",
    "lignite coal": "lignite",
    "nut shell": "nutshell",
    "nutshell": "nutshell",
    "purified bituminous coal": "bituminous coal",
    "reagglomerated bituminous coal": "bituminous coal",
    "sub bituminous coal": "bituminous coal",
    "subbituminous coal": "bituminous coal",
    "wood": "wood",
}


def normalize_ac_raw_materials(df_in: pd.DataFrame) -> None:
    """Canonicalize activated-carbon precursor labels.

    The coconut plant variants (coconut, shell, and husk) are treated as one
    precursor class. Bituminous coal includes its sub-bituminous and purified
    reported variants.
    """
    _ensure_columns(df_in, ("adsorbent_category", "AC_raw_material"), "")
    mask_ac = (
        df_in["adsorbent_category"].astype(str).str.strip().str.casefold()
        == "activated carbon"
    )
    for idx in df_in.index[mask_ac]:
        key = _categorical_key(df_in.at[idx, "AC_raw_material"])
        if not key:
            continue
        df_in.at[idx, "AC_raw_material"] = RAW_MATERIAL_ALIASES.get(key, key)


AC_SUBCATEGORY_ALIASES = {
    "gac": "granular activated carbon",
    "granular activated carbon": "granular activated carbon",
    "granular activated carbon (gac)": "granular activated carbon",
    "pac": "powdered activated carbon",
    "powdered activated carbon": "powdered activated carbon",
    "powdered activated carbon (pac)": "powdered activated carbon",
}


def normalize_adsorbent_subcategories(df_in: pd.DataFrame) -> None:
    """Use full, stable AC subcategory labels instead of PAC/GAC aliases."""
    _ensure_columns(df_in, ("adsorbent_category", "adsorbent_subcategory"), "")
    mask_ac = (
        df_in["adsorbent_category"].astype(str).str.strip().str.casefold()
        == "activated carbon"
    )
    for idx in df_in.index[mask_ac]:
        key = _categorical_key(df_in.at[idx, "adsorbent_subcategory"])
        canonical = AC_SUBCATEGORY_ALIASES.get(key)
        if canonical:
            df_in.at[idx, "adsorbent_subcategory"] = canonical


def infer_adsorbent_category_from_name(df_in: pd.DataFrame) -> None:
    """Infer adsorbent category/subcategory from name fields."""
    _ensure_columns(
        df_in,
        (
            "adsorbent_category",
            "adsorbent_subcategory",
            "name_full",
            "name_commercial",
            "name_abbreviation",
        ),
        "",
    )

    def _norm_name_cell(value: Any) -> str:
        if not isinstance(value, str):
            return ""
        text = unicodedata.normalize("NFKC", value)
        text = text.replace("Ã‚Â®", "").replace("Ã¢â€žÂ¢", "").replace("Ã‚Â©", "")
        text = re.sub(r"\s+", " ", text)
        return text.strip()

    for idx, row in df_in.iterrows():
        if not _blank(row.get("adsorbent_category")):
            continue
        raw_names = [
            row.get("name_full", ""),
            row.get("name_commercial", ""),
            row.get("name_abbreviation", ""),
        ]
        names = [
            _norm_name_cell(name)
            for name in raw_names
            if isinstance(name, str) and name.strip()
        ]
        if not names:
            continue

        names_lower = [name.lower() for name in names]
        text_all = " ".join(names_lower)

        if re.search(r"strong(?:ly)?\s+bas(?:e|ic)\s+anion\s+exchange\s+resin", text_all):
            df_in.at[idx, "adsorbent_category"] = "Ion exchange resin"
            df_in.at[idx, "adsorbent_subcategory"] = "anion strong base"
            continue

        if re.search(r"weak(?:ly)?\s+bas(?:e|ic)\s+anion\s+exchange\s+resin", text_all):
            df_in.at[idx, "adsorbent_category"] = "Ion exchange resin"
            df_in.at[idx, "adsorbent_subcategory"] = "anion weak base"
            continue

        if "anion exchange resin" in text_all:
            df_in.at[idx, "adsorbent_category"] = "Ion exchange resin"
            df_in.at[idx, "adsorbent_subcategory"] = "anion"
            continue

        is_ae_resin = any(re.search(r"\bae\s+resin\b", name) for name in names_lower)
        is_miex = any("miex" in name for name in names_lower)
        if is_ae_resin or is_miex:
            df_in.at[idx, "adsorbent_category"] = "Ion exchange resin"
            df_in.at[idx, "adsorbent_subcategory"] = "anion"
            continue

        is_cdp = (
            "cyclodextrin" in text_all
            or re.search(r"\bcdp\b", text_all)
            or "decafluorobiphenyl polymer" in text_all
        )
        if is_cdp:
            df_in.at[idx, "adsorbent_category"] = "Cyclodextrin polymer"
            df_in.at[idx, "adsorbent_subcategory"] = "unclassified"
            continue

        has_gac = any(re.search(r"\bgac\b", name, flags=re.I) for name in names)
        has_pac = any(re.search(r"\bpac\b", name, flags=re.I) for name in names)
        if has_gac:
            df_in.at[idx, "adsorbent_category"] = "Activated carbon"
            df_in.at[idx, "adsorbent_subcategory"] = "GAC"
            continue
        if has_pac:
            df_in.at[idx, "adsorbent_category"] = "Activated carbon"
            df_in.at[idx, "adsorbent_subcategory"] = "PAC"
            continue

        if (
            re.search(r"\bactivated\s+carbon\b", text_all)
            or re.search(r"\bactivated\s+charcoal\b", text_all)
            or re.search(r"\bactivated\s+carbon\s+fiber\b", text_all)
            or re.search(r"\bacf\d*\b", text_all)
        ):
            df_in.at[idx, "adsorbent_category"] = "Activated carbon"
            if _blank(row.get("adsorbent_subcategory")):
                df_in.at[idx, "adsorbent_subcategory"] = "unclassified"
            continue

        if re.search(r"\bac\b", text_all):
            df_in.at[idx, "adsorbent_category"] = "Activated carbon"
            if _blank(row.get("adsorbent_subcategory")):
                df_in.at[idx, "adsorbent_subcategory"] = "unclassified"
            continue

        if any(name == "resin" for name in names_lower):
            df_in.at[idx, "adsorbent_category"] = "Unknown"
            df_in.at[idx, "adsorbent_subcategory"] = "unclassified"


_NUM_RE = re.compile(
    r"[-+]?(?:\d+(?:\.\d+)?|\.\d+)(?:[eE][-+]?\d+)?"
)
_SEP_RE = re.compile(r"\s*(?:Ã¢â‚¬â€œ|-|Ã¢â‚¬â€|to|Ã‚Â±|\+/-|\+Ã¢Ââ€ž-)\s*")


def _coerce_number(value: Any) -> float | None:
    if _blank(value):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip()
    full_num = _NUM_RE.fullmatch(text)
    if full_num:
        try:
            return float(full_num.group(0))
        except Exception:
            pass

    if "," in text or ";" in text:
        parts = [part.strip() for part in re.split(r"[;,]", text) if part.strip()]
        nums_list: list[float] = []
        for part in parts:
            if _SEP_RE.search(part):
                nums_list = []
                break
            match = _NUM_RE.search(part)
            if match:
                try:
                    nums_list.append(float(match.group(0)))
                except Exception:
                    pass
        if len(nums_list) >= 2:
            return sum(nums_list) / len(nums_list)

    parts = _SEP_RE.split(text)
    nums: list[float] = []
    for part in parts:
        match = _NUM_RE.search(part)
        if match:
            try:
                nums.append(float(match.group(0)))
            except Exception:
                pass
    if not nums:
        return None
    if len(nums) == 1:
        return nums[0]
    return (nums[0] + nums[1]) / 2.0


def infer_ac_subcategory_from_nominal_grain_size(df_in: pd.DataFrame) -> None:
    """
    Infer PAC vs GAC for activated carbons using the standardized nominal
    grain size in mm. Only fills blank/unclassified subcategories.
    """
    _ensure_columns(
        df_in,
        ("adsorbent_category", "adsorbent_subcategory", "Nominal_grain_size_value_mm"),
        "",
    )

    mask_ac = (
        df_in["adsorbent_category"]
        .astype(str)
        .str.strip()
        .str.lower()
        .eq("activated carbon")
    )
    for idx, row in df_in.loc[mask_ac].iterrows():
        sub_raw = str(row.get("adsorbent_subcategory", "")).strip().lower()
        if sub_raw not in {"", "unclassified"}:
            continue

        size_raw = row.get("Nominal_grain_size_value_mm", "")
        if _blank(size_raw):
            continue

        size_mm = None
        if isinstance(size_raw, str) and "&&" in size_raw:
            parts = [part.strip() for part in size_raw.split("&&") if part.strip()]
            nums = [_coerce_number(part) for part in parts]
            nums = [num for num in nums if num is not None]
            if nums:
                size_mm = sum(nums) / len(nums)
        else:
            size_mm = _coerce_number(size_raw)

        if size_mm is None:
            continue
        if size_mm <= 0.18:
            df_in.at[idx, "adsorbent_subcategory"] = "PAC"
        elif size_mm >= 0.20:
            df_in.at[idx, "adsorbent_subcategory"] = "GAC"


_GRAIN_SIZE_SUFFIX_RE = re.compile(
    r"\(\s*[\d<>=,\-Ã¢â‚¬â€œ ]+(?:Ã‚Âµm|um|mm|mesh)\s*\)",
    flags=re.IGNORECASE,
)


def _has_grain_size_suffix(row: pd.Series) -> bool:
    for col in ("name_abbreviation", "name_commercial"):
        if col in row.index:
            value = row[col]
            if isinstance(value, str) and _GRAIN_SIZE_SUFFIX_RE.search(value):
                return True
    return False


def inject_commercial_adsorbent_properties(
    df_in: pd.DataFrame,
    config: AdsorbentProcessingConfig,
) -> None:
    """
    Inject commercial adsorbent properties from online and consolidated
    property workbooks using canonical product-key matching.
    """
    if not (config.enable_online_injection or config.enable_database_injection):
        return

    mapping = load_adsorbent_map(config.mapping_path)

    def _load_props_index(path: Path, sheet: str) -> pd.DataFrame | None:
        if not path:
            return None
        if not os.path.exists(path):
            print(f"[WARN] Commercial properties file not found: {path}")
            return None
        props = pd.read_excel(path, sheet_name=sheet)
        if props.empty:
            return None
        if "Product_key" not in props.columns:
            def _pk_from_props_row(row: pd.Series) -> str:
                for col in ("Name_Commercial", "Name_Abbreviation"):
                    value = row.get(col, "")
                    if isinstance(value, str) and value.strip():
                        key, _, _ = _resolve_mapped_product(value, mapping, col)
                        if key:
                            return key
                return ""

            props["Product_key"] = props.apply(_pk_from_props_row, axis=1)
        props["Product_key"] = props["Product_key"].map(
            lambda value: _resolve_mapped_product(value, mapping, "Product_key")[0]
        )
        props["Product_key"] = props["Product_key"].astype(str).str.strip().str.lower()
        props = props[props["Product_key"] != ""].copy()
        if props.empty:
            return None
        props = (
            props.sort_values("Product_key")
            .drop_duplicates("Product_key", keep="first")
            .set_index("Product_key")
        )
        return props

    props_online = (
        _load_props_index(config.online_properties_path, config.online_properties_sheet)
        if config.enable_online_injection
        else None
    )
    props_db = (
        _load_props_index(config.database_properties_path, config.database_properties_sheet)
        if config.enable_database_injection
        else None
    )
    if props_online is None and props_db is None:
        return

    name_cols = [
        col
        for col in ("name_commercial", "name_full", "name_abbreviation")
        if col in df_in.columns
    ]
    if not name_cols:
        return

    df_in["__product_key"] = ""
    for idx, row in df_in.iterrows():
        for col in name_cols:
            value = row.get(col, "")
            if isinstance(value, str) and value.strip():
                key, _, _ = _resolve_mapped_product(value, mapping, col)
                if key:
                    df_in.at[idx, "__product_key"] = key.lower()
                    break

    for idx, row in df_in.iterrows():
        product_key = str(row.get("__product_key", "")).strip().lower()
        if not product_key:
            continue
        src_db = (
            props_db.loc[product_key]
            if props_db is not None and product_key in props_db.index
            else None
        )
        src_online = (
            props_online.loc[product_key]
            if props_online is not None and product_key in props_online.index
            else None
        )
        if src_db is None and src_online is None:
            continue
        has_injectable_property = any(
            source is not None
            and any(
                source_column in source.index and not _blank(source.get(source_column))
                for source_column in ADSORBENT_PROPERTY_COLUMN_MAP
            )
            for source in (src_db, src_online)
        )
        if has_injectable_property:
            _clear_resolved_commercial_property_warning(df_in, idx)

        skip_grain_size = _has_grain_size_suffix(row)

        def _assign_value(dst_col: str, new_value: Any) -> None:
            if _blank(new_value):
                return
            try:
                if dst_col in df_in.columns and df_in[dst_col].dtype != "object":
                    df_in[dst_col] = df_in[dst_col].astype("object")
            except Exception:
                pass
            df_in.at[idx, dst_col] = new_value

        if src_db is not None:
            for src_col, dst_col in ADSORBENT_PROPERTY_COLUMN_MAP.items():
                if src_col not in src_db.index or dst_col not in df_in.columns:
                    continue
                if dst_col == "Nominal_grain_size_value_mm" and skip_grain_size:
                    continue
                cur_val = df_in.at[idx, dst_col]
                if not _missing_or_unclassified(cur_val):
                    continue
                new_value = src_db.get(src_col, None)
                if _blank(new_value):
                    continue
                _assign_value(dst_col, new_value)

        if src_online is not None:
            for src_col, dst_col in ADSORBENT_PROPERTY_COLUMN_MAP.items():
                if src_col not in src_online.index or dst_col not in df_in.columns:
                    continue
                if dst_col == "Nominal_grain_size_value_mm" and skip_grain_size:
                    continue
                new_value = src_online.get(src_col, None)
                if _blank(new_value):
                    continue
                cur_val = df_in.at[idx, dst_col]
                if src_col == "Adsorbent_category":
                    _assign_value(dst_col, new_value)
                    continue
                if src_col == "Raw_Material":
                    _assign_value(dst_col, new_value)
                    continue
                if src_col == "Adsorbent_subcategory":
                    cat_val = (
                        df_in.at[idx, "adsorbent_category"]
                        if "adsorbent_category" in df_in.columns
                        else ""
                    )
                    cat_val_norm = str(cat_val).strip().lower()
                    if cat_val_norm == "activated carbon":
                        if not _missing_or_unclassified(cur_val):
                            continue
                        _assign_value(dst_col, new_value)
                        continue
                    _assign_value(dst_col, new_value)
                    continue
                if not _missing_or_unclassified(cur_val):
                    continue
                _assign_value(dst_col, new_value)

    df_in.drop(columns=["__product_key"], inplace=True, errors="ignore")


def materialize_adsorbent_identity_fields(df_in: pd.DataFrame) -> None:
    """Add conservative material-identity fields after name normalization.

    Canonical commercial names identify a material across studies. Generic or
    study-defined adsorbent labels remain study-local; absent identifiers are
    explicit unknowns rather than implicitly merged materials.
    """
    _ensure_columns(
        df_in,
        ("study_no", "adsorbent_id", "name_commercial", "source_row_index"),
        "",
    )
    statuses: list[str] = []
    bases: list[str] = []
    identity_keys: list[str] = []
    instance_keys: list[str] = []
    for position, (_, row) in enumerate(df_in.iterrows()):
        study = normalize_product_key(_text(row.get("study_no"))) or "missing-study"
        adsorbent = normalize_product_key(_text(row.get("adsorbent_id")))
        commercial = normalize_product_key(_text(row.get("name_commercial")))
        source_row = _text(row.get("source_row_index")) or str(position)
        instance_key = (
            f"study::{study}|adsorbent::{adsorbent}"
            if adsorbent
            else f"study::{study}|source-row::{source_row}"
        )
        if commercial:
            statuses.append("specific")
            bases.append("canonical name_commercial")
            identity_keys.append(f"specific::{commercial}")
        elif adsorbent:
            statuses.append("generic_study_local")
            bases.append("study_no + adsorbent_id")
            identity_keys.append(instance_key)
        else:
            statuses.append("unknown")
            bases.append("study_no + source_row_index")
            identity_keys.append(instance_key)
        instance_keys.append(instance_key)

    df_in[ADSORBENT_IDENTITY_STATUS_COLUMN] = statuses
    df_in[ADSORBENT_IDENTITY_BASIS_COLUMN] = bases
    df_in[ADSORBENT_IDENTITY_KEY_COLUMN] = identity_keys
    df_in[ADSORBENT_STUDY_INSTANCE_KEY_COLUMN] = instance_keys


def enrich_adsorbent_fields(
    df_in: pd.DataFrame,
    config: AdsorbentProcessingConfig,
) -> None:
    """Apply adsorbent semantic cleanup, inference, naming, and enrichment."""
    preserve_raw_adsorbent_categorical_fields(df_in)
    cleanup_adsorbent_subcategory(df_in)
    infer_adsorbent_category_from_name(df_in)
    infer_ac_raw_material_from_name(df_in)
    standardize_commercial_adsorbent_names(df_in, config.mapping_path)
    materialize_adsorbent_identity_fields(df_in)
    _ensure_columns(df_in, ADSORBENT_PROPERTY_COLUMN_MAP.values(), "")
    inject_commercial_adsorbent_properties(df_in, config)


def materialize_nonionic_resin_iec_defaults(df_in: pd.DataFrame) -> None:
    """Set IEC features to zero for nonionic resins.

    Nonionic resins are modeled with ion-exchange resins in the Resin family,
    but have no ion-exchange capacity by definition.  A numeric zero therefore
    represents the material property more faithfully than a text value such as
    ``not applicable`` or a missing feature.  This deliberately affects only
    the model-facing IEC outputs; reported source fields remain unchanged.
    """
    iec_columns = (
        "ion_exchange_capacity_value_meq/g",
        "ion_exchange_capacity_value_meq/L",
    )
    category_column = next(
        (
            column
            for column in ("adsorbent_category", "Adsorbent_category")
            if column in df_in.columns
        ),
        None,
    )
    if category_column is None:
        return
    _ensure_columns(df_in, iec_columns, "")
    category_key = (
        df_in[category_column]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.casefold()
    )
    mask = category_key.eq("nonionic resin")
    if mask.any():
        df_in.loc[mask, list(iec_columns)] = 0.0


def normalize_polymer_matrix(value: Any, adsorbent_category: Any) -> str:
    raw = _collapse_spaces(_text(value))
    if not raw:
        return ""
    category = _text(adsorbent_category).casefold()
    if category not in {"ion exchange resin", "nonionic resin"}:
        return raw

    key = unicodedata.normalize("NFKC", raw).casefold()
    key = key.replace("â€“", "-").replace("â€”", "-")
    compact = re.sub(r"[^a-z0-9]+", " ", key).strip()

    if "phenol formaldehyde" in compact or "phenolformaldehyde" in compact:
        return "Phenol-formaldehyde"
    if "polyvinylbenzyl chloride" in compact or compact == "pvbc":
        return "Polyvinylbenzyl chloride"
    if "polyvinylpyridine" in compact or compact in {"pvpy", "pvp"}:
        return "Polyvinylpyridine (PVPy)"
    if (
        "polymethacrylate" in compact
        or "poly methyl methacrylate" in compact
        or compact == "pmma"
    ):
        return "Polymethacrylate (PMMA)"
    if (
        "polyacrylic" in compact
        or "acrylic resin" in compact
        or "acrylic polymer" in compact
        or "aliphatic acrylic polymer" in compact
        or "acrylate resin" in compact
    ):
        return "Polyacrylic"

    has_styrene = any(
        token in compact
        for token in ("polystyrene", "styrene", "styrenic", "ps dvb")
    )
    has_dvb = any(
        token in compact
        for token in ("divinylbenzene", "dvb", "cross linked", "crosslinked")
    )
    if has_styrene and has_dvb:
        return "Polystyrene-divinylbenzene (PS-DVB)"
    if (
        has_styrene
        or compact == "ps"
        or "macroporous styrene" in compact
        or "styrene resin" in compact
        or "styrene polymer" in compact
    ):
        return "Polystyrene (PS)"
    return raw


_NON_FUNCTIONAL_GROUP_PHASES = frozenset(
    {
        "alumina",
        "ferric oxide",
        "ferrous oxide",
        "ferrihydrite",
        "goethite",
        "hematite",
        "iron oxide",
        "magnetite",
        "silica",
        "silicon dioxide",
        "sio2",
        "zeolite",
    }
)

_INORGANIC_COMPOUND_PATTERN = re.compile(
    r"\b(?:aluminum|aluminium|calcium|cobalt|copper|ferric|ferrous|iron|magnesium|"
    r"manganese|nickel|sodium|titanium|zinc)(?: (?:[ivxlcdm]+|\d+))? "
    r"(?:di)?(?:oxides?|hydroxides?|carbonates?|sulfates?|sulfides?|phosphates?|chlorides?)\b"
)
_METAL_OXIDE_FORMULA_PATTERN = re.compile(r"^(?:al|co|cu|fe|mn|ni|ti|zn)\d*o\d+$")
_SIMPLE_CARBON_MOIETY_PATTERN = re.compile(
    r"^[\-\s]*(?:c(?:h|f|cl|br|i)\d*)[\-\s]*$",
    flags=re.IGNORECASE,
)

# A resin carries a genuine amino-acid ligand only when the chelating chemistry
# is named: iminodiacetic, aminophosphonic, picolylamine, and similar.  The bare
# phrase "amino acid(s)" is not such a designation.  It appears in this corpus
# only as a paraphrase of the vendor description "complex amine" (Purolite
# PFA694E), so it is routed to ``Amine`` with every other complex-amine spelling
# rather than becoming a second level for one chemistry.
_AMINO_ACID_LIGAND_PATTERN = re.compile(
    r"\b(?:iminodiacetic|iminodiacetate|iminodiacetato|"
    r"aminophosphonic|aminomethylphosphonic|aminophosphonate|"
    r"aminocarboxylic|aminocarboxylate|"
    r"glycine|glycinate|glutamic|glutamate|aspartic|aspartate|"
    r"picolylamine|bispicolylamine|"
    r"thiourea amino acid)\b"
)

FUNCTIONAL_GROUP_INDICATOR_COLUMNS = {
    "Quaternary ammonium": "contains_quaternary_ammonium",
    "Sulfonic acid": "contains_sulfonic_acid",
    "Carboxyl": "contains_carboxyl",
    "Carbonyl": "contains_carbonyl",
    "Hydroxyl": "contains_hydroxyl",
    "Lactone": "contains_lactone",
    "Pyridinic nitrogen": "contains_pyridinic_nitrogen",
    "Pyrrolic nitrogen": "contains_pyrrolic_nitrogen",
    "Amino acid": "contains_amino_acid",
    "Amine": "contains_amine",
    "Amide": "contains_amide",
    "Nitrile": "contains_nitrile",
    "Alkoxy": "contains_alkoxy",
    "Glycosidic bond": "contains_glycosidic_bond",
    "Ester": "contains_ester",
    "Benzyl chloride": "contains_benzyl_chloride",
    "Siloxane": "contains_siloxane",
    "Unspecified functional groups": "contains_unspecified_functional_groups",
}

PORE_CLASS_INDICATOR_COLUMNS = {
    "Microporous": "contains_microporous",
    "Mesoporous": "contains_mesoporous",
    "Macroporous": "contains_macroporous",
    "Gel": "contains_gel",
    "Nonporous": "contains_nonporous",
}

_PORE_CLASS_ALIASES = {
    "micropore": "Microporous",
    "microporous": "Microporous",
    "mesopore": "Mesoporous",
    "mesoporous": "Mesoporous",
    "macropore": "Macroporous",
    "macroporous": "Macroporous",
    "macroreticular": "Macroporous",
    "gel": "Gel",
    "nonporous": "Nonporous",
}


def _is_non_functional_group_label(raw: str, compact: str) -> bool:
    """Return whether a token describes a material phase or simple moiety.

    These labels can be meaningful adsorbent descriptors, but they are not
    chemical functional groups and should not become categorical levels when
    extracted into ``functional_group``.
    """
    if compact in _NON_FUNCTIONAL_GROUP_PHASES:
        return True
    if _INORGANIC_COMPOUND_PATTERN.search(compact):
        return True
    if _METAL_OXIDE_FORMULA_PATTERN.fullmatch(compact):
        return True
    if "nanoparticle" in compact or "mineral" in compact:
        return True

    # Examples include CH2 and CF3 (with optional surrounding bond dashes).
    normalized_raw = raw.replace("–", "-").replace("—", "-")
    return bool(_SIMPLE_CARBON_MOIETY_PATTERN.fullmatch(normalized_raw))


def _canonical_functional_group(token: str) -> str:
    raw = _collapse_spaces(token).strip(" ,;")
    key = unicodedata.normalize("NFKC", raw).casefold()
    key = key.replace("−", "-").replace("–", "-").replace("—", "-")
    key = re.sub(r"\s+", " ", key)
    compact = re.sub(r"[^a-z0-9+]+", " ", key).strip()

    # These labels are modeled as a categorical feature. Canonicalize chemical
    # family names rather than preserving spelling, counterion, Type-I/II, or
    # alkyl-substituent variants as separate sparse levels.
    if (
        "quaternary" in compact
        or "ammonium" in compact
        or "aminium" in compact
        or "permanently cationic" in compact
        or "cationic amine" in compact
        or "n+" in key
        or re.search(r"n\s*\^?\s*\+", key)
        or ("type i" in compact or "type ii" in compact) and "r" in compact
    ):
        return "Quaternary ammonium"
    if "sulfon" in compact or "so3" in compact:
        return "Sulfonic acid"
    if "carbox" in compact or "cooh" in compact:
        return "Carboxyl"
    if "carbonyl" in compact or key in {"c=o", "c = o"}:
        return "Carbonyl"
    if "phenol" in compact or "hydroxyl" in compact or re.search(r"\boh\b", compact):
        return "Hydroxyl"
    if "lactone" in compact:
        return "Lactone"
    if "pyridin" in compact:
        return "Pyridinic nitrogen"
    if "pyrrol" in compact:
        return "Pyrrolic nitrogen"
    if _AMINO_ACID_LIGAND_PATTERN.search(compact):
        return "Amino acid"
    if "amine" in compact or "amino" in compact:
        return "Amine"
    if "amide" in compact:
        return "Amide"
    if "nitrile" in compact:
        return "Nitrile"
    if "alkoxy" in compact:
        return "Alkoxy"
    if key in {"c-o", "c - o", "c-h", "c - h"}:
        return ""
    if "glycosidic" in compact:
        return "Glycosidic bond"
    if "trifluoroethyl" in compact:
        return ""
    if "ester" in compact:
        return "Ester"
    if "benzyl chloride" in compact:
        return "Benzyl chloride"
    if "siloxane" in compact:
        return "Siloxane"
    # Acid/base strength describes ion-exchange behavior, not a chemical
    # functional group.  It belongs in adsorbent classification rather than
    # in the functional-group field or its ML indicators.
    if re.search(r"\bweak(?:ly)?\s+(?:acid|acidic|base|basic)\b", compact):
        return ""
    if "abundant functional groups" in compact:
        return "Unspecified functional groups"
    return "" if _is_non_functional_group_label(raw, compact) else raw


def normalize_functional_groups(value: Any) -> str:
    raw = _text(value)
    if not raw:
        return ""
    tokens = re.split(r"\s*(?:;|&&|\|)\s*", raw)
    normalized: list[str] = []
    seen: set[str] = set()
    for token in tokens:
        canonical = _canonical_functional_group(token)
        key = canonical.casefold()
        if canonical and key not in seen:
            seen.add(key)
            normalized.append(canonical)
    return "; ".join(normalized)


def consolidate_functional_groups_by_identity(df_in: pd.DataFrame) -> None:
    """Resolve reported functional groups per material instead of per study row.

    ``functional_group`` reaches a performance row from the study that reported
    it, so one commercial product can carry different labels in different
    papers.  Amberlite IRA910, for example, was recorded as ``Amine`` in one
    study and ``Quaternary ammonium`` in two others, which lets an indicator
    behave as a tag for those studies rather than as a property of the resin.

    A named material is identified across studies by
    ``adsorbent_identity_key``, so its reported groups are pooled into one union
    per identity -- the rule :func:`_functional_group_union` already applies
    when the consolidated adsorbent database is built, which also brings the
    performance rows back in line with that database.  ``generic_study_local``
    and ``unknown`` identities keep their own values because those labels are
    only meaningful inside one study.

    Every reported value stays in ``functional_group_raw``, and identities whose
    studies disagreed are named in
    ``functional_group_identity_conflict`` for review.
    """
    _ensure_columns(
        df_in,
        (
            "functional_group",
            ADSORBENT_IDENTITY_KEY_COLUMN,
            ADSORBENT_IDENTITY_STATUS_COLUMN,
            FUNCTIONAL_GROUP_IDENTITY_CONFLICT_COLUMN,
        ),
        "",
    )
    identity_keys = df_in[ADSORBENT_IDENTITY_KEY_COLUMN].map(_text)
    statuses = df_in[ADSORBENT_IDENTITY_STATUS_COLUMN].map(
        lambda value: _text(value).casefold()
    )
    shared = statuses.eq("specific") & identity_keys.ne("")
    if not shared.any():
        return

    studies = (
        df_in["study_no"].map(_text)
        if "study_no" in df_in.columns
        else pd.Series("", index=df_in.index)
    )
    positions_by_identity: dict[str, list[Any]] = {}
    for position, identity_key in identity_keys[shared].items():
        positions_by_identity.setdefault(identity_key, []).append(position)

    for identity_key, positions in positions_by_identity.items():
        reported = {
            position: _text(df_in.at[position, "functional_group"])
            for position in positions
        }
        pooled = _functional_group_union(reported.values())
        if not pooled:
            continue
        df_in.loc[positions, "functional_group"] = pooled

        # Count the studies behind each reported label so a disagreement can be
        # judged without reopening every source row.  A blank study label falls
        # back to the row itself, which keeps the count a lower bound.
        studies_by_value: dict[str, set[str]] = {}
        for position, value in reported.items():
            if value:
                studies_by_value.setdefault(value, set()).add(
                    studies.at[position] or f"row::{position}"
                )
        if len(studies_by_value) < 2:
            continue
        detail = ", ".join(
            f"{value} ({len(reporting)} {'study' if len(reporting) == 1 else 'studies'})"
            for value, reporting in sorted(
                studies_by_value.items(), key=lambda item: (-len(item[1]), item[0])
            )
        )
        df_in.loc[positions, FUNCTIONAL_GROUP_IDENTITY_CONFLICT_COLUMN] = (
            f"{identity_key} pooled to '{pooled}' from {detail}"
        )


def materialize_functional_group_ml_features(df_in: pd.DataFrame) -> None:
    """Expand reported functional groups into three-state ML indicators.

    The normalized semicolon-delimited field remains in ``Standardized_Full``
    for auditability.  These flags are the model-facing representation, so a
    row with ``Amine; Amide`` has both ``contains_amine`` and
    ``contains_amide`` set to ``True``.

    For a nonblank parent field, ``False`` means that the group was not included
    in that reported label set; it does not assert physical absence.  Indicators
    remain blank only when the parent field itself is unreported.
    """
    _ensure_columns(df_in, ("functional_group",), "")
    reported = df_in["functional_group"].map(lambda value: bool(_text(value)))
    token_sets = df_in["functional_group"].map(
        lambda value: {
            _collapse_spaces(token).casefold()
            for token in re.split(r"\s*(?:;|&&|\|)\s*", _text(value))
            if _collapse_spaces(token)
        }
    )
    for label, column in FUNCTIONAL_GROUP_INDICATOR_COLUMNS.items():
        key = label.casefold()
        df_in[column] = [
            (key in tokens) if is_reported else ""
            for is_reported, tokens in zip(reported, token_sets)
        ]


def normalize_pore_class(value: Any) -> str:
    """Canonicalize a possibly multi-valued pore-class label."""
    normalized: list[str] = []
    seen: set[str] = set()
    for token in re.split(r"\s*(?:;|&&|\|)\s*", _text(value)):
        raw = _collapse_spaces(token).strip(" ,;")
        if not raw:
            continue
        key = re.sub(r"[^a-z0-9]+", "", raw.casefold())
        canonical = _PORE_CLASS_ALIASES.get(key, raw)
        canonical_key = canonical.casefold()
        if canonical_key not in seen:
            seen.add(canonical_key)
            normalized.append(canonical)
    return "; ".join(normalized)


def materialize_pore_class_ml_features(df_in: pd.DataFrame) -> None:
    """Normalize pore classes and expand them into three-state ML indicators.

    For a nonblank parent field, ``False`` means that the class was not included
    in the reported label set; it does not assert physical absence.  Indicators
    remain blank only when ``pore_class`` itself is unreported.  Macroreticular
    is canonicalized to Macroporous before the indicators are generated.
    """
    _ensure_columns(df_in, ("pore_class",), "")
    df_in["pore_class"] = df_in["pore_class"].map(normalize_pore_class)
    reported = df_in["pore_class"].map(lambda value: bool(_text(value)))

    def pore_classes(value: Any) -> set[str]:
        classes: set[str] = set()
        for token in re.split(r"\s*(?:;|&&|\|)\s*", _text(value)):
            key = re.sub(r"[^a-z0-9]+", "", _collapse_spaces(token).casefold())
            label = _PORE_CLASS_ALIASES.get(key)
            if label:
                classes.add(label)
        return classes

    class_sets = df_in["pore_class"].map(pore_classes)
    for label, column in PORE_CLASS_INDICATOR_COLUMNS.items():
        df_in[column] = [
            (label in classes) if is_reported else ""
            for is_reported, classes in zip(reported, class_sets)
        ]


def normalize_model_facing_adsorbent_categories(df_in: pd.DataFrame) -> None:
    """Apply categorical canonicalization after adsorbent enrichment is complete."""
    normalize_ac_raw_materials(df_in)
    normalize_adsorbent_subcategories(df_in)

    _ensure_columns(df_in, ("adsorbent_category", "functional_group", "polymer_matrix"), "")
    df_in["functional_group"] = df_in["functional_group"].apply(normalize_functional_groups)
    consolidate_functional_groups_by_identity(df_in)
    materialize_functional_group_ml_features(df_in)
    materialize_pore_class_ml_features(df_in)
    df_in["polymer_matrix"] = [
        normalize_polymer_matrix(value, category)
        for value, category in zip(
            df_in["polymer_matrix"], df_in["adsorbent_category"])
    ]


_NUMBER_TOKEN = r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?"


def parse_numeric_values(value: Any) -> list[float]:
    """Parse values/lists/ranges while keeping negative range endpoints valid."""
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return []
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return [float(value)] if math.isfinite(float(value)) else []

    text = unicodedata.normalize("NFKC", str(value))
    text = text.replace("âˆ’", "-").replace("â€“", "-").replace("â€”", "-")
    # An uncertainty is not a second independent observation.
    text = re.sub(
        rf"({_NUMBER_TOKEN})\s*(?:Â±|\+/-)\s*{_NUMBER_TOKEN}",
        r"\1",
        text,
    )
    # Turn the separator in both 8.5-17 and -80--2 into whitespace while
    # preserving the sign of the second endpoint in the latter.
    text = re.sub(r"(?<=\d)\s*-\s*(?=-?\d)", " ", text)
    values: list[float] = []
    for match in re.findall(_NUMBER_TOKEN, text):
        try:
            number = float(match)
        except ValueError:
            continue
        if math.isfinite(number):
            values.append(number)
    return values


def _variant_reason(row: pd.Series) -> str:
    activation = _text(row.get("AC_activation_method")).casefold()
    if "reactivat" in activation:
        return "reactivated variant"

    labels = " | ".join(_text(row.get(col)) for col in IDENTITY_SOURCE_COLUMNS)
    normalized = unicodedata.normalize("NFKC", labels).casefold()
    if re.search(r"(?:^|[\s_-])(ht|at)(?:$|[\s_-])", normalized):
        return "explicit HT/AT modified variant"
    if "high-temperature treated" in normalized or "ammonia-treated" in normalized:
        return "explicit chemically/thermally modified variant"
    return ""


def _resolve_identity(
    row: pd.Series,
    mapping: dict[str, str],
) -> tuple[str, str, str, str]:
    for column in IDENTITY_SOURCE_COLUMNS:
        value = _text(row.get(column))
        canonical = mapping.get(normalize_product_key(value))
        if canonical:
            return canonical, normalize_product_key(canonical), f"map:{column}", value

    commercial = _text(row.get("name_commercial"))
    if commercial:
        canonical = clean_product_display(commercial)
        return canonical, normalize_product_key(canonical), "reported:name_commercial", commercial
    return "", "", "unresolved", ""


def _study_key(row: pd.Series) -> str:
    extraction_dataset = _text(row.get("extraction_dataset"))
    study = _text(row.get("study_no"))
    return f"{extraction_dataset}|{study}" if extraction_dataset else study


def build_audit_dataframe(
    raw: pd.DataFrame,
    mapping: dict[str, str],
    value_converter: Callable[
        [pd.Series, NumericProperty], tuple[list[float], str]
    ],
) -> tuple[pd.DataFrame, dict[int, dict[str, list[float]]]]:
    audit = raw.copy()
    if "source_row_index" not in audit.columns:
        audit.insert(0, "source_row_index", range(len(audit)))

    audit.insert(1, "canonical_adsorbent", "")
    audit.insert(2, "Product_key", "")
    audit.insert(3, "identity_match_method", "")
    audit.insert(4, "identity_matched_value", "")
    audit.insert(5, "included_in_database", False)
    audit.insert(6, "exclusion_reason", "")
    audit.insert(7, "variant_status", "")

    audit["polymer_matrix_normalized"] = [
        normalize_polymer_matrix(value, category)
        for value, category in zip(
            audit.get("polymer_matrix", pd.Series("", index=audit.index)),
            audit.get("adsorbent_category", pd.Series("", index=audit.index)),
        )
    ]
    audit["functional_group_normalized"] = audit.get(
        "functional_group", pd.Series("", index=audit.index)
    ).apply(normalize_functional_groups)
    audit["pore_class_normalized"] = audit.get(
        "pore_class", pd.Series("", index=audit.index)
    ).apply(normalize_pore_class)

    parsed_by_row: dict[int, dict[str, list[float]]] = {}
    conversion_notes: list[str] = []
    for idx, row in audit.iterrows():
        canonical, product_key, method, matched = _resolve_identity(row, mapping)
        variant = _variant_reason(row)
        included = bool(product_key) and not variant
        exclusion = ""
        if not product_key:
            exclusion = (
                "No explicit adsorbent-map match and no reported name_commercial; "
                "generic study labels are not consolidated across studies"
            )
        elif variant:
            exclusion = variant

        audit.at[idx, "canonical_adsorbent"] = canonical
        audit.at[idx, "Product_key"] = product_key
        audit.at[idx, "identity_match_method"] = method
        audit.at[idx, "identity_matched_value"] = matched
        audit.at[idx, "included_in_database"] = included
        audit.at[idx, "exclusion_reason"] = exclusion
        audit.at[idx, "variant_status"] = variant or "base/as-reported"

        parsed_by_row[idx] = {}
        row_notes: list[str] = []
        for prop in NUMERIC_PROPERTIES:
            values, note = value_converter(row, prop)
            parsed_by_row[idx][prop.output_col] = values
            if note:
                row_notes.append(note)
        conversion_notes.append("; ".join(dict.fromkeys(row_notes)))

    audit["numeric_conversion_notes"] = conversion_notes
    for prop in NUMERIC_PROPERTIES:
        audit[f"{prop.output_col}__parsed_values"] = [
            "; ".join(f"{value:.12g}" for value in parsed_by_row[idx][prop.output_col])
            for idx in audit.index
        ]
    return audit, parsed_by_row


def _unique_text(values: Iterable[Any]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = _text(value)
        key = text.casefold()
        if text and key not in seen:
            seen.add(key)
            out.append(text)
    return out


def _joined_unique(values: Iterable[Any]) -> str:
    return "; ".join(_unique_text(values))


def _categorical_consensus(values: Iterable[Any]) -> tuple[str, str]:
    cleaned = [_text(value) for value in values]
    cleaned = [value for value in cleaned if value]
    if not cleaned:
        return "", ""

    display_by_key: dict[str, str] = {}
    counts: Counter[str] = Counter()
    for value in cleaned:
        key = value.casefold()
        display_by_key.setdefault(key, value)
        counts[key] += 1
    ranked = sorted(counts, key=lambda key: (-counts[key], display_by_key[key].casefold()))
    chosen = display_by_key[ranked[0]]
    if len(ranked) == 1:
        return chosen, ""
    detail = ", ".join(f"{display_by_key[key]} ({counts[key]})" for key in ranked)
    return chosen, detail


def _functional_group_union(values: Iterable[Any]) -> str:
    groups: list[str] = []
    seen: set[str] = set()
    for value in values:
        for token in normalize_functional_groups(value).split(";"):
            token = token.strip()
            key = token.casefold()
            if token and key not in seen:
                seen.add(key)
                groups.append(token)
    return "; ".join(groups)


def _choose_canonical_display(group: pd.DataFrame) -> str:
    mapped = group.loc[
        group["identity_match_method"].astype(str).str.startswith("map:"),
        "canonical_adsorbent",
    ]
    candidates = mapped if not mapped.empty else group["canonical_adsorbent"]
    chosen, _ = _categorical_consensus(candidates)
    return chosen


def _format_range(lo: float, hi: float) -> str:
    if math.isclose(lo, hi, rel_tol=1e-12, abs_tol=1e-12):
        return f"{lo:.12g}"
    return f"{lo:.12g}â€“{hi:.12g}"


def build_database_dataframe(
    audit: pd.DataFrame,
    parsed_by_row: dict[int, dict[str, list[float]]],
) -> pd.DataFrame:
    included = audit[audit["included_in_database"]].copy()
    if included.empty:
        return pd.DataFrame(columns=["Name_Commercial", "Product_key"])

    records: list[dict[str, Any]] = []
    for product_key, group in included.groupby("Product_key", sort=True):
        record: dict[str, Any] = {
            "Name_Commercial": _choose_canonical_display(group),
            "Product_key": product_key,
            "Study_count": len({_study_key(row) for _, row in group.iterrows()}),
            "Source_row_count": len(group),
            "Data_sources": _joined_unique(group.get("DOI", pd.Series(dtype=object))),
            "Reported_names": _joined_unique(
                value
                for column in ("name_full", "name_commercial")
                if column in group.columns
                for value in group[column]
            ),
            "Reported_abbreviations": _joined_unique(
                group.get("name_abbreviation", pd.Series(dtype=object))
            ),
            "Reported_vendors": _joined_unique(
                group.get("vendor", pd.Series(dtype=object))
            ),
        }

        categorical_conflicts: list[str] = []
        for output_col, source_col in CATEGORICAL_PROPERTIES:
            chosen, conflict = _categorical_consensus(
                group.get(source_col, pd.Series(dtype=object))
            )
            record[output_col] = chosen
            if conflict:
                categorical_conflicts.append(f"{output_col}: {conflict}")

        record["Functional_group"] = _functional_group_union(
            group.get("functional_group_normalized", pd.Series(dtype=object))
        )
        record["Categorical_conflicts"] = "; ".join(categorical_conflicts)

        quality_flags: list[str] = []
        for note in _unique_text(group["numeric_conversion_notes"]):
            quality_flags.extend(part.strip() for part in note.split(";") if part.strip())

        group_indices = list(group.index)
        numeric_values_by_output: dict[str, list[float]] = {}
        for prop in NUMERIC_PROPERTIES:
            values: list[float] = []
            studies_with_values: set[str] = set()
            for idx in group_indices:
                row_values = parsed_by_row[idx][prop.output_col]
                if row_values:
                    values.extend(row_values)
                    studies_with_values.add(_study_key(audit.loc[idx]))
            numeric_values_by_output[prop.output_col] = values

            if values:
                lo, hi = min(values), max(values)
                record[prop.output_col] = (lo + hi) / 2.0
                record[f"{prop.output_col}_range"] = _format_range(lo, hi)
                record[f"{prop.output_col}_n_values"] = len(values)
                record[f"{prop.output_col}_n_studies"] = len(studies_with_values)
                if prop.output_col in {
                    "element_C_value",
                    "element_N_value",
                    "element_O_value",
                } and any(0 < value < 1 for value in values) and any(
                    value > 20 for value in values
                ):
                    quality_flags.append(
                        f"{prop.output_col}: possible mixed fraction/percentage scale"
                    )
            else:
                record[prop.output_col] = None
                record[f"{prop.output_col}_range"] = ""
                record[f"{prop.output_col}_n_values"] = 0
                record[f"{prop.output_col}_n_studies"] = 0

        elemental_values = {
            column: numeric_values_by_output.get(column, [])
            for column in ("element_C_value", "element_N_value", "element_O_value")
        }
        if any(elemental_values.values()):
            normalized, note = normalize_elemental_fraction_value_lists(elemental_values)
            for column, value in normalized.items():
                if elemental_values[column]:
                    record[column] = value
                    record[f"{column}_range"] = "" if value is None else f"{value:.12g}"
            if note:
                quality_flags.append(note)

        for column in (
            "Total_pore_percentage",
            "Micro_pore_percentage",
            "Meso_pore_percentage",
            "Macro_pore_percentage",
        ):
            values = numeric_values_by_output.get(column, [])
            if not values:
                continue
            value, note = normalize_bounded_fraction_values(values, field=column)
            record[column] = value
            record[f"{column}_range"] = "" if value is None else f"{value:.12g}"
            if note:
                quality_flags.append(note)

        record["Quality_flags"] = "; ".join(dict.fromkeys(quality_flags))
        records.append(record)

    database = pd.DataFrame.from_records(records)
    identity_cols = [
        "Name_Commercial",
        "Product_key",
        "Study_count",
        "Source_row_count",
        "Data_sources",
        "Reported_names",
        "Reported_abbreviations",
        "Reported_vendors",
    ]
    categorical_cols = [
        output_col for output_col, _ in CATEGORICAL_PROPERTIES
    ] + ["Functional_group", "Categorical_conflicts"]
    numeric_cols = [
        column
        for prop in NUMERIC_PROPERTIES
        for column in (
            prop.output_col,
            f"{prop.output_col}_range",
            f"{prop.output_col}_n_values",
            f"{prop.output_col}_n_studies",
        )
    ]
    ordered = identity_cols + categorical_cols + numeric_cols + ["Quality_flags"]
    return database.reindex(columns=ordered)


def _excel_safe(dataframe: pd.DataFrame) -> pd.DataFrame:
    safe = dataframe.copy()
    for column in safe.columns:
        if safe[column].dtype == "object":
            safe[column] = safe[column].map(
                lambda value: (
                    "'" + value
                    if isinstance(value, str) and value.startswith(("=", "+", "-", "@"))
                    else value
                )
            )
    return safe


def _style_worksheet(
    worksheet,
    table_name: str,
    freeze_cell: str = "A2",
) -> None:
    worksheet.freeze_panes = freeze_cell
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 34

    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if worksheet.max_row >= 2 and worksheet.max_column >= 1:
        table = Table(displayName=table_name, ref=worksheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    wide_headers = {
        "Data_sources",
        "Reported_names",
        "Reported_abbreviations",
        "Reported_vendors",
        "Categorical_conflicts",
        "Quality_flags",
        "exclusion_reason",
        "numeric_conversion_notes",
        "article_title",
        "authors",
    }
    for column_cells in worksheet.columns:
        letter = column_cells[0].column_letter
        header = _text(column_cells[0].value)
        sample = [len(_text(cell.value)) for cell in column_cells[: min(80, worksheet.max_row)]]
        width = min(max(max(sample, default=len(header)) + 2, 10), 28)
        if header in wide_headers:
            width = 45
        elif header.endswith("_range"):
            width = 20
        worksheet.column_dimensions[letter].width = width

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def write_workbook(
    database: pd.DataFrame,
    audit: pd.DataFrame,
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        _excel_safe(database).to_excel(
            writer, sheet_name=DATABASE_SHEET, index=False, na_rep=""
        )
        _excel_safe(audit).to_excel(writer, sheet_name=AUDIT_SHEET, index=False, na_rep="")
        _style_worksheet(writer.book[DATABASE_SHEET], "AdsorbentDatabaseTable")
        _style_worksheet(writer.book[AUDIT_SHEET], "ConsolidationAuditTable")


def _values_equal(left: Any, right: Any) -> bool:
    if _blank(left) and _blank(right):
        return True
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        if isinstance(left, bool) or isinstance(right, bool):
            return left == right
        return math.isclose(float(left), float(right), rel_tol=1e-12, abs_tol=1e-12)
    return _collapse_spaces(str(left)) == _collapse_spaces(str(right))


def _resolve_literature_product(
    row: pd.Series,
    mapping: dict[str, str],
) -> tuple[str, str, str]:
    for column in ("Product_key", "Name_Commercial"):
        key, display, method = _resolve_mapped_product(row.get(column), mapping, column)
        if not key:
            continue
        if method.startswith("map:"):
            return key, display, method
        current_display = _text(row.get("Name_Commercial"))
        return key, current_display or display, method
    return "", "", "unresolved"


def _coalesce_literature_database(
    literature: pd.DataFrame,
    mapping: dict[str, str],
) -> tuple[pd.DataFrame, list[dict[str, Any]]]:
    records: list[dict[str, Any]] = []
    audit_records: list[dict[str, Any]] = []
    _ensure_columns(literature, AUTHORITATIVE_PROPERTY_COLUMNS, "")

    for source_idx, row in literature.iterrows():
        key, display, method = _resolve_literature_product(row, mapping)
        if not key:
            audit_records.append(
                {
                    "online_name": "",
                    "resolved_Product_key": "",
                    "match_status": "literature_unresolved",
                    "match_method": method,
                    "literature_name": row.get("Name_Commercial", ""),
                    "action": "excluded",
                    "fields_filled": "",
                    "fields_overridden": "",
                    "unmatched_candidates": "",
                    "vendor": "",
                    "data_sources": row.get("Data_sources", ""),
                    "notes": f"Literature row {source_idx} has no resolvable Product_key",
                }
            )
            continue
        record = row.to_dict()
        record["Product_key"] = key
        if display:
            record["Name_Commercial"] = display
        record["__source_row_index"] = source_idx
        records.append(record)

    if not records:
        return pd.DataFrame(columns=AUTHORITATIVE_PROPERTY_COLUMNS), audit_records

    resolved = pd.DataFrame.from_records(records)
    coalesced: list[dict[str, Any]] = []
    for product_key, group in resolved.groupby("Product_key", sort=True):
        chosen = group.iloc[0].to_dict()
        conflicts: list[str] = []
        source_rows = [str(value) for value in group["__source_row_index"]]
        for _, row in group.iloc[1:].iterrows():
            for column in AUTHORITATIVE_PROPERTY_COLUMNS:
                value = row.get(column)
                if _blank(value):
                    continue
                current = chosen.get(column)
                if _blank(current):
                    chosen[column] = value
                elif not _values_equal(current, value):
                    conflicts.append(f"{column}: {current} <> {value}")
        if len(group) > 1:
            audit_records.append(
                {
                    "online_name": "",
                    "resolved_Product_key": product_key,
                    "match_status": "literature_duplicate_collapsed",
                    "match_method": "adsorbent_map",
                    "literature_name": chosen.get("Name_Commercial", ""),
                    "action": "coalesced_literature_rows",
                    "fields_filled": "",
                    "fields_overridden": "",
                    "unmatched_candidates": "",
                    "vendor": "",
                    "data_sources": chosen.get("Data_sources", ""),
                    "notes": (
                        f"Collapsed literature rows {', '.join(source_rows)}"
                        + (f"; conflicts kept from first row: {'; '.join(conflicts)}" if conflicts else "")
                    ),
                }
            )
        chosen.pop("__source_row_index", None)
        coalesced.append(chosen)

    return pd.DataFrame.from_records(coalesced), audit_records


def _resolve_online_product(
    row: pd.Series,
    mapping: dict[str, str],
) -> tuple[str, str, str]:
    for column in ("Product_key", "Name_Commercial", "Name_Abbreviation"):
        if column not in row.index:
            continue
        key, display, method = _resolve_mapped_product(row.get(column), mapping, column)
        if key:
            return key, display, method
    return "", "", "unresolved"


def _join_field_names(values: Iterable[str]) -> str:
    return "; ".join(dict.fromkeys(value for value in values if value))


def _append_unique_joined(existing: Any, value: Any) -> str:
    parts = _unique_text(str(existing).split(";") if _text(existing) else [])
    for part in _unique_text(str(value).split(";") if _text(value) else []):
        if part.casefold() not in {seen.casefold() for seen in parts}:
            parts.append(part)
    return "; ".join(parts)


def _coalesce_online_properties(
    online: pd.DataFrame,
    mapping: dict[str, str],
) -> tuple[list[pd.Series], list[dict[str, Any]]]:
    coalesced_by_key: dict[str, dict[str, Any]] = {}
    audit_records: list[dict[str, Any]] = []

    for online_idx, row in online.iterrows():
        key, display, method = _resolve_online_product(row, mapping)
        online_name = _text(row.get("Name_Commercial")) or display
        if not key:
            audit_records.append(
                {
                    "online_row_index": online_idx,
                    "online_name": online_name,
                    "resolved_Product_key": "",
                    "match_status": "online_unresolved",
                    "match_method": method,
                    "literature_name": "",
                    "action": "not_applied",
                    "fields_filled": "",
                    "fields_overridden": "",
                    "fields_unchanged": "",
                    "fields_conflicted": "",
                    "unmatched_candidates": "",
                    "vendor": row.get("Vendor", ""),
                    "data_sources": row.get("Data_sources", ""),
                    "vendor_verification_notes": row.get(
                        "vendor_verification_notes", ""
                    ),
                    "online_row_count_for_key": 1,
                    "notes": "Online row has no resolvable product key",
                }
            )
            continue

        entry = coalesced_by_key.get(key)
        if entry is None:
            entry = row.to_dict()
            entry["Product_key"] = key
            entry["__display"] = display
            entry["__match_method"] = method
            entry["__online_names"] = online_name
            entry["__online_row_indices"] = str(online_idx)
            entry["__online_row_count"] = 1
            entry["__fields_conflicted"] = ""
            coalesced_by_key[key] = entry
            continue

        entry["__online_row_count"] = int(entry["__online_row_count"]) + 1
        entry["__online_row_indices"] = _append_unique_joined(
            entry.get("__online_row_indices", ""), online_idx
        )
        entry["__online_names"] = _append_unique_joined(
            entry.get("__online_names", ""), online_name
        )
        if method.startswith("map:") and not str(
            entry.get("__match_method", "")
        ).startswith("map:"):
            entry["__match_method"] = method
            entry["__display"] = display

        for column, new_value in row.items():
            if _blank(new_value):
                continue
            if column in {"Vendor", "Data_sources", "vendor_verification_notes"}:
                entry[column] = _append_unique_joined(entry.get(column, ""), new_value)
                continue
            current = entry.get(column)
            if _blank(current):
                entry[column] = new_value
            elif (
                column in ONLINE_TO_AUTHORITATIVE_COLUMN_MAP
                and not _values_equal(current, new_value)
            ):
                entry["__fields_conflicted"] = _append_unique_joined(
                    entry.get("__fields_conflicted", ""),
                    f"{column}: {current} <> {new_value}",
                )

    for key, entry in coalesced_by_key.items():
        if int(entry.get("__online_row_count", 1)) > 1:
            conflicts = _text(entry.get("__fields_conflicted"))
            notes = (
                f"Coalesced online rows {entry.get('__online_row_indices', '')}; "
                "kept first nonblank value for duplicate-field conflicts"
                if conflicts
                else f"Coalesced online rows {entry.get('__online_row_indices', '')}"
            )
            audit_records.append(
                {
                    "online_row_index": entry.get("__online_row_indices", ""),
                    "online_name": entry.get("__online_names", ""),
                    "resolved_Product_key": key,
                    "match_status": "duplicate_online_key_coalesced",
                    "match_method": entry.get("__match_method", ""),
                    "literature_name": "",
                    "action": "coalesced_online_rows",
                    "fields_filled": "",
                    "fields_overridden": "",
                    "fields_unchanged": "",
                    "fields_conflicted": conflicts,
                    "unmatched_candidates": "",
                    "vendor": entry.get("Vendor", ""),
                    "data_sources": entry.get("Data_sources", ""),
                    "vendor_verification_notes": entry.get(
                        "vendor_verification_notes", ""
                    ),
                    "online_row_count_for_key": entry.get("__online_row_count", ""),
                    "notes": notes,
                }
            )

    return [pd.Series(entry) for entry in coalesced_by_key.values()], audit_records


def write_authoritative_workbook(
    database: pd.DataFrame,
    audit: pd.DataFrame,
    output_path: Path,
    sheet_name: str = AUTHORITATIVE_SHEET,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        _excel_safe(database).to_excel(writer, sheet_name=sheet_name, index=False, na_rep="")
        _excel_safe(audit).to_excel(
            writer, sheet_name=AUTHORITATIVE_AUDIT_SHEET, index=False, na_rep=""
        )
        _style_worksheet(writer.book[sheet_name], "AdsorbentPropertiesTable")
        _style_worksheet(writer.book[AUTHORITATIVE_AUDIT_SHEET], "MergeAuditTable")


def build_authoritative_adsorbent_database(
    config: AdsorbentProcessingConfig,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Merge literature and vendor properties into the DB used for enrichment."""
    output_path = (
        config.authoritative_output_path
        if config.authoritative_output_path is not None
        else config.database_output_path.with_name(
            "commercial_adsorbent_properties_authoritative.xlsx"
        )
    )
    mapping = load_adsorbent_map(config.mapping_path)
    literature = pd.read_excel(config.database_output_path, sheet_name=DATABASE_SHEET)
    base, audit_records = _coalesce_literature_database(literature, mapping)
    _ensure_columns(base, AUTHORITATIVE_PROPERTY_COLUMNS, "")

    online = pd.read_excel(config.online_properties_path, sheet_name=config.online_properties_sheet)
    online_rows, online_audit_records = _coalesce_online_properties(online, mapping)
    audit_records.extend(online_audit_records)
    base_by_key = {str(row["Product_key"]): idx for idx, row in base.iterrows()}
    base_names = base["Name_Commercial"].fillna("").astype(str).tolist()

    for row in online_rows:
        online_name = _text(row.get("__online_names")) or _text(row.get("Name_Commercial"))
        key = _text(row.get("Product_key"))
        display = _text(row.get("__display"))
        method = _text(row.get("__match_method")) or "unresolved"
        fields_filled: list[str] = []
        fields_overridden: list[str] = []
        fields_unchanged: list[str] = []
        notes: list[str] = []

        if key in base_by_key:
            match_status = "matched"
            action = "online_values_applied"
            base_idx = base_by_key[key]
            literature_name = base.at[base_idx, "Name_Commercial"]
            for src_col, dst_col in ONLINE_TO_AUTHORITATIVE_COLUMN_MAP.items():
                if src_col not in row.index or dst_col not in AUTHORITATIVE_PROPERTY_COLUMNS:
                    continue
                new_value = row.get(src_col)
                if _blank(new_value):
                    continue
                old_value = base.at[base_idx, dst_col]
                if _blank(old_value):
                    fields_filled.append(dst_col)
                elif _values_equal(old_value, new_value):
                    fields_unchanged.append(dst_col)
                else:
                    fields_overridden.append(dst_col)
                base.at[base_idx, dst_col] = new_value
                if src_col == "ion_exchange_capacity_value_mmol/g":
                    notes.append("mmol/g source stored as meq/g by existing IEC convention")
        else:
            match_status = "online_only_unmatched"
            action = "not_applied"
            literature_name = ""

        candidates = ""
        if match_status == "online_only_unmatched" and online_name:
            candidates = "; ".join(
                difflib.get_close_matches(online_name, base_names, n=5, cutoff=0.45)
            )
        if match_status == "matched" and not fields_filled and not fields_overridden:
            action = "matched_no_property_changes"

        audit_records.append(
            {
                "online_row_index": row.get("__online_row_indices", ""),
                "online_name": online_name or display,
                "resolved_Product_key": key,
                "match_status": match_status,
                "match_method": method,
                "literature_name": literature_name,
                "action": action,
                "fields_filled": _join_field_names(fields_filled),
                "fields_overridden": _join_field_names(fields_overridden),
                "fields_unchanged": _join_field_names(fields_unchanged),
                "fields_conflicted": row.get("__fields_conflicted", ""),
                "unmatched_candidates": candidates,
                "vendor": row.get("Vendor", ""),
                "data_sources": row.get("Data_sources", ""),
                "vendor_verification_notes": row.get("vendor_verification_notes", ""),
                "online_row_count_for_key": row.get("__online_row_count", ""),
                "notes": "; ".join(notes),
            }
        )

    authoritative = base.reindex(columns=AUTHORITATIVE_PROPERTY_COLUMNS).sort_values(
        "Product_key"
    )
    materialize_nonionic_resin_iec_defaults(authoritative)
    audit = pd.DataFrame.from_records(audit_records).reindex(
        columns=AUTHORITATIVE_AUDIT_COLUMNS
    )
    write_authoritative_workbook(
        authoritative,
        audit,
        output_path,
        config.authoritative_properties_sheet,
    )
    print(f"Saved authoritative adsorbent properties: {output_path}")
    print(f"  {config.authoritative_properties_sheet}: {len(authoritative)} literature-backed adsorbents")
    print(f"  {AUTHORITATIVE_AUDIT_SHEET}: {len(audit)} merge/audit rows")
    return authoritative, audit


def build_consolidated_adsorbent_database(
    config: AdsorbentProcessingConfig,
    value_converter: Callable[[pd.Series, NumericProperty], tuple[list[float], str]],
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Normalize raw adsorbent properties and consolidate them across studies."""
    raw = pd.read_excel(config.raw_input_path, sheet_name=config.raw_input_sheet)
    mapping = load_adsorbent_map(config.mapping_path)
    audit, parsed_by_row = build_audit_dataframe(raw, mapping, value_converter)
    database = build_database_dataframe(audit, parsed_by_row)
    write_workbook(database, audit, config.database_output_path)
    included = int(audit["included_in_database"].sum())
    print(f"Saved adsorbent database: {config.database_output_path}")
    print(f"  Adsorbent_Database: {len(database)} consolidated adsorbents")
    print(f"  Consolidation_Audit: {len(audit)} raw rows ({included} included)")
    return database, audit

