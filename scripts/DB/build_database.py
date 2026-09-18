"""Public command for building the merged and standardized databases.

Run this file rather than invoking ``merge.py`` or ``normalize.py``
directly. ``merge.py`` owns raw-record assembly. ``normalize.py`` then
owns normalization and produces both the consolidated adsorbent-property
database and the standardized performance database.
"""

from __future__ import annotations

import argparse
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
AD_ROOT = (
    SCRIPT_DIR.parent.parent
    if SCRIPT_DIR.name.lower() == "db" and SCRIPT_DIR.parent.name.lower() == "scripts"
    else SCRIPT_DIR.parent
)

DEFAULT_TRAINING_EXTRACTION_OUTPUT_DIR = (
    AD_ROOT / "output" / "Extraction" / "Training" / "LLM_prediction" / "GPT4.1_20260603"
)
DEFAULT_VALIDATION_EXTRACTION_OUTPUT_DIR = (
    AD_ROOT / "output" / "Extraction" / "Validation" / "LLM_prediction" / "GPT4.1_20260705"
)
DEFAULT_APPLICATION_EXTRACTION_OUTPUT_DIR = (
    AD_ROOT / "output" / "Extraction" / "Application" / "GPT4.1_20260705"
)

DEFAULT_MERGE_OUTPUT_DIR = AD_ROOT / "output" / "Merge"
DEFAULT_STUDY_METADATA_XLSX = AD_ROOT.parent / "web of science search" / "within_scope_records.xlsx"
DEFAULT_RAW_PERFORMANCE_FILE = DEFAULT_MERGE_OUTPUT_DIR / "pfas_adsorption_performance_merged_raw.xlsx"
DEFAULT_RAW_ADSORBENT_FILE = DEFAULT_MERGE_OUTPUT_DIR / "adsorbent_merged_raw.xlsx"
DEFAULT_NORMALIZED_ADSORBENT_FILE = DEFAULT_MERGE_OUTPUT_DIR / "adsorbent_merged_normalized.xlsx"
DEFAULT_STANDARDIZED_PERFORMANCE_FILE = DEFAULT_MERGE_OUTPUT_DIR / "pfas_adsorption_performance_database.xlsx"
DEFAULT_DATA_AUDIT_FILE = DEFAULT_MERGE_OUTPUT_DIR / "data_audit.xlsx"
DEFAULT_DATA_AUDIT_SCRIPT = SCRIPT_DIR / "data_audit.py"
DEFAULT_RECORD_REVIEW_DECISIONS_FILE = DEFAULT_MERGE_OUTPUT_DIR / "record_review_decisions.csv"
DEFAULT_ADSORBENT_PROPERTY_REVIEW_DECISIONS_FILE = (
    DEFAULT_MERGE_OUTPUT_DIR / "adsorbent_property_review_decisions.csv"
)
DEFAULT_PERFORMANCE_NORMALIZATION_DECISIONS_FILE = (
    DEFAULT_MERGE_OUTPUT_DIR / "performance_normalization_decisions.csv"
)
DEFAULT_COMM_PROPS_ONLINE_FILE = DEFAULT_MERGE_OUTPUT_DIR / "commercial_adsorbent_properties_Online.xlsx"
DEFAULT_COMM_PROPS_DB_FILE = DEFAULT_MERGE_OUTPUT_DIR / "commercial_adsorbent_properties_database.xlsx"
DEFAULT_COMM_PROPS_AUTHORITATIVE_FILE = (
    DEFAULT_MERGE_OUTPUT_DIR / "commercial_adsorbent_properties_authoritative.xlsx"
)
DEFAULT_PFAS_PROPS_FILE = AD_ROOT.parent.parent / "PFAS" / "pfas_features.xlsx"

# ----------------------- User settings -----------------------

# Empty means all studies discovered under each configured dataset.
# Use a short list while testing, for example: ["study_17"].
SELECTED_STUDY_IDS = []

# Empty means all configured datasets. Keep ["training"] while testing if you
# want one extraction directory only.
SELECTED_DATASET_LABELS: list[str] = []

DATASETS = [
    {
        "label": "training",
        "extraction_output_dir": DEFAULT_TRAINING_EXTRACTION_OUTPUT_DIR,
    },
    {
        "label": "application",
        "extraction_output_dir": DEFAULT_APPLICATION_EXTRACTION_OUTPUT_DIR,
    },
    {
        "label": "validation",
        "extraction_output_dir": DEFAULT_VALIDATION_EXTRACTION_OUTPUT_DIR,
    },
]

REVIEW_DECISION_COLUMNS = ("source_row_index", "decision")
REVIEW_OUTCOME_COLUMN = "review_outcome"
VALID_HUMAN_REVIEW_DECISIONS = {"reliable", "unreliable"}


def _env_path(name: str, default: Path) -> Path:
    value = os.getenv(name)
    return Path(value).expanduser() if value else default


def _env_path_or_none(name: str) -> Path | None:
    value = os.getenv(name)
    return Path(value).expanduser() if value else None


def _split_values(values: list[str] | None) -> list[str] | None:
    if values is None:
        return None

    out: list[str] = []
    for value in values:
        for part in str(value).split(","):
            part = part.strip()
            if part:
                out.append(part)
    return out


def _split_study_ids(values: list[str] | None) -> list[str] | None:
    return _split_values(values)


def _norm_label(value: object) -> str:
    return str(value or "").strip().lower()


def _source_row_key(value: object) -> str:
    """Normalize source-row identifiers read from CSV or Excel."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return "" if text.casefold() in {"", "nan", "none"} else text


def _read_human_review_decisions(path: Path) -> dict[str, str]:
    """Read completed source-record decisions and reject ambiguous entries."""
    if not path.exists() or not path.stat().st_size:
        return {}

    import csv

    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = set(reader.fieldnames or [])
        rows = list(reader)
    missing = [column for column in REVIEW_DECISION_COLUMNS if column not in fieldnames]
    if missing:
        raise ValueError(
            f"Review decisions CSV {path} is missing required columns {missing}. "
            f"Expected at least {list(REVIEW_DECISION_COLUMNS)}."
        )
    if not rows:
        return {}

    decisions_by_source: dict[str, set[str]] = {}
    for row_number, row in enumerate(rows, start=2):
        source_key = _source_row_key(row.get("source_row_index"))
        decision = _norm_label(row.get("decision"))
        if not source_key or not decision:
            continue
        if decision not in VALID_HUMAN_REVIEW_DECISIONS:
            allowed = ", ".join(sorted(VALID_HUMAN_REVIEW_DECISIONS))
            raise ValueError(
                f"Unsupported decision {row.get('decision')!r} in {path} row {row_number}. "
                f"Use one of: {allowed}; leave blank while review is pending."
            )
        decisions_by_source.setdefault(source_key, set()).add(decision)

    conflicts = {
        source_key: sorted(decisions)
        for source_key, decisions in decisions_by_source.items()
        if len(decisions) > 1
    }
    if conflicts:
        raise ValueError(
            "Conflicting completed human-review decisions for source_row_index values: "
            + ", ".join(f"{key} ({'/'.join(values)})" for key, values in conflicts.items())
        )
    return {source_key: next(iter(decisions)) for source_key, decisions in decisions_by_source.items()}


def apply_human_review_decisions(
    standardized_file: Path,
    review_decisions_file: Path,
) -> dict[str, int]:
    """Materialize completed review decisions in every source-record worksheet.

    ``review_outcome`` records a completed ``reliable`` or ``unreliable``
    decision without overriding automated normalization or conversion status.
    """
    decisions = _read_human_review_decisions(review_decisions_file)
    if not standardized_file.exists():
        raise FileNotFoundError(f"Standardized database not found: {standardized_file}")

    workbook = load_workbook(standardized_file)
    found_source_keys: set[str] = set()
    updated_rows = 0
    unreliable_rows = 0
    sheets_updated = 0
    for worksheet in workbook.worksheets:
        headers = {
            str(cell.value).strip(): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }
        source_column = headers.get("source_row_index")
        if source_column is None:
            continue

        outcome_column = headers.get(REVIEW_OUTCOME_COLUMN)
        if outcome_column is None:
            # Do not add review metadata to the model-ready projection, which
            # intentionally omits status fields.  All source-record sheets and
            # the compact review queue include record_status.
            if "record_status" not in headers:
                continue
            outcome_column = worksheet.max_column + 1
            worksheet.cell(row=1, column=outcome_column, value=REVIEW_OUTCOME_COLUMN)

        sheets_updated += 1
        for row_number in range(2, worksheet.max_row + 1):
            source_key = _source_row_key(worksheet.cell(row=row_number, column=source_column).value)
            decision = decisions.get(source_key, "")
            if source_key in decisions:
                found_source_keys.add(source_key)
            unreliable = decision == "unreliable"
            worksheet.cell(row=row_number, column=outcome_column, value=decision)
            updated_rows += 1
            unreliable_rows += int(unreliable)

    workbook.save(standardized_file)
    unmatched = len(set(decisions) - found_source_keys)
    return {
        "completed_decisions": len(decisions),
        "unreliable_decisions": sum(value == "unreliable" for value in decisions.values()),
        "unmatched_decisions": unmatched,
        "worksheets_updated": sheets_updated,
        "rows_updated": updated_rows,
        "unreliable_rows": unreliable_rows,
    }


def _parse_dataset_spec(spec: str) -> dict[str, object]:
    if "=" not in spec:
        raise argparse.ArgumentTypeError(
            "Dataset specs must use LABEL=PATH, for example: training=C:\\path\\to\\Extraction\\Training\\..."
        )
    label, path = spec.split("=", 1)
    label = label.strip()
    path = path.strip()
    if not label or not path:
        raise argparse.ArgumentTypeError("Dataset specs require both LABEL and PATH.")
    return {"label": label, "extraction_output_dir": Path(path).expanduser()}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Merge the canonical extracted chain artifacts, then use normalize.py to "
            "build the consolidated adsorbent and PFAS performance databases."
        )
    )
    parser.add_argument(
        "--extraction-output-dir",
        type=Path,
        default=_env_path_or_none("EXTRACTION_OUTPUT_DIR"),
        help=(
            "Single extraction output directory whose _chain_artifacts folder is the "
            "canonical source of extracted records. Overrides the DATASETS list for "
            "quick one-directory runs."
        ),
    )
    parser.add_argument(
        "--dataset-label",
        default=os.getenv("MERGE_DATASET_LABEL"),
        help="Dataset label used with --extraction-output-dir.",
    )
    parser.add_argument(
        "--dataset",
        action="append",
        type=_parse_dataset_spec,
        help=(
            "Dataset to merge, as LABEL=PATH. Repeat for training, validation, "
            "application, etc. Overrides the DATASETS list."
        ),
    )
    parser.add_argument(
        "--selected-dataset-labels",
        nargs="*",
        help=(
            "Labels to keep from DATASETS or --dataset. Comma-separated values "
            "are also accepted."
        ),
    )
    parser.add_argument(
        "--all-datasets",
        action="store_true",
        help="Use every configured dataset, ignoring SELECTED_DATASET_LABELS.",
    )
    parser.add_argument(
        "--study-ids",
        nargs="*",
        help=(
            "Study IDs to process, for example: --study-ids study_17 study_18. "
            "Comma-separated values are also accepted."
        ),
    )
    parser.add_argument(
        "--all-studies",
        action="store_true",
        help="Process every study folder under _chain_artifacts.",
    )
    parser.add_argument(
        "--study-metadata-xlsx",
        type=Path,
        default=_env_path("STUDY_METADATA_XLSX", DEFAULT_STUDY_METADATA_XLSX),
        help="Workbook with DOI/title/year metadata.",
    )
    parser.add_argument(
        "--study-metadata-sheet",
        default=os.getenv("STUDY_METADATA_SHEET", "within_scope_records"),
        help="Metadata worksheet name.",
    )
    parser.add_argument(
        "--merge-output-dir",
        type=Path,
        default=_env_path("MERGE_OUTPUT_DIR", DEFAULT_MERGE_OUTPUT_DIR),
        help="Directory where merged and standardized workbooks are written.",
    )
    parser.add_argument(
        "--performance-output-file",
        type=Path,
        default=None,
        help="Raw merged performance workbook path.",
    )
    parser.add_argument(
        "--adsorbent-output-file",
        type=Path,
        default=None,
        help="Raw merged adsorbent workbook path.",
    )
    parser.add_argument(
        "--adsorbent-normalized-output-file",
        type=Path,
        default=None,
        help="Normalized adsorbent-property workbook path used for consolidation and enrichment.",
    )
    parser.add_argument(
        "--standardized-output-file",
        type=Path,
        default=None,
        help="Standardized performance database workbook path.",
    )
    parser.add_argument(
        "--adsorbent-map-file",
        type=Path,
        default=SCRIPT_DIR / "adsorbent_map.csv",
        help="Canonical adsorbent-name mapping CSV.",
    )
    parser.add_argument(
        "--adsorbent-database-input-sheet",
        default="Sheet1",
        help="Raw adsorbent worksheet passed to normalize.py.",
    )
    parser.add_argument(
        "--standardizer-script",
        type=Path,
        default=SCRIPT_DIR / "normalize.py",
        help="Path to normalize.py.",
    )
    parser.add_argument(
        "--standardizer-sheet-name",
        default=os.getenv("UNIT_STANDARDIZER_SHEET_NAME", "Sheet1"),
        help="Input sheet passed to normalize.py.",
    )
    parser.add_argument(
        "--comm-props-online-file",
        type=Path,
        default=_env_path_or_none("COMM_PROPS_ONLINE_PATH"),
        help="Commercial adsorbent online properties workbook used by normalize.py.",
    )
    parser.add_argument(
        "--comm-props-db-file",
        type=Path,
        default=_env_path_or_none("COMM_PROPS_DB_PATH"),
        help="Literature-derived commercial adsorbent database workbook.",
    )
    parser.add_argument(
        "--comm-props-authoritative-file",
        type=Path,
        default=_env_path_or_none("COMM_PROPS_AUTHORITATIVE_PATH"),
        help="Authoritative adsorbent properties workbook used to enrich the performance database.",
    )
    parser.add_argument(
        "--pfas-props-file",
        type=Path,
        default=_env_path_or_none("PFAS_PROPS_PATH"),
        help="PFAS properties workbook used by normalize.py.",
    )
    parser.add_argument(
        "--python",
        default=sys.executable,
        help="Python executable used for the normalization subprocess.",
    )
    parser.add_argument(
        "--skip-merge",
        action="store_true",
        help="Use existing raw merged workbooks and run normalize.py only.",
    )
    parser.add_argument(
        "--skip-standardize",
        action="store_true",
        help="Do not build either standardized database.",
    )
    parser.add_argument(
        "--disable-adsorbent-database-injection",
        action="store_true",
        help=(
            "Build the adsorbent database but do not use it to fill missing "
            "adsorbent properties in the performance database."
        ),
    )
    parser.add_argument(
        "--skip-data-audit",
        "--skip-model-data-audit",
        dest="skip_data_audit",
        action="store_true",
        help="Skip the figure-oriented data audit after standardization.",
    )
    parser.add_argument(
        "--data-audit-script",
        "--model-data-audit-script",
        dest="data_audit_script",
        type=Path,
        default=DEFAULT_DATA_AUDIT_SCRIPT,
        help="Path to the figure-oriented model-data audit script.",
    )
    parser.add_argument(
        "--data-audit-output-file",
        "--model-data-audit-output-file",
        dest="data_audit_output_file",
        type=Path,
        default=None,
        help="Optional output workbook path for the post-build model data audit.",
    )
    parser.add_argument(
        "--review-decisions-file",
        type=Path,
        default=None,
        help=(
            "CSV containing completed human-review decisions. Defaults to "
            "record_review_decisions.csv in --merge-output-dir."
        ),
    )
    parser.add_argument(
        "--adsorbent-property-review-decisions-file",
        type=Path,
        default=None,
        help=(
            "CSV of adsorbent-scoped property-normalization approvals. Defaults to "
            "adsorbent_property_review_decisions.csv in --merge-output-dir."
        ),
    )
    parser.add_argument(
        "--performance-normalization-decisions-file",
        type=Path,
        default=None,
        help=(
            "CSV of condition-scoped manual performance-normalization approvals. "
            "Defaults to performance_normalization_decisions.csv in --merge-output-dir."
        ),
    )
    parser.add_argument(
        "--skip-review-decisions",
        action="store_true",
        help="Do not apply completed human-review decisions to the standardized database.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Print the resolved workflow without writing files.",
    )
    return parser.parse_args()


def _resolve_outputs(args: argparse.Namespace) -> tuple[Path, Path, Path, Path]:
    merge_output_dir = args.merge_output_dir.expanduser().resolve()
    performance_file = (
        args.performance_output_file.expanduser().resolve()
        if args.performance_output_file
        else merge_output_dir / DEFAULT_RAW_PERFORMANCE_FILE.name
    )
    adsorbent_file = (
        args.adsorbent_output_file.expanduser().resolve()
        if args.adsorbent_output_file
        else merge_output_dir / DEFAULT_RAW_ADSORBENT_FILE.name
    )
    standardized_file = (
        args.standardized_output_file.expanduser().resolve()
        if args.standardized_output_file
        else merge_output_dir / DEFAULT_STANDARDIZED_PERFORMANCE_FILE.name
    )
    normalized_adsorbent_file = (
        args.adsorbent_normalized_output_file.expanduser().resolve()
        if args.adsorbent_normalized_output_file
        else merge_output_dir / DEFAULT_NORMALIZED_ADSORBENT_FILE.name
    )
    return performance_file, adsorbent_file, normalized_adsorbent_file, standardized_file


def _resolve_review_decisions_file(args: argparse.Namespace) -> Path:
    if args.review_decisions_file:
        return args.review_decisions_file.expanduser().resolve()
    return args.merge_output_dir.expanduser().resolve() / DEFAULT_RECORD_REVIEW_DECISIONS_FILE.name


def _resolve_adsorbent_property_review_decisions_file(args: argparse.Namespace) -> Path:
    if args.adsorbent_property_review_decisions_file:
        return args.adsorbent_property_review_decisions_file.expanduser().resolve()
    return (
        args.merge_output_dir.expanduser().resolve()
        / DEFAULT_ADSORBENT_PROPERTY_REVIEW_DECISIONS_FILE.name
    )


def _resolve_performance_normalization_decisions_file(args: argparse.Namespace) -> Path:
    if args.performance_normalization_decisions_file:
        return args.performance_normalization_decisions_file.expanduser().resolve()
    return (
        args.merge_output_dir.expanduser().resolve()
        / DEFAULT_PERFORMANCE_NORMALIZATION_DECISIONS_FILE.name
    )


def _resolve_path(value: object) -> Path:
    return Path(os.path.expandvars(str(value))).expanduser().resolve()


def _resolve_study_ids(args: argparse.Namespace) -> list[str]:
    if args.all_studies:
        return []
    study_ids = _split_study_ids(args.study_ids)
    if study_ids is not None:
        return study_ids
    return list(SELECTED_STUDY_IDS)


def _resolve_datasets(args: argparse.Namespace) -> list[dict[str, str]]:
    cli_dataset_override = False
    if args.extraction_output_dir:
        cli_dataset_override = True
        raw_datasets = [
            {
                "label": args.dataset_label or "training",
                "extraction_output_dir": args.extraction_output_dir,
            }
        ]
    elif args.dataset:
        cli_dataset_override = True
        raw_datasets = args.dataset
    else:
        raw_datasets = DATASETS

    if args.all_datasets:
        allowed_labels: set[str] = set()
    elif args.selected_dataset_labels is not None:
        allowed_labels = {_norm_label(v) for v in (_split_values(args.selected_dataset_labels) or [])}
    elif cli_dataset_override:
        allowed_labels = set()
    else:
        allowed_labels = {_norm_label(v) for v in SELECTED_DATASET_LABELS}

    datasets: list[dict[str, str]] = []
    for dataset in raw_datasets:
        label = str(dataset.get("label", "")).strip()
        if not label:
            raise SystemExit(f"Dataset entry is missing a label: {dataset}")
        if allowed_labels and _norm_label(label) not in allowed_labels:
            continue

        resolved: dict[str, str] = {"label": label}
        if dataset.get("artifact_root"):
            resolved["artifact_root"] = str(_resolve_path(dataset["artifact_root"]))
        elif dataset.get("extraction_output_dir"):
            resolved["extraction_output_dir"] = str(_resolve_path(dataset["extraction_output_dir"]))
        else:
            raise SystemExit(
                f"Dataset {label!r} needs extraction_output_dir or artifact_root."
            )
        datasets.append(resolved)

    if not datasets:
        raise SystemExit("No datasets selected. Check SELECTED_DATASET_LABELS or --selected-dataset-labels.")
    return datasets


def _resolve_standardizer_inputs(args: argparse.Namespace) -> tuple[Path, Path, Path, Path]:
    merge_output_dir = args.merge_output_dir.expanduser().resolve()
    comm_online_file = (
        args.comm_props_online_file.expanduser().resolve()
        if args.comm_props_online_file
        else merge_output_dir / DEFAULT_COMM_PROPS_ONLINE_FILE.name
    )
    comm_db_file = (
        args.comm_props_db_file.expanduser().resolve()
        if args.comm_props_db_file
        else merge_output_dir / DEFAULT_COMM_PROPS_DB_FILE.name
    )
    comm_authoritative_file = (
        args.comm_props_authoritative_file.expanduser().resolve()
        if args.comm_props_authoritative_file
        else merge_output_dir / DEFAULT_COMM_PROPS_AUTHORITATIVE_FILE.name
    )
    pfas_props_file = (
        args.pfas_props_file.expanduser().resolve()
        if args.pfas_props_file
        else DEFAULT_PFAS_PROPS_FILE
    )
    return comm_online_file, comm_db_file, comm_authoritative_file, pfas_props_file


def _print_plan(
    args: argparse.Namespace,
    datasets: list[dict[str, str]],
    selected_study_ids: list[str],
    performance_file: Path,
    adsorbent_file: Path,
    normalized_adsorbent_file: Path,
    standardized_file: Path,
    comm_online_file: Path,
    comm_db_file: Path,
    comm_authoritative_file: Path,
    pfas_props_file: Path,
    review_decisions_file: Path,
    adsorbent_property_review_decisions_file: Path,
    performance_normalization_decisions_file: Path,
) -> None:
    if not selected_study_ids:
        study_selection = "all studies discovered under _chain_artifacts"
    else:
        study_selection = ", ".join(selected_study_ids)

    print("Post-processing workflow")
    print("  datasets          :")
    for dataset in datasets:
        if "artifact_root" in dataset:
            root_label = "artifact_root"
        else:
            root_label = "extraction_output_dir"
        print(f"    - {dataset['label']} ({root_label}): {dataset[root_label]}")
    print(f"  studies           : {study_selection}")
    print(f"  metadata workbook : {args.study_metadata_xlsx.expanduser().resolve()}")
    print(f"  raw performance   : {performance_file}")
    print(f"  raw adsorbent     : {adsorbent_file}")
    print(f"  normalized adsorbent : {normalized_adsorbent_file}")
    print(f"  literature adsorbent database    : {comm_db_file}")
    print(f"  authoritative adsorbent database : {comm_authoritative_file}")
    print(f"  performance database             : {standardized_file}")
    print(f"  online properties                : {comm_online_file}")
    print(f"  PFAS props                       : {pfas_props_file}")
    print(f"  data audit                       : {'disabled' if args.skip_data_audit else 'enabled'}")
    print(
        "  human-review decisions           : "
        f"{'disabled' if args.skip_review_decisions else review_decisions_file}"
    )
    print(f"  adsorbent-property decisions     : {adsorbent_property_review_decisions_file}")
    print(f"  performance-normalization decisions: {performance_normalization_decisions_file}")
    sys.stdout.flush()


def run_merge(
    args: argparse.Namespace,
    datasets: list[dict[str, str]],
    selected_study_ids: list[str],
    performance_file: Path,
    adsorbent_file: Path,
    normalized_adsorbent_file: Path,
    adsorbent_property_review_decisions_file: Path,
) -> None:
    if str(SCRIPT_DIR) not in sys.path:
        sys.path.insert(0, str(SCRIPT_DIR))

    import merge

    print("\n[1/4] Merging extracted records...", flush=True)
    merge.run(
        datasets=datasets,
        selected_study_ids=selected_study_ids,
        study_metadata_xlsx=str(args.study_metadata_xlsx.expanduser().resolve()),
        study_metadata_sheet=args.study_metadata_sheet,
        performance_output_file=str(performance_file),
        adsorbent_output_file=str(adsorbent_file),
        adsorbent_normalized_output_file=str(normalized_adsorbent_file),
        adsorbent_property_review_decisions_file=str(
            adsorbent_property_review_decisions_file
        ),
    )


def run_standardizer(
    args: argparse.Namespace,
    performance_file: Path,
    normalized_adsorbent_file: Path,
    standardized_file: Path,
    comm_online_file: Path,
    comm_db_file: Path,
    comm_authoritative_file: Path,
    pfas_props_file: Path,
    performance_normalization_decisions_file: Path,
) -> None:
    env = os.environ.copy()
    env["BUILD_DATABASE_ORCHESTRATED"] = "1"
    env["UNIT_STANDARDIZER_INPUT_PATH"] = str(performance_file)
    env["UNIT_STANDARDIZER_OUTPUT_PATH"] = str(standardized_file)
    env["UNIT_STANDARDIZER_SHEET_NAME"] = args.standardizer_sheet_name
    env["UNIT_STANDARDIZER_ADSORBENT_INPUT_PATH"] = str(normalized_adsorbent_file)
    env["UNIT_STANDARDIZER_ADSORBENT_OUTPUT_PATH"] = str(comm_db_file)
    env["UNIT_STANDARDIZER_ADSORBENT_SHEET_NAME"] = (
        args.adsorbent_database_input_sheet
    )
    env["UNIT_STANDARDIZER_ADSORBENT_MAP_PATH"] = str(
        args.adsorbent_map_file.expanduser().resolve()
    )
    env["ENABLE_ADSORBENT_DATABASE_BUILD"] = "1"
    env["COMM_PROPS_ONLINE_PATH"] = str(comm_online_file)
    env["COMM_PROPS_ONLINE_SHEET"] = "Sheet1"
    env["ENABLE_COMM_PROPS_ONLINE_INJECTION"] = "0"
    env["COMM_PROPS_AUTHORITATIVE_PATH"] = str(comm_authoritative_file)
    env["COMM_PROPS_AUTHORITATIVE_SHEET"] = "Adsorbent_Properties"
    env["ENABLE_AUTHORITATIVE_ADSORBENT_DATABASE_BUILD"] = "1"
    env["COMM_PROPS_DB_PATH"] = str(comm_authoritative_file)
    env["COMM_PROPS_DB_SHEET"] = "Adsorbent_Properties"
    env["ENABLE_COMM_PROPS_DB_INJECTION"] = (
        "0" if args.disable_adsorbent_database_injection else "1"
    )
    env["PFAS_PROPS_PATH"] = str(pfas_props_file)
    env["PFAS_PROPS_SHEET"] = "Sheet1"
    env["ENABLE_PFAS_PROPS_INJECTION"] = "1"
    env["PERFORMANCE_NORMALIZATION_DECISIONS_PATH"] = str(
        performance_normalization_decisions_file
    )

    cmd = [args.python, str(args.standardizer_script.expanduser().resolve())]
    print(
        "\n[2/4] Standardizing units and building adsorbent/performance databases...",
        flush=True,
    )
    subprocess.run(cmd, cwd=str(SCRIPT_DIR), env=env, check=True)


def run_data_audit(
    args: argparse.Namespace,
    standardized_file: Path,
    pfas_props_file: Path,
) -> None:
    audit_script = args.data_audit_script.expanduser().resolve()
    if not audit_script.exists():
        raise FileNotFoundError(f"Model data audit script not found: {audit_script}")
    audit_output = (
        args.data_audit_output_file.expanduser().resolve()
        if args.data_audit_output_file
        else standardized_file.with_name(DEFAULT_DATA_AUDIT_FILE.name)
    )
    cmd = [
        args.python,
        str(audit_script),
        "--input-path",
        str(standardized_file),
        "--output-path",
        str(audit_output),
        "--pfas-features-path",
        str(pfas_props_file),
    ]
    print("\n[4/4] Building figure-oriented model-data audit...", flush=True)
    subprocess.run(cmd, cwd=str(audit_script.parent), check=True)


def main() -> None:
    args = parse_args()
    datasets = _resolve_datasets(args)
    selected_study_ids = _resolve_study_ids(args)
    (
        performance_file,
        adsorbent_file,
        normalized_adsorbent_file,
        standardized_file,
    ) = _resolve_outputs(args)
    review_decisions_file = _resolve_review_decisions_file(args)
    adsorbent_property_review_decisions_file = (
        _resolve_adsorbent_property_review_decisions_file(args)
    )
    performance_normalization_decisions_file = (
        _resolve_performance_normalization_decisions_file(args)
    )
    (
        comm_online_file,
        comm_db_file,
        comm_authoritative_file,
        pfas_props_file,
    ) = _resolve_standardizer_inputs(args)

    _print_plan(
        args,
        datasets,
        selected_study_ids,
        performance_file,
        adsorbent_file,
        normalized_adsorbent_file,
        standardized_file,
        comm_online_file,
        comm_db_file,
        comm_authoritative_file,
        pfas_props_file,
        review_decisions_file,
        adsorbent_property_review_decisions_file,
        performance_normalization_decisions_file,
    )
    if args.dry_run:
        return

    if args.skip_merge and args.skip_standardize:
        raise SystemExit("Nothing to do: merge and standardization were both skipped.")

    if not args.skip_merge:
        run_merge(
            args,
            datasets,
            selected_study_ids,
            performance_file,
            adsorbent_file,
            normalized_adsorbent_file,
            adsorbent_property_review_decisions_file,
        )

    if not args.skip_standardize:
        run_standardizer(
            args,
            performance_file,
            normalized_adsorbent_file,
            standardized_file,
            comm_online_file,
            comm_db_file,
            comm_authoritative_file,
            pfas_props_file,
            performance_normalization_decisions_file,
        )

    if not args.skip_standardize and not args.skip_review_decisions:
        review_summary = apply_human_review_decisions(
            standardized_file,
            review_decisions_file,
        )
        print(
            "\n[3/4] Applied human-review decisions: "
            f"{review_summary['unreliable_decisions']} unreliable source decisions, "
            f"{review_summary['unreliable_rows']} worksheet rows flagged."
        )
        if review_summary["unmatched_decisions"]:
            print(
                "[WARN] Completed review decisions did not match a source row in this build: "
                f"{review_summary['unmatched_decisions']}"
            )

    if not args.skip_data_audit:
        if not standardized_file.exists():
            raise FileNotFoundError(
                f"Cannot audit a missing standardized database: {standardized_file}"
            )
        run_data_audit(args, standardized_file, pfas_props_file)

    print("\nDone.")


if __name__ == "__main__":
    main()
