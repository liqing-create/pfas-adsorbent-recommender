from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
# IMPLEMENTATION IMPORTS ------------------------------------------------------
import csv
import base64
import binascii
import hashlib
import json
import re
import io
import sys
from dataclasses import dataclass, asdict
from datetime import datetime
import pandas as pd

try:  # Package execution
    from . import workflow_config as config
except ImportError:  # Direct ``python docling_json_preprocess.py`` execution
    import workflow_config as config

config.add_project_import_paths()

# Project-local imports -------------------------------------------------------
try:
    from .helpers.section_heading_utils import (
        METHODS_LABEL,
        RESULTS_LABEL,
        SECTION_LABELS,
        classify_heading,
        find_pre_methods_results_window,
        is_table_heading_text,
        leading_major_section_number,
        normalize_heading,
    )
    from .helpers.docling_json_common import (
        iter_body_refs as iter_docling_body_refs,
        parse_ref as parse_docling_ref,
        resolve_ref as resolve_docling_ref,
        table_markdown as render_structured_table_markdown,
    )
    from .helpers.table_repair_core import (
        attempt_summary as _core_attempt_summary,
        build_repair_selection_payload as _core_build_repair_selection_payload,
        candidate_grid_repair_reasons as _core_candidate_grid_repair_reasons,
        make_table_quality_decision as _core_make_table_quality_decision,
        make_table_repair_candidate as _core_make_table_repair_candidate,
        normalize_table_repairability_status as _core_normalize_table_repairability_status,
        select_best_table_repair_candidate as _core_select_best_table_repair_candidate,
    )
    from .helpers.pdf_table_extraction import (
        extract_camelot_embedded_caption,
        extract_pymupdf_table_candidate,
        repair_docling_table_with_camelot,
    )
    from .helpers.docx_table_extraction import extract_docx_native_table_candidate
except ImportError:
    from helpers.section_heading_utils import (
    METHODS_LABEL,
    RESULTS_LABEL,
    SECTION_LABELS,
    classify_heading,
    find_pre_methods_results_window,
    is_table_heading_text,
    leading_major_section_number,
    normalize_heading,
)
    from helpers.docling_json_common import (
        iter_body_refs as iter_docling_body_refs,
        parse_ref as parse_docling_ref,
        resolve_ref as resolve_docling_ref,
        table_markdown as render_structured_table_markdown,
    )
    from helpers.table_repair_core import (
        attempt_summary as _core_attempt_summary,
        build_repair_selection_payload as _core_build_repair_selection_payload,
        candidate_grid_repair_reasons as _core_candidate_grid_repair_reasons,
        make_table_quality_decision as _core_make_table_quality_decision,
        make_table_repair_candidate as _core_make_table_repair_candidate,
        normalize_table_repairability_status as _core_normalize_table_repairability_status,
        select_best_table_repair_candidate as _core_select_best_table_repair_candidate,
    )
    from helpers.pdf_table_extraction import (
        extract_camelot_embedded_caption,
        extract_pymupdf_table_candidate,
        repair_docling_table_with_camelot,
    )
    from helpers.docx_table_extraction import extract_docx_native_table_candidate

from chains.llm_usage import predict_with_usage, summarize_llm_usage

# Runtime settings and shared paths are centrally configured in workflow_config.
FORCE_RERUN_LLM = config.FORCE_RERUN_LLM
DEBUG_MODE = config.WRITE_DEBUG_ARTIFACTS
DEBUG_PRINT_MAX_CHARS: Optional[int] = config.DEBUG_PRINT_MAX_CHARS

# Large tables are divided into repeated-header table records before chunking.
LARGE_TABLE_SPLIT_CHAR_THRESHOLD = 10_000
LARGE_TABLE_SPLIT_TARGET_CHARS = 4_500


INPUT_DIR = config.SOURCE_STUDIES_DIR
OUTPUT_DIR = config.PROCESSED_DIR
WITHIN_SCOPE_XLSX = config.STUDY_RECORDS_XLSX
WITHIN_SCOPE_SHEET = "within_scope_records"
EXCEL_SUPPLEMENTARY_TABLE_NAME = "supplementary_table.xlsx"
# The curated study-138 workbook uses four header rows before its analyte data.
# Repeating them in each split keeps the downstream table context intact.
EXCEL_SUPPLEMENTARY_HEADER_ROWS = 4

FIGURE_TABLE_CSV = config.FIGURE_TABLE_INFO_CSV
TABLE_ISSUE_DECISIONS_CSV = config.PROCESSED_DIR / "table_issues_decisions.csv"
TABLE_ISSUE_FIELDNAMES = ["study", "file_type", "table_name", "issue"]
TABLE_ISSUE_FIXED_FILE_COLUMN = "fixed_file"
TABLE_ISSUE_DECISION_FIELDNAMES = [
    *TABLE_ISSUE_FIELDNAMES,
    "Decision",
    TABLE_ISSUE_FIXED_FILE_COLUMN,
]
PROCESSING_REPORT = config.PROCESSING_REPORT_FILE
DOCLING_TEXT_REPAIR_DECISIONS_CSV = Path(__file__).with_name(
    "docling_text_repair_decisions.csv"
)
DOCLING_TEXT_REPAIR_REQUIRED_COLUMNS = {
    "study",
    "source_file",
    "object_ref",
    "match_pattern",
    "replacement",
    "expected_matches",
    "approval_status",
}
# ``action`` is deliberately optional for backwards compatibility with the
# existing repair ledger.  Omitted values retain the original text-repair
# behaviour; object exclusions require an explicit value.
DOCLING_OBJECT_ACTION_COLUMN = "action"
DOCLING_OBJECT_ACTION_REPLACE_TEXT = "replace_text"
DOCLING_OBJECT_ACTION_EXCLUDE = "exclude"
DOCLING_OBJECT_ACTIONS = {
    DOCLING_OBJECT_ACTION_REPLACE_TEXT,
    DOCLING_OBJECT_ACTION_EXCLUDE,
}

LLM_PROVIDER = "openai"
LLM_MODEL_NAME = "gpt-4.1-2025-04-14"

TABLE_GRID_TRIAGE_SCHEMA_VERSION = 2
TABLE_POSITION_ALLOWED_STATUSES = {"clean", "table_repairable", "table_repair_needed", "uncertain"}
TABLE_POSITION_REVIEW_STATUSES = {"table_repair_needed", "uncertain"}
TABLE_REPAIR_ACCEPT_STATUSES = {"clean", "table_repairable"}
TABLE_MARKDOWN_EXPANDABLE_STATUSES = {"table_repairable"}
USE_REPAIRED_TABLE_WHEN_STILL_REVIEW = False
TABLE_REPAIR_BBOX_PAD = 3.0
TABLE_REPAIR_COLUMN_MIN_GAP = 2.0
TABLE_REPAIR_MIN_NONEMPTY_CELLS = 6
EXPAND_MARKDOWN_REPAIRABLE_TABLES = True
MARKDOWN_REPAIRABLE_EXPANSION_MIN_ROWS_ADDED = 1
# Backward-compatible aliases for older internal helper names/artifact semantics.
EXPAND_POSITION_ALIGNED_TABLES = EXPAND_MARKDOWN_REPAIRABLE_TABLES
POSITION_ALIGNED_EXPANSION_MIN_ROWS_ADDED = MARKDOWN_REPAIRABLE_EXPANSION_MIN_ROWS_ADDED
TABLE_REPAIR_OPTIONAL_BLANK_COLUMN_RE = re.compile(r"\b(?:crosslinker)\b", re.I)
# Evidence packet used for structured caption/footnote recovery.
# Deterministic recovery may use internal regex candidates; the LLM-facing
# payload is kept minimal: object_refs, layout, local text, and table_preview
# for tables only.
STRUCTURED_CAPTION_TABLE_BEFORE_CHARS = 900
STRUCTURED_CAPTION_TABLE_AFTER_CHARS = 900
STRUCTURED_CAPTION_FIGURE_BEFORE_CHARS = 900
STRUCTURED_CAPTION_FIGURE_AFTER_CHARS = 900
STRUCTURED_CAPTION_OBJECT_PREVIEW_CHARS = 2_500
STRUCTURED_CAPTION_TABLE_PREVIEW_HEAD_ROWS = 4
STRUCTURED_CAPTION_TABLE_PREVIEW_TAIL_ROWS = 2
STRUCTURED_CAPTION_TABLE_PREVIEW_MIN_ROWS_TO_COMPRESS = 10
STRUCTURED_CAPTION_CANDIDATE_WINDOW_CHARS = 350
STRUCTURED_CAPTION_SCHEMA_VERSION = 12
TABLE_HEADER_ROWS_SCHEMA_VERSION = 1
TABLE_HEADER_LLM_SAMPLE_ROWS = 10
# Caption-driven figure recovery: candidate picture groups are internal only.
# A picture is saved/reported only after Docling direct caption or a nearby
# caption anchor attaches to it.
FIGURE_CAPTION_ASSIGNMENT_MAX_STREAM_GAP = 12
MIN_PICTURE_ASSET_AREA = 4096

LARGE_TABLE_SPLIT_MIN_TABLE_CELL_PROXY = 200
LARGE_TABLE_SPLIT_TARGET_TABLE_CELL_PROXY = 200
# Backward-compatible aliases for older references. The large-table splitter
# now applies these limits with the rendered-table rectangle proxy.
LARGE_TABLE_SPLIT_MIN_DATA_CELL_PROXY = LARGE_TABLE_SPLIT_MIN_TABLE_CELL_PROXY
LARGE_TABLE_SPLIT_TARGET_DATA_CELL_PROXY = LARGE_TABLE_SPLIT_TARGET_TABLE_CELL_PROXY
LARGE_TABLE_SPLIT_PART_SUFFIX = "part"
TABLE_ROW_HEADER_FILL_MAX_COLS = 2
PRESERVE_MANUAL_TABLE_COPY_FILES = True
MANUAL_TABLE_COPY_SUFFIX_RE = re.compile(
    r"\s+-\s+Copy(?:\s*\(\d+\))?$",
    re.I,
)
MANUAL_TABLE_COPY_LEGACY_NAME_RE = re.compile(
    r"\.md\s+-\s+Copy(?:\s*\(\d+\))?$",
    re.I,
)
TARGET_MAIN_LABELS = {
    SECTION_LABELS["abstract"],
    SECTION_LABELS["materials_and_methods"],
    SECTION_LABELS["results_and_discussion"],
}

# Data classes ----------------------------------------------------------------
@dataclass
class StreamObject:
    order_index: int
    object_ref: str
    object_type: str
    label: str
    text: str
    content_layer: str
    page_no: Optional[int]
    bbox: Optional[Dict[str, Any]]
    raw: Dict[str, Any]


# General utilities -----------------------------------------------------------
def _study_sort_key(path: Path | str) -> int:
    name = path.name if isinstance(path, Path) else str(path)
    m = re.search(r"(\d+)$", name)
    return int(m.group(1)) if m else 10**9


def _study_number(study: str) -> Optional[int]:
    m = re.fullmatch(r"study_(\d+)", (study or "").strip(), flags=re.I)
    return int(m.group(1)) if m else None


def _load_within_scope_abstracts(path: Path) -> Dict[int, str]:
    """Load the curated abstract used as the authoritative abstract source."""
    if not path.is_file():
        raise FileNotFoundError(f"within_scope_records.xlsx not found: {path}")
    df = pd.read_excel(path, sheet_name=WITHIN_SCOPE_SHEET)
    columns = {str(c).strip().lower(): c for c in df.columns}
    study_col = columns.get("study_no")
    abstract_col = columns.get("abstract")
    if study_col is None or abstract_col is None:
        raise ValueError(
            f"{path} must contain 'study_no' and 'Abstract' columns "
            f"in sheet '{WITHIN_SCOPE_SHEET}'"
        )

    out: Dict[int, str] = {}
    for _, row in df.iterrows():
        try:
            study_no = int(row[study_col])
        except Exception:
            continue
        value = row[abstract_col]
        if pd.isna(value):
            abstract = ""
        else:
            abstract = re.sub(r"\s+", " ", str(value)).strip()
        out[study_no] = abstract
    return out


def _first_balanced_json_object(text: str) -> str:
    """Extract the first balanced JSON object from surrounding prose/fences."""
    start = text.find("{")
    while start >= 0:
        depth = 0
        in_string = False
        escaped = False
        for i in range(start, len(text)):
            ch = text[i]
            if in_string:
                if escaped:
                    escaped = False
                elif ch == "\\":
                    escaped = True
                elif ch == '"':
                    in_string = False
                continue
            if ch == '"':
                in_string = True
            elif ch == "{":
                depth += 1
            elif ch == "}":
                depth -= 1
                if depth == 0:
                    return text[start : i + 1]
        start = text.find("{", start + 1)
    return ""


def _escape_invalid_json_backslashes(text: str) -> str:
    """Keep valid JSON escapes, but repair raw LaTeX-style backslashes."""
    out: List[str] = []
    in_string = False
    escaped = False
    valid_escape_chars = set('"\\/bfnrtu')
    for ch in text:
        if not in_string:
            out.append(ch)
            if ch == '"':
                in_string = True
            continue
        if escaped:
            if ch not in valid_escape_chars:
                out.append("\\")
            out.append(ch)
            escaped = False
            continue
        out.append(ch)
        if ch == "\\":
            escaped = True
        elif ch == '"':
            in_string = False
    return "".join(out)


def _json_loads_dict_lenient(text: str) -> Dict[str, Any]:
    attempts = [text]
    repaired = _escape_invalid_json_backslashes(text)
    if repaired != text:
        attempts.append(repaired)
    for candidate in attempts:
        try:
            obj = json.loads(candidate)
        except Exception:
            continue
        if isinstance(obj, dict):
            return obj
    return {}


def _safe_json_loads(s: str) -> Dict[str, Any]:
    if not s:
        return {}
    txt = s.strip()
    candidates: List[str] = [txt]
    for match in re.finditer(r"```(?:json)?\s*(.*?)\s*```", txt, flags=re.S | re.I):
        fenced = match.group(1).strip()
        if fenced:
            candidates.append(fenced)
            object_text = _first_balanced_json_object(fenced)
            if object_text:
                candidates.append(object_text)
    object_text = _first_balanced_json_object(txt)
    if object_text:
        candidates.append(object_text)

    seen: set[str] = set()
    for candidate in candidates:
        candidate = candidate.strip()
        if not candidate or candidate in seen:
            continue
        seen.add(candidate)
        obj = _json_loads_dict_lenient(candidate)
        if obj:
            return obj
    return {}


def _sha256_text(text: str) -> str:
    return hashlib.sha256((text or "").encode("utf-8", errors="replace")).hexdigest()


def _sanitize_filename(value: str) -> str:
    value = re.sub(r"[<>:\"/\\|?*]+", "_", value or "")
    value = re.sub(r"\s+", "_", value.strip())
    return value.strip("_") or "object"

def _is_preserved_manual_table_copy(path: Path) -> bool:
    if not PRESERVE_MANUAL_TABLE_COPY_FILES:
        return False
    return bool(
        (
            path.suffix.lower() == ".md"
            and MANUAL_TABLE_COPY_SUFFIX_RE.search(path.stem.strip())
        )
        or MANUAL_TABLE_COPY_LEGACY_NAME_RE.search(path.name.strip())
    )


def _manual_table_copy_target(path: Path) -> Optional[Path]:
    """Return the canonical table path represented by a preserved copy.

    A manually repaired file is conventionally saved beside the generated
    file with a suffix such as `` - Copy`` or `` - Copy (1)``.  Strip only
    that terminal suffix so names such as ``segment02`` and ``part01`` remain
    part of the canonical artifact name.
    """
    if not _is_preserved_manual_table_copy(path):
        return None
    if path.suffix.lower() == ".md":
        match = MANUAL_TABLE_COPY_SUFFIX_RE.search(path.stem.strip())
        if match:
            canonical_stem = path.stem[: match.start()].rstrip()
            if canonical_stem:
                return path.with_name(f"{canonical_stem}{path.suffix}")

    legacy_match = MANUAL_TABLE_COPY_LEGACY_NAME_RE.search(path.name.strip())
    if legacy_match:
        canonical_stem = path.name[: legacy_match.start()].rstrip()
        if canonical_stem:
            return path.with_name(f"{canonical_stem}.md")
    return None


def _iter_table_markdown_files(tables_dir: Path) -> Iterable[Path]:
    try:
        candidates = sorted(tables_dir.iterdir(), key=lambda path: path.name.casefold())
    except OSError:
        return
    for path in candidates:
        if path.is_file() and (
            path.suffix.lower() == ".md" or _is_preserved_manual_table_copy(path)
        ):
            yield path


def _manual_table_split_part_parent(path: Path) -> Optional[Path]:
    """Return the unsplit canonical target for a ``*_partNN`` table path."""
    match = re.fullmatch(r"(.+)_part\d+", path.stem.strip(), re.I)
    if not match:
        return None
    return path.with_name(f"{match.group(1)}{path.suffix}")


def _manual_table_split_part_number(path: Path) -> Optional[int]:
    """Return the numeric suffix of a canonical ``*_partNN`` table path."""
    match = re.fullmatch(r".+_part(\d+)", path.stem.strip(), re.I)
    if not match:
        return None
    try:
        return int(match.group(1))
    except ValueError:
        return None


def _read_valid_manual_table_copy(path: Path) -> Optional[str]:
    """Read a preserved copy only when it contains a usable Markdown table."""
    try:
        text = path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError) as exc:
        print(f"[manual table override skipped] {path}: cannot read ({exc})")
        return None

    pipe_lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip().startswith("|") and line.strip().endswith("|")
    ]
    # Markdown table delimiter cells require at least one hyphen.  Requiring
    # three hyphens here incorrectly rejects compact but valid rows such as
    # ``|-|-:|`` produced by the manual table workflow.
    has_separator = any(
        re.match(r"^\|\s*:?-+\s*(?:\||$)", line) for line in pipe_lines
    )
    if len(pipe_lines) < 3 or not has_separator:
        print(
            f"[manual table override skipped] {path}: "
            "no validated Markdown table body found"
        )
        return None
    return text.rstrip() + "\n"


def _manual_table_copy_preference_key(path: Path) -> Tuple[int, str]:
    """Prefer the most recently edited copy, with a stable name tie-breaker."""
    try:
        modified_ns = path.stat().st_mtime_ns
    except OSError:
        modified_ns = 0
    return modified_ns, path.name.casefold()


def _manual_split_copy_groups(
    usable_by_target: Dict[Path, List[Tuple[Path, str]]],
) -> Dict[Path, List[Tuple[int, Path, Path, str]]]:
    """Group manually maintained ``*_partNN`` copies by their base target.

    This supports a user splitting a formerly unsplit generated table after a
    manual edit increases its column count.  A consolidated `` - Copy`` keeps
    precedence when one exists for the base target.
    """
    groups: Dict[Path, List[Tuple[int, Path, Path, str]]] = {}
    for part_target, usable in usable_by_target.items():
        parent = _manual_table_split_part_parent(part_target)
        part_number = _manual_table_split_part_number(part_target)
        if parent is None or part_number is None or parent in usable_by_target:
            continue
        selected_path, selected_text = max(
            usable,
            key=lambda item: _manual_table_copy_preference_key(item[0]),
        )
        groups.setdefault(parent, []).append(
            (part_number, part_target, selected_path, selected_text)
        )

    valid_groups: Dict[Path, List[Tuple[int, Path, Path, str]]] = {}
    for parent, parts in groups.items():
        parts.sort(key=lambda item: item[0])
        part_numbers = [item[0] for item in parts]
        if len(parts) < 2 or len(set(part_numbers)) != len(part_numbers):
            continue
        valid_groups[parent] = parts
    return valid_groups


def _record_path_matches(path_value: Any, target: Path) -> bool:
    try:
        return Path(str(path_value or "")).name.casefold() == target.name.casefold()
    except (TypeError, ValueError):
        return False


def _resolve_fixed_file_path(tables_dir: Path, fixed_file: str) -> Optional[Path]:
    """Resolve a decisions-file fixed-file value to an existing local file."""
    raw_value = str(fixed_file or "").strip().strip('"')
    if not raw_value:
        return None

    raw_path = Path(raw_value)
    base_candidates = (
        [raw_path]
        if raw_path.is_absolute()
        else [
            tables_dir / raw_path,
            tables_dir.parent / raw_path,
            tables_dir.parent.parent / raw_path,
        ]
    )
    candidates: List[Path] = []
    for base_candidate in base_candidates:
        candidates.append(base_candidate)
        if base_candidate.suffix.lower() != ".md":
            candidates.append(base_candidate.with_name(base_candidate.name + ".md"))
        legacy_match = MANUAL_TABLE_COPY_LEGACY_NAME_RE.search(base_candidate.name)
        if legacy_match:
            canonical_copy_name = (
                base_candidate.name[: legacy_match.start()].rstrip()
                + " - Copy.md"
            )
            candidates.append(base_candidate.with_name(canonical_copy_name))

    seen: set[Path] = set()
    for candidate in candidates:
        if candidate in seen:
            continue
        seen.add(candidate)
        if candidate.is_file():
            return candidate
    return None


def _explicit_manual_copy_candidates(
    *,
    tables_dir: Path,
    table_issue_decisions: Dict[Tuple[str, str, str], Dict[str, str]],
) -> Dict[Path, List[Path]]:
    """Map Fixed decisions with an explicit file to generated table targets."""
    candidates_by_target: Dict[Path, List[Path]] = {}
    study_key = tables_dir.parent.name.casefold()
    for key, decision_info in table_issue_decisions.items():
        if key[0] != study_key:
            continue
        if _table_issue_decision_value(decision_info) != "fixed":
            continue
        fixed_file = str(decision_info.get("fixed_file") or "").strip()
        if not fixed_file:
            continue

        target_name = _sanitize_filename(
            f"{key[1]}_{_table_id_filename_fragment(key[2])}"
        ) + ".md"
        target = tables_dir / target_name
        for fixed_file_value in re.split(r"[;\r\n]+", fixed_file):
            fixed_path = _resolve_fixed_file_path(tables_dir, fixed_file_value)
            if fixed_path is None:
                print(
                    f"[manual table override skipped] fixed_file not found for "
                    f"{key[0]}/{key[1]}/{key[2]}: {fixed_file_value.strip()}"
                )
                continue
            candidates_by_target.setdefault(target, []).append(fixed_path)
    return candidates_by_target


def _apply_manual_table_copy_overrides(
    *,
    tables_dir: Path,
    chunk_records: List[Dict[str, Any]],
    figtab_rows: List[Dict[str, Any]],
    table_issue_decisions: Dict[Tuple[str, str, str], Dict[str, str]],
) -> List[Dict[str, str]]:
    """Replace generated table artifacts with validated preserved repairs.

    The override is applied after all source processing so generated files do
    not overwrite the manual copy.  Chunk and figure/table metadata records
    are updated at the same time; otherwise the Markdown file and downstream
    JSONL inputs would disagree after a rerun.
    """
    explicit_candidates_by_target = _explicit_manual_copy_candidates(
        tables_dir=tables_dir,
        table_issue_decisions=table_issue_decisions,
    )
    explicit_source_paths = {
        str(path).casefold()
        for paths in explicit_candidates_by_target.values()
        for path in paths
    }
    candidates_by_target: Dict[Path, List[Path]] = {}
    for candidate in _iter_table_markdown_files(tables_dir):
        if str(candidate).casefold() in explicit_source_paths:
            continue
        target = _manual_table_copy_target(candidate)
        if target is None:
            continue
        candidates_by_target.setdefault(target, []).append(candidate)
    candidates_by_target.update(explicit_candidates_by_target)

    usable_by_target: Dict[Path, List[Tuple[Path, str]]] = {}
    for target, candidates in candidates_by_target.items():
        usable = [
            (candidate, text)
            for candidate in candidates
            if (text := _read_valid_manual_table_copy(candidate)) is not None
        ]
        if usable:
            usable_by_target[target] = usable

    manual_split_groups = _manual_split_copy_groups(usable_by_target)
    manual_split_part_targets = {
        part_target
        for parts in manual_split_groups.values()
        for _, part_target, _, _ in parts
    }

    # A consolidated manual copy such as ``main_paper_table_2 - Copy.md``
    # supersedes generated ``main_paper_table_2_partNN.md`` artifacts and any
    # older per-part manual copies.  Suppress those part targets so they
    # cannot be applied again later in this loop.
    consolidated_targets = {
        target
        for target in usable_by_target
        if _manual_table_split_part_parent(target) is None
    }
    superseded_part_targets = {
        target
        for target in candidates_by_target
        if _manual_table_split_part_parent(target) in consolidated_targets
    }

    overrides: List[Dict[str, str]] = []
    for parent_target, parts in sorted(
        manual_split_groups.items(), key=lambda item: item[0].name.casefold()
    ):
        part_targets = [part_target for _, part_target, _, _ in parts]
        record_targets = [parent_target, *part_targets]
        for table_path in _iter_table_markdown_files(tables_dir):
            canonical_path = (
                _manual_table_copy_target(table_path)
                if _is_preserved_manual_table_copy(table_path)
                else table_path
            )
            if (
                canonical_path is not None
                and _manual_table_split_part_parent(canonical_path) == parent_target
                and canonical_path not in record_targets
            ):
                record_targets.append(canonical_path)

        matching_chunk_indices = [
            index
            for index, record in enumerate(chunk_records)
            if any(
                _record_path_matches(record.get("artifact_path"), record_target)
                for record_target in record_targets
            )
        ]
        matched_chunk_records = len(matching_chunk_indices)

        matching_figtab_indices = [
            index
            for index, row in enumerate(figtab_rows)
            if any(
                _record_path_matches(row.get("file_path"), record_target)
                for record_target in record_targets
            )
        ]
        matched_figtab_rows = len(matching_figtab_indices)

        # Write the canonical split artifacts before discarding stale generated
        # siblings.  The preserved copies remain intact as the source of truth.
        for _, part_target, _, selected_text in parts:
            part_target.write_text(selected_text, encoding="utf-8")

        removed_split_paths: List[str] = []
        selected_target_set = set(part_targets)
        for table_path in _iter_table_markdown_files(tables_dir):
            if _is_preserved_manual_table_copy(table_path):
                continue
            canonical_path = table_path
            if canonical_path == parent_target or (
                _manual_table_split_part_parent(canonical_path) == parent_target
                and canonical_path not in selected_target_set
            ):
                try:
                    table_path.unlink()
                except OSError as exc:
                    print(
                        f"[manual table split override warning] could not remove "
                        f"superseded table part {table_path}: {exc}"
                    )
                else:
                    removed_split_paths.append(str(table_path))

        if matching_chunk_indices:
            representative_index = next(
                (
                    index
                    for index in matching_chunk_indices
                    if _record_path_matches(
                        chunk_records[index].get("artifact_path"), parent_target
                    )
                ),
                matching_chunk_indices[0],
            )
            representative = chunk_records[representative_index]
            base_table_id = re.sub(
                r"_part\d+$", "", str(representative.get("table_id") or ""), flags=re.I
            )
            physical_table_id = str(
                representative.get("physical_table_id") or base_table_id
            )
            replacement_records: List[Dict[str, Any]] = []
            for part_number, part_target, selected_path, selected_text in parts:
                replacement = dict(representative)
                replacement["text"] = selected_text.strip()
                replacement["markdown"] = selected_text.strip()
                replacement["artifact_path"] = str(part_target)
                replacement["source"] = "manual_table_copy"
                replacement["metadata_source"] = (
                    f"{representative.get('metadata_source') or 'generated'}+manual_table_copy"
                )
                replacement["manual_table_copy_applied"] = True
                replacement["manual_table_copy_path"] = str(selected_path)
                replacement["needs_review"] = False
                replacement["table_id"] = f"{base_table_id}_part{part_number:02d}"
                replacement["physical_table_id"] = physical_table_id
                replacement["split_part"] = part_number
                replacement_records.append(replacement)

            matching_chunk_index_set = set(matching_chunk_indices)
            chunk_records[:] = [
                output_record
                for index, record in enumerate(chunk_records)
                for output_record in (
                    replacement_records
                    if index == representative_index
                    else ([] if index in matching_chunk_index_set else [record])
                )
            ]

        if matching_figtab_indices:
            representative_figtab_index = next(
                (
                    index
                    for index in matching_figtab_indices
                    if _record_path_matches(
                        figtab_rows[index].get("file_path"), parent_target
                    )
                ),
                matching_figtab_indices[0],
            )
            representative_row = figtab_rows[representative_figtab_index]
            first_part_target = part_targets[0]
            first_part_source = parts[0][2]
            representative_row["file_path"] = str(first_part_target)
            representative_row["source"] = "manual_table_copy"
            representative_row["metadata_source"] = (
                f"{representative_row.get('metadata_source') or 'generated'}+manual_table_copy"
            )
            representative_row["needs_review"] = "false"
            representative_row["manual_table_copy_applied"] = "true"
            representative_row["manual_table_copy_path"] = str(first_part_source)
            matching_figtab_index_set = set(matching_figtab_indices)
            figtab_rows[:] = [
                row
                for index, row in enumerate(figtab_rows)
                if index == representative_figtab_index or index not in matching_figtab_index_set
            ]

        if matched_chunk_records == 0:
            source_names = ", ".join(selected_path.name for _, _, selected_path, _ in parts)
            print(
                f"[manual table split override warning] [{source_names}] -> "
                f"{parent_target.name}: no generated chunk record matched"
            )
        print(
            f"[manual table split override] {parent_target.name} "
            f"(parts={len(parts)}, chunk_records={matched_chunk_records}, "
            f"figtab_rows={matched_figtab_rows}, removed_parts={len(removed_split_paths)})"
        )
        overrides.append(
            {
                "source_path": "; ".join(
                    str(selected_path) for _, _, selected_path, _ in parts
                ),
                "target_path": "; ".join(str(part_target) for _, part_target, _, _ in parts),
                "chunk_records": str(matched_chunk_records),
                "figtab_rows": str(matched_figtab_rows),
                "removed_parts": str(len(removed_split_paths)),
            }
        )

    for target, usable in sorted(
        usable_by_target.items(), key=lambda item: item[0].name.casefold()
    ):
        if target in superseded_part_targets or target in manual_split_part_targets:
            continue

        selected_path, selected_text = max(
            usable,
            key=lambda item: _manual_table_copy_preference_key(item[0]),
        )
        if len(usable) > 1:
            selected_names = ", ".join(path.name for path, _ in usable)
            print(
                f"[manual table override] multiple copies for {target.name}; "
                f"selected {selected_path.name} from [{selected_names}]"
            )

        target.write_text(selected_text, encoding="utf-8")
        split_part_targets = [
            part_target
            for part_target in candidates_by_target
            if _manual_table_split_part_parent(part_target) == target
        ]
        for table_path in _iter_table_markdown_files(tables_dir):
            canonical_path = (
                _manual_table_copy_target(table_path)
                if _is_preserved_manual_table_copy(table_path)
                else table_path
            )
            if (
                canonical_path is not None
                and _manual_table_split_part_parent(canonical_path) == target
                and canonical_path not in split_part_targets
            ):
                split_part_targets.append(canonical_path)

        record_targets = [target, *split_part_targets]
        matching_chunk_indices = [
            index
            for index, record in enumerate(chunk_records)
            if any(
                _record_path_matches(record.get("artifact_path"), record_target)
                for record_target in record_targets
            )
        ]
        matched_chunk_records = len(matching_chunk_indices)
        if matching_chunk_indices:
            representative_index = next(
                (
                    index
                    for index in matching_chunk_indices
                    if _record_path_matches(
                        chunk_records[index].get("artifact_path"), target
                    )
                ),
                matching_chunk_indices[0],
            )
            representative = chunk_records[representative_index]
            representative["text"] = selected_text.strip()
            representative["markdown"] = selected_text.strip()
            representative["artifact_path"] = str(target)
            representative["source"] = "manual_table_copy"
            representative["metadata_source"] = (
                f"{representative.get('metadata_source') or 'generated'}+manual_table_copy"
            )
            representative["manual_table_copy_applied"] = True
            representative["manual_table_copy_path"] = str(selected_path)
            representative["needs_review"] = False
            logical_table_id = (
                representative.get("logical_table_id")
                or representative.get("physical_table_id")
            )
            if logical_table_id:
                representative["table_id"] = logical_table_id
            else:
                representative["table_id"] = re.sub(
                    r"_part\d+$", "", str(representative.get("table_id") or ""), flags=re.I
                )
            representative["split_part"] = None
            matching_chunk_index_set = set(matching_chunk_indices)
            chunk_records[:] = [
                record
                for index, record in enumerate(chunk_records)
                if index == representative_index or index not in matching_chunk_index_set
            ]

        matching_figtab_indices = [
            index
            for index, row in enumerate(figtab_rows)
            if any(
                _record_path_matches(row.get("file_path"), record_target)
                for record_target in record_targets
            )
        ]
        matched_figtab_rows = len(matching_figtab_indices)
        if matching_figtab_indices:
            representative_figtab_index = next(
                (
                    index
                    for index in matching_figtab_indices
                    if _record_path_matches(figtab_rows[index].get("file_path"), target)
                ),
                matching_figtab_indices[0],
            )
            representative_row = figtab_rows[representative_figtab_index]
            representative_row["file_path"] = str(target)
            representative_row["source"] = "manual_table_copy"
            representative_row["metadata_source"] = (
                f"{representative_row.get('metadata_source') or 'generated'}+manual_table_copy"
            )
            representative_row["needs_review"] = "false"
            representative_row["manual_table_copy_applied"] = "true"
            representative_row["manual_table_copy_path"] = str(selected_path)
            matching_figtab_index_set = set(matching_figtab_indices)
            figtab_rows[:] = [
                row
                for index, row in enumerate(figtab_rows)
                if index == representative_figtab_index or index not in matching_figtab_index_set
            ]

        removed_split_paths: List[str] = []
        split_target_set = set(split_part_targets)
        for table_path in _iter_table_markdown_files(tables_dir):
            canonical_path = (
                _manual_table_copy_target(table_path)
                if _is_preserved_manual_table_copy(table_path)
                else table_path
            )
            if canonical_path not in split_target_set:
                continue
            try:
                table_path.unlink()
            except OSError as exc:
                print(
                    f"[manual table override warning] could not remove superseded "
                    f"table part {table_path}: {exc}"
                )
            else:
                removed_split_paths.append(str(table_path))

        if matched_chunk_records == 0:
            print(
                f"[manual table override warning] {selected_path.name} -> "
                f"{target.name}: no generated chunk record matched"
            )
        print(
            f"[manual table override] {selected_path.name} -> {target.name} "
            f"(chunk_records={matched_chunk_records}, figtab_rows={matched_figtab_rows}, "
            f"removed_parts={len(removed_split_paths)})"
        )
        overrides.append(
            {
                "source_path": str(selected_path),
                "target_path": str(target),
                "chunk_records": str(matched_chunk_records),
                "figtab_rows": str(matched_figtab_rows),
                "removed_parts": str(len(removed_split_paths)),
            }
        )
    return overrides

def _roman_numeral_to_int(value: str) -> Optional[int]:
    value = str(value or "").upper()
    if not value or not re.fullmatch(
        r"M{0,3}(?:CM|CD|D?C{0,3})(?:XC|XL|L?X{0,3})(?:IX|IV|V?I{0,3})",
        value,
    ):
        return None
    total = 0
    previous = 0
    values = {"I": 1, "V": 5, "X": 10, "L": 50, "C": 100, "D": 500, "M": 1000}
    for ch in reversed(value):
        current = values[ch]
        if current < previous:
            total -= current
        else:
            total += current
            previous = current
    return total


def _table_id_filename_fragment(table_id: str) -> str:
    fragment = str(table_id or "").strip().lower().replace(" ", "_")
    match = re.match(r"^(table)_(?P<num>[ivxlcdm]+)(?P<suffix>(?:_.+)?)$", fragment, flags=re.I)
    if not match:
        return fragment
    value = _roman_numeral_to_int(match.group("num"))
    if value is None:
        return fragment
    return f"{match.group(1)}_{value}{match.group('suffix') or ''}"

def _iter_nested_values(obj: Any) -> Iterable[Any]:
    """Yield nested scalar/dict/list values from a Docling object."""
    yield obj
    if isinstance(obj, dict):
        for value in obj.values():
            yield from _iter_nested_values(value)
    elif isinstance(obj, list):
        for value in obj:
            yield from _iter_nested_values(value)


def _decode_docling_picture_png(pic: Dict[str, Any]) -> Optional[bytes]:
    """Extract embedded Docling picture bytes and normalize to PNG when possible.

    Docling picture images are commonly serialized as an ImageRef such as:
    {"image": {"uri": "data:image/png;base64,..."}}. PNG payloads are written
    directly; other image/* payloads are converted with Pillow when available.
    """
    for value in _iter_nested_values(pic):
        if not isinstance(value, str):
            continue
        s = value.strip()
        if not s:
            continue

        mime = ""
        payload = s
        m = re.match(r"^data:(?P<mime>image/[^;,]+);base64,(?P<data>.+)$", s, flags=re.I | re.S)
        if m:
            mime = (m.group("mime") or "").lower()
            payload = m.group("data") or ""
        elif not re.fullmatch(r"[A-Za-z0-9+/=\r\n]+", s) or len(s) < 128:
            continue

        try:
            data = base64.b64decode(re.sub(r"\s+", "", payload), validate=True)
        except (binascii.Error, ValueError):
            continue

        if data.startswith(b"\x89PNG\r\n\x1a\n"):
            return data
        if mime and not mime.startswith("image/"):
            continue

        try:
            from PIL import Image  # type: ignore

            output = io.BytesIO()
            with Image.open(io.BytesIO(data)) as img:
                img.save(output, format="PNG")
            return output.getvalue()
        except Exception:
            continue
    return None


def _write_picture_png_asset(
    *,
    study: str,
    source_file: str,
    object_id: str,
    pic: Dict[str, Any],
) -> Tuple[str, str]:
    """Save an embedded figure PNG and return (file_stem, file_path)."""
    file_stem = _sanitize_filename(f"{source_file}_{object_id.lower().replace(' ', '_')}")
    png_bytes = _decode_docling_picture_png(pic)
    if not png_bytes:
        return file_stem, ""

    figures_dir = OUTPUT_DIR / study / "figures"
    figures_dir.mkdir(parents=True, exist_ok=True)
    figure_file = figures_dir / f"{file_stem}.png"
    figure_file.write_bytes(png_bytes)
    return figure_file.stem, str(figure_file)


def _picture_pixel_size(pic: Dict[str, Any]) -> Optional[Tuple[int, int]]:
    """Return embedded picture dimensions when available.

    Used only as a guardrail against decorative images/logos/icons. Unknown
    dimensions are treated as usable so PDF figures are not discarded merely
    because an image backend was unavailable.
    """
    data = _decode_docling_picture_png(pic)
    if not data:
        return None
    try:
        from PIL import Image  # type: ignore

        with Image.open(io.BytesIO(data)) as img:
            return int(img.width), int(img.height)
    except Exception:
        return None


def _picture_is_candidate_asset(pic: Dict[str, Any]) -> bool:
    size = _picture_pixel_size(pic)
    if not size:
        return True
    width, height = size
    return width * height >= MIN_PICTURE_ASSET_AREA

_DOCLING_LIGATURE_REPLACEMENTS = {
    "/uniFB00": "ff",
    "/uniFB01": "fi",
    "/uniFB02": "fl",
    "/uniFB03": "ffi",
    "/uniFB04": "ffl",
    "\ue103": "fi",
    "\ue104": "fl",
    "\ue09d": "ft",
}


@dataclass(frozen=True)
class DoclingTextRepairDecision:
    """One approved, document-scoped repair or object-exclusion decision."""

    study: str
    source_file: str
    object_ref: str
    match_pattern: str
    replacement: str
    expected_matches: int
    row_number: int
    action: str = DOCLING_OBJECT_ACTION_REPLACE_TEXT
    source_pdf_page: str = ""
    verification_note: str = ""


def _load_docling_text_repair_decisions(
    path: Optional[Path] = None,
) -> List[DoclingTextRepairDecision]:
    """Load approved text repairs and object exclusions from the CSV ledger."""
    csv_path = path or DOCLING_TEXT_REPAIR_DECISIONS_CSV
    if not csv_path.is_file():
        raise FileNotFoundError(
            "Docling text-repair decision ledger is missing: "
            f"{csv_path}"
        )

    with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        actual_columns = set(reader.fieldnames or [])
        missing_columns = DOCLING_TEXT_REPAIR_REQUIRED_COLUMNS - actual_columns
        if missing_columns:
            raise ValueError(
                f"{csv_path} is missing required columns: "
                f"{sorted(missing_columns)}"
            )

        decisions: List[DoclingTextRepairDecision] = []
        for row_number, row in enumerate(reader, start=2):
            approval_status = str(row.get("approval_status") or "").strip().casefold()
            if not approval_status:
                continue
            if approval_status != "approved":
                continue

            action = (
                str(row.get(DOCLING_OBJECT_ACTION_COLUMN) or "")
                .strip()
                .casefold()
                or DOCLING_OBJECT_ACTION_REPLACE_TEXT
            )
            if action not in DOCLING_OBJECT_ACTIONS:
                raise ValueError(
                    f"{csv_path}:{row_number} has an unsupported "
                    f"{DOCLING_OBJECT_ACTION_COLUMN}: {action!r}; expected one of "
                    f"{sorted(DOCLING_OBJECT_ACTIONS)}"
                )

            required_columns = {"study", "source_file", "object_ref"}
            if action == DOCLING_OBJECT_ACTION_REPLACE_TEXT:
                required_columns.update({"match_pattern", "replacement"})
            required_values = {
                column: str(row.get(column) or "").strip()
                for column in required_columns
            }
            empty_columns = [
                column for column, value in required_values.items() if not value
            ]
            if empty_columns:
                raise ValueError(
                    f"{csv_path}:{row_number} has empty required values for "
                    f"{sorted(empty_columns)}"
                )
            match_pattern = ""
            replacement = ""
            expected_matches = 0
            if action == DOCLING_OBJECT_ACTION_REPLACE_TEXT:
                try:
                    expected_matches = int(
                        str(row.get("expected_matches") or "").strip()
                    )
                except ValueError as exc:
                    raise ValueError(
                        f"{csv_path}:{row_number} has an invalid expected_matches value"
                    ) from exc
                if expected_matches < 1:
                    raise ValueError(
                        f"{csv_path}:{row_number} expected_matches must be at least 1"
                    )

                match_pattern = required_values["match_pattern"]
                replacement = required_values["replacement"]
                try:
                    compiled_pattern = re.compile(match_pattern)
                except re.error as exc:
                    raise ValueError(
                        f"{csv_path}:{row_number} has an invalid match_pattern: "
                        f"{match_pattern!r}"
                    ) from exc
                if compiled_pattern.search(""):
                    raise ValueError(
                        f"{csv_path}:{row_number} match_pattern must not match an empty string"
                    )

            decisions.append(
                DoclingTextRepairDecision(
                    study=required_values["study"],
                    source_file=required_values["source_file"],
                    object_ref=required_values["object_ref"],
                    match_pattern=match_pattern,
                    replacement=replacement,
                    expected_matches=expected_matches,
                    row_number=row_number,
                    action=action,
                    source_pdf_page=str(row.get("source_pdf_page") or "").strip(),
                    verification_note=str(row.get("verification_note") or "").strip(),
                )
            )
    return decisions


def _apply_approved_docling_object_exclusions(
    *,
    study: str,
    tables_dir: Path,
    chunk_records: List[Dict[str, Any]],
    object_records: List[Dict[str, Any]],
    figtab_rows: List[Dict[str, Any]],
) -> List[Dict[str, str]]:
    """Remove centrally approved table objects from downstream artifacts.

    Exclusions are intentionally applied after table extraction and manual
    overrides.  This preserves the source object's normal continuation and
    caption handling for its siblings, then removes only the exact object
    named in the ledger from every downstream table artifact.
    """
    decisions = [
        decision
        for decision in _load_docling_text_repair_decisions()
        if decision.study == study
        and decision.action == DOCLING_OBJECT_ACTION_EXCLUDE
    ]
    if not decisions:
        return []

    seen_targets: Dict[Tuple[str, str], int] = {}
    for decision in decisions:
        target = (decision.source_file, decision.object_ref)
        previous_row = seen_targets.get(target)
        if previous_row is not None:
            raise ValueError(
                "Duplicate approved Docling object exclusions for "
                f"{study}/{decision.source_file}/{decision.object_ref}: "
                f"ledger rows {previous_row} and {decision.row_number}"
            )
        seen_targets[target] = decision.row_number

    try:
        resolved_tables_dir = tables_dir.resolve()
    except OSError:
        resolved_tables_dir = tables_dir

    applied: List[Dict[str, str]] = []
    for decision in decisions:
        matching_objects = [
            record
            for record in object_records
            if record.get("source_file") == decision.source_file
            and record.get("object_ref") == decision.object_ref
        ]
        if len(matching_objects) != 1:
            raise ValueError(
                "Docling object-exclusion target must match exactly one "
                f"document object for {study}/{decision.source_file}/"
                f"{decision.object_ref}; matched {len(matching_objects)} "
                f"(decision ledger row {decision.row_number})"
            )
        if matching_objects[0].get("object_type") != "table":
            raise ValueError(
                "Docling object-exclusion target must be a table for "
                f"{study}/{decision.source_file}/{decision.object_ref} "
                f"(decision ledger row {decision.row_number})"
            )

        matching_chunk_indices = [
            index
            for index, record in enumerate(chunk_records)
            if record.get("file_type") == "table"
            and record.get("source_file") == decision.source_file
            and record.get("object_ref") == decision.object_ref
        ]
        matching_figtab_indices = [
            index
            for index, row in enumerate(figtab_rows)
            if row.get("type") == "table"
            and row.get("source_file") == decision.source_file
            and row.get("object_ref") == decision.object_ref
        ]

        artifact_paths = {
            str(chunk_records[index].get("artifact_path") or "").strip()
            for index in matching_chunk_indices
        }
        artifact_paths.update(
            str(figtab_rows[index].get("file_path") or "").strip()
            for index in matching_figtab_indices
        )
        artifact_paths.discard("")

        matching_chunk_index_set = set(matching_chunk_indices)
        chunk_records[:] = [
            record
            for index, record in enumerate(chunk_records)
            if index not in matching_chunk_index_set
        ]
        matching_figtab_index_set = set(matching_figtab_indices)
        figtab_rows[:] = [
            row
            for index, row in enumerate(figtab_rows)
            if index not in matching_figtab_index_set
        ]

        removed_artifacts = 0
        for raw_path in sorted(artifact_paths, key=str.casefold):
            artifact_path = Path(raw_path)
            try:
                resolved_artifact_path = artifact_path.resolve()
            except OSError:
                resolved_artifact_path = artifact_path
            if resolved_artifact_path.parent != resolved_tables_dir:
                print(
                    f"[docling-object-exclusion warning] not removing artifact "
                    f"outside the study tables directory: {artifact_path}"
                )
                continue
            if not artifact_path.is_file():
                continue
            try:
                artifact_path.unlink()
            except OSError as exc:
                print(
                    f"[docling-object-exclusion warning] could not remove "
                    f"{artifact_path}: {exc}"
                )
            else:
                removed_artifacts += 1

        print(
            f"[docling-object-exclusion] {study}/{decision.source_file}/"
            f"{decision.object_ref}: removed {len(matching_chunk_indices)} "
            f"chunk record(s), {len(matching_figtab_indices)} figure/table "
            f"metadata record(s), and {removed_artifacts} artifact(s)."
        )
        applied.append(
            {
                "source_file": decision.source_file,
                "object_ref": decision.object_ref,
                "chunk_records": str(len(matching_chunk_indices)),
                "figtab_rows": str(len(matching_figtab_indices)),
                "removed_artifacts": str(removed_artifacts),
            }
        )
    return applied


def _apply_approved_docling_text_repairs(
    doc: Dict[str, Any],
    *,
    study: str,
    source_file: str,
) -> int:
    """Apply only exact, source-PDF-approved repairs to in-memory text.

    The raw Docling JSON is intentionally never rewritten.  Keeping decisions
    in a separate ledger makes each change reviewable and reusable whenever
    preprocessing is rerun.
    """
    decisions = [
        decision
        for decision in _load_docling_text_repair_decisions()
        if decision.study == study
        and decision.source_file == source_file
        and decision.action == DOCLING_OBJECT_ACTION_REPLACE_TEXT
    ]
    if not decisions:
        return 0

    text_objects = {
        str(item.get("self_ref") or ""): item
        for item in doc.get("texts", [])
        if isinstance(item, dict) and isinstance(item.get("text"), str)
    }
    table_objects = {
        str(item.get("self_ref") or ""): item
        for item in doc.get("tables", [])
        if isinstance(item, dict)
    }
    applied_count = 0
    for decision in decisions:
        text_target = text_objects.get(decision.object_ref)
        table_target = table_objects.get(decision.object_ref)
        if text_target is None and table_target is None:
            raise ValueError(
                "Docling text-repair target is missing for "
                f"{study}/{source_file}: {decision.object_ref} "
                f"(decision ledger row {decision.row_number})"
            )

        pattern = re.compile(decision.match_pattern)
        match_count = 0
        if text_target is not None:
            repaired_text, match_count = pattern.subn(
                decision.replacement,
                str(text_target["text"]),
            )
            text_target["text"] = repaired_text
        else:
            table_data = table_target.get("data") or {}
            grid = table_data.get("grid") if isinstance(table_data, dict) else None
            if not isinstance(grid, list):
                raise ValueError(
                    "Docling text-repair table target has no grid for "
                    f"{study}/{source_file}: {decision.object_ref} "
                    f"(decision ledger row {decision.row_number})"
                )
            for row in grid:
                if not isinstance(row, list):
                    continue
                for cell in row:
                    if not isinstance(cell, dict) or not isinstance(cell.get("text"), str):
                        continue
                    repaired_text, cell_matches = pattern.subn(
                        decision.replacement,
                        cell["text"],
                    )
                    cell["text"] = repaired_text
                    match_count += cell_matches
        if match_count != decision.expected_matches:
            raise ValueError(
                "Docling text-repair match count changed for "
                f"{study}/{source_file}/{decision.object_ref}: "
                f"{decision.match_pattern!r} matched {match_count} time(s), "
                f"expected {decision.expected_matches} "
                f"(decision ledger row {decision.row_number})."
            )
        applied_count += match_count

    print(
        f"[docling-text-repair] {study}/{source_file}: applied "
        f"{applied_count} approved repair(s) from "
        f"{DOCLING_TEXT_REPAIR_DECISIONS_CSV.name}."
    )
    return applied_count


_HTML_BREAK_RE = re.compile(
    r"(?:<\s*br\s*/?\s*>|&lt;\s*br\s*/?\s*&gt;)",
    re.I,
)


def _remove_html_breaks(text: str) -> str:
    return _HTML_BREAK_RE.sub(" ", text or "")

_MD_TABLE_SEPARATOR_CELL_RE = re.compile(r":?-{3,}:?")


def _repair_docling_symbol_artifacts(text: str) -> str:
    s = text or ""
    s = _remove_html_breaks(s)
    for token, replacement in _DOCLING_LIGATURE_REPLACEMENTS.items():
        s = s.replace(token, replacement)

    s = s.replace("\u00c2\u00b5", "\u00b5")
    s = re.sub(r"\u013e(?=\s*mol\b)", "\u00b5", s, flags=re.I)
    s = s.replace("\u00bc", "=")
    s = s.replace("\u00fe", "+")
    s = re.sub(r"/C14\s*C\b", "\u00b0C", s)
    s = re.sub(r"(?<![A-Za-z0-9/])/C0\s*(?=\d)", "-", s)
    s = re.sub(
        r"\b(\d+(?:\.\d+)?)\s*([Ee])\s*([+\-])\s*(\d+)\b",
        r"\1\2\3\4",
        s,
    )
    return s


def _split_markdown_pipe_row(line: str) -> List[str]:
    parts: List[str] = []
    buf: List[str] = []
    escaped = False
    for ch in line:
        if ch == "\\" and not escaped:
            buf.append(ch)
            escaped = True
            continue
        if ch == "|" and not escaped:
            parts.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
        escaped = False
    parts.append("".join(buf))
    return parts


def _clean_rendered_markdown_table(markdown: str) -> str:
    lines: List[str] = []
    for line in (markdown or "").splitlines():
        if not line.lstrip().startswith("|"):
            lines.append(_clean_text(line))
            continue

        parts = _split_markdown_pipe_row(line)
        if len(parts) <= 2:
            lines.append(_clean_text(line))
            continue

        for idx in range(1, len(parts) - 1):
            cell = parts[idx].strip()
            cleaned = (
                cell
                if _MD_TABLE_SEPARATOR_CELL_RE.fullmatch(cell)
                else _clean_text(cell)
            )
            parts[idx] = f" {cleaned} "
        lines.append("|".join(parts))
    return "\n".join(lines).strip()


def _clean_text(text: str) -> str:
    s = text or ""
    s = _repair_docling_symbol_artifacts(s)

    # Docling may emit inline equation wrappers such as <eq>K_{d}</eq>.
    # These are not Markdown-native and can make some editors behave
    # inconsistently around nearby pipe tables. Normalize them to inline math.
    def _eq_tag_to_inline_math(match: re.Match[str]) -> str:
        inner = re.sub(r"\s+", " ", match.group(1) or "").strip()
        return f"${inner}$" if inner else ""

    s = re.sub(
        r"<\s*eq\s*>(.*?)<\s*/\s*eq\s*>",
        _eq_tag_to_inline_math,
        s,
        flags=re.I | re.S,
    )

    # Docling uses GLYPH<0> for minus signs in many units and exponents.
    s = re.sub(r"(?:GLYPH<0>|GLYPH&lt;0&gt;)", "-", s)
    s = re.sub(r"GLYPH<\d+>", "", s)
    s = re.sub(r"-\s+(?=\d)", "-", s)
    s = re.sub(
        r"\b(\d+(?:\.\d+)?[Ee][+\-]\d+)\s+e\s+(\d+(?:\.\d+)?[Ee]\s*(?:[+\-]\s*)?\d+)\b",
        r"\1 - \2",
        s,
    )
    s = re.sub(
        r"(?<![A-Za-z0-9.+\-])(\d+(?:\.\d+)?)\s+e\s+(\d+(?:\.\d+)?)(?![A-Za-z])",
        r"\1 - \2",
        s,
    )
    s = s.replace("\u00a0", " ")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\s+([,.;:)\]])", r"\1", s)
    s = re.sub(r"([([])\s+", r"\1", s)
    return s.strip()


def _looks_like_ocr_formula_junk(text: str) -> bool:
    s = (text or "").strip()
    if not s:
        return True
    if "\\begin{array" in s or "\\end{array" in s:
        return True
    if s.count("&") >= 8:
        return True
    if s.count("\\\\") >= 6:
        return True
    if len(s) >= 4_000:
        return True
    if len(s) >= 600 and re.search(r"(?:[A-Za-z]\s){20,}[A-Za-z]", s):
        return True
    if s.count("\\") >= 40 and len(s) < 1_500:
        return True
    return False


def _clean_formula_text(text: str) -> str:
    s = _clean_text(text)
    if _looks_like_ocr_formula_junk(s):
        return ""
    s = re.sub(r"\\text\s*\{\s*([^}]*)\s*\}", r"\1", s)
    s = re.sub(r"\\begin\{matrix\}(?:\s|\\)+\\end\{matrix\}", "", s, flags=re.DOTALL)
    s = re.sub(r"(?:\\quad\s*){6,}", " ", s)
    s = re.sub(r"(?:\\qquad\s*){4,}", " ", s)
    if not re.search(r"\\begin\{(?:align\*?|aligned|array|tabular)\}", s):
        s = re.sub(r"\s*&\s*", " ", s)
    s = re.sub(r"(?:[A-Za-z]\s){10,}[A-Za-z]", " ", s)
    s = re.sub(r"(?:\\\\\s*){5,}", r"\\\\", s)
    s = re.sub(r"\s{2,}", " ", s).strip()
    return "" if _looks_like_ocr_formula_junk(s) else s


def _page_bbox(obj: Dict[str, Any]) -> Tuple[Optional[int], Optional[Dict[str, Any]]]:
    prov = (obj.get("prov") or [])
    if not prov:
        return None, None
    first = prov[0] or {}
    return first.get("page_no"), first.get("bbox")


def _load_json(path: Path) -> Dict[str, Any]:
    with path.open("r", encoding="utf-8", errors="replace") as f:
        obj = json.load(f)
    if not isinstance(obj, dict):
        raise ValueError(f"JSON root is not an object: {path}")
    return obj


# Docling object resolution ---------------------------------------------------
def _parse_ref(ref: str) -> Tuple[str, int]:
    return parse_docling_ref(ref)


def _resolve_ref(doc: Dict[str, Any], ref: str) -> Dict[str, Any]:
    return resolve_docling_ref(doc, ref)


def _docling_ref_text_parts_with_script(
    doc: Dict[str, Any],
    ref: str,
) -> List[Tuple[str, str]]:
    """Return ordered (text, script) fragments under a Docling ref."""
    try:
        collection, _ = _parse_ref(ref)
        obj = _resolve_ref(doc, ref)
    except Exception:
        return []

    if collection == "texts":
        text = _clean_text(str(obj.get("text") or obj.get("orig") or ""))
        script = str(((obj.get("formatting") or {}).get("script") or "baseline")).lower()
        return [(text, script)] if text else []

    parts: List[Tuple[str, str]] = []
    for child_ref in obj.get("children") or []:
        child = str((child_ref or {}).get("$ref") or "")
        if child:
            parts.extend(_docling_ref_text_parts_with_script(doc, child))
    return parts


def _join_docling_inline_text_parts(parts: List[Tuple[str, str]]) -> str:
    """Join DOCX/Docling inline-run fragments without crossing block lines."""
    out = ""
    for raw, script in parts:
        text = _clean_text(raw)
        if not text:
            continue
        script = (script or "baseline").lower()
        if not out:
            out = text
            continue

        no_space = False
        if script in {"sub", "super"}:
            no_space = True
        elif re.match(r"^[.,;:%)\]]", text):
            no_space = True
        elif out.endswith(("(", "[", "/", "-", "−", "–", "—")):
            no_space = True
        elif re.fullmatch(r"-?\d+", text) and re.search(r"[A-Za-zµμ]$", out):
            no_space = True

        out = f"{out}{text}" if no_space else f"{out} {text}"

    out = re.sub(r"\s+([.,;:%)\]])", r"\1", out)
    out = re.sub(r"([([])\s+", r"\1", out)
    out = re.sub(r"\s{2,}", " ", out)
    return _clean_text(out)


def _docling_inline_group_text(doc: Dict[str, Any], group_ref: str) -> str:
    return _join_docling_inline_text_parts(_docling_ref_text_parts_with_script(doc, group_ref))


def _docling_text_refs_under_ref(doc: Dict[str, Any], ref: str) -> List[str]:
    """Return descendant text refs under ref, preserving Docling child order."""
    try:
        collection, _ = _parse_ref(ref)
        obj = _resolve_ref(doc, ref)
    except Exception:
        return []
    if collection == "texts":
        return [ref]
    refs: List[str] = []
    for child_ref in obj.get("children") or []:
        child = str((child_ref or {}).get("$ref") or "")
        if child:
            refs.extend(_docling_text_refs_under_ref(doc, child))
    return refs


def _docling_inline_text_ref_to_group_ref(doc: Dict[str, Any]) -> Dict[str, str]:
    """Map each text run in an inline group to the group's logical text block."""
    mapping: Dict[str, str] = {}
    for group in doc.get("groups") or []:
        if not isinstance(group, dict):
            continue
        if str(group.get("label") or "") != "inline":
            continue
        group_ref = str(group.get("self_ref") or "")
        if not group_ref:
            continue
        for text_ref in _docling_text_refs_under_ref(doc, group_ref):
            mapping.setdefault(text_ref, group_ref)
    return mapping


def _page_bbox_for_logical_text_block(
    doc: Dict[str, Any],
    group_ref: str,
) -> Tuple[Optional[int], Optional[Dict[str, Any]]]:
    """Use the first child text layout as the logical inline block layout."""
    for text_ref in _docling_text_refs_under_ref(doc, group_ref):
        try:
            page_no, bbox = _page_bbox(_resolve_ref(doc, text_ref))
        except Exception:
            continue
        if page_no is not None or bbox is not None:
            return page_no, bbox
    try:
        return _page_bbox(_resolve_ref(doc, group_ref))
    except Exception:
        return None, None

def _logical_text_from_docling_ref(doc: Dict[str, Any], ref: str) -> str:
    try:
        collection, _ = _parse_ref(ref)
        obj = _resolve_ref(doc, ref)
    except Exception:
        return ""
    if collection == "groups" and str(obj.get("label") or "") == "inline":
        return _docling_inline_group_text(doc, ref)
    return _clean_text(str(obj.get("text") or obj.get("orig") or ""))


def _flatten_body(doc: Dict[str, Any]) -> List[StreamObject]:
    """Flatten Docling body refs into a semantic stream.

    DOCX conversions often split one printed paragraph/caption into many
    formatted text runs inside an ``inline`` group, for example ``K`` +
    subscript ``d``. Caption recovery should operate on that logical paragraph,
    not on the individual formatting runs, so inline groups are emitted as one
    text StreamObject and their child text refs are suppressed.
    """
    out: List[StreamObject] = []
    order = 0
    inline_group_by_text_ref = _docling_inline_text_ref_to_group_ref(doc)
    emitted_inline_groups: set[str] = set()

    def append_stream_object(
        *,
        ref: str,
        object_type: str,
        label: str,
        text: str,
        content_layer: str,
        page_no: Optional[int],
        bbox: Optional[Dict[str, Any]],
        raw: Dict[str, Any],
    ) -> None:
        nonlocal order
        out.append(
            StreamObject(
                order_index=order,
                object_ref=ref,
                object_type=object_type,
                label=label,
                text=text,
                content_layer=content_layer,
                page_no=page_no,
                bbox=bbox,
                raw=raw,
            )
        )
        order += 1

    def emit_inline_group(group_ref: str) -> None:
        if group_ref in emitted_inline_groups:
            return
        try:
            group = _resolve_ref(doc, group_ref)
        except Exception:
            return
        text = _docling_inline_group_text(doc, group_ref)
        emitted_inline_groups.add(group_ref)
        if not text:
            return
        page_no, bbox = _page_bbox_for_logical_text_block(doc, group_ref)
        append_stream_object(
            ref=group_ref,
            object_type="text",
            label="inline",
            text=text,
            content_layer=str(group.get("content_layer") or ""),
            page_no=page_no,
            bbox=bbox,
            raw=group,
        )

    for ref in iter_docling_body_refs(doc):
        try:
            collection, _ = _parse_ref(ref)
            obj = _resolve_ref(doc, ref)
        except Exception:
            continue

        label = str(obj.get("label") or "")
        if collection == "groups" and label == "inline":
            emit_inline_group(ref)
            continue

        if collection == "texts" and ref in inline_group_by_text_ref:
            emit_inline_group(inline_group_by_text_ref[ref])
            continue

        object_type = {"texts": "text", "tables": "table", "pictures": "picture"}.get(collection, collection)
        if object_type == "table" and label == "document_index":
            continue
        page_no, bbox = _page_bbox(obj)
        text = _clean_text(str(obj.get("text") or obj.get("orig") or "")) if object_type == "text" else ""
        append_stream_object(
            ref=ref,
            object_type=object_type,
            label=label,
            text=text,
            content_layer=str(obj.get("content_layer") or ""),
            page_no=page_no,
            bbox=bbox,
            raw=obj,
        )
    return out

def _text_from_ref(doc: Dict[str, Any], ref_obj: Dict[str, Any]) -> str:
    ref = str((ref_obj or {}).get("$ref") or "")
    return _logical_text_from_docling_ref(doc, ref)

_PARENT_CAPTION_RE = re.compile(
    r"^\s*(?:Table|Tab\.?|Figure|Fig\.?|Scheme|Exhibit)\b",
    flags=re.I,
)

def _linked_caption_text(doc: Dict[str, Any], obj: Dict[str, Any]) -> str:
    parts = [_text_from_ref(doc, r) for r in (obj.get("captions") or [])]
    parts = [p for p in parts if p]
    return " ".join(parts).strip() if parts else ""


def _linked_caption_refs(doc: Dict[str, Any], obj: Dict[str, Any]) -> List[str]:
    refs: List[str] = []
    for ref_obj in obj.get("captions") or []:
        ref = str((ref_obj or {}).get("$ref") or "")
        if ref and _text_from_ref(doc, ref_obj):
            refs.append(ref)
    return list(dict.fromkeys(refs))


def _direct_child_refs(doc: Dict[str, Any], parent_ref: str) -> List[str]:
    try:
        parent = _resolve_ref(doc, parent_ref)
    except Exception:
        return []
    refs: List[str] = []
    for child in parent.get("children") or []:
        ref = str((child or {}).get("$ref") or "")
        if ref:
            refs.append(ref)
    return refs


def _stream_object_matches_caption_kind(
    doc: Dict[str, Any],
    so: StreamObject,
    caption_object_type: str,
) -> bool:
    if caption_object_type == "table":
        return so.object_type == "table" and not _is_empty_layout_table(doc, so.raw)
    return so.object_type == "picture" and _picture_is_candidate_asset(so.raw)


def _parent_caption_locally_plausible(
    doc: Dict[str, Any],
    stream: List[StreamObject],
    stream_index: int,
    parent_ref: str,
    caption_object_type: str,
) -> bool:
    if stream_index < 0 or stream_index >= len(stream):
        return False

    current_ref = stream[stream_index].object_ref
    child_refs = _direct_child_refs(doc, parent_ref)
    if current_ref not in child_refs:
        return False

    current_pos = child_refs.index(current_ref)
    index_by_ref = {so.object_ref: idx for idx, so in enumerate(stream)}

    for child_ref in child_refs[:current_pos]:
        child_idx = index_by_ref.get(child_ref)
        if child_idx is not None:
            child_so = stream[child_idx]
            if _stream_object_matches_caption_kind(doc, child_so, caption_object_type):
                return False

        child_text = _logical_text_from_docling_ref(doc, child_ref)
        if child_text and _is_valid_caption_text(child_text, caption_object_type):
            return False

    return True


def _parent_caption_text_and_ref(
    doc: Dict[str, Any],
    obj: Dict[str, Any],
    *,
    object_type: str = "",
    stream: Optional[List[StreamObject]] = None,
    stream_index: Optional[int] = None,
) -> Tuple[str, str]:
    parent_ref = str(((obj.get("parent") or {}).get("$ref") or ""))
    if not parent_ref:
        return "", ""
    parent_text = _logical_text_from_docling_ref(doc, parent_ref)
    caption_object_type = "table" if object_type == "table" else "figure" if object_type else ""
    if caption_object_type:
        if not _is_valid_caption_text(parent_text, caption_object_type):
            return "", ""
    elif not _PARENT_CAPTION_RE.match(parent_text):
        return "", ""

    if stream is not None and stream_index is not None and caption_object_type:
        if not _parent_caption_locally_plausible(
            doc,
            stream,
            stream_index,
            parent_ref,
            caption_object_type,
        ):
            return "", ""

    return parent_text, parent_ref

def _caption_text(doc: Dict[str, Any], obj: Dict[str, Any]) -> str:
    linked_caption = _linked_caption_text(doc, obj)
    if linked_caption:
        return linked_caption
    return ""


def _caption_text_refs(doc: Dict[str, Any], obj: Dict[str, Any]) -> List[str]:
    return _linked_caption_refs(doc, obj)


def _linked_text(doc: Dict[str, Any], obj: Dict[str, Any], key: str) -> Tuple[str, List[str]]:
    parts: List[str] = []
    refs: List[str] = []
    for ref_obj in obj.get(key) or []:
        ref = str((ref_obj or {}).get("$ref") or "")
        text = _text_from_ref(doc, ref_obj)
        if text:
            parts.append(text)
            if ref:
                refs.append(ref)
    return "\n".join(parts).strip(), refs


def _neighbor_identifier_caption(
    stream: List[StreamObject],
    stream_index: int,
    *,
    object_type: str,
    before: bool = True,
) -> Tuple[str, str]:
    """Return an immediately adjacent identifier-only caption and its text ref.

    Docling can split a printed caption so that the identifier (for example
    "Table 1") is a standalone text object adjacent to the table while the
    title/description is linked through the table's ``captions`` field. Treat
    only the immediately adjacent same-page text object as eligible here; wider
    caption discovery remains the job of regex/LLM fallback.
    """
    j = stream_index - 1 if before else stream_index + 1
    if j < 0 or j >= len(stream):
        return "", ""

    current = stream[stream_index]
    neighbor = stream[j]
    if (
        current.page_no is not None
        and neighbor.page_no is not None
        and current.page_no != neighbor.page_no
    ):
        return "", ""
    if not _usable_local_text_object(neighbor):
        return "", ""

    text = _clean_text(neighbor.text)
    if not _caption_text_is_identifier_only(text, object_type):
        return "", ""
    return text, neighbor.object_ref


def _caption_text_with_adjacent_identifier(
    doc: Dict[str, Any],
    stream: List[StreamObject],
    stream_index: int,
) -> str:
    so = stream[stream_index]
    object_type = "table" if so.object_type == "table" else "figure"

    linked_caption = _linked_caption_text(doc, so.raw)
    if linked_caption:
        if _object_id_from_caption(linked_caption, fallback=""):
            return linked_caption

        # Tables normally have captions above; figures normally have captions below.
        directions = (True, False) if so.object_type == "table" else (False, True)
        for before in directions:
            identifier, _ = _neighbor_identifier_caption(
                stream,
                stream_index,
                object_type=object_type,
                before=before,
            )
            if identifier:
                return _clean_text(f"{identifier} {linked_caption}")
        return linked_caption
    if so.object_type == "table":
        adjacent_candidate = _adjacent_table_caption_candidate(doc, stream, stream_index)
        adjacent_caption = _clean_text(str((adjacent_candidate or {}).get("text") or ""))
        if adjacent_caption and _is_valid_caption_text(adjacent_caption, "table"):
            return adjacent_caption
    parent_caption, _ = _parent_caption_text_and_ref(
        doc,
        so.raw,
        object_type=object_type,
        stream=stream,
        stream_index=stream_index,
    )
    return parent_caption

def _caption_text_refs_with_adjacent_identifier(
    doc: Dict[str, Any],
    stream: List[StreamObject],
    stream_index: int,
) -> List[str]:
    so = stream[stream_index]
    object_type = "table" if so.object_type == "table" else "figure"

    linked_refs = _linked_caption_refs(doc, so.raw)
    linked_caption = _linked_caption_text(doc, so.raw)
    if linked_caption:
        refs = list(linked_refs)
        if not _object_id_from_caption(linked_caption, fallback=""):
            directions = (True, False) if so.object_type == "table" else (False, True)
            for before in directions:
                _, ref = _neighbor_identifier_caption(
                    stream,
                    stream_index,
                    object_type=object_type,
                    before=before,
                )
                if ref:
                    refs.append(ref)
                    break
        return list(dict.fromkeys(refs))

    if so.object_type == "table":
        adjacent_candidate = _adjacent_table_caption_candidate(doc, stream, stream_index)
        adjacent_caption = _clean_text(str((adjacent_candidate or {}).get("text") or ""))
        if adjacent_caption and _is_valid_caption_text(adjacent_caption, "table"):
            return list(
                dict.fromkeys(
                    str(ref)
                    for ref in (adjacent_candidate or {}).get("refs") or []
                    if ref
                )
            )
    _, parent_ref = _parent_caption_text_and_ref(
        doc,
        so.raw,
        object_type=object_type,
        stream=stream,
        stream_index=stream_index,
    )
    return [parent_ref] if parent_ref else []

def _local_text_join(*parts: str, max_chars: int, keep_tail: bool = False) -> str:
    """Join local evidence snippets, de-duplicating exact text blocks."""
    out: List[str] = []
    seen: set[str] = set()
    for part in parts:
        text = _clean_text(part)
        if not text:
            continue
        key = text.casefold()
        if key in seen:
            continue
        seen.add(key)
        out.append(text)
    joined = "\n".join(out).strip()
    if len(joined) <= max_chars:
        return joined
    if keep_tail:
        return "...[truncated]\n" + joined[-max_chars:].lstrip()
    return joined[:max_chars].rstrip() + "\n...[truncated]"

def _canonical_object_id(kind_raw: str, num_raw: str) -> str:
    kind_key = kind_raw.replace(".", "").lower()
    if kind_key == "fig":
        kind = "Figure"
    elif kind_key in {"tab", "tabel"}:
        kind = "Table"
    else:
        kind = kind_raw.replace(".", "").capitalize()

    num = re.sub(r"[\s._-]+", "", str(num_raw or "").upper())
    for dash in ("‐", "-", "‒", "–", "—", "―"):
        num = num.replace(dash, "")
    if re.fullmatch(r"\d+S", num):
        num = "S" + num[:-1]
    if re.fullmatch(r"(?:SI|SM|SF|ST)\d+", num):
        num = "S" + re.sub(r"^(?:SI|SM|SF|ST)", "", num)
    return f"{kind} {num}"

def _object_id_from_caption(caption: str, fallback: str) -> str:
    m = re.match(
        r"^\s*(Table|Tab\.?|Tabel|Exhibit|Figure|Fig\.?|Scheme)\s*\.?\s*"
        r"((?:S|SI|SM|SF|ST|A)?\s*[.\-]?\s*\d+(?:\s*[A-Za-z]|[._\-]\d+)?|\d+\s*S|[IVXLCDM]+)\b",
        caption or "",
        flags=re.I,
    )
    if not m:
        return fallback
    return _canonical_object_id(m.group(1), m.group(2))

def _valid_caption_object_id(caption: str, object_type: str) -> str:
    """Return a canonical id only when caption has the expected numbered anchor."""
    object_id = _object_id_from_caption(caption, fallback="")
    if not object_id:
        return ""
    label = object_id.split()[0].lower()
    if object_type == "table":
        return object_id if label in {"table", "exhibit"} else ""
    return object_id if label in {"figure", "scheme", "exhibit"} else ""


def _is_valid_caption_text(caption: str, object_type: str) -> bool:
    return bool(_valid_caption_object_id(caption, object_type))


def _is_real_figure_object_id(object_id: str) -> bool:
    return bool(re.match(r"^(?:Figure|Scheme|Exhibit)\s+\S+", _clean_text(object_id or ""), flags=re.I))


def _logical_table_group_key(table_id: str) -> str:
    return re.sub(r"_(\d+)$", "", table_id or "").strip()

def _uncaptioned_table_id(table_counter: int) -> str:
    """Return a fallback id that cannot collide with printed Table N ids.

    Captionless tables are real table objects, but they should not consume the
    article's numbered table namespace. For example, an uncaptioned
    nomenclature/list table before printed Table 1 should become
    "Uncaptioned Table 1", leaving "Table 1" for the printed captioned table.
    """
    try:
        n = int(table_counter)
    except Exception:
        n = 0
    return f"Uncaptioned Table {max(1, n)}"

# Caption / table-note candidate heuristics ----------------------------------
_CAPTION_CONTEXT_PREFIX_WORDS = ("Supplementary", "Supplemental", "Supporting", "Appendix")
_CAPTION_CONTEXT_PREFIX_PAT = rf"(?:(?:{'|'.join(_CAPTION_CONTEXT_PREFIX_WORDS)})\s+)?"
_CAPTION_SUPP_NUM_PREFIX_PAT = r"(?:SM|SI|SF|ST|S|A)"
_CAPTION_ROMAN_NUM_PAT = r"[IVXLCDM]+\b"
_CAPTION_NUM_PAT = (
    rf"(?:{_CAPTION_SUPP_NUM_PREFIX_PAT}\s*[.\-]?\s*\d+|"
    rf"\d+S\b|\d+(?:[._\-]\d+)?|{_CAPTION_ROMAN_NUM_PAT})"
)
_CAPTION_LABEL_NUM_SEP_PAT = (
    rf"(?:\s+|\s*[-\u2013\u2014]\s*|"
    rf"(?={_CAPTION_SUPP_NUM_PREFIX_PAT}\s*[.\-]?\s*\d+|\d+))"
)
_CAPTION_LINE_RE = re.compile(
    rf"""^\s*
        (?:\#{{1,6}}\s*)?
        (?:\d{{1,6}}\s+)?
        {_CAPTION_CONTEXT_PREFIX_PAT}
        (?P<kind>Table|Tab\.?|Tabel|Exhibit|Figure|Fig\.?|Scheme)
        \.?
        {_CAPTION_LABEL_NUM_SEP_PAT}
        (?P<num>{_CAPTION_NUM_PAT})
        \s*\.?\s*
        (?P<delim>[:|]|[-\u2013\u2014])?
        \s*
        (?P<desc>.*\S.*|)
        \s*$
    """,
    re.IGNORECASE | re.VERBOSE,
)
_CAPTION_PROSE_START_RE = re.compile(
    r"^(?:shows?|lists?|summari[sz]es?|presents?|reports?|contains?|provides?|"
    r"illustrates?|depicts?|compares?|demonstrates?|describes?|represents?)\b",
    re.IGNORECASE,
)
_TABLE_FOOTNOTE_MARKER_RE = re.compile(
    r"^(?:Note\s*:|Notes\s*:|Abbreviations?\s*:|"
    r"\(?[a-z]\)|[a-z]\.|[*\u2020\u2021\u00a7])",
    re.IGNORECASE,
)
_TABLE_ABBREVIATION_NOTE_RE = re.compile(
    r"^[A-Za-z0-9][A-Za-z0-9()/,\s-]{0,120}\s+-\s+\S"
)


def _caption_candidate_from_text(text: str, object_type: str) -> Optional[Dict[str, str]]:
    """Return a normalized caption candidate from one local text line/block."""
    s = _clean_text(text)
    if not s or len(s) > 2_000:
        return None
    m = _CAPTION_LINE_RE.match(s)
    if not m:
        return None

    kind_raw = m.group("kind") or ""
    kind_key = kind_raw.replace(".", "").lower()
    allowed = {"exhibit"}
    if object_type == "table":
        allowed.update({"table", "tab", "tabel"})
    else:
        allowed.update({"figure", "fig", "scheme"})
    if kind_key not in allowed:
        return None

    # Avoid body prose such as "Table 1 shows ..." being accepted as a caption
    # when no printed-caption delimiter is present. Captions like
    # "Table 1 Experimental conditions" are still allowed.
    desc = (m.group("desc") or "").strip()
    has_printed_delim = bool((m.group("delim") or "").strip())
    if not has_printed_delim:
        try:
            has_printed_delim = s[m.end("num") :].lstrip().startswith(".")
        except Exception:
            has_printed_delim = False
    if not has_printed_delim and _CAPTION_PROSE_START_RE.match(desc):
        return None

    return {
        "object_id": _canonical_object_id(kind_raw, m.group("num") or ""),
        "kind": kind_key,
        "text": s,
        "desc": desc,
    }


def _caption_candidate_is_identifier_only(candidate: Dict[str, Any]) -> bool:
    if _clean_text(str(candidate.get("desc") or "")):
        return False
    text_id = _object_id_from_caption(str(candidate.get("text") or ""), fallback="")
    return bool(text_id and text_id == _clean_text(str(candidate.get("object_id") or "")))


def _caption_text_is_identifier_only(text: str, object_type: str) -> bool:
    candidate = _caption_candidate_from_text(text, object_type)
    return bool(candidate and _caption_candidate_is_identifier_only(candidate))

def _is_caption_continuation_line(line: str, object_type: str) -> bool:
    s = _clean_text(line)
    if not s or len(s) > 700:
        return False
    if _caption_candidate_from_text(s, object_type):
        return False
    if _is_table_footnote_candidate_line(s):
        return False
    if re.match(r"^(?:SI\s+\d|References?|Acknowledg(?:e)?ments?|Methods?|Results?)\b", s, re.I):
        return False
    # Single letter footnote markers after a table are not caption continuations.
    if re.fullmatch(r"[a-z]", s, re.I):
        return False
    # A continuation may start with punctuation in DOCX conversions, e.g.
    # "Fig. S1" + ". Amount of PFAS ...".
    return bool(re.search(r"[A-Za-z]", s))


def _merge_caption_identifier_with_continuation(
    candidate: Dict[str, Any],
    lines: List[str],
    index: int,
    object_type: str,
) -> Tuple[Dict[str, Any], int]:
    if not _caption_candidate_is_identifier_only(candidate):
        return candidate, index

    merged_parts = [_clean_text(lines[index])]
    j = index + 1
    # Treat a newline as a paragraph boundary. Identifier-only captions may
    # take exactly one immediately following title/description line, e.g.
    # "Fig. S2" + ". SEM Micrographs ..." or "Table S3" + title. If a
    # caption has multiple sentences, they are expected to be in that same
    # continuation text object/paragraph rather than in later lines.
    if j < len(lines):
        nxt = _clean_text(lines[j])
        if _is_caption_continuation_line(nxt, object_type):
            merged_parts.append(nxt)
            j += 1

    if len(merged_parts) == 1:
        return candidate, index

    merged = dict(candidate)
    text = " ".join(merged_parts)
    text = re.sub(r"\s+([.,;:])", r"\1", text)
    text = re.sub(r"^(Figure|Fig\.|Scheme)\s+(S?\d+)\s+\.", r"\1 \2.", text, flags=re.I)
    merged["text"] = _clean_text(text)
    merged["desc"] = _clean_text(" ".join(merged_parts[1:]))
    merged["continued_caption"] = True
    return merged, j - 1


def _find_caption_candidates(text: str, object_type: str) -> List[Dict[str, Any]]:
    candidates: List[Dict[str, Any]] = []
    seen: set[Tuple[str, str]] = set()
    lines = (text or "").splitlines()
    i = 0
    while i < len(lines):
        candidate = _caption_candidate_from_text(lines[i], object_type)
        if not candidate:
            i += 1
            continue
        candidate, i = _merge_caption_identifier_with_continuation(candidate, lines, i, object_type)
        key = (str(candidate["object_id"]), str(candidate["text"]))
        if key not in seen:
            seen.add(key)
            candidates.append(candidate)
        i += 1
    return candidates

def _normalize_merged_caption_text(parts: List[str]) -> str:
    text = " ".join(_clean_text(part) for part in parts if _clean_text(part))
    text = re.sub(r"\s+([.,;:%)\]])", r"\1", text)
    text = re.sub(r"([([])\s+", r"\1", text)
    text = re.sub(r"(?<=[A-Za-z]/[A-Za-z])\s+(?=\d+\b)", "", text)
    text = re.sub(
        r"^(Figure|Fig\.|Scheme)\s+(S?\d+)\s+\.",
        r"\1 \2.",
        text,
        flags=re.I,
    )
    return _clean_text(text)


def _merge_adjacent_caption_block_candidate(
    candidate: Dict[str, Any],
    lines: List[str],
    index: int,
    object_type: str,
) -> Dict[str, Any]:
    """Merge only identifier-only captions with one adjacent caption-like line.
    Deterministic tier-2 recovery should not stitch arbitrary neighboring text
    objects into a caption. Full DOCX inline-run captions are handled separately
    by _previous_inline_group_caption_candidate(), which uses Docling group
    structure rather than line proximity.
    """
    if not _caption_candidate_is_identifier_only(candidate):
        return candidate
    j = index + 1
    while j < len(lines) and not _clean_text(lines[j]):
        j += 1
    if j >= len(lines):
        return candidate
    nxt = _clean_text(lines[j])
    if (
        _caption_candidate_from_text(nxt, object_type)
        or _is_table_footnote_candidate_line(nxt)
        or not _is_caption_continuation_line(nxt, object_type)
    ):
        return candidate

    text = _normalize_merged_caption_text([lines[index], nxt])
    reparsed = _caption_candidate_from_text(text, object_type)
    merged = dict(reparsed or candidate)
    merged["text"] = text
    merged["desc"] = _clean_text(nxt)
    merged["continued_caption"] = True
    return merged


def _find_adjacent_caption_block_candidates(
    lines: List[str],
    object_type: str,
) -> List[Dict[str, Any]]:
    candidates: List[Dict[str, Any]] = []
    seen: set[Tuple[str, str]] = set()
    i = 0
    while i < len(lines):
        candidate = _caption_candidate_from_text(lines[i], object_type)
        if not candidate:
            i += 1
            continue
        candidate = _merge_adjacent_caption_block_candidate(candidate, lines, i, object_type)
        key = (str(candidate.get("object_id") or ""), str(candidate.get("text") or ""))
        if key not in seen:
            seen.add(key)
            candidates.append(candidate)
        i += 1
    return candidates

def _adjacent_table_caption_candidate(
    doc: Dict[str, Any],
    stream: List[StreamObject],
    stream_index: int,
) -> Optional[Dict[str, Any]]:
    """Return a conservative caption candidate from immediately above a table.
    The input stream has already normalized Docling inline groups into logical
    text blocks, so this function should not stitch formatting-run fragments.
    Deterministic recovery accepts simple cases only: a complete caption block,
    or identifier-only + one immediately adjacent caption-like block.
    """
    if stream_index < 0 or stream_index >= len(stream):
        return None
    if stream[stream_index].object_type != "table":
        return None

    previous_items: List[Tuple[str, str]] = []
    char_total = 0
    for j in range(stream_index - 1, -1, -1):
        so = stream[j]
        if so.object_type == "text":
            if not _usable_local_text_object(so):
                continue
            txt = _clean_text(so.text)
            if not txt:
                continue
            previous_items.append((txt, so.object_ref))
            char_total += len(txt) + 1
            if _caption_candidate_from_text(txt, "table"):
                break
            if char_total >= STRUCTURED_CAPTION_CANDIDATE_WINDOW_CHARS:
                break
            continue
        if so.object_type == "picture" and char_total < STRUCTURED_CAPTION_CANDIDATE_WINDOW_CHARS:
            continue

        if previous_items:
            break
        return None

    if not previous_items:
        return None
    line_items = [(text, ref) for text, ref in reversed(previous_items) if _clean_text(text)]
    lines = [_clean_text(text) for text, _ in line_items]
    candidates = _find_adjacent_caption_block_candidates(lines, "table")
    non_identifier = [c for c in candidates if not _caption_candidate_is_identifier_only(c)]
    if len(non_identifier) != 1:
        return None
    candidate = dict(non_identifier[0])
    candidate["source"] = "adjacent_before"

    caption_norm = re.sub(r"\s+", " ", _clean_text(str(candidate.get("text") or ""))).casefold()
    refs: List[str] = []
    for line, ref in line_items:
        line_norm = re.sub(r"\s+", " ", _clean_text(line)).casefold()
        if line_norm and (line_norm in caption_norm or caption_norm in line_norm):
            refs.append(ref)
    candidate["refs"] = refs or [ref for _, ref in line_items]

    return candidate


def _looks_like_orphan_table_cell_text(text: str) -> bool:
    """Heuristic for text Docling leaked out of a table body as a text block."""
    s = _clean_text(text)
    if not s or len(s) > 160:
        return False
    if re.search(r"[.;:]\s*$", s):
        return False
    if re.search(r"\b(?:mg|g|kg|L|ml|mL|ppm|ppb|PFAS|PFOA|PFOS|GAC|DOC|pH)\b", s):
        return True
    return bool(re.search(r"\d", s) and not re.search(r"\b(?:the|this|these|those|therefore|however)\b", s, re.I))


def _is_table_footnote_candidate_line(text: str) -> bool:
    s = _clean_text(text)
    if not s:
        return False
    return bool(
        _TABLE_FOOTNOTE_MARKER_RE.match(s)
        or (_TABLE_ABBREVIATION_NOTE_RE.match(s) and ("," in s or ";" in s))
    )


def _leading_table_footnote_candidates(text: str) -> List[str]:
    """Footnote candidates safe enough for deterministic attachment.

    The line may be separated from the table by short text objects that are
    likely leaked table cells, but not by normal prose.
    """
    candidates: List[str] = []
    started = False
    for line in (text or "").splitlines():
        s = _clean_text(line)
        if not s:
            continue
        if _is_table_footnote_candidate_line(s):
            started = True
            candidates.append(s)
            continue
        if started:
            break
        if _looks_like_orphan_table_cell_text(s):
            continue
        break
    return candidates


def _norm_consumed_text(text: str) -> str:
    return re.sub(r"\s+", " ", _clean_text(text)).strip().lower()


def _remember_consumed_text(consumed_text_norms: set[str], text: str) -> None:
    norm = _norm_consumed_text(text)
    if norm:
        consumed_text_norms.add(norm)
    for line in (text or "").splitlines():
        line_norm = _norm_consumed_text(line)
        if line_norm:
            consumed_text_norms.add(line_norm)


def _is_consumed_text_piece(text: str, consumed_text_norms: set[str]) -> bool:
    norm = _norm_consumed_text(text)
    if not norm:
        return False
    if norm in consumed_text_norms:
        return True
    if len(norm) >= 20 and any(norm in known for known in consumed_text_norms):
        return True
    return False


def _is_caption_like_text(text: str) -> bool:
    return bool(
        _caption_candidate_from_text(text, "table")
        or _caption_candidate_from_text(text, "figure")
    )


# Table markdown and quality checks ------------------------------------------
def _escape_md_cell(value: Any) -> str:
    return str(value or "").replace("|", r"\|").strip()


def _split_md_row(line: str) -> List[str]:
    s = (line or "").strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    cells: List[str] = []
    buf: List[str] = []
    escaped = False
    for ch in s:
        if escaped:
            buf.append(ch)
            escaped = False
            continue
        if ch == "\\":
            buf.append(ch)
            escaped = True
            continue
        if ch == "|":
            cells.append("".join(buf).strip())
            buf = []
        else:
            buf.append(ch)
    cells.append("".join(buf).strip())
    return cells


def _cells_to_md_row(cells: List[str]) -> str:
    return "| " + " | ".join(cells or []) + " |"



def _md_separator_cell_for_width(cell: str, width: int) -> str:
    raw = (cell or "").replace(" ", "")
    left_align = raw.startswith(":")
    right_align = raw.endswith(":")
    target_width = max(width, 3 + int(left_align) + int(right_align))
    dash_count = max(3, target_width - int(left_align) - int(right_align))
    return (
        (":" if left_align else "")
        + ("-" * dash_count)
        + (":" if right_align else "")
    )


def _align_markdown_pipe_table_lines(lines: List[str]) -> List[str]:
    """Pad Markdown pipe-table cells so raw .md files are human-readable."""
    if len(lines) < 2:
        return lines

    rows = [_split_md_row(line) for line in lines]
    width = max((len(row) for row in rows), default=0)
    if width <= 1:
        return lines

    rows = [row + [""] * (width - len(row)) for row in rows]
    col_widths = [
        max(3, max(len(row[col_idx]) for row in rows))
        for col_idx in range(width)
    ]

    out: List[str] = []
    for row in rows:
        is_separator = bool(
            len(row) > 1
            and all(
                _MD_TABLE_SEPARATOR_CELL_RE.fullmatch(cell.replace(" ", ""))
                for cell in row
            )
        )
        rendered_cells: List[str] = []
        for col_idx, cell in enumerate(row):
            rendered = (
                _md_separator_cell_for_width(cell, col_widths[col_idx])
                if is_separator
                else cell
            )
            rendered_cells.append(rendered.ljust(col_widths[col_idx]))
        out.append("| " + " | ".join(rendered_cells) + " |")
    return out


def _pretty_print_markdown_pipe_tables(block: str) -> str:
    """Align each Markdown pipe table inside a block without changing values."""
    lines = (block or "").rstrip().splitlines()
    if not lines:
        return ""

    out: List[str] = []
    i = 0
    while i < len(lines):
        if (
            lines[i].lstrip().startswith("|")
            and i + 1 < len(lines)
            and _is_md_separator_row(lines[i + 1])
        ):
            j = i + 2
            while j < len(lines) and lines[j].lstrip().startswith("|"):
                j += 1
            out.extend(_align_markdown_pipe_table_lines(lines[i:j]))
            i = j
            continue

        out.append(lines[i])
        i += 1

    return "\n".join(out).rstrip()

def _pad_or_trim_cells(cells: List[str], n_cols: int) -> List[str]:
    cells = list(cells or [])
    if len(cells) < n_cols:
        return cells + [""] * (n_cols - len(cells))
    if len(cells) > n_cols:
        return cells[:n_cols]
    return cells


NUMERICISH_TABLE_CELL_PAT = re.compile(
    r"^[\d\s,.;:+\-−–—×xX*/^().%Ee\[\]]+$"
)


def _is_blank_table_cell(cell: str) -> bool:
    return not (cell or "").strip()


def _numericish_table_cell_for_fill(cell: str) -> bool:
    s = re.sub(r"<[^>]+>", "", cell or "").strip()
    s = re.sub(r"\s+", " ", s)
    return bool(s and NUMERICISH_TABLE_CELL_PAT.fullmatch(s))


def _fillable_leading_row_header_columns(
    rows: List[List[str]],
    max_cols: int,
) -> set[int]:
    """Find leading text-like columns where blank cells likely mean merged headers."""
    if not rows:
        return set()

    n_cols = max((len(row) for row in rows), default=0)
    if n_cols <= 0:
        return set()

    padded = [_pad_or_trim_cells(row, n_cols) for row in rows]
    fillable: set[int] = set()

    for col in range(min(max_cols, n_cols)):
        values: List[str] = []
        blanks_after_label = 0
        seen_label = False

        for row in padded:
            has_later_content = any(
                not _is_blank_table_cell(cell) for cell in row[col + 1 :]
            )
            if not has_later_content:
                continue

            if not _is_blank_table_cell(row[col]):
                seen_label = True
                values.append(row[col].strip())
            elif seen_label:
                blanks_after_label += 1

        if not values or blanks_after_label <= 0:
            continue

        non_numeric = sum(
            not _numericish_table_cell_for_fill(value) for value in values
        )
        if non_numeric * 2 >= len(values):
            fillable.add(col)

    return fillable


def _fill_leading_row_header_cells(
    data_rows: List[List[str]],
    max_cols: int = TABLE_ROW_HEADER_FILL_MAX_COLS,
) -> List[List[str]]:
    """Fill sparse merged-cell row headers before splitting a large table."""
    if not data_rows:
        return data_rows

    n_cols = max((len(row) for row in data_rows), default=0)
    if n_cols <= 0:
        return data_rows

    rows = [_pad_or_trim_cells(list(row), n_cols) for row in data_rows]
    max_cols = min(max_cols, n_cols)
    fill_cols = _fillable_leading_row_header_columns(rows, max_cols)
    if not fill_cols:
        return rows

    last = [""] * max_cols
    for row in rows:
        for col in range(max_cols):
            if col not in fill_cols:
                continue

            if not _is_blank_table_cell(row[col]):
                new_value = row[col].strip()
                if last[col] and new_value != last[col]:
                    for child_col in range(col + 1, max_cols):
                        last[child_col] = ""
                last[col] = new_value
                continue

            has_later_content = any(
                not _is_blank_table_cell(cell) for cell in row[col + 1 :]
            )
            parent_available = col == 0 or not _is_blank_table_cell(row[col - 1])
            if last[col] and has_later_content and parent_available:
                row[col] = last[col]

    return rows


def _strip_cell_emphasis(text: str) -> str:
    s = _clean_text(text)
    s = re.sub(r"^(?:\*{1,3}|_{1,3})\s*", "", s)
    s = re.sub(r"\s*(?:\*{1,3}|_{1,3})$", "", s)
    return s.strip()


def _is_md_separator_row(line: str) -> bool:
    cells = _split_md_row(line)
    return bool(
        len(cells) > 1
        and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)
    )


def _embedded_table_caption_from_md_row(line: str) -> Optional[str]:
    cells = [_strip_cell_emphasis(cell) for cell in _split_md_row(line)]
    cells = [cell for cell in cells if cell]
    if not cells:
        return None
    first_candidate: Optional[Dict[str, str]] = None
    for width in range(1, min(4, len(cells)) + 1):
        joined = _clean_text(" ".join(cells[:width]))
        candidate = _caption_candidate_from_text(joined, "table")
        if candidate and _is_valid_caption_text(str(candidate.get("text") or ""), "table"):
            first_candidate = candidate
            break
    if not first_candidate:
        return None
    if not _caption_candidate_is_identifier_only(first_candidate):
        return str(first_candidate.get("text") or "").strip()
    descriptions = [
        cell
        for cell in cells[1:]
        if not _caption_candidate_from_text(cell, "table")
        and not _is_table_footnote_candidate_line(cell)
    ]
    if descriptions:
        return _clean_text(f"{cells[0]} {max(descriptions, key=len)}")
    return str(first_candidate.get("text") or "").strip()


def _embedded_table_caption_marker_from_md_row(line: str) -> bool:
    """Detect a split in-table caption without reconstructing its text.

    Docling can distribute an embedded caption across header cells, for
    example ``Table | Class # | ... | S8.``. That row is strong evidence of a
    caption boundary, but it is not safe to turn the cells into caption text.
    Camelot is responsible for recovering the caption itself.
    """
    cells = [_strip_cell_emphasis(cell) for cell in _split_md_row(line)]
    cells = [cell for cell in cells if cell]
    for idx, cell in enumerate(cells[:-1]):
        if not re.match(r"^\s*(?:Table|Tab\.?|Tabel|Exhibit)\b", cell, re.IGNORECASE):
            continue
        for following in cells[idx + 1:]:
            # Ignore intervening header labels. They are table structure, not
            # caption prose; only use the separated identifier as the marker.
            candidate = _caption_candidate_from_text(f"Table {following}", "table")
            if candidate and _is_valid_caption_text(str(candidate.get("text") or ""), "table"):
                return True
    return False


def _table_has_embedded_caption_marker(
    doc: Dict[str, Any],
    table_obj: Dict[str, Any],
) -> bool:
    """Return whether the table's early rows contain a split caption marker."""
    try:
        table_body = _table_body_to_markdown(doc, table_obj)
    except Exception:
        return False
    for line in table_body.splitlines()[:3]:
        if line.lstrip().startswith("|") and _embedded_table_caption_marker_from_md_row(line):
            return True
    return False


def _camelot_embedded_caption_recovery(
    *,
    pdf_path: Optional[Path],
    table_obj: Dict[str, Any],
) -> Tuple[str, Dict[str, Any]]:
    """Return a validated Camelot caption candidate and its diagnostic payload."""
    if not pdf_path or extract_camelot_embedded_caption is None:
        return "", {}
    try:
        payload = extract_camelot_embedded_caption(
            pdf_path=pdf_path,
            table_obj=table_obj,
            bbox_pad=TABLE_REPAIR_BBOX_PAD,
            min_nonempty_cells=TABLE_REPAIR_MIN_NONEMPTY_CELLS,
        )
    except Exception as exc:
        return "", {"status": "failed", "error": str(exc)}
    candidate_caption = _clean_text(str(payload.get("caption_candidate") or ""))
    if _valid_caption_object_id(candidate_caption, "table"):
        return candidate_caption, payload
    return "", payload


def _promote_embedded_table_caption(table_body: str) -> Tuple[str, str]:
    lines = (table_body or "").splitlines()
    if len(lines) < 3:
        return "", table_body
    for idx, line in enumerate(lines[:3]):
        if not line.lstrip().startswith("|"):
            continue
        if idx + 2 >= len(lines):
            return "", table_body
        caption = _embedded_table_caption_from_md_row(line)
        if (
            caption
            and _is_md_separator_row(lines[idx + 1])
            and lines[idx + 2].lstrip().startswith("|")
        ):
            new_lines = [
                *lines[:idx],
                lines[idx + 2],
                lines[idx + 1],
                *lines[idx + 3:],
            ]
            return caption, "\n".join(new_lines).strip()
    return "", table_body


def _docling_table_grid(table_obj: Dict[str, Any]) -> List[List[str]]:
    data = table_obj.get("data") or {}
    try:
        n_rows = int(data.get("num_rows") or 0)
        n_cols = int(data.get("num_cols") or 0)
    except Exception:
        n_rows = n_cols = 0
    if n_rows <= 0 or n_cols <= 0:
        return []
    grid = [["" for _ in range(n_cols)] for _ in range(n_rows)]
    for cell in data.get("table_cells") or []:
        try:
            r0 = int(cell.get("start_row_offset_idx"))
            c0 = int(cell.get("start_col_offset_idx"))
        except Exception:
            continue
        if 0 <= r0 < n_rows and 0 <= c0 < n_cols:
            txt = _clean_text(str(cell.get("text") or ""))
            if txt and not grid[r0][c0]:
                grid[r0][c0] = txt
            elif txt:
                grid[r0][c0] = (grid[r0][c0] + " " + txt).strip()
    return grid


def _docling_table_grid_preserve_breaks(table_obj: Dict[str, Any]) -> List[List[str]]:
    """Return the Docling grid while preserving intra-cell line breaks.

    DOCX tables often encode collapsed records as line-separated values inside
    one Docling cell, e.g. "PFOA\nPFOS" paired with "414\n538". The normal
    markdown renderer collapses those into spaces, but expansion needs the
    original positional separators.
    """
    data = table_obj.get("data") or {}
    try:
        n_rows = int(data.get("num_rows") or 0)
        n_cols = int(data.get("num_cols") or 0)
    except Exception:
        n_rows = n_cols = 0
    if n_rows <= 0 or n_cols <= 0:
        return []
    grid = [["" for _ in range(n_cols)] for _ in range(n_rows)]
    for cell in data.get("table_cells") or []:
        try:
            r0 = int(cell.get("start_row_offset_idx"))
            c0 = int(cell.get("start_col_offset_idx"))
        except Exception:
            continue
        if 0 <= r0 < n_rows and 0 <= c0 < n_cols:
            raw = str(cell.get("text") or "")
            parts = [_clean_text(part) for part in re.split(r"\r?\n+", raw)]
            txt = "\n".join(part for part in parts if part)
            if txt and not grid[r0][c0]:
                grid[r0][c0] = txt
            elif txt:
                grid[r0][c0] = (grid[r0][c0] + "\n" + txt).strip()
    return grid


def _docling_table_has_multiline_body_cells(table_obj: Dict[str, Any]) -> bool:
    for cell in ((table_obj.get("data") or {}).get("table_cells") or []):
        if cell.get("column_header"):
            continue
        raw = str(cell.get("text") or "")
        parts = [_clean_text(part) for part in re.split(r"\r?\n+", raw)]
        if len([part for part in parts if part]) > 1:
            return True
    return False


def _table_body_to_markdown(doc: Dict[str, Any], table_obj: Dict[str, Any]) -> str:
    if _docling_table_has_multiline_body_cells(table_obj):
        expanded_body, _expansion_payload = _expanded_position_aligned_table_body(table_obj)
        if expanded_body:
            return expanded_body

    try:
        rendered = render_structured_table_markdown(table_obj, doc)
        if rendered.strip():
            return _clean_rendered_markdown_table(rendered)
    except Exception:
        pass

    grid = _docling_table_grid(table_obj)
    if not grid:
        return ""
    header = grid[0]
    lines = [
        "| " + " | ".join(_escape_md_cell(x) for x in header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in grid[1:]:
        lines.append("| " + " | ".join(_escape_md_cell(x) for x in row) + " |")
    return "\n".join(lines).strip()


def _iter_descendant_docling_refs(
    doc: Dict[str, Any],
    obj: Dict[str, Any],
    wanted_collections: set[str],
) -> Iterable[str]:
    """Yield nested Docling refs under an object's children."""
    seen: set[str] = set()

    def visit(ref_obj: Any) -> Iterable[str]:
        ref = str((ref_obj or {}).get("$ref") if isinstance(ref_obj, dict) else ref_obj or "")
        if not ref or ref in seen:
            return
        seen.add(ref)
        try:
            collection, _ = _parse_ref(ref)
            child = _resolve_ref(doc, ref)
        except Exception:
            return

        if collection in wanted_collections:
            yield ref
        for child_ref in child.get("children") or []:
            yield from visit(child_ref)

    for child_ref in obj.get("children") or []:
        yield from visit(child_ref)


def _table_nonempty_cell_texts(table_obj: Dict[str, Any]) -> List[str]:
    texts: List[str] = []
    data = table_obj.get("data") or {}
    for cell in data.get("table_cells") or []:
        text = _clean_text(str(cell.get("text") or ""))
        if text:
            texts.append(text)
    for row in data.get("grid") or []:
        for cell in row or []:
            if not isinstance(cell, dict):
                continue
            text = _clean_text(str(cell.get("text") or ""))
            if text:
                texts.append(text)
    return list(dict.fromkeys(texts))


def _table_descendant_texts(doc: Dict[str, Any], table_obj: Dict[str, Any]) -> List[str]:
    texts: List[str] = []
    for ref in _iter_descendant_docling_refs(doc, table_obj, {"texts"}):
        text = _text_from_ref(doc, {"$ref": ref})
        if text:
            texts.append(text)
    return list(dict.fromkeys(texts))


def _table_has_candidate_picture_child(doc: Dict[str, Any], table_obj: Dict[str, Any]) -> bool:
    for ref in _iter_descendant_docling_refs(doc, table_obj, {"pictures"}):
        try:
            if _picture_is_candidate_asset(_resolve_ref(doc, ref)):
                return True
        except Exception:
            continue
    return False


def _is_empty_layout_table(doc: Dict[str, Any], table_obj: Dict[str, Any]) -> bool:
    """Return True for DOCX layout/spacing tables with no semantic payload.

    Empty layout tables should not enter caption recovery, table grouping,
    table_counter, table_segment_indices, or markdown output. Tables with
    picture children are preserved so figure-container tables can still be
    processed by _figure_like_table_info().
    """
    if _table_nonempty_cell_texts(table_obj):
        return False
    if _table_descendant_texts(doc, table_obj):
        return False
    if _table_has_candidate_picture_child(doc, table_obj):
        return False
    return True

def _figure_like_table_info(doc: Dict[str, Any], table_obj: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    """Detect DOCX rich-cell tables that are actually figure containers."""
    picture_refs: List[str] = []
    seen_picture_refs: set[str] = set()
    for ref in _iter_descendant_docling_refs(doc, table_obj, {"pictures"}):
        if ref in seen_picture_refs:
            continue
        try:
            pic = _resolve_ref(doc, ref)
        except Exception:
            continue
        if not _picture_is_candidate_asset(pic):
            continue
        seen_picture_refs.add(ref)
        picture_refs.append(ref)
    if not picture_refs:
        return None

    text_refs = list(_iter_descendant_docling_refs(doc, table_obj, {"texts"}))
    text_ref_texts = [_text_from_ref(doc, {"$ref": ref}) for ref in text_refs]
    text_ref_texts = [text for text in text_ref_texts if text]
    caption_source = "\n".join(_table_nonempty_cell_texts(table_obj) or text_ref_texts)
    candidates = [
        candidate
        for candidate in _find_caption_candidates(caption_source, "figure")
        if not _caption_candidate_is_identifier_only(candidate)
    ]
    if len(candidates) != 1:
        return None

    caption = _clean_text(str(candidates[0].get("text") or ""))
    object_id = _valid_caption_object_id(caption, "figure")
    if not object_id:
        return None
    return {
        "object_id": object_id,
        "caption": caption,
        "picture_refs": picture_refs,
        "text_refs": text_refs,
    }


def _compose_table_block(table_body: str, caption: str = "", footnote: str = "") -> str:
    lines: List[str] = []
    if caption.strip():
        lines.extend([caption.strip(), ""])
    if table_body.strip():
        lines.append(table_body.strip())
    if footnote.strip():
        lines.extend(["", footnote.strip()])
    return "\n".join(lines).rstrip() + "\n"


def _markdown_from_text_grid(rows: List[List[str]]) -> str:
    if not rows:
        return ""
    width = max((len(row) for row in rows), default=0)
    if width <= 0:
        return ""
    normalized = [list(row) + [""] * (width - len(row)) for row in rows]
    lines = [
        "| " + " | ".join(_escape_md_cell(x) for x in normalized[0][:width]) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    for row in normalized[1:]:
        lines.append("| " + " | ".join(_escape_md_cell(x) for x in row[:width]) + " |")
    return "\n".join(lines).strip()


def _excel_cell_to_text(value: Any) -> str:
    """Convert an Excel cell value to stable, markdown-safe text."""
    if value is None:
        return ""
    if isinstance(value, (datetime,)):
        return value.isoformat(sep=" ")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _clean_text(str(value))


def _excel_sheet_to_markdown_rows(sheet: Any) -> Tuple[str, List[List[str]]]:
    """Read one openpyxl worksheet, preserving merged header values."""
    max_row = int(sheet.max_row or 0)
    max_col = int(sheet.max_column or 0)
    if max_row <= 0 or max_col <= 0:
        return "", []

    grid = [
        [_excel_cell_to_text(sheet.cell(row=row, column=col).value) for col in range(1, max_col + 1)]
        for row in range(1, max_row + 1)
    ]
    for merged_range in sheet.merged_cells.ranges:
        top_left = _excel_cell_to_text(
            sheet.cell(row=merged_range.min_row, column=merged_range.min_col).value
        )
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                grid[row - 1][col - 1] = top_left

    def _row_has_text(row: List[str]) -> bool:
        return any(str(cell or "").strip() for cell in row)

    while grid and not _row_has_text(grid[-1]):
        grid.pop()
    while grid and not _row_has_text(grid[0]):
        grid.pop(0)
    if not grid:
        return "", []

    width = max(
        (idx + 1 for row in grid for idx, cell in enumerate(row) if str(cell or "").strip()),
        default=0,
    )
    grid = [row[:width] for row in grid]
    caption = grid[0][0].strip() if grid and grid[0] else ""
    if _is_valid_caption_text(caption, "table"):
        grid = grid[1:]
        while grid and not _row_has_text(grid[0]):
            grid.pop(0)
    return caption, grid


def _split_excel_table_by_proxy(
    *,
    rows: List[List[str]],
    caption: str,
    header_row_count: int = EXCEL_SUPPLEMENTARY_HEADER_ROWS,
) -> List[str]:
    """Render an Excel grid into markdown parts within enumeration's proxy budget."""
    if not rows:
        return []
    width = max((len(row) for row in rows), default=0)
    if width <= 0:
        return []
    normalized = [list(row) + [""] * (width - len(row)) for row in rows]
    header_count = min(max(1, int(header_row_count)), max(1, len(normalized) - 1))
    header_rows = normalized[:header_count]
    data_rows = normalized[header_count:]
    separator = _cells_to_md_row(["---"] * width)
    repeated_header = [_cells_to_md_row(row) for row in header_rows] + [separator]
    header_proxy = max(0, len(repeated_header) - 2) * width

    parts: List[str] = []
    current: List[str] = []
    current_proxy = header_proxy
    for row in data_rows:
        padded_row = _cells_to_md_row(row)
        trial_body = "\n".join([*repeated_header, *current, padded_row])
        trial_block = _compose_table_block(trial_body, caption=caption)
        if current and (
            current_proxy + width > LARGE_TABLE_SPLIT_TARGET_TABLE_CELL_PROXY
            or len(trial_block) > LARGE_TABLE_SPLIT_TARGET_CHARS
        ):
            parts.append(_compose_table_block("\n".join([*repeated_header, *current]), caption=caption))
            current = [padded_row]
            current_proxy = header_proxy + width
        else:
            current.append(padded_row)
            current_proxy += width
    if current:
        parts.append(_compose_table_block("\n".join([*repeated_header, *current]), caption=caption))

    return parts or [_compose_table_block("\n".join(repeated_header), caption=caption)]


def _process_excel_supplementary_tables(
    *,
    study: str,
    excel_path: Path,
    chunk_records: List[Dict[str, Any]],
    object_records: List[Dict[str, Any]],
    figtab_rows: List[Dict[str, Any]],
) -> int:
    """Convert the curated supplementary workbook into canonical table markdown.

    Only the explicitly curated ``supplementary_table.xlsx`` is read.  In
    particular, this does not fall back to ``supplementary_table_original.xlsx``.
    """
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to preprocess supplementary Excel tables") from exc

    workbook = openpyxl.load_workbook(excel_path, read_only=False, data_only=True)
    tables_dir = OUTPUT_DIR / study / "tables"
    tables_dir.mkdir(parents=True, exist_ok=True)
    table_count = 0
    # Keep the downstream source label consistent with Docling-extracted
    # supplementary tables; ``source`` and ``metadata_source`` identify the
    # Excel origin.
    source_file = "supplementary_material"
    metadata_source = excel_path.name

    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        caption, rows = _excel_sheet_to_markdown_rows(sheet)
        split_blocks = _split_excel_table_by_proxy(rows=rows, caption=caption)
        if not split_blocks:
            continue
        fallback_id = f"Uncaptioned Table {sheet_index}"
        table_id = _object_id_from_caption(caption, fallback=fallback_id)
        object_ref = f"#/excel_tables/{_sanitize_filename(sheet.title)}"
        order_index = 10_000 + sheet_index
        logical_table_id = table_id
        for part_index, block in enumerate(split_blocks, start=1):
            output_table_id = (
                table_id
                if len(split_blocks) == 1
                else f"{table_id}_{LARGE_TABLE_SPLIT_PART_SUFFIX}{part_index:02d}"
            )
            table_file = tables_dir / (
                f"{_sanitize_filename(source_file + '_' + _table_id_filename_fragment(output_table_id))}.md"
            )
            output_block = _pretty_print_markdown_pipe_tables(_remove_html_breaks(block))
            table_file.write_text(output_block.rstrip() + "\n", encoding="utf-8")
            record = {
                "study_folder": study,
                "source_file": source_file,
                "file_type": "table",
                "object_type": "table",
                "object_ref": object_ref,
                "table_id": output_table_id,
                "logical_table_id": logical_table_id,
                "physical_table_id": table_id,
                "physical_table_segment": 1,
                "split_part": part_index if len(split_blocks) > 1 else None,
                "caption_text": caption,
                "footnote_text": "",
                "text": output_block.strip(),
                "markdown": output_block.strip(),
                "page_no": None,
                "bbox": None,
                "source": "excel",
                "metadata_source": metadata_source,
                "caption_valid": _is_valid_caption_text(caption, "table"),
                "caption_repair_status": "not_attempted",
                "repair_reasons": [],
                "initial_position_status": "not_applicable",
                "post_repair_position_status": "not_applicable",
                "initial_table_quality_status": "not_assessed",
                "initial_table_recommended_action": "use",
                "initial_table_quality_decision": {},
                "final_table_quality_status": "not_assessed",
                "final_table_recommended_action": "use",
                "final_table_quality_decision": {},
                "needs_review": False,
                "artifact_path": str(table_file),
                "order_start": order_index,
                "order_end": order_index,
            }
            chunk_records.append(record)
            figtab_rows.append(
                {
                    "study_folder": study,
                    "source_file": source_file,
                    "file_type": source_file,
                    "type": "table",
                    "object_ref": object_ref,
                    "object_id": output_table_id,
                    "caption": caption,
                    "footnote": "",
                    "page_no": "",
                    "bbox": "{}",
                    "standardized_file_name": table_file.stem,
                    "file_path": str(table_file),
                    "source": "excel",
                    "metadata_source": metadata_source,
                    "needs_review": "false",
                }
            )
        object_records.append(
            {
                "study_folder": study,
                "source_file": source_file,
                "order_index": order_index,
                "object_ref": object_ref,
                "object_type": "table",
                "label": "table",
                "text": caption,
                "content_layer": "body",
                "page_no": None,
                "bbox": None,
            }
        )
        table_count += len(split_blocks)

    workbook.close()
    return table_count


def _md_cell_preserve_breaks(value: Any) -> str:
    parts = [
        _clean_text(part)
        for part in re.split(r"\r?\n+", str(value or ""))
        if _clean_text(part)
    ]
    if not parts:
        return ""
    return _escape_md_cell(" ".join(parts))


def _markdown_from_text_grid_preserve_breaks(rows: List[List[str]]) -> str:
    if not rows:
        return ""
    width = max((len(row) for row in rows), default=0)
    if width <= 0:
        return ""
    normalized = [list(row) + [""] * (width - len(row)) for row in rows]
    lines = [
        "| " + " | ".join(_md_cell_preserve_breaks(x) for x in normalized[0][:width]) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    for row in normalized[1:]:
        lines.append("| " + " | ".join(_md_cell_preserve_breaks(x) for x in row[:width]) + " |")
    return "\n".join(lines).strip()


def _docling_grid_text_rows(table_obj: Dict[str, Any]) -> List[List[str]]:
    data = table_obj.get("data") or {}
    raw_grid = data.get("grid") or []
    rows: List[List[str]] = []
    if raw_grid:
        width = max((len(row) for row in raw_grid), default=0)
        for raw_row in raw_grid:
            row: List[str] = []
            for col_idx in range(width):
                cell = raw_row[col_idx] if col_idx < len(raw_row) else {}
                if isinstance(cell, dict):
                    row.append(_clean_text(str(cell.get("text") or "")))
                else:
                    row.append(_clean_text(str(cell or "")))
            rows.append(row)
        return rows
    return _docling_table_grid(table_obj)


def _cell_units_for_position_expansion(cell: str) -> List[str]:
    """Split one collapsed cell into position-aligned units.

    Prefer Docling/DOCX line breaks. When no line breaks remain, fall back to a
    conservative whitespace split for simple token lists. The fallback is only
    used inside rows already classified as position_aligned.
    """
    text = str(cell or "").strip()
    if not text:
        return []
    line_parts = [_clean_text(part) for part in re.split(r"\r?\n+", text)]
    line_parts = [part for part in line_parts if part]
    if len(line_parts) > 1:
        return line_parts

    cleaned = _clean_text(re.sub(r"\s+", " ", text))
    if not cleaned:
        return []
    # Simple fallback: split plain lists such as "F100 F200 F816" or
    # "653 646 598". Avoid splitting formula-like strings unless line breaks
    # were preserved, because formulas can contain meaningful adjacent tokens.
    if re.search(r"[()]", cleaned):
        return [cleaned]
    parts = cleaned.split()
    return parts if len(parts) > 1 else [cleaned]


def _md_cell_from_expanded_value(value: str) -> str:
    return _escape_md_cell(re.sub(r"\s+", " ", _clean_text(value)).strip())


def _expanded_position_aligned_table_body(table_obj: Dict[str, Any]) -> Tuple[str, Dict[str, Any]]:
    """Expand a position_aligned compacted table into one record per row.

    This is intentionally deterministic. It preserves one or more original
    header rows and zips line-separated cell values by position.
    """
    grid = _docling_table_grid_preserve_breaks(table_obj)
    if not grid or len(grid) < 2:
        return "", {"status": "skipped", "reason": "empty_or_too_short_grid"}

    header_count = max(1, min(_table_header_row_count(table_obj), len(grid) - 1))
    header_rows = [
        [_md_cell_from_expanded_value(cell.replace("\n", " ")) for cell in row]
        for row in grid[:header_count]
    ]
    expanded_rows: List[List[str]] = []
    rows_expanded = 0

    for row in grid[header_count:]:
        units_by_cell = [_cell_units_for_position_expansion(cell) for cell in row]
        nonempty_lengths = [len(units) for units in units_by_cell if units]
        if not nonempty_lengths:
            continue
        candidate_counts: Dict[int, int] = {}
        for length in nonempty_lengths:
            if length > 1:
                candidate_counts[length] = candidate_counts.get(length, 0) + 1
        target_len = 0
        for length, count in candidate_counts.items():
            if count >= 2 and (not target_len or count > candidate_counts[target_len] or length > target_len):
                target_len = length

        if target_len <= 1:
            expanded_rows.append([_md_cell_from_expanded_value(cell.replace("\n", " ")) for cell in row])
            continue

        shared_prefix_idx: Optional[int] = None
        plus_one_indices = [
            idx for idx, units in enumerate(units_by_cell) if len(units) == target_len + 1
        ]
        if len(plus_one_indices) == 1:
            first_nonempty_idx = next(
                (idx for idx, units in enumerate(units_by_cell) if units),
                None,
            )
            if plus_one_indices[0] == first_nonempty_idx:
                shared_prefix_idx = plus_one_indices[0]

        allowed_lengths = {0, 1, target_len}
        if shared_prefix_idx is not None:
            allowed_lengths.add(target_len + 1)

        # If a nonempty cell has a conflicting intermediate count, do not guess.
        if any(len(units) not in allowed_lengths for units in units_by_cell):
            expanded_rows.append([_md_cell_from_expanded_value(cell.replace("\n", " ")) for cell in row])
            continue

        for i in range(target_len):
            out_row: List[str] = []
            for idx, units in enumerate(units_by_cell):
                if not units:
                    out_row.append("")
                elif len(units) == target_len:
                    out_row.append(_md_cell_from_expanded_value(units[i]))
                elif idx == shared_prefix_idx and len(units) == target_len + 1:
                    out_row.append(_md_cell_from_expanded_value(f"{units[0]} {units[i + 1]}"))
                else:
                    # Shared row-level value. Keep it explicit in every expanded row.
                    out_row.append(_md_cell_from_expanded_value(units[0]))
            expanded_rows.append(out_row)
        rows_expanded += target_len - 1

    if rows_expanded < POSITION_ALIGNED_EXPANSION_MIN_ROWS_ADDED:
        return "", {"status": "skipped", "reason": "no_rows_expanded"}

    width = max(
        max((len(row) for row in header_rows), default=0),
        max((len(row) for row in expanded_rows), default=0),
    )
    normalized_headers = [row + [""] * (width - len(row)) for row in header_rows]
    lines = ["| " + " | ".join(normalized_headers[0][:width]) + " |"]
    lines.append("| " + " | ".join("---" for _ in range(width)) + " |")
    for row in normalized_headers[1:]:
        lines.append("| " + " | ".join(row[:width]) + " |")
    for row in expanded_rows:
        padded = [*row, *([""] * max(0, width - len(row)))]
        lines.append("| " + " | ".join(padded[:width]) + " |")

    return "\n".join(lines).strip(), {
        "status": "expanded",
        "header_rows": header_count,
        "original_data_rows": max(0, len(grid) - header_count),
        "expanded_data_rows": len(expanded_rows),
        "rows_added": rows_expanded,
    }


def _table_header_row_count(table_obj: Dict[str, Any]) -> int:
    header_end = 1
    for cell in ((table_obj.get("data") or {}).get("table_cells") or []):
        if cell.get("column_header"):
            try:
                header_end = max(header_end, int(cell.get("end_row_offset_idx") or 1))
            except Exception:
                pass
    return header_end


def _table_data_cell_proxy(table_obj: Dict[str, Any]) -> int:
    return sum(
        1
        for cell in ((table_obj.get("data") or {}).get("table_cells") or [])
        if not cell.get("column_header") and _clean_text(str(cell.get("text") or ""))
    )


def _markdown_table_cell_proxy_from_lines(pipe_lines: List[str]) -> int:
    # The formatter pads all rows to the widest row before the record reaches
    # enumeration.  Use that final width for the early large-table decision as
    # well, rather than undercounting a narrow first header row.
    columns = max((len(_split_md_row(line)) for line in pipe_lines), default=0)
    return max(0, len(pipe_lines) - 2) * columns


def _enumeration_table_proxy_warnings(
    study: str,
    chunk_records: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    """Print actionable warnings for table chunks enumeration would reject."""
    warnings: List[Dict[str, Any]] = []
    table_record_count = 0
    for record in chunk_records:
        if str(record.get("file_type") or "").strip().lower() != "table":
            continue
        table_record_count += 1
        pipe_lines = _markdown_table_pipe_lines(str(record.get("text") or ""))
        columns = max((len(_split_md_row(line)) for line in pipe_lines), default=0)
        cell_proxy = max(0, len(pipe_lines) - 2) * columns
        if cell_proxy <= LARGE_TABLE_SPLIT_TARGET_TABLE_CELL_PROXY:
            continue
        artifact_path = str(record.get("artifact_path") or "")
        warnings.append(
            {
                "table_name": Path(artifact_path).name or str(record.get("chunk_id") or "unknown"),
                "artifact_path": artifact_path,
                "chunk_id": str(record.get("chunk_id") or ""),
                "pipe_rows": len(pipe_lines),
                "columns": columns,
                "cell_proxy": cell_proxy,
                "limit": LARGE_TABLE_SPLIT_TARGET_TABLE_CELL_PROXY,
            }
        )

    if warnings:
        print(
            f"[TABLE PROXY WARNING] {study}: {len(warnings)} of {table_record_count} table "
            f"chunk(s) exceed enumeration's {LARGE_TABLE_SPLIT_TARGET_TABLE_CELL_PROXY}-cell limit."
        )
        for warning in warnings:
            print(
                f"[TABLE PROXY WARNING] {study}/{warning['table_name']}: "
                f"{warning['cell_proxy']} cells "
                f"({warning['pipe_rows']} pipe rows × {warning['columns']} columns); "
                f"limit {warning['limit']}."
            )
    return warnings


def _markdown_table_pipe_lines(block: str) -> List[str]:
    return [
        line.rstrip()
        for line in (block or "").splitlines()
        if line.lstrip().startswith("|") and "|" in line.lstrip()[1:]
    ]


def _split_large_table_block(
    *,
    table_obj: Dict[str, Any],
    full_block: str,
    caption: str,
    footnote: str,
    study: str = "",
    source_file: str = "",
    table_id: str = "",
) -> List[str]:
    pipe_lines = _markdown_table_pipe_lines(full_block)
    table_cell_proxy = _markdown_table_cell_proxy_from_lines(pipe_lines)
    if (
        len(full_block) <= LARGE_TABLE_SPLIT_CHAR_THRESHOLD
        and table_cell_proxy <= LARGE_TABLE_SPLIT_MIN_TABLE_CELL_PROXY
    ):
        return [full_block]

    separator_index = next(
        (
            i
            for i, line in enumerate(pipe_lines)
            if len(_split_md_row(line)) > 1
            and all(
                re.fullmatch(r":?-{3,}:?", cell.replace(" ", ""))
                for cell in _split_md_row(line)
            )
        ),
        None,
    )
    if separator_index is None or separator_index == 0:
        return [full_block]

    physical_rows = pipe_lines[:separator_index] + pipe_lines[separator_index + 1 :]
    header_decision = _resolve_table_header_rows_for_block(
        study=study,
        source_file=source_file,
        table_id=table_id,
        table_obj=table_obj,
        caption_text=caption,
        physical_rows=physical_rows,
        force_llm=True,
        suffix="_split",
    )
    header_count = min(max(1, int(header_decision.get("header_count") or 1)), len(physical_rows) - 1)
    header_row_cells = [_split_md_row(row) for row in physical_rows[:header_count]]
    data_row_cells = physical_rows[header_count:]
    if len(data_row_cells) < 2:
        return [full_block]

    data_row_cells = _fill_leading_row_header_cells(
        [_split_md_row(row) for row in data_row_cells]
    )

    # The final Markdown formatter pads every row to the widest row.  Budget
    # splits against that same width, rather than the first header row alone;
    # otherwise a 6-column header expanded by 8- or 10-column data rows can
    # silently exceed enumeration's 200-cell guard after preprocessing.
    column_count = max(
        (len(row) for row in [*header_row_cells, *data_row_cells]),
        default=0,
    )
    if column_count < 2:
        return [full_block]

    def _padded_md_row(cells: List[str]) -> str:
        return _cells_to_md_row([*cells, *([""] * max(0, column_count - len(cells)))])

    header_rows = [_padded_md_row(row) for row in header_row_cells]
    data_rows = [_padded_md_row(row) for row in data_row_cells]
    separator = _cells_to_md_row(["---"] * column_count)
    repeated_header = [header_rows[0], separator, *header_rows[1:]]
    repeated_header_body_proxy = max(0, len(repeated_header) - 2) * column_count
    parts: List[str] = []
    current: List[str] = []
    current_cell_proxy = repeated_header_body_proxy
    for row in data_rows:
        row_cell_proxy = column_count
        trial_body = "\n".join([*repeated_header, *current, row])
        trial_block = _compose_table_block(trial_body, caption=caption, footnote=footnote)
        if current and (
            len(trial_block) > LARGE_TABLE_SPLIT_TARGET_CHARS
            or current_cell_proxy + row_cell_proxy
            > LARGE_TABLE_SPLIT_TARGET_TABLE_CELL_PROXY
        ):
            body = "\n".join([*repeated_header, *current])
            parts.append(_compose_table_block(body, caption=caption, footnote=footnote))
            current = [row]
            current_cell_proxy = repeated_header_body_proxy + row_cell_proxy
        else:
            current.append(row)
            current_cell_proxy += row_cell_proxy
    if current:
        body = "\n".join([*repeated_header, *current])
        parts.append(_compose_table_block(body, caption=caption, footnote=footnote))

    return parts if len(parts) > 1 else [full_block]



def _normalized_md_cells(cells: List[str]) -> List[str]:
    return [_clean_text(str(cell or "")).casefold() for cell in cells]


def _table_header_rows_for_carryover(
    table_obj: Dict[str, Any],
    header_count: Optional[int] = None,
) -> List[List[str]]:
    grid = _docling_table_grid(table_obj)
    if not grid:
        return []
    resolved_header_count = _table_header_row_count(table_obj) if header_count is None else header_count
    header_count = min(max(1, int(resolved_header_count or 1)), len(grid))
    rows = [list(row) for row in grid[:header_count]]
    width = max((len(row) for row in rows), default=0)
    if width < 2:
        return []
    return [row + [""] * (width - len(row)) for row in rows]


def _table_columns_compatible_with_header(table_obj: Dict[str, Any], header_rows: List[List[str]]) -> bool:
    if not header_rows or not header_rows[0]:
        return False
    return _table_column_count(table_obj) == len(header_rows[0])


def _table_has_own_nonempty_header_rows(
    table_obj: Dict[str, Any],
    carried_header_rows: Optional[List[List[str]]] = None,
) -> bool:
    grid = _docling_table_grid(table_obj)
    if not grid:
        return False

    first_row = list(grid[0])
    if not any(_clean_text(cell) for cell in first_row):
        return False

    header_cells = [
        cell
        for cell in ((table_obj.get("data") or {}).get("table_cells") or [])
        if cell.get("column_header") and _clean_text(str(cell.get("text") or ""))
    ]

    if carried_header_rows:
        carried_first = list(carried_header_rows[0] or [])
        if _normalized_md_cells(first_row[: len(carried_first)]) == _normalized_md_cells(carried_first):
            return True
        carried_first_cell = _clean_text(carried_first[0] if carried_first else "")
        first_cell = _clean_text(first_row[0] if first_row else "")
        if (
            carried_first_cell
            and first_cell
            and first_cell.casefold() == carried_first_cell.casefold()
            and (header_cells or _header_context_score(first_row) >= 1)
        ):
            return True
        if _first_markdown_row_looks_like_data_against_header(first_row, carried_header_rows):
            return False

    if header_cells:
        return True

    candidate_rows = [list(row) for row in grid[: min(3, len(grid))]]
    if _header_context_score(first_row) >= 2:
        return True
    return any(
        _header_context_score(row) >= 2 and _header_context_score(first_row) >= 1
        for row in candidate_rows[1:]
    )


def _body_pipe_lines(table_body: str) -> List[str]:
    return [line.rstrip() for line in (table_body or "").splitlines() if line.lstrip().startswith("|")]


def _body_already_has_header(table_body: str, header_rows: List[List[str]]) -> bool:
    pipe_lines = _body_pipe_lines(table_body)
    if not pipe_lines or not header_rows:
        return False
    if len(pipe_lines) < len(header_rows):
        return False
    for idx, header in enumerate(header_rows):
        if _normalized_md_cells(_split_md_row(pipe_lines[idx])) != _normalized_md_cells(header):
            return False
    return True


def _inject_carried_table_header(table_body: str, header_rows: List[List[str]]) -> str:
    if not table_body.strip() or not header_rows:
        return table_body
    if _body_already_has_header(table_body, header_rows):
        return table_body

    pipe_lines = _body_pipe_lines(table_body)
    separator_index = next(
        (
            i
            for i, line in enumerate(pipe_lines)
            if len(_split_md_row(line)) > 1
            and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in _split_md_row(line))
        ),
        None,
    )
    if separator_index is None:
        data_rows = pipe_lines
    else:
        data_rows = pipe_lines[:separator_index] + pipe_lines[separator_index + 1 :]

    header_width = len(header_rows[0])
    normalized_header = [_cells_to_md_row(row + [""] * (header_width - len(row))) for row in header_rows]
    separator = _cells_to_md_row(["---"] * header_width)
    return "\n".join([normalized_header[0], separator, *normalized_header[1:], *data_rows]).strip()

_NUM_PAT = r"(?<![A-Za-z])[-+−–—]?\d+(?:\.\d+)?(?:\s*[–—−-]\s*\d+(?:\.\d+)?)?(?![A-Za-z])"
_NUM_RE = re.compile(_NUM_PAT)
_PLUS_MINUS_PAT = r"(?:\u00b1|\u00c2\u00b1|\u0105|\uff71|\+/-)"
_EXCLUDED_HEADER_RE = re.compile(r"\b(ref(?:erence)?s?|citation|source|doi|author|study|year|formula|cas|transition)\b", re.I)
_EXCLUDED_ROW_LABEL_RE = re.compile(r"\b(?:molecular|chemical)\s+(?:formula|structure)\b|\bcas(?:\s+number)?\b|\bmass\s+transition\b", re.I)
_VALUE_NUM_PAT = r"[-+−–—]?(?:\d+(?:\.\d+)?|\.\d+)(?:[eE][-+]?\d+)?"
_VALUE_TOKEN_PAT = rf"{_VALUE_NUM_PAT}(?:\s*{_PLUS_MINUS_PAT}\s*{_VALUE_NUM_PAT})?"
_VALUE_ONLY_RE = re.compile(
    rf"^\s*{_VALUE_TOKEN_PAT}(?:\s*[,;]\s*{_VALUE_TOKEN_PAT})*\s*%?\s*$"
)
_VALUE_AT_SEGMENT_START_RE = re.compile(
    rf"(?:(?<=^)|(?<=[,;:\[(])\s*)({_VALUE_TOKEN_PAT})(?!\s*/\s*[A-Za-z])"
)
_PACKED_VALUE_NUM_PAT = r"[-+\u2212]?(?:\d+(?:\.\d+)?|\.\d+)(?:[eE][-+]?\d+)?"
_PACKED_VALUE_TOKEN_PAT = (
    rf"{_PACKED_VALUE_NUM_PAT}"
    rf"(?:(?:\s*[\u2013\u2014]\s*|-(?!\s)){_PACKED_VALUE_NUM_PAT})?"
    rf"(?:\s*{_PLUS_MINUS_PAT}\s*{_PACKED_VALUE_NUM_PAT})?"
    r"%?"
)
_PACKED_NUMERIC_VALUES_ONLY_RE = re.compile(
    rf"^\s*[\[(]?\s*{_PACKED_VALUE_TOKEN_PAT}"
    rf"(?:\s+{_PACKED_VALUE_TOKEN_PAT})*"
    rf"\s*[\])]?\s*$"
)
_RANGE_LIKE_HEADER_RE = re.compile(
    r"(?:%|\b(?:range|distribution|size|diameter|radius|length|width|height|"
    r"thickness|particle|pore|mesh|fraction|recovery|removal|yield|"
    r"efficiency|percent(?:age)?)\b)",
    re.I,
)
_PACKED_RANGE_DELIM_PAT = r"(?:[\u2012\u2013\u2014\u2212-]|\bto\b)"
_SINGLE_PACKED_NUMERIC_RANGE_RE = re.compile(
    rf"^\s*[\[(]?\s*{_PACKED_VALUE_NUM_PAT}\s*"
    rf"{_PACKED_RANGE_DELIM_PAT}\s*"
    rf"{_PACKED_VALUE_NUM_PAT}\s*%?\s*[\])]?\s*$",
    re.I,
)


def _norm_cell(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z]+", " ", value or "").strip().lower()


def _duplicate_text_cell(cell: str) -> bool:
    toks = _norm_cell(cell).split()
    if len(toks) < 2 or len(toks) % 2:
        return False
    mid = len(toks) // 2
    return toks[:mid] == toks[mid:] and any(re.search(r"[a-z]", t) for t in toks[:mid])


def _independent_numeric_count(cell: str) -> int:
    text = _clean_text(re.sub(r"<[^>]+>", " ", cell or ""))
    if not re.search(r"\d", text):
        return 0

    # Pure numeric cells may contain several real values separated by comma or
    # semicolon. Count those directly.
    if _VALUE_ONLY_RE.fullmatch(text):
        return len(re.findall(_VALUE_TOKEN_PAT, text))
    return len(list(_VALUE_AT_SEGMENT_START_RE.finditer(text)))

def _looks_like_single_numeric_range(cell: str, header_cell: str = "") -> bool:
    text = _clean_text(re.sub(r"<[^>]+>", " ", cell or ""))
    if not _SINGLE_PACKED_NUMERIC_RANGE_RE.fullmatch(text):
        return False

    # En/em dashes, "to", and spaced ASCII hyphens are explicit range markers.
    if re.search(r"[\u2012\u2013\u2014]|\bto\b|\d\s+-\s+\d", text, re.I):
        return True

    # Ambiguous forms like "0.42 -0.84" are ranges in range-like columns, but
    # could be two adjacent values elsewhere.
    header_text = _clean_text(header_cell or "")
    return bool(_RANGE_LIKE_HEADER_RE.search(header_text))

def _packed_numeric_value_count(cell: str) -> int:
    """Count whitespace-packed numeric values, excluding units and labels."""
    text = _clean_text(re.sub(r"<[^>]+>", " ", cell or ""))
    if not re.search(r"\d", text):
        return 0
    if not _PACKED_NUMERIC_VALUES_ONLY_RE.fullmatch(text):
        return 0
    return len(re.findall(_PACKED_VALUE_TOKEN_PAT, text))


_DENSE_VALUE_UNIT_EXPONENT_RE = re.compile(
    r"\b(?:m|cm|mm|um|nm|kg|g|mg|ug|mol|mmol|umol|l|ml|eq)\s*[-+]?\s*\d+\b",
    re.I,
)
_DENSE_VALUE_PARAMETER_RE = re.compile(
    r"\b(?:r\s*\^?\s*2|1\s*/\s*n|k\s*f|k\s*d)\b",
    re.I,
)
_DENSE_VALUE_PLACEHOLDER_RE = re.compile(r"\b(?:n\s*/?\s*a|n\.?a\.?|nd)\b", re.I)
_DENSE_VALUE_NUM_PAT = r"[-+\u2212]?(?:\d+(?:\.\d+)?|\.\d+)(?:[eE][-+]?\d+)?"
_DENSE_VALUE_CLUSTER_RE = re.compile(
    rf"(?<![A-Za-z0-9/])"
    rf"(?:[<>]=?|[\u2264\u2265~\u2248\u223c])?\s*"
    rf"{_DENSE_VALUE_NUM_PAT}"
    rf"(?:\s*(?:[-\u2013\u2014]|to)\s*{_DENSE_VALUE_NUM_PAT})?"
    rf"(?:\s*{_PLUS_MINUS_PAT}\s*{_DENSE_VALUE_NUM_PAT})?"
    rf"\s*%?"
    rf"(?!\s*/\s*[A-Za-z])"
)
_DENSE_VALUE_UNIT_AFTER_RE = re.compile(
    r"^\s*(?:%|[A-Za-z/]*\s*)?(?:mg|ug|g|kg|ng|mol|mmol|umol|l|ml|eq|mm|cm|nm|um|m|mesh|ppm|ppb|ppt|%)\b",
    re.I,
)

_DENSE_VALUE_BARE_SEPARATOR_RE = re.compile(r"^\s+$")


def _trim_dense_value_span(text: str, start: int, end: int) -> Tuple[int, int]:
    while start < end and text[start].isspace():
        start += 1
    while end > start and text[end - 1].isspace():
        end -= 1
    return start, end


def _longest_bare_numeric_cluster_run(text: str, spans: List[Tuple[int, int]]) -> int:
    ordered = sorted(spans)
    if not ordered:
        return 0
    longest = run = 1
    prev_end = ordered[0][1]
    for start, end in ordered[1:]:
        gap = text[prev_end:start]
        if _DENSE_VALUE_BARE_SEPARATOR_RE.fullmatch(gap or ""):
            run += 1
            longest = max(longest, run)
        else:
            run = 1
        prev_end = end
    return longest

def _dense_value_normalized_text(cell: str) -> str:
    text = _clean_text(re.sub(r"<[^>]+>", " ", cell or ""))
    text = re.sub(r"[*\u2020\u2021\u00a7]+", " ", text)
    text = _DENSE_VALUE_PARAMETER_RE.sub(" ", text)
    text = _DENSE_VALUE_UNIT_EXPONENT_RE.sub(" ", text)
    return text


def _dense_value_cluster_count(cell: str) -> int:
    """Count likely collapsed value clusters, requiring bare value separators."""
    text = _dense_value_normalized_text(cell)
    if not re.search(r"\d|\bn\s*/?\s*a\b|\bnd\b", text, re.I):
        return 0
    placeholder_spans = [
        _trim_dense_value_span(text, *match.span())
        for match in _DENSE_VALUE_PLACEHOLDER_RE.finditer(text)
    ]
    numeric_matches = list(_DENSE_VALUE_CLUSTER_RE.finditer(text))
    numeric_spans: List[Tuple[int, int]] = []
    strong_count = 0
    for match in numeric_matches:
        start, end = _trim_dense_value_span(text, *match.span())
        if any(start < p_end and end > p_start for p_start, p_end in placeholder_spans):
            continue
        token = text[start:end]
        after = text[end : end + 16]
        is_strong = bool(
            re.search(r"\.\d|[%<>\u2264\u2265~\u2248\u223c]|[-\u2013\u2014]\s*\d|\bto\b", token, re.I)
            or re.search(_PLUS_MINUS_PAT, token)
            or _DENSE_VALUE_UNIT_AFTER_RE.search(after)
        )
        numeric_spans.append((start, end))
        if is_strong:
            strong_count += 1

    spans = sorted([*placeholder_spans, *numeric_spans])
    if not spans:
        return 0
    bare_run = _longest_bare_numeric_cluster_run(text, spans)
    if bare_run < 2:
        return 0
    remainder = text
    for start, end in reversed(spans):
        remainder = remainder[:start] + " " + remainder[end:]
    remainder = re.sub(
        r"\b(?:min|max|mean|median|avg|average|approx(?:imately)?|about|by|weight|basis|carbon|surface|area|diameter|porosity|purity|matrix|functional|group|capacity|exchange|number|size|screen|through|mesh|moisture|pore|volume|particle|iodine|effective|range|value|values|and|or|of|the|to|x|l|d|o|i|dvb|gel|polyamine|polyacrylic|styrene|macroreticular|crosslinked|aromatic|polymer)\b",
        " ",
        remainder,
        flags=re.I,
    )
    remainder = re.sub(r"[\s,;:()\[\]{}%/\\+<>=\-\u2013\u2014\u00d7\u00b1\.]+", "", remainder)
    has_unexplained_words = bool(re.search(r"[A-Za-z]{3,}", remainder))

    count = len(spans)
    if count < 4:
        return count if not has_unexplained_words else 0
    if bare_run >= min(4, count) and (strong_count >= 2 or not has_unexplained_words):
        return count
    if bare_run >= min(4, count) and count >= 5 and strong_count >= 1:
        return count
    return 0


def _suspicious_numeric_value_count(cell: str, header_cell: str = "") -> int:
    """Count values only when they look like collapsed adjacent values."""
    if _looks_like_single_numeric_range(cell, header_cell):
        return 1
    packed_count = _packed_numeric_value_count(cell)
    if packed_count:
        return packed_count
    return _dense_value_cluster_count(cell)

def _has_neighboring_value_collapse(row: List[str], header: List[str]) -> bool:
    """Detect multiple whitespace-packed numeric values in one table cell."""
    for col_idx, cell in enumerate(row):
        header_cell = header[col_idx] if col_idx < len(header) else ""
        if _EXCLUDED_HEADER_RE.search(header_cell or ""):
            continue

        packed_count = _packed_numeric_value_count(cell)
        if packed_count >= 2:
            if _looks_like_single_numeric_range(cell, header_cell):
                continue
            return True

    return False


def _has_dense_value_cell_collapse(row: List[str], header: List[str]) -> bool:
    """Detect record-dense value cells where repair is worth attempting."""
    for col_idx, cell in enumerate(row):
        header_cell = header[col_idx] if col_idx < len(header) else ""
        if _EXCLUDED_HEADER_RE.search(header_cell or ""):
            continue
        if _dense_value_cluster_count(cell) >= 4:
            return True
    return False


def _looks_like_row_label_value_glued(cell: str) -> bool:
    """Detect a row label meshed with its first numeric value, e.g. PFBA 5.86±0.50."""
    text = _clean_text(re.sub(r"<[^>]+>", " ", cell or ""))
    if not re.search(r"[A-Za-z]{2,}", text):
        return False
    if _EXCLUDED_ROW_LABEL_RE.search(text):
        return False
    value_pat = rf"{_NUM_PAT}\s*{_PLUS_MINUS_PAT}\s*{_NUM_PAT}"
    if re.search(value_pat, text):
        return True
    return bool(re.search(r"[A-Za-z]{2,}\s+[-+\u2212\u2013\u2014]?\d+\.\d+", text))


_SIMPLE_SPAN_VALUE_RE = re.compile(
    r"(?:[-+\u2212\u2013\u2014]?\d+(?:\.\d+)?%?|n\.?a\.?|N\.?A\.?)$"
)
_ROW_LABEL_SIMPLE_VALUE_RE = re.compile(
    r"^(.+?)\s+([-+\u2212\u2013\u2014]?\d+(?:\.\d+)?%?|n\.?a\.?|N\.?A\.?)$"
)


def _split_compacted_cell_to_width(text: str, width: int) -> Optional[List[str]]:
    if width <= 1:
        return None
    cleaned = _clean_text(re.sub(r"<[^>]+>", " ", text or ""))
    if not cleaned:
        return None
    parts = cleaned.split()
    if len(parts) != width:
        return None
    if all(_SIMPLE_SPAN_VALUE_RE.fullmatch(part) for part in parts):
        return parts
    return None


def _docling_body_cell_span_width(cell: Dict[str, Any]) -> int:
    try:
        start = int(cell.get("start_col_offset_idx"))
        end = int(cell.get("end_col_offset_idx"))
        if end > start:
            return end - start
    except Exception:
        pass
    try:
        return int(cell.get("col_span") or 1)
    except Exception:
        return 1


def _has_docling_compacted_body_span(table_obj: Dict[str, Any]) -> bool:
    for cell in ((table_obj.get("data") or {}).get("table_cells") or []):
        if cell.get("column_header"):
            continue
        span = _docling_body_cell_span_width(cell)
        if span <= 1:
            continue
        text = _clean_text(str(cell.get("text") or ""))
        if _packed_numeric_value_count(text) >= min(3, max(2, span)):
            return True
    return False


def _md_table_rows(block: str) -> List[List[str]]:
    rows = [_split_md_row(line) for line in (block or "").splitlines() if line.lstrip().startswith("|")]
    return [row for row in rows if row]

def _md_meaningful_body_rows(block: str) -> List[List[str]]:
    """Return non-empty Markdown body rows, excluding header/separator rows."""
    rows = _md_table_rows(block)
    if len(rows) <= 2:
        return []
    meaningful: List[List[str]] = []
    for row in rows[2:]:
        if any(_clean_text(re.sub(r"<[^>]+>", " ", cell or "")) for cell in row):
            meaningful.append(row)
    return meaningful


def _docling_table_shape(table_obj: Dict[str, Any]) -> Tuple[int, int, int]:
    data = table_obj.get("data") or {}
    try:
        n_rows = int(data.get("num_rows") or 0)
    except Exception:
        n_rows = 0
    try:
        n_cols = int(data.get("num_cols") or 0)
    except Exception:
        n_cols = 0
    return n_rows, n_cols, len(data.get("table_cells") or [])


def _table_bbox_dimensions(table_obj: Dict[str, Any]) -> Tuple[float, float]:
    bbox = ((table_obj.get("prov") or [{}])[0] or {}).get("bbox") or {}
    try:
        width = float(bbox.get("r")) - float(bbox.get("l"))
        height = float(bbox.get("t")) - float(bbox.get("b"))
    except Exception:
        return 0.0, 0.0
    return max(0.0, width), max(0.0, height)


def _docling_has_adjacent_column_x_order_inversion(table_obj: Dict[str, Any]) -> bool:
    """Detect rows where assigned adjacent columns contradict visual x-order."""
    data = table_obj.get("data") or {}
    try:
        n_cols = int(data.get("num_cols") or 0)
    except Exception:
        n_cols = 0
    if n_cols < 2:
        return False

    width, _ = _table_bbox_dimensions(table_obj)
    min_gap = max(20.0, width * 0.04)
    header_by_row: Dict[int, Dict[int, float]] = {}
    body_by_row: Dict[int, Dict[int, float]] = {}

    for cell in data.get("table_cells") or []:
        if not _clean_text(str(cell.get("text") or "")):
            continue
        if int(cell.get("row_span") or 1) != 1 or int(cell.get("col_span") or 1) != 1:
            continue
        bbox = cell.get("bbox") or {}
        try:
            row_idx = int(cell.get("start_row_offset_idx"))
            col_idx = int(cell.get("start_col_offset_idx"))
            left = float(bbox.get("l"))
        except Exception:
            continue
        if row_idx < 0 or col_idx < 0 or col_idx >= n_cols:
            continue

        target = header_by_row if cell.get("column_header") else body_by_row
        target.setdefault(row_idx, {})[col_idx] = left

    if not body_by_row:
        return False

    header_normal_pairs: set[int] = set()
    for row in header_by_row.values():
        for col_idx in range(n_cols - 1):
            left = row.get(col_idx)
            right = row.get(col_idx + 1)
            if left is not None and right is not None and right > left + min_gap:
                header_normal_pairs.add(col_idx)

    for col_idx in range(n_cols - 1):
        comparable = 0
        inverted = 0
        normal = 0
        for row in body_by_row.values():
            left = row.get(col_idx)
            right = row.get(col_idx + 1)
            if left is None or right is None:
                continue
            if left > right + min_gap:
                comparable += 1
                inverted += 1
            elif right > left + min_gap:
                comparable += 1
                normal += 1

        required_inversions = max(2, (comparable + 3) // 4)
        if (
            comparable >= 3
            and inverted >= required_inversions
            and (col_idx in header_normal_pairs or normal >= 2)
        ):
            return True

    return False


def _docling_structural_table_corruption_reasons(
    doc: Dict[str, Any],
    table_obj: Dict[str, Any],
    block: str,
) -> List[str]:
    """Detect table-grid failures where Docling produced no usable body cells.

    This is intentionally about structural evidence, not specific table IDs:
    a captioned/large table region has a Docling table object, but the emitted
    grid has no meaningful data rows, a single dense collapsed cell, or only a
    single spanning header cell. In those cases, Docling column metadata is not
    trustworthy and PDF-based repair should be attempted directly.
    """
    reasons: List[str] = []

    def add(reason: str) -> None:
        if reason not in reasons:
            reasons.append(reason)

    markdown_rows = _md_table_rows(block)
    body_rows = _md_meaningful_body_rows(block)
    n_rows, n_cols, cell_count = _docling_table_shape(table_obj)
    width, height = _table_bbox_dimensions(table_obj)
    visually_large = width >= 200.0 and height >= 80.0

    data = table_obj.get("data") or {}
    cells = data.get("table_cells") or []
    max_col_span = 1
    for cell in cells:
        max_col_span = max(max_col_span, _docling_body_cell_span_width(cell))

    if _docling_has_adjacent_column_x_order_inversion(table_obj):
        add("adjacent_column_x_order_inversion")

    all_md_cells = [cell for row in markdown_rows for cell in row if (cell or "").strip()]
    max_md_cell_chars = max((len(_clean_text(cell)) for cell in all_md_cells), default=0)
    max_md_cell_numbers = max((_independent_numeric_count(cell) for cell in all_md_cells), default=0)

    descendant_texts = _table_descendant_texts(doc, table_obj)
    descendant_char_count = sum(len(text) for text in descendant_texts)

    grid_rows = _docling_grid_text_rows(table_obj)
    header_count = min(max(1, _table_header_row_count(table_obj)), len(grid_rows)) if grid_rows else 0
    body_grid_rows = grid_rows[header_count:] if header_count else []
    header_nonempty = sum(
        1
        for row in grid_rows[:header_count]
        for cell in row
        if _clean_text(str(cell or ""))
    )
    body_nonempty = sum(
        1
        for row in body_grid_rows
        for cell in row
        if _clean_text(str(cell or ""))
    )
    if (
        header_count >= 1
        and len(grid_rows) > header_count
        and body_grid_rows
        and body_nonempty == 0
        and header_nonempty >= max(2, min(4, n_cols))
    ):
        add("header_only_blank_body_rows")

    if markdown_rows and not body_rows and n_rows <= 1:
        severe_grid_collapse = cell_count <= 2 or n_cols <= 1 or visually_large
        if severe_grid_collapse:
            add("markdown_has_no_meaningful_data_rows")

    if n_rows <= 1 and n_cols <= 1 and cell_count <= 1:
        dense_descendants = len(descendant_texts) >= 20 or descendant_char_count >= 500
        dense_markdown_cell = max_md_cell_chars >= 500 or max_md_cell_numbers >= 10
        if dense_descendants or dense_markdown_cell:
            add("single_cell_dense_text_collapse")

    if n_rows <= 1 and n_cols >= 3 and cell_count <= 2 and max_col_span >= min(3, n_cols):
        add("single_spanning_header_only_table")

    return reasons


def _header_width(header_rows: Optional[List[List[str]]]) -> int:
    return max((len(row) for row in (header_rows or []) if row), default=0)


_ROW_LABEL_CODE_RE = re.compile(r"^[A-Za-z]{1,6}[-_ ]?\d{1,4}[A-Za-z]?$")
_CATEGORICAL_CODE_COLUMN_RE = re.compile(
    r"\b(?:matrix|stream|sample|condition|group|phase|adsorbent|sorbent|"
    r"material|compound|analyte|pfas|type|category|descriptor|treatment|"
    r"water|medium)\b",
    re.I,
)
_VALUE_METRIC_HEADER_RE = re.compile(
    r"\b(?:q\s*e|q\s*t|q\s*m|k\s*d|log\s*k|d\s*r|removal|recovery|"
    r"concentration|conc\.?|capacity|dose|rate|percent|%|value|mean|sd|"
    r"std|error|r\s*2|k\s*\d*)\b",
    re.I,
)


def _looks_like_row_label_code(value: str) -> bool:
    s = _clean_text(value)
    if not s or len(s) > 40:
        return False
    return bool(_ROW_LABEL_CODE_RE.fullmatch(s))


def _row_label_leakage_suspected(row: List[str], header: List[str]) -> bool:
    """Detect leaked row IDs without flagging legitimate coded category columns."""
    if len(row) < 2:
        return False
    if not (_looks_like_row_label_code(row[0]) and _looks_like_row_label_code(row[1])):
        return False
    if _clean_text(row[0]).casefold() == _clean_text(row[1]).casefold():
        return False

    second_header = _clean_text(header[1] if len(header) > 1 else "")
    if second_header and _CATEGORICAL_CODE_COLUMN_RE.search(second_header):
        return False

    # A leaked row label usually lands where a value/blank column should be.
    # If the second header is a named non-metric dimension, treat coded values
    # like SC1/RT1 as legitimate categories rather than corruption.
    if second_header and not (
        _VALUE_METRIC_HEADER_RE.search(second_header)
        or _looks_like_row_label_code(second_header)
    ):
        return False

    return sum(_independent_numeric_count(cell) > 0 for cell in row[2:]) >= 2


def _first_markdown_row_looks_like_data_against_header(
    first_row: List[str],
    expected_header_rows: Optional[List[List[str]]],
) -> bool:
    if not first_row or not expected_header_rows:
        return False
    header = list(expected_header_rows[0] or [])
    if not header:
        return False
    if _normalized_md_cells(first_row[: len(header)]) == _normalized_md_cells(header):
        return False

    first_header = _clean_text(header[0] if header else "")
    first_cell = _clean_text(first_row[0] if first_row else "")
    numeric_cells_after_first = sum(
        _independent_numeric_count(cell) > 0 for cell in first_row[1:]
    )
    if (
        re.search(r"\b(?:cdp|sample|adsorb|sorbent|material|compound|pfas|analyte|id|name)\b", first_header, re.I)
        and first_cell
        and first_cell.casefold() != first_header.casefold()
        and numeric_cells_after_first >= 2
    ):
        return True

    return bool(
        first_cell
        and numeric_cells_after_first >= max(3, min(6, max(1, len(first_row) // 3)))
    )


def _continuation_header_grid_corruption_reasons(
    *,
    table_body: str,
    expected_header_rows: Optional[List[List[str]]],
    header_carryover_blocked: bool,
) -> List[str]:
    """Detect continuation fragments whose schema nearly matches but lost columns/header.

    Large multi-page tables normally keep the same schema across physical
    fragments. A later fragment that cannot accept carried headers because it is
    off by only one or two columns is more likely to be a corrupted extraction
    grid than a true new table schema.
    """
    if not header_carryover_blocked or not expected_header_rows:
        return []
    expected_width = _header_width(expected_header_rows)
    if expected_width <= 1:
        return []
    rows = _md_table_rows(table_body)
    if not rows:
        return []

    reasons: List[str] = []
    observed_width = len(rows[0])
    delta = abs(observed_width - expected_width)
    if observed_width != expected_width:
        reasons.append("continuation_header_missing_due_to_column_mismatch")
    if 1 <= delta <= 2:
        reasons.append("continuation_near_miss_column_count")
    if _first_markdown_row_looks_like_data_against_header(rows[0], expected_header_rows):
        reasons.append("continuation_starts_with_data_row_instead_of_header")
    return reasons


def _markdown_grid_corruption_reasons(block: str) -> List[str]:
    """Detect Markdown-grid corruption that should bypass Markdown-only repair."""
    rows = _md_table_rows(block)
    if len(rows) < 2:
        return []

    reasons: List[str] = []

    def add(reason: str) -> None:
        if reason not in reasons:
            reasons.append(reason)

    header = rows[0]
    data_rows = rows[2:] if len(rows) > 2 else []

    if len(data_rows) >= 3:
        nonempty_data_row_indices = [
            idx
            for idx, row in enumerate(data_rows)
            if any(_clean_text(re.sub(r"<[^>]+>", " ", cell or "")) for cell in row)
        ]
        if nonempty_data_row_indices == [0]:
            first_body = data_rows[0]
            nonempty_cells = [
                _clean_text(re.sub(r"<[^>]+>", " ", cell or ""))
                for cell in first_body
                if _clean_text(re.sub(r"<[^>]+>", " ", cell or ""))
            ]
            alpha_cells = sum(bool(re.search(r"[A-Za-z]", cell)) for cell in nonempty_cells)
            numeric_cells = sum(_independent_numeric_count(cell) > 0 for cell in nonempty_cells)
            if (
                len(nonempty_cells) >= 3
                and alpha_cells >= max(2, len(nonempty_cells) // 2)
                and numeric_cells <= max(1, len(nonempty_cells) // 5)
            ):
                add("markdown_header_only_blank_body_rows")

    # The Markdown renderer sometimes promotes the first data row into the
    # header row when a continuation fragment has no repeated header.
    if _looks_like_row_label_code(header[0] if header else ""):
        if sum(_independent_numeric_count(cell) > 0 for cell in header[1:]) >= 3:
            add("data_row_used_as_markdown_header")

    for row in data_rows:
        if not row:
            continue
        # Label leakage: a row/sample/material identifier appears in a neighbor
        # value column even though the row already has a row-label-like first
        # cell. Legitimate coded categorical columns such as Matrix=SC1/RT1 are
        # excluded by header semantics.
        if _row_label_leakage_suspected(row, header):
            add("row_label_leakage")

        if _has_neighboring_value_collapse(row, header):
            add("neighboring_value_collapse")

        if _has_dense_value_cell_collapse(row, header):
            add("dense_value_cell_collapsed_records")

    return reasons


def _markdown_width_compatible_with_header(block: str, header_rows: Optional[List[List[str]]]) -> bool:
    expected_width = _header_width(header_rows)
    if expected_width <= 1:
        return False
    rows = _md_table_rows(block)
    if not rows:
        return False
    return len(rows[0]) == expected_width


def _inject_carried_header_into_table_block(
    block: str,
    header_rows: Optional[List[List[str]]],
    *,
    caption: str = "",
    footnote: str = "",
) -> Tuple[str, bool]:
    """Inject carried headers into a table block after PDF repair when safe."""
    if not header_rows:
        return block, False
    pipe_body = "\n".join(_body_pipe_lines(block)).strip()
    if not pipe_body:
        return block, False
    if _body_already_has_header(pipe_body, header_rows):
        return block, False
    if not _markdown_width_compatible_with_header(pipe_body, header_rows):
        return block, False
    injected_body = _inject_carried_table_header(pipe_body, header_rows)
    return _compose_table_block(injected_body, caption=caption, footnote=footnote).rstrip(), True



def _markdown_table_shape(block: str) -> Tuple[int, int]:
    rows = _md_table_rows(block)
    if not rows:
        return 0, 0
    return max(0, len(rows) - 2), max((len(row) for row in rows), default=0)


def _drop_column_from_rows(rows: List[List[str]], drop_idx: int) -> List[List[str]]:
    out: List[List[str]] = []
    for row in rows:
        if drop_idx < len(row):
            out.append(row[:drop_idx] + row[drop_idx + 1 :])
        else:
            out.append(list(row))
    return out


def _project_header_rows_for_candidate_width(
    header_rows: Optional[List[List[str]]],
    observed_width: int,
) -> Tuple[Optional[List[List[str]]], Optional[int], str]:
    """Return exact or projected header rows for a candidate table width.

    Exact-width candidates receive the carried header. For one-column-short
    candidates, allow dropping a known optional/blank-like column such as the
    Crosslinker column in Table S6. This prevents rejecting a clean 17-column
    repair merely because a completely blank/non-extractable column was omitted.
    """
    expected_width = _header_width(header_rows)
    if not header_rows or expected_width <= 1 or observed_width <= 1:
        return None, None, "none"
    if observed_width == expected_width:
        return header_rows, None, "carried_header"
    if observed_width == expected_width - 1:
        first_header = header_rows[0]
        for idx, cell in enumerate(first_header):
            if TABLE_REPAIR_OPTIONAL_BLANK_COLUMN_RE.search(_clean_text(cell)):
                return _drop_column_from_rows(header_rows, idx), idx, "projected_header_drop_optional_blank_column"
    return None, None, "none"


def _prepare_repair_candidate_block(
    block: str,
    header_rows: Optional[List[List[str]]],
    *,
    caption: str = "",
    footnote: str = "",
) -> Tuple[str, Dict[str, Any]]:
    """Normalize a repair candidate and inject/project carried headers if safe."""
    pipe_body = "\n".join(_body_pipe_lines(block)).strip()
    if not pipe_body:
        return block.strip(), {"header_action": "none", "dropped_header_col_idx": None}
    observed_rows = _md_table_rows(pipe_body)
    observed_width = max((len(row) for row in observed_rows), default=0)

    if header_rows and _body_already_has_header(pipe_body, header_rows):
        prepared = _compose_table_block(pipe_body, caption=caption, footnote=footnote).rstrip()
        return prepared, {"header_action": "already_had_carried_header", "dropped_header_col_idx": None}

    candidate_header_rows, dropped_idx, header_action = _project_header_rows_for_candidate_width(
        header_rows,
        observed_width,
    )
    if candidate_header_rows:
        injected_body = _inject_carried_table_header(pipe_body, candidate_header_rows)
        prepared = _compose_table_block(injected_body, caption=caption, footnote=footnote).rstrip()
        return prepared, {"header_action": header_action, "dropped_header_col_idx": dropped_idx}

    prepared = _compose_table_block(pipe_body, caption=caption, footnote=footnote).rstrip()
    return prepared, {"header_action": "none", "dropped_header_col_idx": None}


def _candidate_grid_repair_reasons(block: str) -> List[str]:
    return _core_candidate_grid_repair_reasons(block)


def _make_table_repair_candidate(
    *,
    name: str,
    block: str,
    payload: Optional[Dict[str, Any]],
    expected_header_rows: Optional[List[List[str]]],
    caption: str,
    footnote: str,
    docling_bbox: Optional[Dict[str, Any]] = None,
    known_outside_texts: Optional[List[str]] = None,
    reference_block: str = "",
    extra_repair_reasons: Optional[List[str]] = None,
) -> Dict[str, Any]:
    return _core_make_table_repair_candidate(
        name=name,
        block=block,
        payload=payload,
        expected_header_rows=expected_header_rows,
        caption_text=caption,
        footnote_text=footnote,
        docling_bbox=docling_bbox,
        known_outside_texts=known_outside_texts,
        reference_block=reference_block,
        extra_repair_reasons=extra_repair_reasons,
    )


def _select_best_table_repair_candidate(candidates: List[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    return _core_select_best_table_repair_candidate(candidates)


def _make_table_quality_decision(
    candidate: Optional[Dict[str, Any]],
    *,
    repairability_status: Optional[str] = None,
    phase: str = "",
    default_corrupt_action: str = "repair",
) -> Dict[str, Any]:
    return _core_make_table_quality_decision(
        candidate,
        repairability_status=repairability_status,
        phase=phase,
        default_corrupt_action=default_corrupt_action,
    )


def repair_docling_table_with_pdf_tools(
    *,
    pdf_path: str | Path,
    table_obj: Dict[str, Any],
    docling_markdown: str,
    caption_text: str = "",
    footnote_text: str = "",
    expected_header_rows: Optional[List[List[str]]] = None,
    route: str = "pdf_repair_unguided",
    bbox_pad: float = 3.0,
    column_min_gap: float = 2.0,
    min_nonempty_cells: int = 6,
    prefer_lattice: bool = False,
    known_outside_texts: Optional[List[str]] = None,
    docling_original_repair_reasons: Optional[List[str]] = None,
) -> Dict[str, Any]:
    """Generate PDF extraction candidates, then select via shared repair core."""
    prov = (table_obj.get("prov") or [{}])[0]
    try:
        page_no = int(prov.get("page_no") or 0)
    except Exception:
        page_no = 0

    candidates: List[Dict[str, Any]] = []
    attempts: List[Dict[str, Any]] = []
    docling_bbox = prov.get("bbox") or {}
    docling_original_repair_reasons = list(docling_original_repair_reasons or [])

    def add_candidate(name: str, block: str, payload: Optional[Dict[str, Any]]) -> None:
        if not str(block or "").strip():
            return
        candidates.append(_make_table_repair_candidate(
            name=name,
            block=block,
            payload=payload or {},
            expected_header_rows=expected_header_rows,
            caption=caption_text,
            footnote=footnote_text,
            docling_bbox=docling_bbox,
            known_outside_texts=known_outside_texts,
            reference_block=docling_markdown,
            extra_repair_reasons=(
                docling_original_repair_reasons
                if name == "docling_original"
                else None
            ),
        ))

    add_candidate("docling_original", docling_markdown, {"status": "original"})

    camelot_payload = repair_docling_table_with_camelot(
        pdf_path=pdf_path,
        table_obj=table_obj,
        docling_markdown=docling_markdown,
        caption_text=caption_text,
        footnote_text=footnote_text,
        bbox_pad=bbox_pad,
        column_min_gap=column_min_gap,
        min_nonempty_cells=min_nonempty_cells,
        prefer_lattice=False,
        use_docling_columns=not prefer_lattice,
    )
    stream_attempts = camelot_payload.get("stream_attempts") or []
    if not stream_attempts:
        stream_attempt = camelot_payload.get("stream_attempt")
        if isinstance(stream_attempt, dict):
            stream_attempts = [stream_attempt]
    for stream_attempt in stream_attempts:
        if not isinstance(stream_attempt, dict):
            continue
        stream_mode = str(stream_attempt.get("stream_mode") or "").strip()
        stream_name = f"camelot_stream_{stream_mode}" if stream_mode else "camelot_stream"
        attempts.append(_core_attempt_summary(stream_name, stream_attempt))
        add_candidate(
            stream_name,
            str(stream_attempt.get("block") or ""),
            stream_attempt,
        )

    camelot_strategy = str(camelot_payload.get("strategy") or "")
    if camelot_strategy != "stream":
        attempts.append(_core_attempt_summary(
            f"camelot_{camelot_strategy or 'repair'}",
            camelot_payload,
        ))
        add_candidate(
            f"camelot_{camelot_strategy or 'repair'}",
            str(camelot_payload.get("block") or ""),
            camelot_payload,
        )

    if camelot_payload.get("strategy") == "stream" or (
        prefer_lattice and camelot_payload.get("strategy") != "lattice"
    ):
        lattice_payload = repair_docling_table_with_camelot(
            pdf_path=pdf_path,
            table_obj=table_obj,
            docling_markdown=docling_markdown,
            caption_text=caption_text,
            footnote_text=footnote_text,
            bbox_pad=bbox_pad,
            column_min_gap=column_min_gap,
            min_nonempty_cells=min_nonempty_cells,
            prefer_lattice=True,
        )
        attempts.append(_core_attempt_summary(
            f"camelot_{lattice_payload.get('strategy') or 'repair'}",
            lattice_payload,
        ))
        add_candidate(
            f"camelot_{lattice_payload.get('strategy') or 'repair'}",
            str(lattice_payload.get("block") or ""),
            lattice_payload,
        )

    if page_no:
        for strategy in ("lines", "text"):
            pymupdf_payload = extract_pymupdf_table_candidate(
                pdf_path=pdf_path,
                page_no=page_no,
                strategy=strategy,
                table_bbox=docling_bbox,
                bbox_pad=bbox_pad,
            )
            attempts.append(_core_attempt_summary(f"pymupdf_{strategy}", pymupdf_payload))
            add_candidate(f"pymupdf_{strategy}", str(pymupdf_payload.get("block") or ""), pymupdf_payload)

    return _core_build_repair_selection_payload(
        route=route,
        candidates=candidates,
        attempts=attempts,
        source_path=str(pdf_path),
        source_path_key="pdf_path",
        page_no=page_no,
    )


def repair_docling_table_with_docx_tools(
    *,
    docx_path: Path | str,
    table_id: str,
    docling_markdown: str,
    caption_text: str = "",
    footnote_text: str = "",
    expected_header_rows: Optional[List[List[str]]] = None,
    table_ordinal_hint: Optional[int] = None,
    preserve_line_breaks: bool = False,
    expand_multiline_cells: bool = True,
    docling_original_repair_reasons: Optional[List[str]] = None,
) -> Dict[str, Any]:
    """Generate DOCX extraction candidates, then select via shared repair core."""
    path = Path(docx_path)
    route = "docx_repair_native"
    candidates: List[Dict[str, Any]] = []
    attempts: List[Dict[str, Any]] = []
    docling_original_repair_reasons = list(docling_original_repair_reasons or [])

    def add_candidate(name: str, block: str, payload: Optional[Dict[str, Any]]) -> None:
        if not str(block or "").strip():
            return
        candidates.append(_make_table_repair_candidate(
            name=name,
            block=block,
            payload=payload or {},
            expected_header_rows=expected_header_rows,
            caption=caption_text,
            footnote=footnote_text,
            known_outside_texts=[caption_text, footnote_text],
            reference_block=docling_markdown,
            extra_repair_reasons=(
                docling_original_repair_reasons
                if name == "docling_original"
                else None
            ),
        ))

    add_candidate("docling_original", docling_markdown, {"status": "original"})

    docx_payload = extract_docx_native_table_candidate(
        docx_path=path,
        table_id=table_id,
        docling_markdown=docling_markdown,
        caption_text=caption_text,
        footnote_text=footnote_text,
        expected_header_rows=expected_header_rows,
        table_ordinal_hint=table_ordinal_hint,
        preserve_line_breaks=preserve_line_breaks,
        expand_multiline_cells=expand_multiline_cells,
    )
    attempts.append(_core_attempt_summary("docx_native", docx_payload))
    add_candidate("docx_native", str(docx_payload.get("block") or ""), docx_payload)

    if docx_payload.get("status") in {"skipped", "failed", "not_found"}:
        return {
            "status": docx_payload.get("status") or "skipped",
            "route": route,
            "accepted": False,
            "reason": docx_payload.get("reason"),
            "error": docx_payload.get("error"),
            "requested_table_id": docx_payload.get("requested_table_id") or table_id,
            "requested_ordinal": docx_payload.get("requested_ordinal") or table_ordinal_hint,
            "tables_scanned": docx_payload.get("tables_scanned", 0),
            "docx_path": str(path),
            "attempts": attempts,
            "candidates": candidates,
            "generated_at": datetime.now().isoformat(timespec="seconds"),
        }

    payload = _core_build_repair_selection_payload(
        route=route,
        candidates=candidates,
        attempts=attempts,
        source_path=str(path),
        source_path_key="docx_path",
        extra={
            "requested_table_id": docx_payload.get("requested_table_id") or table_id,
            "requested_ordinal": docx_payload.get("requested_ordinal") or table_ordinal_hint,
            "tables_scanned": docx_payload.get("tables_scanned", 0),
        },
    )
    payload["generated_at"] = datetime.now().isoformat(timespec="seconds")
    return payload

def _write_table_repair_bundle_artifacts(
    *,
    study: str,
    source_file: str,
    physical_table_id: str,
    table_obj: Dict[str, Any],
    reasons: List[str],
    initial_status: str,
    repair_status: str,
    repair_payload: Dict[str, Any],
    candidates: List[Dict[str, Any]],
    selected_candidate: Optional[Dict[str, Any]],
) -> Dict[str, Any]:
    """Write one navigable repair bundle instead of many unrelated artifacts."""
    art_dir = OUTPUT_DIR / study / "artifacts"
    art_dir.mkdir(parents=True, exist_ok=True)
    stem = _sanitize_filename(f"{source_file}_{physical_table_id}_repair")
    page_no, bbox = _page_bbox(table_obj)

    selected_name = str((selected_candidate or {}).get("name") or "")
    selected_clean = bool(selected_candidate and selected_candidate.get("accepted"))
    summary_candidates: List[Dict[str, Any]] = []
    for candidate in candidates:
        summary_candidates.append({
            k: v
            for k, v in candidate.items()
            if k not in {"block"}
        })

    payload = {
        "status": "accepted" if selected_clean else "needs_manual_review",
        "selection_mode": "accepted" if selected_clean else "provisional_review" if selected_name else "none",
        "selected_candidate": selected_name or None,
        "selected_candidate_accepted": selected_clean,
        "selected_candidate_is_provisional": bool(selected_name and not selected_clean),
        "reasons": reasons,
        "source_file": source_file,
        "table_id": physical_table_id,
        "page_no": page_no,
        "bbox": bbox,
        "initial_table_grid_status": initial_status,
        "post_repair_table_grid_status": repair_status,
        "initial_table_quality_decision": repair_payload.get("initial_table_quality_decision"),
        "post_repair_quality_decision": (
            repair_payload.get("post_repair_quality_decision")
            or repair_payload.get("post_expansion_quality_decision")
        ),
        "repair_payload": {k: v for k, v in repair_payload.items() if k not in {"block", "raw_block"}},
        "candidates": summary_candidates,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }

    md_lines = [f"# {physical_table_id} repair candidates", ""]
    if selected_clean:
        selection_message = f"Selected accepted candidate: `{selected_name}`"
    elif selected_name:
        selection_message = (
            "No candidate accepted; provisional review candidate: "
            f"`{selected_name}`"
        )
    else:
        selection_message = "No candidate accepted; no provisional review candidate available"
    md_lines.extend([selection_message, ""])

    attempts = list((repair_payload or {}).get("attempts") or [])
    if attempts:
        md_lines.extend(["## Extraction attempts", ""])
        for attempt in attempts:
            name = str(attempt.get("name") or "unknown")
            status = str(attempt.get("status") or "unknown")
            strategy = str(attempt.get("strategy") or "")
            reason = attempt.get("reason") or attempt.get("error") or ""
            produced_block = bool(attempt.get("produced_block"))
            details: List[str] = [
                f"status=`{status}`",
                f"produced_block=`{produced_block}`",
            ]
            if attempt.get("attempted") is not None:
                details.insert(0, f"attempted=`{bool(attempt.get('attempted'))}`")
            if strategy:
                details.append(f"strategy=`{strategy}`")
            stream_mode = str(attempt.get("stream_mode") or "")
            if stream_mode:
                details.append(f"stream_mode=`{stream_mode}`")
            if reason:
                details.append(f"reason=`{reason}`")
            for key in (
                "tables_found",
                "best_nonempty_cells",
                "shape",
                "bbox",
                "pymupdf_clip",
                "pymupdf_clip_source",
                "camelot_area",
                "docling_columns",
                "docling_columns_source",
            ):
                if key not in attempt:
                    continue
                value = attempt.get(key)
                if value is None or value == "":
                    continue
                details.append(f"{key}=`{value}`")
            md_lines.append(f"- `{name}`: " + ", ".join(details))
        md_lines.append("")

    for candidate in candidates:
        name = str(candidate.get("name") or "candidate")
        if name == selected_name:
            selected_flag = " selected" if selected_clean else " provisional_review_candidate"
        else:
            selected_flag = ""
        candidate_is_selected = name == selected_name
        validation_status = "eligible_for_auto_selection" if candidate.get("accepted") else "needs_repair"
        selection_status = (
            "selected_for_output"
            if candidate_is_selected and selected_clean
            else "provisional_review_choice"
            if candidate_is_selected
            else "not_selected"
        )
        md_lines.extend([
            f"## {name}{selected_flag}",
            "",
            f"- candidate validation: `{validation_status}`",
            f"- selection: `{selection_status}`",
            f"- shape: `{candidate.get('shape')}`",
            f"- header_action: `{candidate.get('header_action')}`",
            f"- reasons: `{candidate.get('reasons')}`",
            "",
        ])
        block = str(candidate.get("block") or "").strip()
        md_lines.extend([block if block else "_No Markdown block available._", ""])

    summary_path = art_dir / f"{stem}_summary.json"
    candidates_path = art_dir / f"{stem}_candidates.md"
    summary_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    candidates_path.write_text("\n".join(md_lines).rstrip() + "\n", encoding="utf-8")
    payload["artifact_path"] = str(summary_path)
    payload["candidate_comparison_path"] = str(candidates_path)
    return payload

def _looks_like_row_label_simple_value_glued(cell: str) -> bool:
    text = _clean_text(re.sub(r"<[^>]+>", " ", cell or ""))
    match = _ROW_LABEL_SIMPLE_VALUE_RE.match(text)
    if not match:
        return False
    label = match.group(1).strip()
    if not re.search(r"[A-Za-z]{2,}", label):
        return False
    if _EXCLUDED_ROW_LABEL_RE.search(label):
        return False
    return True


def _has_row_label_value_glue_from_rows(rows: List[List[str]]) -> bool:
    if len(rows) < 3:
        return False
    for row in rows[2:]:
        if not row or _EXCLUDED_ROW_LABEL_RE.search(row[0] or ""):
            continue
        independent_counts = [_independent_numeric_count(c) for c in row]
        if row and _looks_like_row_label_value_glued(row[0]) and sum(independent_counts[1:]) >= 2:
            return True
    return False


def _has_repeated_row_label_simple_value_glue_from_rows(rows: List[List[str]]) -> bool:
    if len(rows) < 5:
        return False
    candidate_rows = 0
    data_like_rows = 0
    for row in rows[2:]:
        if not row or _EXCLUDED_ROW_LABEL_RE.search(row[0] or ""):
            continue
        if sum(_independent_numeric_count(c) for c in row[1:]) < 2:
            continue
        data_like_rows += 1
        if _looks_like_row_label_simple_value_glued(row[0]):
            candidate_rows += 1
    if candidate_rows < 3:
        return False
    required = max(3, min(5, (data_like_rows + 3) // 4))
    return candidate_rows >= required


def _has_row_label_value_glue(block: str) -> bool:
    return _has_row_label_value_glue_from_rows(_md_table_rows(block))


def _has_repeated_row_label_simple_value_glue(block: str) -> bool:
    return _has_repeated_row_label_simple_value_glue_from_rows(_md_table_rows(block))


def _repair_docling_spanned_body_cells(table_obj: Dict[str, Any]) -> Tuple[str, Dict[str, Any]]:
    rows = _docling_grid_text_rows(table_obj)
    if not rows:
        return "", {"status": "skipped", "reason": "empty_grid"}
    width = max((len(row) for row in rows), default=0)
    rows = [list(row) + [""] * (width - len(row)) for row in rows]
    changed: List[Dict[str, Any]] = []

    for cell in ((table_obj.get("data") or {}).get("table_cells") or []):
        if cell.get("column_header"):
            continue
        span = _docling_body_cell_span_width(cell)
        if span <= 1:
            continue
        parts = _split_compacted_cell_to_width(str(cell.get("text") or ""), span)
        if not parts:
            continue
        try:
            row_idx = int(cell.get("start_row_offset_idx"))
            col_idx = int(cell.get("start_col_offset_idx"))
        except Exception:
            continue
        if row_idx < 0 or row_idx >= len(rows) or col_idx < 0:
            continue
        while col_idx + span > len(rows[row_idx]):
            rows[row_idx].append("")
        text = _clean_text(str(cell.get("text") or ""))
        covered = rows[row_idx][col_idx : col_idx + span]
        if any(_clean_text(value) not in {"", text} for value in covered):
            continue
        rows[row_idx][col_idx : col_idx + span] = parts
        changed.append({"row": row_idx + 1, "col": col_idx + 1, "span": span, "text": text})

    if not changed:
        return "", {"status": "skipped", "reason": "no_deterministic_span_splits"}
    return _markdown_from_text_grid(rows), {
        "status": "repaired",
        "strategy": "docling_body_colspan_split",
        "changed_cells": changed,
    }


def _is_suspected_merged_table(block: str) -> bool:
    rows = _md_table_rows(block)
    if len(rows) == 2:
        nonempty_header_cells = [cell for cell in rows[0] if (cell or "").strip()]
        if len(nonempty_header_cells) == 1:
            cell = nonempty_header_cells[0]
            if len(_clean_text(cell)) >= 500 and _independent_numeric_count(cell) >= 10:
                return True
            if _independent_numeric_count(cell) >= 25:
                return True
        return False
    if len(rows) < 3:
        return False
    if _has_repeated_row_label_simple_value_glue_from_rows(rows):
        return True
    headers = rows[0]
    excluded = {i for i, h in enumerate(headers) if _EXCLUDED_HEADER_RE.search(h or "")}
    for row in rows[2:]:
        if row and _EXCLUDED_ROW_LABEL_RE.search(row[0] or ""):
            continue
        raw_counts = [_independent_numeric_count(c) for c in row]
        if row and _looks_like_row_label_value_glued(row[0]) and sum(raw_counts[1:]) >= 2:
            return True
        cells = [
            (c, headers[i] if i < len(headers) else "")
            for i, c in enumerate(row)
            if i not in excluded and (c or "").strip()
        ]
        if any(_duplicate_text_cell(c) for c, _header_cell in cells):
            return True
        counts = [_suspicious_numeric_value_count(c, header_cell) for c, header_cell in cells]
        if sum(c >= 3 for c in counts) >= 2:
            return True
        if max(counts, default=0) >= 4 and sum(c >= 2 for c in counts) >= 2:
            return True
    return False


# LLM table triage ------------------------------------------------------------
_TABLE_POSITION_CHAIN = None


def _get_table_position_chain():
    global _TABLE_POSITION_CHAIN
    if _TABLE_POSITION_CHAIN is None:
        try:
            from .chains.table_compaction_triage_chain import create_table_compaction_triage_chain  # type: ignore
        except Exception:
            try:
                from chains.table_compaction_triage_chain import create_table_compaction_triage_chain  # type: ignore
            except Exception as exc:
                raise RuntimeError(f"Cannot import table compaction triage chain: {exc}")
        _TABLE_POSITION_CHAIN = create_table_compaction_triage_chain(
            provider=LLM_PROVIDER,
            model_name=LLM_MODEL_NAME,
        )
    return _TABLE_POSITION_CHAIN


def _normalize_position_status(value: Any) -> str:
    return _core_normalize_table_repairability_status(value)


def _position_triage_cache_path(study: str, source_file: str, table_id: str, suffix: str = "") -> Path:
    stem = _sanitize_filename(f"{source_file}_{table_id}{suffix}_position_triage")
    return OUTPUT_DIR / study / "artifacts" / f"{stem}.json"


def _classify_table_position_with_cache(
    *,
    study: str,
    source_file: str,
    table_id: str,
    caption_text: str,
    table_markdown: str,
    suffix: str = "",
) -> Dict[str, Any]:
    meta = {
        "schema_version": TABLE_GRID_TRIAGE_SCHEMA_VERSION,
        "provider": LLM_PROVIDER,
        "model_name": LLM_MODEL_NAME,
        "caption_sha256": _sha256_text(caption_text),
        "table_sha256": _sha256_text(table_markdown),
    }
    path = _position_triage_cache_path(study, source_file, table_id, suffix=suffix)
    if not FORCE_RERUN_LLM and path.exists():
        try:
            obj = _load_json(path)
            if obj.get("meta") == meta:
                obj["status"] = _normalize_position_status(obj.get("status"))
                return obj
        except Exception:
            pass

    if predict_with_usage is None:
        payload = {"status": "uncertain", "meta": meta, "error": "predict_with_usage unavailable"}
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        return payload

    try:
        chain = _get_table_position_chain()
        call = predict_with_usage(
            chain,
            provider=LLM_PROVIDER,
            model_name=LLM_MODEL_NAME,
            study_id=study,
            source_file=source_file,
            table_id=table_id,
            caption_text=caption_text or "",
            table_markdown=table_markdown or "",
        )
        parsed = _safe_json_loads(call.text)
        status = _normalize_position_status(parsed.get("status"))
        payload = {
            "status": status,
            "meta": meta,
            "llm_call": call.usage.metadata_dict(),
            "generated_at": datetime.now().isoformat(timespec="seconds"),
        }
        print(
            f"[LLM] {study}/{source_file} table position triage {table_id}: status={status} "
            f"prompt={call.usage.prompt_tokens} completion={call.usage.completion_tokens} total={call.usage.total_tokens}",
            file=sys.stderr,
        )
    except Exception as exc:
        payload = {"status": "uncertain", "meta": meta, "error": f"LLM table triage failed: {exc}"}
        print(f"⚠️ LLM table triage failed → {study}/{source_file}/{table_id}: {exc}", file=sys.stderr)

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload




# LLM table header row resolution ---------------------------------------------
_TABLE_HEADER_ROWS_CHAIN = None
_TABLE_HEADER_DESCRIPTOR_RE = re.compile(
    r"^(?:"
    r"adsorbents?|sorbents?|materials?|matrix|matrices|media|medium|date|time|"
    r"sample|condition|conditions|unit|units|analytes?|compounds?|pfas|"
    r"experiment|replicate|solvent|reagents?|dose|concentration|initial|final|"
    r"water|groundwater|parameter|parameters|components?"
    r")\b",
    re.I,
)


def _get_table_header_rows_chain():
    global _TABLE_HEADER_ROWS_CHAIN
    if _TABLE_HEADER_ROWS_CHAIN is None:
        try:
            from .chains.table_header_rows_chain import create_table_header_rows_chain  # type: ignore
        except Exception:
            try:
                from chains.table_header_rows_chain import create_table_header_rows_chain  # type: ignore
            except Exception as exc:
                raise RuntimeError(f"Cannot import table header rows chain: {exc}")
        _TABLE_HEADER_ROWS_CHAIN = create_table_header_rows_chain(
            provider=LLM_PROVIDER,
            model_name=LLM_MODEL_NAME,
        )
    return _TABLE_HEADER_ROWS_CHAIN


def _table_header_rows_cache_path(study: str, source_file: str, table_id: str, suffix: str = "") -> Path:
    stem = _sanitize_filename(f"{source_file}_{table_id}{suffix}_table_header_rows")
    return OUTPUT_DIR / study / "artifacts" / f"{stem}.json"


def _table_physical_rows_from_block(table_block: str) -> List[str]:
    pipe_lines = _body_pipe_lines(table_block)
    if not pipe_lines:
        return []
    separator_index = next(
        (
            i
            for i, line in enumerate(pipe_lines)
            if len(_split_md_row(line)) > 1
            and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in _split_md_row(line))
        ),
        None,
    )
    if separator_index is None:
        return pipe_lines
    return pipe_lines[:separator_index] + pipe_lines[separator_index + 1 :]


def _header_context_score(cells: List[str]) -> int:
    nonempty = [_clean_text(cell) for cell in cells if _clean_text(cell)]
    if not nonempty:
        return 0
    score = 0
    first = nonempty[0]
    if _TABLE_HEADER_DESCRIPTOR_RE.search(first):
        score += 2
    numeric_cells = sum(_independent_numeric_count(cell) > 0 for cell in nonempty)
    alpha_cells = sum(bool(re.search(r"[A-Za-z]", cell)) for cell in nonempty)
    if alpha_cells >= max(2, len(nonempty) // 2) and numeric_cells <= max(1, len(nonempty) // 5):
        score += 1
    short_label_cells = sum(bool(re.fullmatch(r"[A-Za-z][A-Za-z0-9+\-/(). ]{0,18}", cell)) for cell in nonempty)
    if len(nonempty) >= 4 and short_label_cells >= len(nonempty) * 0.6:
        score += 1
    return score


def _table_header_llm_suspicious(
    table_obj: Dict[str, Any],
    physical_rows: List[str],
    docling_header_count: int,
    *,
    force_llm: bool = False,
) -> bool:
    if force_llm:
        return True
    if docling_header_count > 1:
        return False
    if len(physical_rows) <= 2:
        return False
    first_width = len(_split_md_row(physical_rows[0]))
    if first_width < 4 and _table_data_cell_proxy(table_obj) < LARGE_TABLE_SPLIT_MIN_DATA_CELL_PROXY:
        return False
    lookahead = physical_rows[docling_header_count : min(len(physical_rows), docling_header_count + 5)]
    if any(_header_context_score(_split_md_row(row)) >= 2 for row in lookahead):
        return True
    if first_width >= 6 and any(_header_context_score(_split_md_row(row)) >= 1 for row in lookahead):
        return True
    return False


def _numbered_header_sample(physical_rows: List[str]) -> str:
    sample = physical_rows[:TABLE_HEADER_LLM_SAMPLE_ROWS]
    return "\n".join(f"ROW {idx}: {row}" for idx, row in enumerate(sample, start=1)).strip()


def _normalize_llm_header_rows(raw: Dict[str, Any], sample_count: int) -> List[int]:
    values = raw.get("header_rows") if isinstance(raw, dict) else []
    if not isinstance(values, list):
        return []
    rows: List[int] = []
    for value in values:
        try:
            idx = int(value)
        except Exception:
            continue
        if 1 <= idx <= sample_count and idx not in rows:
            rows.append(idx)
    rows = sorted(rows)
    if not rows:
        return []
    expected = list(range(1, rows[-1] + 1))
    if rows != expected:
        return []
    # Require at least one sampled non-header row for contrast; otherwise the
    # model may simply be marking every row it saw.
    if rows[-1] >= sample_count:
        return []
    return rows


def _resolve_table_header_rows_for_block(
    *,
    study: str,
    source_file: str,
    table_id: str,
    table_obj: Dict[str, Any],
    caption_text: str,
    physical_rows: List[str],
    force_llm: bool = False,
    suffix: str = "",
) -> Dict[str, Any]:
    if not physical_rows:
        return {"source": "none", "header_count": 0, "header_rows": []}

    max_header = max(1, len(physical_rows) - 1)
    docling_header_count = min(max(1, _table_header_row_count(table_obj)), max_header)
    fallback = {
        "source": "docling",
        "header_count": docling_header_count,
        "header_rows": list(range(1, docling_header_count + 1)),
    }

    if not _table_header_llm_suspicious(
        table_obj,
        physical_rows,
        docling_header_count,
        force_llm=force_llm,
    ):
        return fallback

    sample_rows = physical_rows[:TABLE_HEADER_LLM_SAMPLE_ROWS]
    sample_count = len(sample_rows)
    if sample_count <= 1:
        return fallback

    numbered_rows = _numbered_header_sample(physical_rows)
    meta = {
        "schema_version": TABLE_HEADER_ROWS_SCHEMA_VERSION,
        "provider": LLM_PROVIDER,
        "model_name": LLM_MODEL_NAME,
        "caption_sha256": _sha256_text(caption_text),
        "sample_sha256": _sha256_text(numbered_rows),
        "docling_header_count": docling_header_count,
    }
    cache_table_id = table_id or "unknown_table"
    path = _table_header_rows_cache_path(study or "unknown_study", source_file or "unknown_source", cache_table_id, suffix=suffix)
    if not FORCE_RERUN_LLM and path.exists():
        try:
            cached = _load_json(path)
            if cached.get("meta") == meta and isinstance(cached.get("header_count"), int):
                return cached
        except Exception:
            pass

    payload: Dict[str, Any] = {
        **fallback,
        "meta": meta,
        "llm_header_rows": [],
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }

    if predict_with_usage is None:
        payload["error"] = "predict_with_usage unavailable"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        return payload

    rendered_prompt = ""
    raw_response = ""
    try:
        chain = _get_table_header_rows_chain()
        try:
            rendered_prompt = chain.prompt.format(
                study_id=study,
                source_file=source_file,
                table_id=table_id,
                caption_text=caption_text or "",
                docling_header_rows=str(docling_header_count),
                numbered_rows=numbered_rows,
            )
        except Exception:
            rendered_prompt = numbered_rows
        call = predict_with_usage(
            chain,
            provider=LLM_PROVIDER,
            model_name=LLM_MODEL_NAME,
            study_id=study,
            source_file=source_file,
            table_id=table_id,
            caption_text=caption_text or "",
            docling_header_rows=str(docling_header_count),
            numbered_rows=numbered_rows,
        )
        raw_response = call.text
        llm_rows = _normalize_llm_header_rows(_safe_json_loads(raw_response), sample_count)
        if llm_rows:
            payload.update(
                {
                    "source": "llm",
                    "header_count": llm_rows[-1],
                    "header_rows": llm_rows,
                    "llm_header_rows": llm_rows,
                    "llm_call": call.usage.metadata_dict(),
                }
            )
        else:
            payload["llm_header_rows"] = []
            payload["llm_raw_response_empty_or_invalid"] = True
        print(
            f"[LLM] {study}/{source_file} table header rows {table_id}: "
            f"source={payload.get('source')} header_count={payload.get('header_count')} "
            f"prompt={call.usage.prompt_tokens} completion={call.usage.completion_tokens} total={call.usage.total_tokens}",
            file=sys.stderr,
        )
    except Exception as exc:
        payload["error"] = f"LLM table header rows failed: {exc}"
        print(f"LLM table header rows failed -> {study}/{source_file}/{table_id}: {exc}", file=sys.stderr)

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    if DEBUG_MODE and study and source_file and table_id:
        dbg_dir = OUTPUT_DIR / study / "debug"
        dbg_dir.mkdir(parents=True, exist_ok=True)
        stem = _sanitize_filename(f"{source_file}_{table_id}{suffix}_table_header_rows")
        if rendered_prompt:
            (dbg_dir / f"{stem}_prompt.txt").write_text(rendered_prompt, encoding="utf-8")
        if raw_response:
            (dbg_dir / f"{stem}_raw.txt").write_text(raw_response, encoding="utf-8")

    return payload


# LLM structured caption / footnote resolution --------------------------------
_STRUCTURED_CAPTION_CHAIN = None


def _get_structured_caption_chain():
    global _STRUCTURED_CAPTION_CHAIN
    if _STRUCTURED_CAPTION_CHAIN is None:
        try:
            from .chains.structured_caption_resolution_chain import create_structured_caption_resolution_chain  # type: ignore
        except Exception:
            try:
                from chains.structured_caption_resolution_chain import create_structured_caption_resolution_chain  # type: ignore
            except Exception as exc:
                raise RuntimeError(f"Cannot import structured caption resolution chain: {exc}")
        _STRUCTURED_CAPTION_CHAIN = create_structured_caption_resolution_chain(
            provider=LLM_PROVIDER,
            model_name=LLM_MODEL_NAME,
        )
    return _STRUCTURED_CAPTION_CHAIN


def _structured_caption_cache_path(study: str, source_file: str) -> Path:
    stem = _sanitize_filename(f"{source_file}_structured_caption_resolution")
    return OUTPUT_DIR / study / "artifacts" / f"{stem}.json"


def _truncate_text(value: str, max_chars: int) -> str:
    txt = value or ""
    return txt if len(txt) <= max_chars else txt[:max_chars].rstrip() + " ...[truncated]"


def _usable_local_text_object(so: StreamObject) -> bool:
    return (
        so.object_type == "text"
        and so.content_layer != "furniture"
        and so.label not in {"page_header", "page_footer"}
        and bool(_clean_text(so.text))
    )

def _raw_char_bounded_local_text(
    doc: Dict[str, Any],
    stream: List[StreamObject],
    stream_index: int,
    *,
    before: bool,
    max_chars: int,
) -> str:
    """Return nearby raw text for repair-contamination checks.

    Unlike the LLM local-text helper, this intentionally keeps furniture,
    headers/footers, citations, and ordinary body text. The text is not used as
    table content; it is only outside-grid evidence when a PDF repair candidate
    pulls nearby text into the table body.
    """
    indices = (
        range(stream_index - 1, -1, -1)
        if before
        else range(stream_index + 1, len(stream))
    )
    pieces: List[str] = []
    char_total = 0

    for j in indices:
        so = stream[j]
        text = _clean_text(so.text)
        if not text and so.object_type in {"table", "picture"}:
            text = _clean_text(_caption_text_with_adjacent_identifier(doc, stream, j))
        if not text:
            continue
        pieces.append(text)
        char_total += len(text) + 1
        if char_total >= max_chars:
            break

    if before:
        pieces.reverse()
        joined = "\n".join(pieces).strip()
        if len(joined) > max_chars:
            joined = "...[truncated]\n" + joined[-max_chars:].lstrip()
        return joined

    joined = "\n".join(pieces).strip()
    if len(joined) > max_chars:
        joined = joined[:max_chars].rstrip() + "\n...[truncated]"
    return joined

def _char_bounded_local_text(
    doc: Dict[str, Any],
    stream: List[StreamObject],
    stream_index: int,
    *,
    before: bool,
    max_chars: int,
) -> str:
    """Return nearby body text using a character budget, not object counts.

    This intentionally takes the tail of the previous text stream and the head
    of the following text stream. One Docling text object may contain a whole
    section, so per-object truncation is the wrong unit here.
    """
    current = stream[stream_index]
    indices = (
        range(stream_index - 1, -1, -1)
        if before
        else range(stream_index + 1, len(stream))
    )
    pieces: List[str] = []
    char_total = 0

    for j in indices:
        so = stream[j]
        # Page numbers are useful layout metadata, but captions often sit
        # immediately across a page break. Keep walking the stream until the
        # character budget is exhausted.
        if not _usable_local_text_object(so):
            # Tables/pictures are not treated as boundaries, because Docling
            # sometimes splits one printed table into several table objects. If
            # a nearby object has an actual linked caption, keep that printed
            # caption text as local evidence for continuation fragments.
            if so.object_type in {"table", "picture"}:
                text = _clean_text(_caption_text_with_adjacent_identifier(doc, stream, j))
                if not text:
                    continue
            else:
                continue
        else:
            text = _clean_text(so.text)
        pieces.append(text)
        char_total += len(text) + 1
        if char_total >= max_chars:
            break

    if before:
        pieces.reverse()
        joined = "\n".join(pieces).strip()
        if len(joined) > max_chars:
            joined = "...[truncated]\n" + joined[-max_chars:].lstrip()
        return joined

    joined = "\n".join(pieces).strip()
    if len(joined) > max_chars:
        joined = joined[:max_chars].rstrip() + "\n...[truncated]"
    return joined


def _table_markdown_preview(table_markdown: str) -> str:
    """Compact long markdown tables for caption/footnote recovery."""
    text = (table_markdown or "").strip()
    if not text:
        return ""

    lines = text.splitlines()
    pipe_indices = [i for i, line in enumerate(lines) if line.lstrip().startswith("|")]
    if len(pipe_indices) < STRUCTURED_CAPTION_TABLE_PREVIEW_MIN_ROWS_TO_COMPRESS:
        return _truncate_text(text, STRUCTURED_CAPTION_OBJECT_PREVIEW_CHARS)

    pipe_set = set(pipe_indices)
    pipe_lines = [lines[i] for i in pipe_indices]
    head = pipe_lines[:STRUCTURED_CAPTION_TABLE_PREVIEW_HEAD_ROWS]
    tail = pipe_lines[-STRUCTURED_CAPTION_TABLE_PREVIEW_TAIL_ROWS:]
    preview_lines = [
        *head,
        "| ... middle rows omitted for caption recovery ... |",
        *tail,
    ]
    preview = "\n".join(preview_lines)
    return _truncate_text(preview, STRUCTURED_CAPTION_OBJECT_PREVIEW_CHARS)




def _prompt_attr(value: Any) -> str:
    text = _clean_text(str(value or ""))
    return (
        text.replace("&", "&amp;")
        .replace('"', "&quot;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _prompt_object_refs(seed: Dict[str, Any]) -> List[str]:
    refs = [str(ref or "").strip() for ref in (seed.get("object_refs") or [])]
    if not refs and seed.get("object_ref"):
        refs = [str(seed.get("object_ref") or "").strip()]
    return [ref for ref in refs if ref]


def _llm_requests_from_seeds(seeds: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    requests: List[Dict[str, Any]] = []
    seen: set[Tuple[str, ...]] = set()
    for seed in seeds:
        refs = _prompt_object_refs(seed)
        if not refs:
            continue
        key = tuple(refs)
        if key in seen:
            continue
        seen.add(key)
        requests.append(
            {
                "object_refs": refs,
                "object_type": str(seed.get("object_type") or ""),
                "missing_fields": list(seed.get("missing_fields") or []),
            }
        )
    return requests


def _prompt_length_for_stream_object(doc: Dict[str, Any], so: StreamObject) -> int:
    if so.object_type == "text":
        if not _usable_local_text_object(so):
            return 0
        return len(_clean_text(so.text)) + 1
    if so.object_type == "table":
        try:
            preview = _table_markdown_preview(_table_body_to_markdown(doc, so.raw))
        except Exception:
            preview = ""
        return len(preview) + 160
    if so.object_type == "picture":
        return 120
    return 0


def _llm_window_for_seed(
    doc: Dict[str, Any],
    stream: List[StreamObject],
    seed: Dict[str, Any],
    index_by_ref: Dict[str, int],
) -> Optional[Tuple[int, int]]:
    indices = [index_by_ref[ref] for ref in _prompt_object_refs(seed) if ref in index_by_ref]
    if not indices:
        return None
    first = min(indices)
    last = max(indices)
    before_chars = (
        STRUCTURED_CAPTION_TABLE_BEFORE_CHARS
        if str(seed.get("object_type") or "") == "table"
        else STRUCTURED_CAPTION_FIGURE_BEFORE_CHARS
    )
    after_chars = (
        STRUCTURED_CAPTION_TABLE_AFTER_CHARS
        if str(seed.get("object_type") or "") == "table"
        else STRUCTURED_CAPTION_FIGURE_AFTER_CHARS
    )

    start = first
    total = 0
    for j in range(first - 1, -1, -1):
        size = _prompt_length_for_stream_object(doc, stream[j])
        if size <= 0:
            start = j
            continue
        if total and total + size > before_chars:
            break
        total += size
        start = j
        if total >= before_chars:
            break

    end = last
    total = 0
    for j in range(last + 1, len(stream)):
        size = _prompt_length_for_stream_object(doc, stream[j])
        if size <= 0:
            end = j
            continue
        if total and total + size > after_chars:
            break
        total += size
        end = j
        if total >= after_chars:
            break

    return start, end


def _merge_llm_windows(windows: List[Tuple[int, int]]) -> List[Tuple[int, int]]:
    if not windows:
        return []
    merged: List[Tuple[int, int]] = []
    for start, end in sorted(windows):
        if not merged or start > merged[-1][1] + 1:
            merged.append((start, end))
        else:
            merged[-1] = (merged[-1][0], max(merged[-1][1], end))
    return merged


def _context_resolution_caption(doc: Dict[str, Any], so: StreamObject, context_resolutions: Dict[str, Dict[str, Any]]) -> str:
    resolution = context_resolutions.get(so.object_ref) or {}
    caption = _clean_text(str(resolution.get("caption_text") or ""))
    if caption:
        return caption
    if so.object_type in {"table", "picture"}:
        return _clean_text(_caption_text(doc, so.raw))
    return ""


def _stream_object_marked_snippet(
    doc: Dict[str, Any],
    so: StreamObject,
    requested_by_ref: Dict[str, Dict[str, Any]],
    context_resolutions: Dict[str, Dict[str, Any]],
) -> str:
    if so.object_type == "text":
        if not _usable_local_text_object(so):
            return ""
        return _clean_text(so.text)

    if so.object_type not in {"table", "picture"}:
        return ""

    seed = requested_by_ref.get(so.object_ref)
    status = "unresolved" if seed else "context"
    missing = ",".join(str(x) for x in (seed or {}).get("missing_fields", []) if x)
    caption = _context_resolution_caption(doc, so, context_resolutions)
    if caption and not seed:
        status = "known"

    object_kind = "TABLE_OBJECT" if so.object_type == "table" else "PICTURE_OBJECT"
    attrs = [
        f'ref="{_prompt_attr(so.object_ref)}"',
        f'status="{status}"',
    ]
    if missing:
        attrs.append(f'missing="{_prompt_attr(missing)}"')
    if so.page_no is not None:
        attrs.append(f'page="{_prompt_attr(so.page_no)}"')
    marker = f"<!-- {object_kind} {' '.join(attrs)} -->"
    end_marker = f"<!-- END_{object_kind} -->"
    lines = [marker]
    if caption:
        label = "Known caption" if status == "known" else "Context caption"
        lines.append(f"{label}: {caption}")

    if so.object_type == "table":
        body = _table_markdown_preview(_table_body_to_markdown(doc, so.raw))
        lines.append(body or "[table preview unavailable]")
    else:
        lines.append("[picture content omitted]")
    lines.append(end_marker)
    return "\n".join(lines).strip()


def _marked_snippets_from_seeds(
    *,
    doc: Dict[str, Any],
    stream: List[StreamObject],
    seeds: List[Dict[str, Any]],
    context_resolutions: Dict[str, Dict[str, Any]],
) -> str:
    index_by_ref = {so.object_ref: idx for idx, so in enumerate(stream)}
    requested_by_ref: Dict[str, Dict[str, Any]] = {}
    for seed in seeds:
        for ref in _prompt_object_refs(seed):
            requested_by_ref[ref] = seed

    windows = [
        window
        for seed in seeds
        if (window := _llm_window_for_seed(doc, stream, seed, index_by_ref)) is not None
    ]
    snippets: List[str] = []
    for snippet_index, (start, end) in enumerate(_merge_llm_windows(windows), start=1):
        lines = [f"=== SNIPPET {snippet_index} ==="]
        last_page: Optional[int] = None
        for so in stream[start : end + 1]:
            if (
                last_page is not None
                and so.page_no is not None
                and so.page_no != last_page
            ):
                lines.append(f"<!-- PAGE BREAK: {last_page} to {so.page_no} -->")
            if so.page_no is not None:
                last_page = so.page_no
            rendered = _stream_object_marked_snippet(
                doc,
                so,
                requested_by_ref,
                context_resolutions,
            )
            if rendered:
                lines.append(rendered)
        snippets.append("\n\n".join(lines).strip())
    return "\n\n".join(snippets).strip()


def _picture_text_preview(doc: Dict[str, Any], picture_obj: Dict[str, Any]) -> str:
    parts: List[str] = []
    for child in picture_obj.get("children") or []:
        text = _text_from_ref(doc, child)
        if text:
            parts.append(text)
    return _truncate_text("\n".join(parts), STRUCTURED_CAPTION_OBJECT_PREVIEW_CHARS)


def _local_evidence_packet(
    *,
    doc: Dict[str, Any],
    stream: List[StreamObject],
    stream_index: int,
) -> Dict[str, Any]:
    so = stream[stream_index]
    is_table = so.object_type == "table"
    before_chars = (
        STRUCTURED_CAPTION_TABLE_BEFORE_CHARS
        if is_table
        else STRUCTURED_CAPTION_FIGURE_BEFORE_CHARS
    )
    after_chars = (
        STRUCTURED_CAPTION_TABLE_AFTER_CHARS
        if is_table
        else STRUCTURED_CAPTION_FIGURE_AFTER_CHARS
    )
    text_before = _char_bounded_local_text(
        doc,
        stream,
        stream_index,
        before=True,
        max_chars=before_chars,
    )
    text_after = _char_bounded_local_text(
        doc,
        stream,
        stream_index,
        before=False,
        max_chars=after_chars,
    )
    # The LLM fallback should see the closest printed text even when Docling
    # stores that text as children of the object rather than as body-stream
    # siblings. Keep this simple by folding linked captions/footnotes into the
    # ordinary before/after text fields instead of adding extra prompt keys.
    linked_caption = _caption_text_with_adjacent_identifier(doc, stream, stream_index)
    if linked_caption:
        if is_table:
            text_before = _local_text_join(
                text_before,
                linked_caption,
                max_chars=before_chars,
                keep_tail=True,
            )
        else:
            text_after = _local_text_join(
                linked_caption,
                text_after,
                max_chars=after_chars,
                keep_tail=False,
            )

    if is_table:
        linked_footnote, _ = _linked_text(doc, so.raw, "footnotes")
        if linked_footnote:
            text_after = _local_text_join(
                linked_footnote,
                text_after,
                max_chars=after_chars,
                keep_tail=False,
            )

    object_type_for_candidates = "table" if is_table else "figure"
    caption_candidates: List[Dict[str, Any]] = []
    if is_table:
        adjacent_candidate = _adjacent_table_caption_candidate(doc, stream, stream_index)
        if adjacent_candidate:
            caption_candidates.append(adjacent_candidate)
    for candidate in _find_caption_candidates(text_before, object_type_for_candidates):
        caption_candidates.append({"source": "before", **candidate})
    if is_table:
        caption_candidates.extend(
            _immediate_after_caption_candidates(stream, stream_index, object_type_for_candidates)
        )
    else:
        for candidate in _find_caption_candidates(text_after, object_type_for_candidates):
            caption_candidates.append({"source": "after", **candidate})
    # caption_candidates and leading_table_footnote_candidates are internal
    # deterministic evidence. They are intentionally not sent to the LLM.
    packet: Dict[str, Any] = {
        "text_before": text_before,
        "text_after": text_after,
        "caption_candidates": caption_candidates,
    }
    if is_table:
        table_body = _table_body_to_markdown(doc, so.raw)
        packet["table_preview"] = _table_markdown_preview(table_body)
        packet["leading_table_footnote_candidates"] = _leading_table_footnote_candidates(text_after)
    return packet


def _preferred_caption_candidate(
    candidates: List[Dict[str, Any]],
    *,
    object_type: str,
) -> Optional[Dict[str, Any]]:
    if not candidates:
        return None

    adjacent = [c for c in candidates if c.get("source") == "adjacent_before"]
    adjacent_unique = {(str(c.get("object_id") or ""), str(c.get("text") or "")) for c in adjacent}
    if len(adjacent_unique) == 1:
        return adjacent[0]

    # Tables usually have captions above. Do not deterministically attach an
    # after-caption to an upstream table; rare below-table captions or mixed-side
    # evidence should be handled by the LLM.
    if object_type == "table":
        preferred = [c for c in candidates if c.get("source") == "before"]
        if not preferred:
            return None
    else:
        preferred = [c for c in candidates if c.get("source") == "after"] or candidates

    # Use deterministic routing only when the preferred side has one plausible
    # printed caption. Multiple plausible captions should go to the LLM.
    unique = {(str(c.get("object_id") or ""), str(c.get("text") or "")) for c in preferred}
    if len(unique) != 1:
        return None
    return preferred[0]


def _regex_resolution_from_seed(seed: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    """Cheap deterministic resolver for unambiguous local caption/note cases."""
    object_ref = str(seed.get("object_ref") or "")
    object_type = str(seed.get("object_type") or "")
    missing_fields = set(seed.get("missing_fields") or [])
    candidate = _preferred_caption_candidate(
        list(seed.get("caption_candidates") or []),
        object_type=object_type,
    )

    object_id: Optional[str] = None
    caption_text: Optional[str] = None
    if candidate and ({"object_id", "caption_text", "table_id", "figure_id"} & missing_fields):
        object_id = str(candidate.get("object_id") or "") or None
        if not _caption_candidate_is_identifier_only(candidate):
            caption_text = str(candidate.get("text") or "") or None

    footnote_text: Optional[str] = None
    if object_type == "table" and ({"footnote_text", "table_footnote"} & missing_fields):
        footnotes = [str(x or "").strip() for x in seed.get("leading_table_footnote_candidates") or []]
        footnotes = [x for x in footnotes if x]
        # Immediate clean Note/Abbreviations patterns are safe to attach. If
        # several are consecutive, keep them together; this is common for tables.
        if footnotes:
            footnote_text = "\n".join(footnotes)

    if not object_id and not caption_text and not footnote_text:
        return None
    object_refs = [str(ref or "").strip() for ref in (seed.get("object_refs") or [])]
    if not object_refs and object_ref:
        object_refs = [object_ref]

    return {
        "object_ref": object_ref or (object_refs[0] if object_refs else ""),
        "object_refs": object_refs,
        "object_type": "table" if object_type == "table" else "picture",
        "status": "resolved",
        "object_id": object_id,
        "caption_text": caption_text,
        "footnote_text": footnote_text,
        "source": "structured_caption_regex",
    }

def _expand_resolution_for_object_refs(resolution: Optional[Dict[str, Any]]) -> List[Dict[str, Any]]:
    if not resolution:
        return []
    refs = [str(ref or "").strip() for ref in (resolution.get("object_refs") or [])]
    if not refs and resolution.get("object_ref"):
        refs = [str(resolution.get("object_ref") or "").strip()]
    refs = [ref for ref in refs if ref]
    if not refs:
        return []
    out: List[Dict[str, Any]] = []
    for ref in refs:
        item = dict(resolution)
        item["object_ref"] = ref
        item["object_refs"] = refs
        out.append(item)
    return out

def _debug_print_and_write(study: str, source_file: str, stem: str, text: str) -> None:
    if not DEBUG_MODE:
        return
    debug_dir = OUTPUT_DIR / study / "debug"
    try:
        debug_dir.mkdir(parents=True, exist_ok=True)
        (debug_dir / f"{source_file}_structured_caption_{stem}.txt").write_text(
            text or "",
            encoding="utf-8",
            errors="replace",
        )
    except Exception as exc:
        print(
            f"⚠️ Failed to write structured caption debug file ({stem}): {exc}",
            file=sys.stderr,
        )

    shown = text or ""
    if DEBUG_PRINT_MAX_CHARS is not None and len(shown) > DEBUG_PRINT_MAX_CHARS:
        shown = shown[:DEBUG_PRINT_MAX_CHARS].rstrip() + "\n...[truncated by DEBUG_PRINT_MAX_CHARS]"
    bar = "=" * 80
    print(
        f"\n{bar}\n[DEBUG structured caption {stem}] {study}/{source_file}\n{bar}\n{shown}\n",
        file=sys.stderr,
    )

def _object_resolution_seed(
    *,
    doc: Dict[str, Any],
    stream: List[StreamObject],
    stream_index: int,
    table_counter_hint: int,
) -> Optional[Dict[str, Any]]:
    so = stream[stream_index]
    if so.object_type not in {"table", "picture"}:
        return None
    if so.object_type == "table" and _is_empty_layout_table(doc, so.raw):
        return None
    object_type = "table" if so.object_type == "table" else "figure"
    linked_caption_raw = _caption_text_with_adjacent_identifier(doc, stream, stream_index)
    object_id = _valid_caption_object_id(linked_caption_raw, object_type)
    linked_caption = linked_caption_raw if object_id else ""
    evidence = _local_evidence_packet(doc=doc, stream=stream, stream_index=stream_index)

    footnote = ""
    if so.object_type == "table":
        footnote, _ = _linked_text(doc, so.raw, "footnotes")
        if _table_footnote_is_clearly_invalid(footnote):
            footnote = ""
        if not footnote:
            footnote, _ = _collect_following_table_footnote(stream, stream_index)
    missing_fields: List[str] = []
    if so.object_type == "table":
        if not object_id:
            missing_fields.extend(["table_id", "caption_text"])
        if not footnote and evidence.get("leading_table_footnote_candidates"):
            missing_fields.append("table_footnote")
    else:
        if not object_id:
            missing_fields.extend(["figure_id", "caption_text"])

    # Do not call regex/LLM recovery for fully resolved objects with no nearby
    # missing field evidence.
    if not missing_fields:
        return None

    return {
        "object_ref": so.object_ref,
        "object_refs": [so.object_ref],
        "object_type": "table" if so.object_type == "table" else "figure",
        "layout": "1",
        "missing_fields": missing_fields,
        "deterministic_caption": linked_caption or None,
        **evidence,
    }

def _seed_has_caption_label_evidence(seed: Dict[str, Any], object_type: str) -> bool:
    candidates = [c for c in (seed.get("caption_candidates") or []) if isinstance(c, dict)]
    if object_type == "table":
        if any(str(c.get("source") or "") != "after" for c in candidates):
            return True
    elif candidates:
        return True
    deterministic_caption = str(seed.get("deterministic_caption") or "").strip()
    if deterministic_caption and _caption_candidate_from_text(deterministic_caption, object_type):
        return True

    return False

def _picture_seed_has_caption_label_evidence(seed: Dict[str, Any]) -> bool:
    return _seed_has_caption_label_evidence(seed, "figure")


def _seed_missing_caption(seed: Dict[str, Any]) -> bool:
    missing = set(seed.get("missing_fields") or [])
    return bool({"table_id", "figure_id", "object_id", "caption_text"} & missing)


def _seed_allowed_for_llm(
    seed: Dict[str, Any],
    deterministic_resolution: Optional[Dict[str, Any]] = None,
) -> bool:
    """Gate LLM calls so it resolves anchors instead of inventing captions.

    Captions require a numbered caption-label anchor. Figures without such an
    anchor are discarded. Tables without such an anchor are not sent for caption
    recovery; footnote-only LLM recovery is allowed only when the table already
    has a valid caption/id from Docling or deterministic recovery.
    """
    object_type = str(seed.get("object_type") or "")
    if not _seed_missing_caption(seed):
        return True
    if _seed_has_caption_label_evidence(seed, object_type):
        return True

    # No caption anchor nearby. Do not ask the LLM to invent captions.
    # For tables this also suppresses footnote-only recovery because there is no
    # reliable logical table identity to attach the note to.
    return False

def _bbox_edges(bbox: Optional[Dict[str, Any]]) -> Optional[Tuple[float, float, float, float]]:
    if not bbox:
        return None
    try:
        l = float(bbox.get("l"))
        r = float(bbox.get("r"))
        t = float(bbox.get("t"))
        b = float(bbox.get("b"))
    except Exception:
        return None
    if r <= l or t <= b:
        return None
    return l, r, t, b


def _overlap_ratio_1d(a0: float, a1: float, b0: float, b1: float) -> float:
    overlap = max(0.0, min(a1, b1) - max(a0, b0))
    denom = max(1e-6, min(abs(a1 - a0), abs(b1 - b0)))
    return overlap / denom


def _picture_layout_relation(a: StreamObject, b: StreamObject) -> str:
    """Return side_by_side/stacked/empty for physically adjacent pictures."""
    if a.page_no is None or b.page_no is None or a.page_no != b.page_no:
        return ""
    ae = _bbox_edges(a.bbox)
    be = _bbox_edges(b.bbox)
    if not ae or not be:
        return ""
    al, ar, at, ab = ae
    bl, br, bt, bb = be
    aw, ah = ar - al, at - ab
    bw, bh = br - bl, bt - bb
    if min(aw, ah, bw, bh) <= 8:
        return ""
    area_ratio = min(aw * ah, bw * bh) / max(aw * ah, bw * bh, 1e-6)
    if area_ratio < 0.10:
        return ""

    vertical_overlap = _overlap_ratio_1d(ab, at, bb, bt)
    horizontal_gap = max(0.0, max(al, bl) - min(ar, br))
    if vertical_overlap >= 0.60 and horizontal_gap <= max(24.0, 0.35 * max(aw, bw)):
        return "side_by_side"

    horizontal_overlap = _overlap_ratio_1d(al, ar, bl, br)
    vertical_gap = max(0.0, max(ab, bb) - min(at, bt))
    if horizontal_overlap >= 0.60 and vertical_gap <= max(24.0, 0.35 * max(ah, bh)):
        return "stacked"
    return ""


def _layout_from_stream_indices(stream: List[StreamObject], indices: List[int]) -> str:
    if len(indices) <= 1:
        return "1"
    objs = [stream[i] for i in indices]
    edges = [_bbox_edges(obj.bbox) for obj in objs]
    edges = [e for e in edges if e]
    if len(edges) != len(objs):
        return f"{len(indices)}*1"

    side_by_side_pairs = 0
    stacked_pairs = 0
    for i in range(len(objs)):
        for j in range(i + 1, len(objs)):
            rel = _picture_layout_relation(objs[i], objs[j])
            if rel == "side_by_side":
                side_by_side_pairs += 1
            elif rel == "stacked":
                stacked_pairs += 1
    if side_by_side_pairs >= stacked_pairs and side_by_side_pairs > 0:
        return f"1*{len(indices)}"
    return f"{len(indices)}*1"

def _candidate_picture_stream_indices(stream: List[StreamObject]) -> List[int]:
    return [
        i
        for i, so in enumerate(stream)
        if so.object_type == "picture" and _picture_is_candidate_asset(so.raw)
    ]

def _stream_between_allows_picture_group(stream: List[StreamObject], start: int, end: int) -> bool:
    """Fallback grouping signal when Docling has no picture bboxes.

    Allow short, panel-like labels and tiny noncandidate picture fragments
    between adjacent picture objects. The fragments can be separators/rules
    created by DOCX conversion; they connect groups but are not output as
    figure assets.
    """
    if end <= start or end - start > 8:
        return False
    for so in stream[start + 1 : end]:
        if so.object_type == "picture" and not _picture_is_candidate_asset(so.raw):
            continue
        if so.object_type != "text":
            return False
        if so.content_layer == "furniture" or so.label in {"page_header", "page_footer"}:
            continue
        text = _clean_text(so.text)
        if not text:
            continue
        if _caption_candidate_from_text(text, "figure") or _caption_candidate_from_text(text, "table"):
            return False
        if len(text) > 80 or re.search(r"[.;:]$", text):
            return False
        if len(text.split()) > 6:
            return False
    return True


def _picture_stream_relation(stream: List[StreamObject], i: int, j: int) -> str:
    if j <= i:
        return ""
    if _bbox_edges(stream[i].bbox) and _bbox_edges(stream[j].bbox):
        return ""
    if _stream_between_allows_picture_group(stream, i, j):
        return "stream_adjacent"
    return ""


def _picture_layout_groups(stream: List[StreamObject]) -> List[List[int]]:
    """Group physically adjacent picture objects on the same page."""
    picture_indices = _candidate_picture_stream_indices(stream)
    parent = {i: i for i in picture_indices}

    def find(x: int) -> int:
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(a: int, b: int) -> None:
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[rb] = ra

    for pos, i in enumerate(picture_indices):
        for j in picture_indices[pos + 1:]:
            if stream[i].page_no is not None and stream[j].page_no is not None and stream[i].page_no != stream[j].page_no:
                continue
            if _picture_layout_relation(stream[i], stream[j]):
                union(i, j)
                continue
            if pos + 1 < len(picture_indices) and j == picture_indices[pos + 1] and _picture_stream_relation(stream, i, j):
                union(i, j)

    groups_by_root: Dict[int, List[int]] = {}
    for i in picture_indices:
        groups_by_root.setdefault(find(i), []).append(i)
    return [sorted(group) for group in groups_by_root.values() if len(group) > 1]


def _direct_caption_resolution_for_group(
    *,
    doc: Dict[str, Any],
    stream: List[StreamObject],
    indices: List[int],
    source: str,
    footnote_text: Optional[str] = None,
) -> List[Dict[str, Any]]:
    captions: List[Tuple[str, str]] = []
    for idx in indices:
        object_type = "table" if stream[idx].object_type == "table" else "figure"
        caption = _caption_text_with_adjacent_identifier(doc, stream, idx)
        if object_type == "table" and not _valid_caption_object_id(caption, "table"):
            try:
                embedded_caption, _ = _promote_embedded_table_caption(
                    _table_body_to_markdown(doc, stream[idx].raw)
                )
            except Exception:
                embedded_caption = ""
            if _valid_caption_object_id(embedded_caption, "table"):
                caption = embedded_caption
        object_id = _valid_caption_object_id(caption, object_type)
        if caption and object_id:
            key = (object_id, caption)
            if key not in captions:
                captions.append(key)
    refs = [stream[idx].object_ref for idx in indices]
    object_type = "picture" if all(ref.startswith("#/pictures/") for ref in refs) else "table"

    if object_type == "table":
        # Continuation pages often repeat the same table identifier while
        # appending "(continued)" to the printed caption.  The identifier is
        # the stable identity here; requiring the complete caption string to
        # match incorrectly turns an already-captioned table group into an
        # unresolved group and lets nearby captions leak into it.
        table_ids = {object_id for object_id, _caption in captions}
        if len(table_ids) != 1:
            return []
        object_id = next(iter(table_ids))
        caption = next(
            (
                candidate_caption
                for candidate_id, candidate_caption in captions
                if candidate_id == object_id
                and not re.search(r"\(\s*continued\s*\)\s*\.?\s*$", candidate_caption, re.I)
            ),
            captions[0][1],
        )
    else:
        if len(captions) != 1:
            return []
        object_id, caption = captions[0]

    return [
        {
            "object_ref": ref,
            "object_refs": refs,
            "object_type": object_type,
            "status": "resolved",
            "object_id": object_id,
            "caption_text": caption,
            "footnote_text": footnote_text or None,
            "source": source,
        }
        for ref in refs
    ]

def _all_candidate_picture_groups(stream: List[StreamObject]) -> List[List[int]]:
    grouped = _picture_layout_groups(stream)
    grouped_indices = {idx for group in grouped for idx in group}
    singletons = [[idx] for idx in _candidate_picture_stream_indices(stream) if idx not in grouped_indices]
    return sorted([*grouped, *singletons], key=lambda group: (min(group), max(group)))


def _caption_candidates_at_stream_index(stream: List[StreamObject], idx: int, object_type: str) -> List[Dict[str, Any]]:
    so = stream[idx]
    if not _usable_local_text_object(so):
        return []
    lines = [_clean_text(so.text)]
    # Allow exactly one adjacent text object for identifier/title splits such as
    # "Fig. S2" + ". SEM Micrographs ...". This mirrors the existing caption
    # continuation rule and avoids swallowing later body prose.
    if idx + 1 < len(stream) and _usable_local_text_object(stream[idx + 1]):
        if so.page_no is None or stream[idx + 1].page_no is None or so.page_no == stream[idx + 1].page_no:
            lines.append(_clean_text(stream[idx + 1].text))
    out: List[Dict[str, Any]] = []
    for candidate in _find_caption_candidates("\n".join(lines), object_type):
        item = dict(candidate)
        item["stream_index"] = idx
        out.append(item)
    return out


def _immediate_after_caption_candidates(
    stream: List[StreamObject],
    stream_index: int,
    object_type: str,
) -> List[Dict[str, Any]]:
    """Return after-side caption evidence only when no prose/object intervenes."""
    for idx in range(stream_index + 1, len(stream)):
        so = stream[idx]
        if so.object_type == "text":
            if not _usable_local_text_object(so):
                continue
            candidates = _caption_candidates_at_stream_index(stream, idx, object_type)
            if not candidates:
                return []

            # A table caption immediately followed by another table is normally
            # an above-caption for that next table, not a below-caption for the
            # previous table.
            if object_type == "table":
                for nxt in range(idx + 1, len(stream)):
                    next_so = stream[nxt]
                    if next_so.object_type == "text":
                        if not _usable_local_text_object(next_so):
                            continue
                        break
                    if next_so.object_type == "table":
                        return []
                    if next_so.object_type == "picture":
                        break

            out: List[Dict[str, Any]] = []
            for candidate in candidates:
                item = dict(candidate)
                item["source"] = "after_immediate"
                out.append(item)
            return out

        if so.object_type in {"table", "picture"}:
            return []

    return []


def _caption_anchor_candidates(stream: List[StreamObject], object_type: str) -> List[Dict[str, Any]]:
    candidates: List[Dict[str, Any]] = []
    seen: set[Tuple[int, str, str]] = set()
    for idx, so in enumerate(stream):
        if so.object_type != "text":
            continue
        for candidate in _caption_candidates_at_stream_index(stream, idx, object_type):
            key = (idx, str(candidate.get("object_id") or ""), str(candidate.get("text") or ""))
            if key not in seen:
                seen.add(key)
                candidates.append(candidate)
    return candidates


def _intervening_text_allows_figure_caption_assignment(stream: List[StreamObject], start: int, end: int) -> bool:
    if end < start:
        return False
    if end - start > FIGURE_CAPTION_ASSIGNMENT_MAX_STREAM_GAP:
        return False
    for so in stream[start + 1 : end]:
        if so.object_type != "text":
            return False
        if so.content_layer == "furniture" or so.label in {"page_header", "page_footer"}:
            continue
        text = _clean_text(so.text)
        if not text:
            continue
        if _caption_candidate_from_text(text, "figure") or _caption_candidate_from_text(text, "table"):
            return False
        # Permit short panel/legend labels between the picture group and caption,
        # but reject normal body prose.
        if len(text) > 80 or re.search(r"[.;:]$", text):
            return False
        if len(text.split()) > 6:
            return False
    return True


def _nearest_picture_group_for_caption(
    stream: List[StreamObject],
    groups: List[List[int]],
    caption_index: int,
) -> Optional[List[int]]:
    # Figure captions are usually below figures: prefer the nearest preceding
    # candidate group, then fall back to a following group for caption-above cases.
    preceding = [group for group in groups if max(group) < caption_index]
    preceding.sort(key=lambda group: caption_index - max(group))
    for group in preceding:
        if _intervening_text_allows_figure_caption_assignment(stream, max(group), caption_index):
            return group

    following = [group for group in groups if min(group) > caption_index]
    following.sort(key=lambda group: min(group) - caption_index)
    for group in following:
        if _intervening_text_allows_figure_caption_assignment(stream, caption_index, min(group)):
            return group
    return None


def _figure_caption_anchor_resolutions(doc: Dict[str, Any], stream: List[StreamObject]) -> List[Dict[str, Any]]:
    """Assign figure captions from caption anchors to adjacent picture groups.

    This is the JSON-era equivalent of the markdown-era caption-first behavior:
    picture groups are only candidates. They become figure outputs only after a
    Docling direct caption or nearby figure/scheme/exhibit caption attaches.
    """
    groups = _all_candidate_picture_groups(stream)
    if not groups:
        return []

    # Do not override Docling direct captions.
    direct_refs = {
        stream[idx].object_ref
        for group in groups
        for idx in group
        if _is_valid_caption_text(_caption_text_with_adjacent_identifier(doc, stream, idx), "figure")
    }
    assigned_refs: set[str] = set()
    resolutions: List[Dict[str, Any]] = []

    for candidate in _caption_anchor_candidates(stream, "figure"):
        caption_index = int(candidate.get("stream_index") or -1)
        if caption_index < 0:
            continue
        group = _nearest_picture_group_for_caption(stream, groups, caption_index)
        if not group:
            continue
        refs = [stream[idx].object_ref for idx in group]
        if any(ref in direct_refs or ref in assigned_refs for ref in refs):
            continue
        object_id = str(candidate.get("object_id") or "").strip()
        caption = str(candidate.get("text") or "").strip()
        if not object_id or not caption or _caption_candidate_is_identifier_only(candidate):
            continue
        for ref in refs:
            resolutions.append(
                {
                    "object_ref": ref,
                    "object_refs": refs,
                    "object_type": "picture",
                    "status": "resolved",
                    "object_id": object_id,
                    "caption_text": caption,
                    "footnote_text": None,
                    "source": "structured_caption_anchor_regex",
                }
            )
        assigned_refs.update(refs)
    return resolutions

def _picture_group_caption_resolutions(doc: Dict[str, Any], stream: List[StreamObject]) -> List[Dict[str, Any]]:
    resolutions: List[Dict[str, Any]] = []
    for group in _picture_layout_groups(stream):
        resolutions.extend(
            _direct_caption_resolution_for_group(
                doc=doc,
                stream=stream,
                indices=group,
                source="docling_picture_group",
           )
        )
    return resolutions


def _picture_is_inside_table_region(
    picture: StreamObject,
    table: StreamObject,
) -> bool:
    """Return whether a picture is physically embedded in a table region."""
    if picture.page_no is not None and table.page_no is not None:
        if picture.page_no != table.page_no:
            return False
    picture_edges = _bbox_edges(picture.bbox)
    table_edges = _bbox_edges(table.bbox)
    if not picture_edges or not table_edges:
        return False
    picture_left, picture_right, picture_top, picture_bottom = picture_edges
    table_left, table_right, table_top, table_bottom = table_edges
    horizontal_coverage = _overlap_ratio_1d(
        picture_left,
        picture_right,
        table_left,
        table_right,
    )
    vertical_coverage = _overlap_ratio_1d(
        picture_bottom,
        picture_top,
        table_bottom,
        table_top,
    )
    return horizontal_coverage >= 0.80 and vertical_coverage >= 0.80


def _intervening_text_is_table_leak(stream: List[StreamObject], start: int, end: int) -> bool:
    for so in stream[start + 1 : end]:
        if so.object_type == "picture":
            if _picture_is_candidate_asset(so.raw):
                if (
                    _picture_is_inside_table_region(so, stream[start])
                    or _picture_is_inside_table_region(so, stream[end])
                ):
                    continue
                return False
            continue
        if so.object_type != "text":
            return False
        if so.content_layer == "furniture" or so.label in {"page_header", "page_footer"}:
            continue
        text = _clean_text(so.text)
        if not text:
            continue
        if _caption_candidate_from_text(text, "table"):
            return False
        if _looks_like_table_footnote(text):
            return False
        if not _looks_like_orphan_table_cell_text(text):
            return False
    return True


def _intervening_real_figure_is_table_boundary(
    doc: Dict[str, Any],
    stream: List[StreamObject],
    start: int,
    end: int,
) -> bool:
    """Return whether a captioned figure separates two table objects.

    A real figure with its own caption is a hard document boundary for table
    continuation.  This deliberately does not treat every picture as a
    boundary: uncaptioned fragments can be part of a compound figure or a
    table-layout artifact.
    """
    for picture_index in range(start + 1, end):
        so = stream[picture_index]
        if so.object_type != "picture":
            continue
        caption = _caption_text_with_adjacent_identifier(
            doc,
            stream,
            picture_index,
        )
        if _is_valid_caption_text(caption, "figure"):
            return True
    return False


def _table_column_count(table_obj: Dict[str, Any]) -> int:
    grid = _docling_table_grid(table_obj)
    return max((len(row) for row in grid), default=0)

def _table_fragment_groups_with_doc(doc: Dict[str, Any], stream: List[StreamObject]) -> List[List[int]]:
    groups: List[List[int]] = []
    table_indices = [
        i
        for i, so in enumerate(stream)
        if so.object_type == "table" and not _is_empty_layout_table(doc, so.raw)
    ]
    used: set[int] = set()
    for idx in table_indices:
        if idx in used:
            continue
        group = [idx]
        base_cols = _table_column_count(stream[idx].raw)
        if base_cols <= 0:
            continue
        first_object_id = _resolved_table_id_for_index(doc, stream, idx)
        prev = idx
        for nxt in table_indices:
            if nxt <= idx or nxt in used:
                continue
            next_object_id = _resolved_table_id_for_index(doc, stream, nxt)
            if next_object_id and next_object_id != first_object_id:
                break
            if _intervening_real_figure_is_table_boundary(doc, stream, prev, nxt):
                break
            if not _intervening_text_is_table_leak(stream, prev, nxt):
                break
            next_cols = _table_column_count(stream[nxt].raw)
            if next_cols != base_cols and not first_object_id.lower().startswith("exhibit "):
                break
            group.append(nxt)
            prev = nxt
        if len(group) > 1:
            used.update(group)
            groups.append(group)
    return groups


def _table_caption_boundary_status(
    doc: Dict[str, Any],
    stream: List[StreamObject],
    idx: int,
    *,
    caption_resolutions: Optional[Dict[str, Dict[str, Any]]] = None,
    pdf_path: Optional[Path] = None,
    use_camelot: bool = False,
) -> str:
    """Return whether this table appears to start its own captioned table.

    ``captionless_checked`` means we tried the available embedded-caption
    detectors and did not find a valid table caption. This lets continuation
    grouping relax column-shape requirements without treating unchecked hidden
    captions as safe boundaries.
    """
    split_caption_marker = _table_has_embedded_caption_marker(doc, stream[idx].raw)
    if split_caption_marker:
        # A cached/regex caption may have been inherited from the preceding
        # table. Check the PDF evidence first so that such a label cannot hide
        # a new caption embedded in this table's header row.
        candidate_caption, _payload = _camelot_embedded_caption_recovery(
            pdf_path=pdf_path,
            table_obj=stream[idx].raw,
        )
        if candidate_caption:
            return "has_embedded_caption"
        # Keep a likely new table separate even when Camelot is unavailable or
        # cannot recover the text. The processing stage will retain it as an
        # unresolved/uncaptioned table instead of silently calling it S7.
        return "embedded_caption_unresolved"

    resolution = (caption_resolutions or {}).get(stream[idx].object_ref) or {}
    refs = [str(ref or "") for ref in (resolution.get("object_refs") or []) if ref]
    resolution_is_single_object = not refs or refs == [stream[idx].object_ref]
    if (
        resolution_is_single_object
        and str(resolution.get("status") or "").strip().lower() == "resolved"
    ):
        object_id = str(resolution.get("object_id") or "")
        caption_text = str(resolution.get("caption_text") or "")
        if _valid_caption_object_id(object_id, "table") or _valid_caption_object_id(caption_text, "table"):
            return "has_caption"

    caption = _caption_text_with_adjacent_identifier(doc, stream, idx)
    if _valid_caption_object_id(caption, "table"):
        return "has_caption"
    try:
        embedded_caption, _ = _promote_embedded_table_caption(
            _table_body_to_markdown(doc, stream[idx].raw)
        )
    except Exception:
        embedded_caption = ""
    if _valid_caption_object_id(embedded_caption, "table"):
        return "has_caption"

    if use_camelot and pdf_path and extract_camelot_embedded_caption is not None:
        try:
            payload = extract_camelot_embedded_caption(
                pdf_path=pdf_path,
                table_obj=stream[idx].raw,
                bbox_pad=TABLE_REPAIR_BBOX_PAD,
                min_nonempty_cells=TABLE_REPAIR_MIN_NONEMPTY_CELLS,
            )
            candidate_caption = _clean_text(str(payload.get("caption_candidate") or ""))
        except Exception:
            return "captionless_unchecked"
        if _valid_caption_object_id(candidate_caption, "table"):
            return "has_caption"
        if str(payload.get("status") or "").strip().lower() in {"extracted", "not_found"}:
            return "captionless_checked"
        return "captionless_unchecked"

    return "captionless_unchecked"




def _resolved_table_id_for_index(
    doc: Dict[str, Any],
    stream: List[StreamObject],
    idx: int,
    caption_resolutions: Optional[Dict[str, Dict[str, Any]]] = None,
) -> str:
    resolution = (caption_resolutions or {}).get(stream[idx].object_ref) or {}
    object_id = _normalize_resolution_object_id(resolution.get("object_id"))
    if _valid_caption_object_id(object_id or "", "table"):
        return object_id or ""
    caption_text = _clean_text(str(resolution.get("caption_text") or ""))
    object_id = _valid_caption_object_id(caption_text, "table")
    if object_id:
        return object_id
    caption = _caption_text_with_adjacent_identifier(doc, stream, idx)
    object_id = _valid_caption_object_id(caption, "table")
    if object_id:
        return object_id
    try:
        embedded_caption, _ = _promote_embedded_table_caption(
            _table_body_to_markdown(doc, stream[idx].raw)
        )
    except Exception:
        embedded_caption = ""
    return _valid_caption_object_id(embedded_caption, "table")


def _table_has_own_caption_evidence(doc: Dict[str, Any], stream: List[StreamObject], idx: int) -> bool:
    return _table_caption_boundary_status(doc, stream, idx) == "has_caption"


def _continuation_column_evidence_is_strong(
    expected_cols: int,
    next_cols: int,
    group_len: int,
    *,
    group_col_counts: Optional[List[int]] = None,
    same_logical_table: bool = False,
    page_numbers_missing: bool = False,
) -> bool:
    if expected_cols <= 0 or next_cols <= 0:
        return False
    if min(expected_cols, next_cols) < 2:
        return False
    if next_cols == expected_cols:
        return True
    if same_logical_table and group_col_counts:
        observed_cols = [cols for cols in group_col_counts if cols > 0]
        if observed_cols:
            widest_observed = max(observed_cols)
            if min(widest_observed, next_cols) >= 5 and abs(next_cols - widest_observed) == 1:
                return True
    if page_numbers_missing and min(expected_cols, next_cols) >= 3:
        return True
    # Some first pages omit a leading class/index column that appears on later
    # continuation pages, as in study_51 Table S3.
    return group_len == 1 and abs(next_cols - expected_cols) == 1 and min(expected_cols, next_cols) >= 5


def _table_continuation_output_groups(
    doc: Dict[str, Any],
    stream: List[StreamObject],
    *,
    study_dir: Optional[Path] = None,
    source_file: str = "",
    caption_resolutions: Optional[Dict[str, Dict[str, Any]]] = None,
) -> List[List[int]]:
    """Group Docling table objects that form one printed table in reading order.

    Page boundaries are weak layout signals, not table boundaries. Structural
    checks below determine whether a later object is a continuation.
    """
    groups: List[List[int]] = []
    table_indices = [
        i
        for i, so in enumerate(stream)
        if so.object_type == "table" and not _is_empty_layout_table(doc, so.raw)
    ]
    used: set[int] = set()
    pdf_path = _source_pdf_path(study_dir, source_file, doc) if study_dir and source_file else None
    caption_boundary_cache: Dict[int, str] = {}

    def caption_boundary_status(idx: int) -> str:
        if idx not in caption_boundary_cache:
            caption_boundary_cache[idx] = _table_caption_boundary_status(
                doc,
                stream,
                idx,
                caption_resolutions=caption_resolutions,
                pdf_path=pdf_path,
                use_camelot=bool(pdf_path),
            )
        return caption_boundary_cache[idx]

    for idx in table_indices:
        if idx in used:
            continue
        group = [idx]
        expected_cols = _table_column_count(stream[idx].raw)
        if expected_cols <= 0:
            continue
        group_col_counts = [expected_cols]
        prev = idx
        for nxt in table_indices:
            if nxt <= idx or nxt in used:
                continue
            prev_page = stream[prev].page_no
            next_page = stream[nxt].page_no
            boundary_status = caption_boundary_status(nxt)
            group_id = _resolved_table_id_for_index(
                doc,
                stream,
                group[0],
                caption_resolutions=caption_resolutions,
            )
            next_id = _resolved_table_id_for_index(
                doc,
                stream,
                nxt,
                caption_resolutions=caption_resolutions,
            )
            same_logical_table = bool(group_id and next_id and group_id == next_id)
            if boundary_status in {"has_embedded_caption", "embedded_caption_unresolved"}:
                break
            if boundary_status == "has_caption":
                if not same_logical_table:
                    break
            if _intervening_real_figure_is_table_boundary(doc, stream, prev, nxt):
                break
            if not _intervening_text_is_table_leak(stream, prev, nxt):
                break
            next_cols = _table_column_count(stream[nxt].raw)
            strong_column_evidence = _continuation_column_evidence_is_strong(
                expected_cols,
                next_cols,
                len(group),
                group_col_counts=group_col_counts,
                same_logical_table=same_logical_table,
                page_numbers_missing=prev_page is None and next_page is None,
            )
            checked_captionless_boundary = caption_boundary_status(nxt) == "captionless_checked"
            if not strong_column_evidence and not checked_captionless_boundary:
                break
            group.append(nxt)
            group_col_counts.append(next_cols)
            expected_cols = max(expected_cols, next_cols)
            prev = nxt
        if len(group) > 1:
            used.update(group)
            groups.append(group)
    return groups


def _table_group_caption_resolutions(doc: Dict[str, Any], stream: List[StreamObject]) -> List[Dict[str, Any]]:
    """Conservatively propagate a caption across Docling-split table fragments.

    This does not attempt to repair leaked table cells. It only assigns the same
    printed caption to nearby captionless fragments when structural signals are
    compatible.
    """
    resolutions: List[Dict[str, Any]] = []
    for group in _table_fragment_groups_with_doc(doc, stream):
        group_footnote, _ = _linked_text(doc, stream[group[-1]].raw, "footnotes")
        if not group_footnote:
            group_footnote, _ = _collect_following_table_footnote(stream, group[-1])
        resolutions.extend(
            _direct_caption_resolution_for_group(
                doc=doc,
                stream=stream,
                indices=group,
                source="docling_table_group",
                footnote_text=group_footnote,
            )
        )
    return resolutions


def _deterministic_group_caption_resolutions(doc: Dict[str, Any], stream: List[StreamObject]) -> List[Dict[str, Any]]:
    return [
        *_picture_group_caption_resolutions(doc, stream),
        *_figure_caption_anchor_resolutions(doc, stream),
        *_table_group_caption_resolutions(doc, stream),
    ]

def _table_group_resolution_seed(
    *,
    doc: Dict[str, Any],
    stream: List[StreamObject],
    indices: List[int],
) -> Optional[Dict[str, Any]]:
    refs = [stream[idx].object_ref for idx in indices]
    if len(refs) <= 1:
        return None
    first_idx = min(indices)
    last_idx = max(indices)
    text_before = _char_bounded_local_text(
        doc,
        stream,
        first_idx,
        before=True,
        max_chars=STRUCTURED_CAPTION_TABLE_BEFORE_CHARS,
    )
    text_after = _char_bounded_local_text(
        doc,
        stream,
        last_idx,
        before=False,
        max_chars=STRUCTURED_CAPTION_TABLE_AFTER_CHARS,
    )
    if not text_before and not text_after:
        return None

    caption_candidates: List[Dict[str, Any]] = []
    for source_name, source_text in (("before", text_before), ("after", text_after)):
        for candidate in _find_caption_candidates(source_text, "table"):
            caption_candidates.append({"source": source_name, **candidate})

    table_previews: List[str] = []
    for n, idx in enumerate(indices, start=1):
        body = _table_body_to_markdown(doc, stream[idx].raw)
        preview = _table_markdown_preview(body)
        if preview:
            table_previews.append(f"[table fragment {n}]\n{preview}")

    leading_footnote_candidates = _leading_table_footnote_candidates(text_after)
    missing_fields = ["table_id", "caption_text"]
    if leading_footnote_candidates:
        missing_fields.append("table_footnote")

    return {
        "object_ref": refs[0],
        "object_refs": refs,
        "object_type": "table",
        "layout": f"{len(refs)}*1",
        "missing_fields": missing_fields,
        "text_before": text_before,
        "text_after": text_after,
        "caption_candidates": caption_candidates,
        "table_preview": "\n\n".join(table_previews),
        "leading_table_footnote_candidates": leading_footnote_candidates,
    }


def _seed_layout(seed: Dict[str, Any]) -> str:
    refs = list(seed.get("object_refs") or [])
    layout = str(seed.get("layout") or "").strip()
    if layout:
        return layout
    return "1" if len(refs) <= 1 else f"{len(refs)}*1"

def _llm_block_from_seed(seed: Dict[str, Any]) -> Dict[str, Any]:
    """Build the minimal prompt-facing unresolved block."""
    object_refs = [str(ref or "").strip() for ref in (seed.get("object_refs") or [])]
    if not object_refs:
        object_ref = str(seed.get("object_ref") or "").strip()
        object_refs = [object_ref] if object_ref else []
    block: Dict[str, Any] = {
        "object_refs": object_refs,
        "layout": _seed_layout({**seed, "object_refs": object_refs}),
        "text_before": str(seed.get("text_before") or ""),
        "text_after": str(seed.get("text_after") or ""),
    }
    if str(seed.get("object_type") or "") == "table":
        table_preview = str(seed.get("table_preview") or "").strip()
        if table_preview:
            block["table_preview"] = table_preview
    return block


def _llm_blocks_from_seeds(seeds: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    blocks = [_llm_block_from_seed(seed) for seed in seeds]
    return [block for block in blocks if block.get("object_refs")]


def _refs_from_llm_item(item: Dict[str, Any]) -> List[str]:
    raw_refs = item.get("object_refs")
    if raw_refs is None and item.get("object_ref") is not None:
        raw_refs = [item.get("object_ref")]
    if not isinstance(raw_refs, list):
        return []
    refs: List[str] = []
    for ref in raw_refs:
        ref_s = str(ref or "").strip()
        if ref_s and ref_s not in refs:
            refs.append(ref_s)
    return refs

def _grouped_object_id_for_ref(object_id: str, object_ref: str, resolution: Dict[str, Any], object_type: str) -> str:
    refs = [str(ref or "") for ref in (resolution.get("object_refs") or [])]
    if object_type != "table" or len(refs) <= 1 or object_ref not in refs:
        return object_id
    if re.search(r"_\d+$", object_id or ""):
        return object_id
    return f"{object_id}_{refs.index(object_ref) + 1}"

def _normalize_resolution_object_id(value: Any) -> Optional[str]:
    txt = str(value or "").strip()
    if not txt or txt.lower() == "null":
        return None
    m = re.match(
        r"^\s*(Table|Tab\.?|Tabel|Exhibit|Figure|Fig\.?|Scheme)\s*\.?\s*"
        r"((?:S|SI|SM|SF|ST|A)?\s*[.\-]?\s*\d+(?:\s*[A-Za-z]|[._\-]\d+)?|\d+\s*S|[IVXLCDM]+)\b",
        txt,
        flags=re.I,
    )
    if not m:
        return txt
    return _canonical_object_id(m.group(1), m.group(2))


def _caption_with_identifier(object_id: Optional[str], caption: Optional[str]) -> Optional[str]:
    caption_text = _clean_text(str(caption or ""))
    if not caption_text:
        return None
    normalized_id = _normalize_resolution_object_id(object_id)
    if not normalized_id:
        return caption_text
    if _object_id_from_caption(caption_text, fallback="") == normalized_id:
        return caption_text
    return _clean_text(f"{normalized_id} {caption_text}")


def _clean_structured_caption_llm_response(
    parsed: Any,
    requested_objects: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    """Normalize grouped object_refs responses into per-object resolutions."""
    if not isinstance(parsed, dict):
        return []
    requested_refs = {
        str(ref or "")
        for obj in requested_objects
        for ref in (obj.get("object_refs") or [])
    }
    clean: List[Dict[str, Any]] = []

    for item in parsed.get("figures") or []:
        if not isinstance(item, dict):
            continue
        refs = _refs_from_llm_item(item)
        refs = [ref for ref in refs if ref in requested_refs and ref.startswith("#/pictures/")]
        if not refs:
            continue
        object_id = _normalize_resolution_object_id(item.get("figure_id"))
        caption = _clean_text(str(item.get("caption_text") or "")) or None
        if not object_id and not caption:
            continue
        caption = _caption_with_identifier(object_id, caption)
        for ref in refs:
            clean.append(
                {
                    "object_ref": ref,
                    "object_refs": refs,
                    "object_type": "picture",
                    "status": "resolved",
                    "object_id": object_id,
                    "caption_text": caption,
                    "footnote_text": None,
                    "source": "structured_caption_llm",
                }
            )

    raw_table_items = [item for item in (parsed.get("tables") or []) if isinstance(item, dict)]
    table_key_to_ref_groups: Dict[Tuple[str, str], List[Tuple[str, ...]]] = {}
    for item in raw_table_items:
        refs = _refs_from_llm_item(item)
        refs = [ref for ref in refs if ref in requested_refs and ref.startswith("#/tables/")]
        object_id = _normalize_resolution_object_id(item.get("table_id")) or ""
        caption = _clean_text(str(item.get("caption_text") or ""))
        if refs and (object_id or caption):
            table_key_to_ref_groups.setdefault((object_id, caption.casefold()), []).append(tuple(refs))
    duplicate_caption_keys = {
        key
        for key, groups in table_key_to_ref_groups.items()
        if len({group for group in groups}) > 1
    }

    for item in raw_table_items:
        refs = _refs_from_llm_item(item)
        refs = [ref for ref in refs if ref in requested_refs and ref.startswith("#/tables/")]
        if not refs:
            continue
        object_id = _normalize_resolution_object_id(item.get("table_id"))
        caption = _clean_text(str(item.get("caption_text") or "")) or None
        footnote = _clean_text(str(item.get("table_footnote") or ""))
        duplicate_key = ((object_id or ""), _clean_text(str(item.get("caption_text") or "")).casefold())
        if duplicate_key in duplicate_caption_keys:
            continue
        if not object_id and not caption and not footnote:
            continue
        caption = _caption_with_identifier(object_id, caption)
        for ref in refs:
            clean.append(
                {
                    "object_ref": ref,
                    "object_refs": refs,
                    "object_type": "table",
                    "status": "resolved",
                    "object_id": object_id,
                    "caption_text": caption,
                    "footnote_text": footnote or None,
                    "source": "structured_caption_llm",
                }
            )
    return clean

def _resolution_covers_seed(seed: Dict[str, Any], resolution: Optional[Dict[str, Any]]) -> bool:
    if not resolution:
        return False
    missing = set(seed.get("missing_fields") or [])
    if {"table_id", "figure_id", "object_id", "caption_text"} & missing:
        if not resolution.get("object_id") or not resolution.get("caption_text"):
            return False
    if {"table_footnote", "footnote_text"} & missing:
        if not resolution.get("footnote_text"):
            return False
    return True


def _merge_caption_resolutions(*resolution_lists: Iterable[Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    merged: Dict[str, Dict[str, Any]] = {}
    for resolutions in resolution_lists:
        for r in resolutions or []:
            if not isinstance(r, dict) or not r.get("object_ref"):
                continue
            ref = str(r.get("object_ref"))
            if ref not in merged:
                merged[ref] = dict(r)
                continue
            dst = merged[ref]
            for key in ("object_id", "caption_text", "footnote_text", "object_refs"):
                if r.get(key):
                    dst[key] = r.get(key)
            if r.get("status"):
                dst["status"] = r.get("status")
            # LLM is higher-authority for ambiguous caption routing; regex is
            # retained if it was the only contributor.
            if r.get("source") == "structured_caption_llm" or not dst.get("source"):
                dst["source"] = r.get("source")
    return merged


def _resolve_structured_captions_with_cache(
    *,
    study: str,
    source_file: str,
    doc: Dict[str, Any],
    stream: List[StreamObject],
) -> Dict[str, Dict[str, Any]]:
    """Resolve missing/suspicious object metadata using deterministic evidence, then LLM."""
    deterministic_resolutions = _deterministic_group_caption_resolutions(doc, stream)
    deterministic_by_ref = {
        str(r.get("object_ref") or ""): r
        for r in deterministic_resolutions
        if isinstance(r, dict) and r.get("object_ref")
    }
    picture_groups = _picture_layout_groups(stream)
    picture_group_by_idx = {idx: tuple(group) for group in picture_groups for idx in group}
    processed_picture_groups: set[Tuple[int, ...]] = set()
    table_groups = _table_fragment_groups_with_doc(doc, stream)
    table_group_by_idx = {idx: tuple(group) for group in table_groups for idx in group}
    processed_table_groups: set[Tuple[int, ...]] = set()

    table_counter = 0
    llm_seeds: List[Dict[str, Any]] = []
    regex_resolutions: List[Dict[str, Any]] = [*deterministic_resolutions]
    def consider_seed(seed: Optional[Dict[str, Any]], deterministic_resolution: Optional[Dict[str, Any]] = None) -> None:
        if not seed:
            return
        regex_resolution = None if deterministic_resolution else _regex_resolution_from_seed(seed)
        if regex_resolution:
            regex_resolutions.extend(_expand_resolution_for_object_refs(regex_resolution))
        resolution = deterministic_resolution or regex_resolution
        if not _resolution_covers_seed(seed, resolution):
            if _seed_allowed_for_llm(seed, resolution):
                llm_seeds.append(seed)

    for idx, so in enumerate(stream):
        if so.object_type == "table" and _is_empty_layout_table(doc, so.raw):
            continue
        if so.object_type == "table":
            table_counter += 1
        if so.object_type not in {"table", "picture"}:
            continue

        picture_group = picture_group_by_idx.get(idx) if so.object_type == "picture" else None
        if picture_group and picture_group not in processed_picture_groups:
            processed_picture_groups.add(picture_group)
            group_deterministic = next(
                (
                    deterministic_by_ref.get(stream[i].object_ref)
                    for i in picture_group
                    if stream[i].object_ref in deterministic_by_ref
                ),
                None,
            )
            if not group_deterministic:
                first_idx = min(picture_group)
                seed = _object_resolution_seed(
                    doc=doc,
                    stream=stream,
                    stream_index=first_idx,
                    table_counter_hint=table_counter,
                )
                if seed:
                    seed["object_refs"] = [stream[i].object_ref for i in picture_group]
                    seed["layout"] = _layout_from_stream_indices(stream, list(picture_group))
                    if _picture_seed_has_caption_label_evidence(seed):
                        consider_seed(seed)
            continue
        if picture_group:
            continue
        if so.object_type == "picture":
            if not _picture_is_candidate_asset(so.raw):
                continue
            seed = _object_resolution_seed(
                doc=doc,
                stream=stream,
                stream_index=idx,
                table_counter_hint=table_counter,
            )
            if seed and _picture_seed_has_caption_label_evidence(seed):
                consider_seed(seed, deterministic_by_ref.get(so.object_ref))
            continue

        table_group = table_group_by_idx.get(idx) if so.object_type == "table" else None
        if table_group and table_group not in processed_table_groups:
            processed_table_groups.add(table_group)
            group_deterministic = next(
                (deterministic_by_ref.get(stream[i].object_ref) for i in table_group if stream[i].object_ref in deterministic_by_ref),
                None,
            )
            if not group_deterministic:
                group_seed = _table_group_resolution_seed(
                    doc=doc,
                    stream=stream,
                    indices=list(table_group),
                )
                consider_seed(group_seed)
            continue
        if table_group:
            continue

        seed = _object_resolution_seed(
            doc=doc,
            stream=stream,
            stream_index=idx,
            table_counter_hint=table_counter,
        )
        if not seed:
            continue
        consider_seed(seed, deterministic_by_ref.get(so.object_ref))

    if not llm_seeds:
        return _merge_caption_resolutions(regex_resolutions)

    llm_objects = _llm_requests_from_seeds(llm_seeds)
    if not llm_objects:
        return _merge_caption_resolutions(regex_resolutions)

    regex_context = _merge_caption_resolutions(regex_resolutions)
    snippets_text = _marked_snippets_from_seeds(
        doc=doc,
        stream=stream,
        seeds=llm_seeds,
        context_resolutions=regex_context,
    )
    if not snippets_text:
        return _merge_caption_resolutions(regex_resolutions)
    meta = {
        "schema_version": STRUCTURED_CAPTION_SCHEMA_VERSION,
        "provider": LLM_PROVIDER,
        "model_name": LLM_MODEL_NAME,
        "snippets_sha256": _sha256_text(snippets_text),
        "requested_objects_sha256": _sha256_text(json.dumps(llm_objects, ensure_ascii=False, sort_keys=True)),
        "regex_resolutions_sha256": _sha256_text(json.dumps(regex_resolutions, ensure_ascii=False, sort_keys=True)),
    }
    path = _structured_caption_cache_path(study, source_file)
    if not FORCE_RERUN_LLM and path.exists():
        try:
            cached = _load_json(path)
            if cached.get("meta") == meta:
                cached_resolutions = [
                    r for r in (cached.get("resolutions") or [])
                    if isinstance(r, dict) and r.get("object_ref")
                ]
                return _merge_caption_resolutions(regex_resolutions, cached_resolutions)
        except Exception:
            pass

    if predict_with_usage is None:
        payload = {
            "meta": meta,
            "resolutions": [],
            "regex_resolutions": regex_resolutions,
            "error": "predict_with_usage unavailable",
        }
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        return _merge_caption_resolutions(regex_resolutions)

    try:
        chain = _get_structured_caption_chain()
        try:
            rendered_prompt = chain.prompt.format(
                study_id=study,
                source_file=source_file,
                snippets_text=snippets_text,
            )
        except Exception as exc:
            rendered_prompt = (
                f"[DEBUG] Failed to render structured caption prompt: {exc}\n\n"
                f"snippets_text:\n{snippets_text}"
            )
        _debug_print_and_write(study, source_file, "snippets", snippets_text)
        _debug_print_and_write(study, source_file, "prompt", rendered_prompt)

        call = predict_with_usage(
            chain,
            provider=LLM_PROVIDER,
            model_name=LLM_MODEL_NAME,
            study_id=study,
            source_file=source_file,
            snippets_text=snippets_text,
        )
        _debug_print_and_write(study, source_file, "raw_response", call.text)

        parsed = _safe_json_loads(call.text)
        clean_resolutions = _clean_structured_caption_llm_response(parsed, llm_objects)
        payload = {
            "meta": meta,
            "study": study,
            "source_file": source_file,
            "resolutions": clean_resolutions,
            "regex_resolutions": regex_resolutions,
            "requested_objects": llm_objects,
            "prompt_snippets": snippets_text,
            "llm_call": call.usage.metadata_dict(),
            "generated_at": datetime.now().isoformat(timespec="seconds"),
        }
        if DEBUG_MODE:
            payload["debug_rendered_prompt"] = rendered_prompt
            payload["debug_raw_response"] = call.text
            payload["debug_parsed_response"] = parsed
        print(
            f"[LLM] {study}/{source_file} structured caption resolution: "
            f"objects={len(llm_objects)} regex_resolutions={len(regex_resolutions)} "
            f"llm_resolutions={len(clean_resolutions)} "
            f"prompt={call.usage.prompt_tokens} completion={call.usage.completion_tokens} total={call.usage.total_tokens}",
            file=sys.stderr,
        )
    except Exception as exc:
        payload = {
            "meta": meta,
            "study": study,
            "source_file": source_file,
            "resolutions": [],
            "regex_resolutions": regex_resolutions,
            "requested_objects": llm_objects,
            "prompt_snippets": snippets_text,
            "error": f"structured caption resolution failed: {exc}",
            "generated_at": datetime.now().isoformat(timespec="seconds"),
        }
        print(f"⚠️ Structured caption resolution failed → {study}/{source_file}: {exc}", file=sys.stderr)

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return _merge_caption_resolutions(regex_resolutions, payload.get("resolutions", []))

def _apply_caption_resolution(
    *,
    object_ref: str,
    object_type: str,
    deterministic_object_id: str,
    deterministic_caption: str,
    deterministic_footnote: str = "",
    resolutions: Dict[str, Dict[str, Any]],
) -> Tuple[str, str, str, str]:
    """Return (object_id, caption_text, footnote_text, metadata_source)."""
    r = resolutions.get(object_ref) or {}
    status = str(r.get("status") or "").strip().lower()
    if status == "resolved":
        object_id = _normalize_resolution_object_id(r.get("object_id")) or deterministic_object_id
        object_id = _grouped_object_id_for_ref(object_id, object_ref, r, object_type)
        caption = _clean_text(str(r.get("caption_text") or "")) or deterministic_caption
        footnote = _clean_text(str(r.get("footnote_text") or "")) or deterministic_footnote
        metadata_source = str(r.get("source") or "structured_caption_llm")
        return object_id, caption, footnote, metadata_source
    return deterministic_object_id, deterministic_caption, deterministic_footnote, "docling_json"


# Footnote / caption helpers --------------------------------------------------
def _looks_like_table_footnote(text: str) -> bool:
    return _is_table_footnote_candidate_line(text)


def _table_footnote_is_clearly_invalid(text: str) -> bool:
    s = _clean_text(text)
    if not s:
        return False
    if _object_id_from_caption(s, fallback=""):
        return True
    if re.match(r"^(?:references?|bibliography|acknowledg(?:e)?ments?|appendix|supplementary\s+materials?)\b", s, re.I):
        return True
    if re.match(r"^https?://|^doi\s*:", s, re.I):
        return True
    return False

def _collect_following_table_footnote(stream: List[StreamObject], table_index: int) -> Tuple[str, List[str]]:
    table_obj = stream[table_index]
    parts: List[str] = []
    refs: List[str] = []
    for nxt in stream[table_index + 1 : table_index + 5]:
        if nxt.object_type != "text":
            break
        if nxt.page_no is not None and table_obj.page_no is not None and nxt.page_no != table_obj.page_no:
            break
        if nxt.label == "footnote" or _looks_like_table_footnote(nxt.text):
            parts.append(nxt.text)
            refs.append(nxt.object_ref)
            continue
        break
    return "\n".join(parts).strip(), refs


# Main prose sectioning -------------------------------------------------------
_SECTION_END_RE = re.compile(
    r"^(?:conclusions?(?:\s+(?:and|&)\s+perspectives)?|"
    r"(?:environmental\s+)?implications?|"
    r"appendix(?:\s+.*)?|supplementary\s+(?:data|materials?))$",
    re.I,
)
_CONCLUSIONS_RE = re.compile(
    r"^conclusions?(?:\s+(?:and|&)\s+perspectives)?$",
    re.I,
)
_BACKMATTER_RE = re.compile(
    r"^(?:"
    r"notes?|acknowledg(?:e)?ments?|"
    r"references?|funding|data\s+availability(?:\s+statement)?|associated\s+content|"
    r"declaration\s+of\s+(?:competing\s+)?interests?|competing\s+interests?|"
    r"conflicts?\s+of\s+interest|credit\s+authorship(?:\s+contribution\s+statement)?|"
    r"author(?:ship)?\s+contributions?|contribution\s+statement|ethics\s+statement|"
    r"consent\s+to\s+participate|consent\s+for\s+publication|"
    r"availability\s+of\s+data\s+and\s+materials|publisher'?s\s+note"
    r")$",
    re.I,
)


def _is_retained_section_stop(text: str) -> bool:
    norm = normalize_heading(text)
    return bool(_SECTION_END_RE.fullmatch(norm) or _BACKMATTER_RE.fullmatch(norm))


def _is_conclusions_section(text: str) -> bool:
    """Return whether *text* is a Conclusions boundary that may be followed by a retained section.

    Some publishers place Experimental Section after Conclusions. Conclusions
    should remain excluded from the cleaned main-paper records, but it should
    not permanently stop the scan before a later Materials/Methods section.
    """
    return bool(_CONCLUSIONS_RE.fullmatch(normalize_heading(text)))


def _is_supplementary_backmatter_stop(text: str) -> bool:
    norm = normalize_heading(text)
    return bool(_BACKMATTER_RE.fullmatch(norm))


def _excel_abstract_record(study: str, abstract_text: str) -> Dict[str, Any]:
    text = re.sub(r"\s+", " ", abstract_text or "").strip()
    return {
        "study_folder": study,
        "source_file": "main_paper",
        "file_type": "main_paper",
        "object_type": "curated_abstract",
        "object_refs": [],
        "text": text,
        "markdown": f"## ABSTRACT\n\n{text}",
        "section_label": SECTION_LABELS["abstract"],
        "metadata_source": "within_scope_records.xlsx",
        "page_start": None,
        "page_end": None,
        "order_start": -1,
        "order_end": -1,
    }


def _markdown_for_text_record(text: str, *, is_heading: bool = False, level: int = 2) -> str:
    txt = _clean_text(text)
    if not txt:
        return ""
    if is_heading:
        lvl = max(1, min(6, int(level or 2)))
        return f"{'#' * lvl} {txt}"
    return txt


def _emit_text_record(
    *,
    study: str,
    source_file: str,
    file_type: str,
    obj: StreamObject,
    section_label: Optional[str] = None,
    section_path: Optional[List[str]] = None,
    original_file_type: Optional[str] = None,
    is_heading: bool = False,
) -> Dict[str, Any]:
    text_value = _clean_formula_text(obj.text) if obj.label == "formula" else obj.text
    markdown = _markdown_for_text_record(text_value, is_heading=is_heading, level=2)
    if not is_heading and obj.label == "list_item" and markdown:
        markdown = f"- {markdown}"
    elif not is_heading and obj.label == "formula" and markdown:
        markdown = f"$$\n{markdown}\n$$"
    rec: Dict[str, Any] = {
        "study_folder": study,
        "source_file": source_file,
        "file_type": file_type,
        "object_type": "text",
        "object_refs": [obj.object_ref],
        "text": text_value,
        "markdown": markdown,
        "page_start": obj.page_no,
        "page_end": obj.page_no,
        "order_start": obj.order_index,
        "order_end": obj.order_index,
    }
    if section_label:
        rec["section_label"] = section_label
    if section_path:
        rec["section_path"] = section_path
    if original_file_type:
        rec["original_file_type"] = original_file_type
    return rec


# Study/source processing -----------------------------------------------------
def _source_document_kind(doc: Dict[str, Any]) -> str:
    """Return the file type Docling actually converted, when known."""
    origin = doc.get("origin") or {}
    filename = str(origin.get("filename") or "").strip().lower()
    mimetype = str(origin.get("mimetype") or "").strip().lower()

    if mimetype == "application/pdf" or filename.endswith(".pdf"):
        return "pdf"
    if (
        filename.endswith(".docx")
        or mimetype
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        return "docx"
    return ""

def _source_pdf_path(study_dir: Path, source_file: str, doc: Dict[str, Any]) -> Optional[Path]:
    origin = doc.get("origin") or {}
    filename = str(origin.get("filename") or "").strip()
    candidates: List[Path] = []
    if filename and filename.lower().endswith(".pdf"):
        candidates.append(study_dir / filename)
    candidates.append(study_dir / f"{source_file}.pdf")
    for path in candidates:
        if path.is_file():
            return path
    return None

def _source_docx_path(study_dir: Path, source_file: str, doc: Dict[str, Any]) -> Optional[Path]:
    origin = doc.get("origin") or {}
    filename = str(origin.get("filename") or "").strip()
    candidates: List[Path] = []
    if filename and filename.lower().endswith(".docx"):
        candidates.append(study_dir / filename)
    candidates.append(study_dir / f"{source_file}.docx")
    for path in candidates:
        if path.is_file():
            return path
    return None


def _write_table_manual_review_artifacts(
    *,
    study: str,
    source_file: str,
    physical_table_id: str,
    table_obj: Dict[str, Any],
    docling_block: str,
    repair_payload: Dict[str, Any],
    pdf_path: Optional[Path],
    reasons: List[str],
    initial_status: str,
    repair_status: str,
    expected_header_rows: Optional[List[List[str]]] = None,
    caption: str = "",
    footnote: str = "",
    candidate_payloads: Optional[List[Dict[str, Any]]] = None,
    selected_candidate: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    page_no, _bbox = _page_bbox(table_obj)
    candidates: List[Dict[str, Any]] = []
    if candidate_payloads is not None:
        candidates = candidate_payloads
    else:
        candidates.append(_make_table_repair_candidate(
            name="docling_original",
            block=docling_block,
            payload={"status": "original"},
            expected_header_rows=expected_header_rows,
            caption=caption,
            footnote=footnote,
        ))
        repaired_block = str(repair_payload.get("block") or "").strip()
        if repaired_block:
            candidates.append(_make_table_repair_candidate(
                name=f"camelot_{repair_payload.get('strategy') or 'repair'}",
                block=repaired_block,
                payload=repair_payload,
                expected_header_rows=expected_header_rows,
                caption=caption,
                footnote=footnote,
            ))
        if pdf_path and page_no:
            for strategy in ("lines", "text"):
                payload = extract_pymupdf_table_candidate(
                    pdf_path=pdf_path,
                    page_no=page_no,
                    strategy=strategy,
                    table_bbox=_bbox,
                    bbox_pad=TABLE_REPAIR_BBOX_PAD,
                )
                block = str(payload.get("block") or "").strip()
                if block:
                    candidates.append(_make_table_repair_candidate(
                        name=f"pymupdf_{strategy}",
                        block=block,
                        payload=payload,
                        expected_header_rows=expected_header_rows,
                        caption=caption,
                        footnote=footnote,
                    ))

    bundle = _write_table_repair_bundle_artifacts(
        study=study,
        source_file=source_file,
        physical_table_id=physical_table_id,
        table_obj=table_obj,
        reasons=reasons,
        initial_status=initial_status,
        repair_status=repair_status,
        repair_payload=repair_payload,
        candidates=candidates,
        selected_candidate=selected_candidate,
    )

    return bundle


def _process_table(
    *,
    study: str,
    study_dir: Path,
    source_file: str,
    doc: Dict[str, Any],
    stream: List[StreamObject],
    stream_index: int,
    table_counter: int,
    consumed_text_refs: set[str],
    consumed_text_norms: set[str],
    suspected_rows: List[Dict[str, str]],
    figtab_rows: List[Dict[str, Any]],
    caption_resolutions: Dict[str, Dict[str, Any]],
    table_segment_indices: Dict[str, int],
    carried_header_rows: Optional[List[List[str]]] = None,
    expected_header_rows: Optional[List[List[str]]] = None,
    header_carryover_blocked: bool = False,
) -> List[Dict[str, Any]]:
    so = stream[stream_index]
    table_obj = so.raw
    deterministic_caption_raw = _caption_text_with_adjacent_identifier(doc, stream, stream_index)
    deterministic_caption = (
        deterministic_caption_raw
        if _is_valid_caption_text(deterministic_caption_raw, "table")
        else ""
    )
    fallback_id = _uncaptioned_table_id(table_counter)
    deterministic_table_id = _object_id_from_caption(deterministic_caption, fallback=fallback_id)
    deterministic_footnote, foot_refs = _linked_text(doc, table_obj, "footnotes")
    if _table_footnote_is_clearly_invalid(deterministic_footnote):
        deterministic_footnote, foot_refs = "", []
    if not deterministic_footnote:
        deterministic_footnote, foot_refs = _collect_following_table_footnote(stream, stream_index)

    table_id, caption, footnote, metadata_source = _apply_caption_resolution(
        object_ref=so.object_ref,
        object_type="table",
        deterministic_object_id=deterministic_table_id,
        deterministic_caption=deterministic_caption,
        deterministic_footnote=deterministic_footnote,
        resolutions=caption_resolutions,
    )
    if _is_valid_caption_text(caption, "table"):
        table_id = _object_id_from_caption(caption, fallback=table_id)
    if caption:
        consumed_text_refs.update(_caption_text_refs_with_adjacent_identifier(doc, stream, stream_index))
    _remember_consumed_text(consumed_text_norms, caption)
    _remember_consumed_text(consumed_text_norms, footnote)
    if footnote == deterministic_footnote:
        consumed_text_refs.update(foot_refs)

    table_body = _table_body_to_markdown(doc, table_obj)
    if carried_header_rows:
        table_body = _inject_carried_table_header(table_body, carried_header_rows)
    embedded_caption, table_body = _promote_embedded_table_caption(table_body)
    if embedded_caption and (not caption or _caption_text_is_identifier_only(caption, "table")):
        caption = embedded_caption
        table_id = _object_id_from_caption(caption, fallback=table_id)
        metadata_source = f"{metadata_source}+embedded_table_caption"
        _remember_consumed_text(consumed_text_norms, embedded_caption)
    embedded_caption_marker = _table_has_embedded_caption_marker(doc, table_obj)
    source_kind = _source_document_kind(doc)
    pdf_path = _source_pdf_path(study_dir, source_file, doc)
    docx_path = _source_docx_path(study_dir, source_file, doc)
    caption_was_missing_or_invalid = not _is_valid_caption_text(caption, "table")
    caption_repair_payload: Dict[str, Any] = {}
    camelot_caption_overrode_resolution = False
    if caption_was_missing_or_invalid or embedded_caption_marker:
        if source_kind == "pdf" and pdf_path and extract_camelot_embedded_caption is not None:
            candidate_caption, caption_repair_payload = _camelot_embedded_caption_recovery(
                pdf_path=pdf_path,
                table_obj=table_obj,
            )
            candidate_table_id = _object_id_from_caption(candidate_caption, fallback="")
            if candidate_table_id and (
                caption_was_missing_or_invalid or candidate_table_id != table_id
            ):
                caption = candidate_caption
                table_id = candidate_table_id
                metadata_source = f"{metadata_source}+camelot_embedded_caption"
                _remember_consumed_text(consumed_text_norms, caption)
                camelot_caption_overrode_resolution = not caption_was_missing_or_invalid
            elif embedded_caption_marker and not candidate_table_id:
                # Do not preserve a preceding table's valid-looking caption
                # when a new embedded caption was detected but not recovered.
                caption = ""
                table_id = fallback_id
                metadata_source = "embedded_caption_unresolved"
        else:
            caption_repair_payload = {"status": "skipped", "reason": "source PDF not found or not source document"}
            if embedded_caption_marker and not caption_was_missing_or_invalid:
                caption = ""
                table_id = fallback_id
                metadata_source = "embedded_caption_unresolved"

        if caption_was_missing_or_invalid or embedded_caption_marker:
            art_dir = OUTPUT_DIR / study / "artifacts"
            art_dir.mkdir(parents=True, exist_ok=True)
            stem = _sanitize_filename(f"{source_file}_{table_id}_camelot_caption_candidate")
            (art_dir / f"{stem}.json").write_text(
                json.dumps({k: v for k, v in caption_repair_payload.items() if k != "raw_block"}, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            if str(caption_repair_payload.get("raw_block") or "").strip():
                (art_dir / f"{stem}.md").write_text(
                    str(caption_repair_payload.get("raw_block")).rstrip() + "\n",
                    encoding="utf-8",
                )

    caption_valid = _is_valid_caption_text(caption, "table")
    logical_table_id = table_id
    segment_index = table_segment_indices.get(logical_table_id, 0) + 1
    table_segment_indices[logical_table_id] = segment_index
    physical_table_id = (
        logical_table_id
        if segment_index == 1
        else f"{logical_table_id}_segment{segment_index:02d}"
    )
    docling_block = _compose_table_block(table_body, caption=caption, footnote=footnote)
    final_block = docling_block
    source = "docling"
    initial_status = "not_flagged"
    repair_status = "not_attempted"
    repair_payload: Dict[str, Any] = {}
    manual_review_artifacts_written = False
    repair_reasons: List[str] = []
    if caption_was_missing_or_invalid:
        repair_reasons.append("caption_missing_or_invalid")
    if camelot_caption_overrode_resolution:
        repair_reasons.append("embedded_caption_overrode_inherited_caption")
        source = "docling_camelot_embedded_caption_override"
    elif caption_valid and caption_was_missing_or_invalid:
        source = "docling_camelot_caption_recovery"

    span_repaired_body, span_repair_payload = _repair_docling_spanned_body_cells(table_obj)
    deterministic_repair_clean = False
    if span_repaired_body:
        span_repaired_block = _compose_table_block(span_repaired_body, caption=caption, footnote=footnote)
        span_repair_payload["block"] = span_repaired_block.rstrip()
        if (
            not _is_suspected_merged_table(span_repaired_block)
            and not _markdown_grid_corruption_reasons(span_repaired_block)
        ):
            final_block = span_repaired_block.rstrip() + "\n"
            source = "docling_body_colspan_split"
            repair_status = "clean"
            deterministic_repair_clean = True
            repair_reasons.append("body_spanned_cell_split")
        else:
            repair_reasons.append("body_spanned_cell_split_incomplete")

    row_label_glue_suspected = (
        False if deterministic_repair_clean else _has_row_label_value_glue(docling_block)
    )
    repeated_simple_row_label_glue_suspected = (
        False if deterministic_repair_clean else _has_repeated_row_label_simple_value_glue(docling_block)
    )
    row_label_pdf_repair_suspected = row_label_glue_suspected or repeated_simple_row_label_glue_suspected
    structural_corruption_reasons = (
        []
        if deterministic_repair_clean
        else _docling_structural_table_corruption_reasons(doc, table_obj, docling_block)
    )
    continuation_corruption_reasons = (
        []
        if deterministic_repair_clean
        else _continuation_header_grid_corruption_reasons(
            table_body=table_body,
            expected_header_rows=expected_header_rows,
            header_carryover_blocked=header_carryover_blocked,
        )
    )
    markdown_grid_corruption_reasons = (
        [] if deterministic_repair_clean else _markdown_grid_corruption_reasons(docling_block)
    )
    docling_compacted_body_span_suspected = (
        False if deterministic_repair_clean else _has_docling_compacted_body_span(table_obj)
    )
    deterministic_table_repair_reasons = [
        *structural_corruption_reasons,
        *continuation_corruption_reasons,
        *markdown_grid_corruption_reasons,
    ]
    deterministic_table_repair_suspected = bool(deterministic_table_repair_reasons)
    docling_quality_extra_reasons = list(deterministic_table_repair_reasons)
    if row_label_glue_suspected:
        docling_quality_extra_reasons.append("row_label_value_glue")
    if repeated_simple_row_label_glue_suspected:
        docling_quality_extra_reasons.append("repeated_row_label_simple_value_glue")
    if docling_compacted_body_span_suspected:
        docling_quality_extra_reasons.append("docling_compacted_body_span")

    docling_quality_candidate = _make_table_repair_candidate(
        name="docling_original",
        block=docling_block,
        payload={"status": "original"},
        expected_header_rows=expected_header_rows,
        caption=caption,
        footnote=footnote,
        extra_repair_reasons=[] if deterministic_repair_clean else docling_quality_extra_reasons,
    )
    initial_quality_decision = _make_table_quality_decision(
        docling_quality_candidate,
        phase="docling_initial",
        default_corrupt_action="repair",
    )
    docling_quality_reasons = list(initial_quality_decision.get("repair_reasons") or [])
    body_suspected_merged = (
        False
        if deterministic_repair_clean
        else (
            _is_suspected_merged_table(docling_block)
            or docling_compacted_body_span_suspected
            or deterministic_table_repair_suspected
            or bool(docling_quality_reasons)
        )
    )
    if body_suspected_merged:
        repair_reasons.append("table_grid_suspected_corrupt")
    for reason in docling_quality_reasons:
        if reason not in repair_reasons:
            repair_reasons.append(reason)

    if body_suspected_merged:
        if deterministic_table_repair_suspected:
            initial_status = "table_repair_needed"
            repair_status = "table_repair_needed"
            needs_pdf_table_repair = True
            repair_reasons.append("deterministic_table_repair_trigger")
            initial_quality_decision = _make_table_quality_decision(
                docling_quality_candidate,
                repairability_status=initial_status,
                phase="docling_initial",
                default_corrupt_action="repair",
            )
        else:
            triage = _classify_table_position_with_cache(
                study=study,
                source_file=source_file,
                table_id=physical_table_id,
                caption_text=caption,
                table_markdown=docling_block,
            )
            initial_status = _normalize_position_status(triage.get("status"))
            needs_pdf_table_repair = initial_status in TABLE_POSITION_REVIEW_STATUSES
            initial_quality_decision = _make_table_quality_decision(
                docling_quality_candidate,
                repairability_status=initial_status,
                phase="docling_initial",
                default_corrupt_action="repair",
            )
        if row_label_pdf_repair_suspected:
            repair_status = "table_repair_needed"
            needs_pdf_table_repair = True
            repair_reasons.append("row_label_value_glue_requires_pdf_repair")
            initial_quality_decision = _make_table_quality_decision(
                docling_quality_candidate,
                repairability_status="table_repair_needed",
                phase="docling_initial",
                default_corrupt_action="repair",
            )

        if (
            initial_status in TABLE_MARKDOWN_EXPANDABLE_STATUSES
            and EXPAND_MARKDOWN_REPAIRABLE_TABLES
            and not row_label_pdf_repair_suspected
            and not deterministic_table_repair_suspected
        ):
            expanded_body, expansion_payload = _expanded_position_aligned_table_body(table_obj)
            expansion_payload.update({
                "source_table_grid_status": initial_status,
                "source_table_quality_decision": initial_quality_decision,
                "accepted": False,
            })
            if expanded_body:
                expanded_block = _compose_table_block(expanded_body, caption=caption, footnote=footnote)
                expansion_extra_reasons = _markdown_grid_corruption_reasons(expanded_block)
                if _is_suspected_merged_table(expanded_block):
                    expansion_extra_reasons.append("table_grid_suspected_corrupt")
                expansion_candidate = _make_table_repair_candidate(
                    name="docling_markdown_repairable_expansion",
                    block=expanded_block,
                    payload=expansion_payload,
                    expected_header_rows=expected_header_rows,
                    caption=caption,
                    footnote=footnote,
                    extra_repair_reasons=expansion_extra_reasons,
                )
                expansion_quality_decision = _make_table_quality_decision(
                    expansion_candidate,
                    phase="docling_markdown_expansion",
                    default_corrupt_action="repair",
                )
                expansion_suspected = expansion_quality_decision.get("quality_status") != "clean"
                expansion_payload["post_expansion_suspected_corrupt"] = expansion_suspected
                expansion_payload["post_expansion_quality_decision"] = expansion_quality_decision
                expansion_payload["block"] = expanded_block.rstrip()
                repair_payload = {
                    "status": "accepted" if not expansion_suspected else "table_repair_needed",
                    "route": "docling_markdown_repairable_expansion",
                    "accepted": not expansion_suspected,
                    "block": expanded_block.rstrip(),
                    "initial_table_quality_decision": initial_quality_decision,
                    "post_repair_quality_decision": expansion_quality_decision,
                    "expansion_payload": {k: v for k, v in expansion_payload.items() if k != "block"},
                }
                if not expansion_suspected:
                    final_block = expanded_block.rstrip() + "\n"
                    source = "docling_markdown_repairable_expansion"
                    repair_status = "clean"
                    expansion_payload["accepted"] = True
                    repair_payload["status"] = "accepted"
                    repair_payload["accepted"] = True
                    repair_payload["expansion_payload"]["accepted"] = True
                else:
                    repair_status = "uncertain"
                    repair_payload["status"] = "uncertain"
                    needs_pdf_table_repair = True
            else:
                repair_status = initial_status
                needs_pdf_table_repair = True

            art_dir = OUTPUT_DIR / study / "artifacts"
            art_dir.mkdir(parents=True, exist_ok=True)
            stem = _sanitize_filename(f"{source_file}_{physical_table_id}_markdown_repairable_expansion")
            (art_dir / f"{stem}.json").write_text(
                json.dumps(expansion_payload, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            if str(expansion_payload.get("block") or "").strip():
                (art_dir / f"{stem}.md").write_text(
                    str(expansion_payload.get("block")).rstrip() + "\n",
                    encoding="utf-8",
                )

        if needs_pdf_table_repair:
            pdf_repair_route = "pdf_repair_docling_guided" if expected_header_rows else "pdf_repair_unguided"
            known_outside_table_texts: List[str] = []

            def _add_known_outside_table_text(outside_text: str) -> None:
                cleaned = _clean_text(outside_text)
                if cleaned and cleaned not in known_outside_table_texts:
                    known_outside_table_texts.append(cleaned)
            for outside_text in (caption, footnote):
                _add_known_outside_table_text(outside_text)
            for outside_text in (
                _raw_char_bounded_local_text(doc, stream, stream_index, before=True, max_chars=350),
                _raw_char_bounded_local_text(doc, stream, stream_index, before=False, max_chars=350),
            ):
                _add_known_outside_table_text(outside_text)
            try:
                if source_kind == "docx" and docx_path:
                    repair_route = "docx_repair_native"
                    repair_payload = repair_docling_table_with_docx_tools(
                        docx_path=docx_path,
                        table_id=physical_table_id,
                        docling_markdown=docling_block,
                        caption_text=caption,
                        footnote_text=footnote,
                        expected_header_rows=expected_header_rows,
                        docling_original_repair_reasons=docling_quality_reasons,
                    )
                elif source_kind == "pdf" and pdf_path:
                    repair_route = pdf_repair_route
                    repair_payload = repair_docling_table_with_pdf_tools(
                        pdf_path=pdf_path,
                        table_obj=table_obj,
                        docling_markdown=docling_block,
                        caption_text=caption,
                        footnote_text=footnote,
                        expected_header_rows=expected_header_rows,
                        route=pdf_repair_route,
                        bbox_pad=TABLE_REPAIR_BBOX_PAD,
                        column_min_gap=TABLE_REPAIR_COLUMN_MIN_GAP,
                        min_nonempty_cells=TABLE_REPAIR_MIN_NONEMPTY_CELLS,
                        prefer_lattice=row_label_pdf_repair_suspected or deterministic_table_repair_suspected,
                        known_outside_texts=known_outside_table_texts,
                        docling_original_repair_reasons=docling_quality_reasons,
                    )
                elif docx_path:
                    # Fallback for older JSONs whose origin metadata is missing.
                    # Prefer a same-stem DOCX over a sidecar PDF because DOCX
                    # table objects usually have no page/bbox provenance.
                    repair_route = "docx_repair_native"
                    repair_payload = repair_docling_table_with_docx_tools(
                        docx_path=docx_path,
                        table_id=physical_table_id,
                        docling_markdown=docling_block,
                        caption_text=caption,
                        footnote_text=footnote,
                        expected_header_rows=expected_header_rows,
                        docling_original_repair_reasons=docling_quality_reasons,
                    )
                elif pdf_path:
                    repair_route = pdf_repair_route
                    repair_payload = repair_docling_table_with_pdf_tools(
                        pdf_path=pdf_path,
                        table_obj=table_obj,
                        docling_markdown=docling_block,
                        caption_text=caption,
                        footnote_text=footnote,
                        expected_header_rows=expected_header_rows,
                        route=pdf_repair_route,
                        bbox_pad=TABLE_REPAIR_BBOX_PAD,
                        column_min_gap=TABLE_REPAIR_COLUMN_MIN_GAP,
                        min_nonempty_cells=TABLE_REPAIR_MIN_NONEMPTY_CELLS,
                        prefer_lattice=row_label_pdf_repair_suspected or deterministic_table_repair_suspected,
                        known_outside_texts=known_outside_table_texts,
                        docling_original_repair_reasons=docling_quality_reasons,
                    )
                else:
                    repair_route = pdf_repair_route
                    repair_payload = {
                        "status": "skipped",
                        "route": repair_route,
                        "reason": "source PDF/DOCX not found",
                        "accepted": False,
                        "candidates": [],
                    }
            except Exception as exc:

                repair_payload = {
                    "status": "failed",
                    "route": repair_route,
                    "reason": "table repair raised exception",
                    "error": f"{type(exc).__name__}: {exc}",
                    "accepted": False,
                    "candidates": [],
                }

            repair_payload["initial_table_quality_decision"] = initial_quality_decision
            repair_payload["docling_quality_candidate"] = {
                k: v for k, v in docling_quality_candidate.items() if k != "block"
            }
            repair_candidates = list(repair_payload.get("candidates") or [])
            selected_repair_candidate = repair_payload.get("selected_candidate_payload")
            if (
                selected_repair_candidate
                and docling_quality_reasons
                and str(selected_repair_candidate.get("name") or "") == "docling_original"
            ):
                repair_payload["selected_original_after_repair_trigger"] = True
                selected_repair_candidate = _make_table_repair_candidate(
                    name="docling_original",
                    block=str(selected_repair_candidate.get("block") or docling_block),
                    payload={"status": "original"},
                    expected_header_rows=expected_header_rows,
                    caption=caption,
                    footnote=footnote,
                    extra_repair_reasons=docling_quality_reasons,
                )
                repair_payload["selected_candidate_payload"] = selected_repair_candidate
                repair_payload["selected_candidate_quality_decision"] = _make_table_quality_decision(
                    selected_repair_candidate,
                    phase=f"{repair_payload.get('route') or 'repair'}_selection",
                    default_corrupt_action="manual_review",
                )
                selected_clean = bool(selected_repair_candidate.get("accepted"))
                repair_payload["selected_candidate"] = selected_repair_candidate.get("name")
                repair_payload["selected_candidate_accepted"] = selected_clean
                repair_payload["accepted"] = selected_clean
                repair_payload["status"] = "accepted" if selected_clean else "table_repair_needed"
                repair_payload["block"] = str(selected_repair_candidate.get("block") or "").rstrip()
                repair_payload["selected_candidate_shape"] = selected_repair_candidate.get("shape")
                repair_payload["selected_candidate_header_action"] = selected_repair_candidate.get("header_action")
                repair_payload["selected_candidate_dropped_header_col_idx"] = selected_repair_candidate.get("dropped_header_col_idx")
                repair_payload["selected_candidate_quality"] = selected_repair_candidate.get("quality")

                replaced_original_candidate = False
                for idx, candidate in enumerate(repair_candidates):
                    if str(candidate.get("name") or "") == "docling_original":
                        repair_candidates[idx] = selected_repair_candidate
                        replaced_original_candidate = True
                        break
                if not replaced_original_candidate:
                    repair_candidates.insert(0, selected_repair_candidate)
                repair_payload["candidates"] = repair_candidates
            if selected_repair_candidate:
                selected_name = str(selected_repair_candidate.get("name") or "table_repair_candidate")
                repaired_block = str(
                    selected_repair_candidate.get("block")
                    or repair_payload.get("block")
                    or ""
                ).strip()

                # Post-selection validation: a repair candidate can be better
                # than the original table but still not clean.  Do not accept it
                # merely because the extraction/selection layer picked it.
                post_repair_quality_decision = _make_table_quality_decision(
                    selected_repair_candidate,
                    phase="post_repair_selection",
                    default_corrupt_action="manual_review",
                )
                post_repair_reasons = list(post_repair_quality_decision.get("repair_reasons") or [])
                repair_payload["post_repair_selected_candidate"] = selected_name
                repair_payload["post_repair_grid_reasons"] = post_repair_reasons
                repair_payload["post_repair_decision_reasons"] = post_repair_reasons
                repair_payload["post_repair_quality_decision"] = post_repair_quality_decision
                repair_payload["block"] = repaired_block

                if post_repair_reasons:
                    try:
                        post_repair_triage = _classify_table_position_with_cache(
                            study=study,
                            source_file=source_file,
                            table_id=physical_table_id,
                            caption_text=caption,
                            table_markdown=repaired_block,
                            suffix=f"_post_repair_{_sanitize_filename(selected_name)}",
                        )
                    except Exception as exc:
                        post_repair_triage = {
                            "status": "uncertain",
                            "error": f"{type(exc).__name__}: {exc}",
                        }
                    repair_status = _normalize_position_status(post_repair_triage.get("status"))
                    repair_payload["post_repair_triage"] = post_repair_triage
                    post_repair_quality_decision = _make_table_quality_decision(
                        selected_repair_candidate,
                        repairability_status=repair_status,
                        phase="post_repair_selection",
                        default_corrupt_action="manual_review",
                    )
                    repair_payload["post_repair_quality_decision"] = post_repair_quality_decision
                else:
                    repair_status = "clean"
                    post_repair_quality_decision = _make_table_quality_decision(
                        selected_repair_candidate,
                        repairability_status=repair_status,
                        phase="post_repair_selection",
                        default_corrupt_action="manual_review",
                    )
                    repair_payload["post_repair_quality_decision"] = post_repair_quality_decision

                repair_payload["post_repair_selected_position_status"] = repair_status

                if (
                    repair_status in TABLE_MARKDOWN_EXPANDABLE_STATUSES
                    and EXPAND_MARKDOWN_REPAIRABLE_TABLES
                    and not row_label_pdf_repair_suspected
                ):
                    post_expanded_body, post_expansion_payload = _expanded_position_aligned_table_body(table_obj)
                    post_expansion_payload.update({
                        "source_table_grid_status": repair_status,
                        "source_selected_candidate": selected_name,
                        "source_table_quality_decision": post_repair_quality_decision,
                        "accepted": False,
                    })

                    if post_expanded_body:
                        post_expanded_block = _compose_table_block(
                            post_expanded_body,
                            caption=caption,
                            footnote=footnote,
                        )
                        post_expansion_extra_reasons = _markdown_grid_corruption_reasons(post_expanded_block)
                        if _is_suspected_merged_table(post_expanded_block):
                            post_expansion_extra_reasons.append("table_grid_suspected_corrupt")
                        expansion_candidate = _make_table_repair_candidate(
                            name="post_repair_docling_expansion",
                            block=post_expanded_block,
                            payload=post_expansion_payload,
                            expected_header_rows=expected_header_rows,
                            caption=caption,
                            footnote=footnote,
                            extra_repair_reasons=post_expansion_extra_reasons,
                        )
                        post_expansion_quality_decision = _make_table_quality_decision(
                            expansion_candidate,
                            phase="post_repair_expansion",
                            default_corrupt_action="manual_review",
                        )
                        post_expansion_reasons = list(
                            post_expansion_quality_decision.get("repair_reasons") or []
                        )
                        post_expansion_payload["post_expansion_reasons"] = post_expansion_reasons
                        post_expansion_payload["post_expansion_quality_decision"] = post_expansion_quality_decision
                        post_expansion_payload["block"] = post_expanded_block.rstrip()

                        if post_expansion_reasons:
                            try:
                                post_expansion_triage = _classify_table_position_with_cache(
                                    study=study,
                                    source_file=source_file,
                                    table_id=physical_table_id,
                                    caption_text=caption,
                                    table_markdown=post_expanded_block,
                                    suffix=f"_post_repair_expansion_{_sanitize_filename(selected_name)}",
                                )
                            except Exception as exc:
                                post_expansion_triage = {
                                    "status": "uncertain",
                                    "error": f"{type(exc).__name__}: {exc}",
                                }
                            post_expansion_status = _normalize_position_status(
                                post_expansion_triage.get("status")
                            )
                            post_expansion_payload["post_expansion_triage"] = post_expansion_triage
                            post_expansion_quality_decision = _make_table_quality_decision(
                                expansion_candidate,
                                repairability_status=post_expansion_status,
                                phase="post_repair_expansion",
                                default_corrupt_action="manual_review",
                            )
                            post_expansion_payload["post_expansion_quality_decision"] = post_expansion_quality_decision
                        else:
                            post_expansion_status = "clean"
                            post_expansion_quality_decision = _make_table_quality_decision(
                                expansion_candidate,
                                repairability_status=post_expansion_status,
                                phase="post_repair_expansion",
                                default_corrupt_action="manual_review",
                            )
                            post_expansion_payload["post_expansion_quality_decision"] = post_expansion_quality_decision

                        repair_payload["post_expansion_quality_decision"] = post_expansion_quality_decision
                        if post_expansion_status == "clean":
                            repair_candidates.append(expansion_candidate)
                            selected_repair_candidate = expansion_candidate
                            repair_payload["selected_candidate"] = "post_repair_docling_expansion"
                            repair_payload["selected_candidate_payload"] = expansion_candidate
                            repair_payload["block"] = post_expanded_block.rstrip()
                            repaired_block = post_expanded_block.rstrip()
                            selected_name = "post_repair_docling_expansion"
                            repair_status = "clean"
                            post_expansion_payload["accepted"] = True
                        else:
                            repair_status = (
                                post_expansion_status
                                if post_expansion_status in TABLE_POSITION_REVIEW_STATUSES
                                else "table_repair_needed"
                            )
                    else:
                        repair_status = "table_repair_needed"

                    art_dir = OUTPUT_DIR / study / "artifacts"
                    art_dir.mkdir(parents=True, exist_ok=True)
                    stem = _sanitize_filename(
                        f"{source_file}_{physical_table_id}_post_repair_expansion"
                    )
                    (art_dir / f"{stem}.json").write_text(
                        json.dumps(post_expansion_payload, ensure_ascii=False, indent=2),
                        encoding="utf-8",
                    )
                    if str(post_expansion_payload.get("block") or "").strip():
                        (art_dir / f"{stem}.md").write_text(
                            str(post_expansion_payload.get("block")).rstrip() + "\n",
                            encoding="utf-8",
                        )

                if repair_status == "clean":
                    repair_payload["accepted"] = True
                    final_block = repaired_block.rstrip() + "\n"
                    source = selected_name
                else:
                    repair_payload["accepted"] = False
                    repair_payload["post_repair_review_reason"] = "selected_candidate_still_not_clean"
                    final_block = repaired_block.rstrip() + "\n"
                    source = f"{selected_name}_post_repair_review"

                repair_payload["post_repair_position_status"] = repair_status
            else:
                repair_status = "table_repair_needed"
                repair_payload["post_repair_position_status"] = repair_status
                repair_payload["accepted"] = False
                # Repair failure/manual review must not suppress the canonical
                # table artifact. Keep the best table body already available
                # (expanded/span-repaired if one was accepted earlier; otherwise
                # the original Docling block) and mark it as needing review.
                if not final_block.strip():
                    final_block = docling_block.rstrip() + "\n"
                if source == "docling":
                    source = "docling_unrepaired_review"

            _write_table_manual_review_artifacts(
                study=study,
                source_file=source_file,
                physical_table_id=physical_table_id,
                table_obj=table_obj,
                docling_block=docling_block,
                repair_payload=repair_payload,
                pdf_path=pdf_path,
                reasons=repair_reasons,
                initial_status=initial_status,
                repair_status=repair_status,
                expected_header_rows=expected_header_rows,
                caption=caption,
                footnote=footnote,
                candidate_payloads=repair_candidates,
                selected_candidate=selected_repair_candidate,
            )
            manual_review_artifacts_written = True

    repair_payload.setdefault("initial_table_quality_decision", initial_quality_decision)
    repair_payload.setdefault("docling_quality_candidate", {
        k: v for k, v in docling_quality_candidate.items() if k != "block"
    })

    if not caption_valid and not manual_review_artifacts_written:
        _write_table_manual_review_artifacts(
            study=study,
            source_file=source_file,
            physical_table_id=physical_table_id,
            table_obj=table_obj,
            docling_block=docling_block,
            repair_payload=repair_payload,
            pdf_path=pdf_path,
            reasons=repair_reasons or ["caption_missing_or_invalid"],
            initial_status=initial_status,
            repair_status=repair_status,
        )
    if not caption_valid:
        suspected_rows.append(
            {
                "study": study,
                "file_type": source_file,
                "table_name": _logical_table_group_key(logical_table_id),
                "issue": "caption_missing_or_invalid",
            }
        )
    if repair_status in TABLE_POSITION_REVIEW_STATUSES:
        suspected_rows.append(
            {
                "study": study,
                "file_type": source_file,
                "table_name": _logical_table_group_key(logical_table_id),
                "issue": repair_status,
            }
        )
    if not final_block.strip():
        final_block = docling_block.rstrip() + "\n"
        if source == "docling":
            source = "docling_empty_output_fallback"

    final_quality_decision = (
        repair_payload.get("post_expansion_quality_decision")
        or repair_payload.get("post_repair_quality_decision")
    )
    if not final_quality_decision:
        if repair_status == "clean" or (
            repair_status == "not_attempted"
            and initial_quality_decision.get("quality_status") == "clean"
        ):
            final_repairability_status = "clean"
        elif repair_status in TABLE_MARKDOWN_EXPANDABLE_STATUSES:
            final_repairability_status = repair_status
        elif repair_status in TABLE_POSITION_REVIEW_STATUSES:
            final_repairability_status = repair_status
        else:
            final_repairability_status = "table_repair_needed"
        final_quality_decision = _make_table_quality_decision(
            docling_quality_candidate,
            repairability_status=final_repairability_status,
            phase="final_table",
            default_corrupt_action="manual_review",
        )

    try:
        split_blocks = _split_large_table_block(
            table_obj=table_obj,
            full_block=final_block,
            caption=caption,
            footnote=footnote,
            study=study,
            source_file=source_file,
            table_id=physical_table_id,
        )
    except Exception as exc:
        repair_reasons.append(f"table_split_failed:{type(exc).__name__}")
        repair_status = "table_repair_needed"
        split_blocks = [final_block.rstrip() + "\n"]

    if not split_blocks:
        split_blocks = [final_block.rstrip() + "\n"]

    # Write canonical table md file(s).
    tables_dir = OUTPUT_DIR / study / "tables"
    tables_dir.mkdir(parents=True, exist_ok=True)
    page_no, bbox = _page_bbox(table_obj)
    records: List[Dict[str, Any]] = []
    for part_index, block in enumerate(split_blocks, start=1):
        output_table_id = (
            physical_table_id
            if len(split_blocks) == 1
            else f"{physical_table_id}_{LARGE_TABLE_SPLIT_PART_SUFFIX}{part_index:02d}"
        )
        table_file = tables_dir / (
            f"{_sanitize_filename(source_file + '_' + _table_id_filename_fragment(output_table_id))}.md"
        )
        output_block = _pretty_print_markdown_pipe_tables(_remove_html_breaks(block))
        table_file.write_text(output_block.rstrip() + "\n", encoding="utf-8")
        record = {
            "study_folder": study,
            "source_file": source_file,
            "file_type": "table",
            "object_type": "table",
            "object_ref": so.object_ref,
            "table_id": output_table_id,
            "logical_table_id": logical_table_id,
            "physical_table_id": physical_table_id,
            "physical_table_segment": segment_index,
            "split_part": part_index if len(split_blocks) > 1 else None,
            "caption_text": caption,
            "footnote_text": footnote,
            "text": output_block.strip(),
            "markdown": output_block.strip(),
            "page_no": page_no,
            "bbox": bbox,
            "source": source,
            "metadata_source": metadata_source,
            "caption_valid": caption_valid,
            "caption_repair_status": str(caption_repair_payload.get("status") or "not_attempted"),
            "repair_reasons": repair_reasons,
            "initial_position_status": initial_status,
            "post_repair_position_status": repair_status,
            "initial_table_quality_status": initial_quality_decision.get("quality_status"),
            "initial_table_recommended_action": initial_quality_decision.get("recommended_action"),
            "initial_table_quality_decision": initial_quality_decision,
            "final_table_quality_status": final_quality_decision.get("quality_status"),
            "final_table_recommended_action": final_quality_decision.get("recommended_action"),
            "final_table_quality_decision": final_quality_decision,
            "needs_review": (repair_status in TABLE_POSITION_REVIEW_STATUSES or not caption_valid),
            "artifact_path": str(table_file),
            "order_start": so.order_index,
            "order_end": so.order_index,
        }
        records.append(record)
        figtab_rows.append(
            {
                "study_folder": study,
                "source_file": source_file,
                "file_type": source_file,
                "type": "table",
                "object_ref": so.object_ref,
                "object_id": output_table_id,
                "caption": caption,
                "footnote": footnote,
                "page_no": page_no or "",
                "bbox": json.dumps(bbox or {}, ensure_ascii=False),
                "standardized_file_name": table_file.stem,
                "file_path": str(table_file),
                "source": source,
                "metadata_source": metadata_source,
                "needs_review": str(record["needs_review"]).lower(),
            }
        )
    return records

def _picture_asset_id_for_ref(object_id: str, object_ref: str, resolutions: Dict[str, Dict[str, Any]]) -> str:
    r = resolutions.get(object_ref) or {}
    refs = [str(ref or "") for ref in (r.get("object_refs") or [])]
    if len(refs) <= 1 or object_ref not in refs:
        return object_id
    return f"{object_id}_panel{refs.index(object_ref) + 1:02d}"

def _process_figure_like_table_metadata(
    *,
    study: str,
    source_file: str,
    doc: Dict[str, Any],
    so: StreamObject,
    figtab_rows: List[Dict[str, Any]],
    consumed_text_refs: set[str],
    consumed_text_norms: set[str],
) -> bool:
    info = _figure_like_table_info(doc, so.raw)
    if not info:
        return False

    object_id = str(info.get("object_id") or "")
    caption = str(info.get("caption") or "")
    picture_refs = [str(ref or "") for ref in (info.get("picture_refs") or []) if ref]
    if not object_id or not caption or not picture_refs:
        return False

    consumed_text_refs.update(str(ref or "") for ref in (info.get("text_refs") or []) if ref)
    _remember_consumed_text(consumed_text_norms, caption)

    for index, picture_ref in enumerate(picture_refs, start=1):
        try:
            pic = _resolve_ref(doc, picture_ref)
        except Exception:
            continue
        asset_object_id = object_id if len(picture_refs) == 1 else f"{object_id}_panel{index:02d}"
        figure_stem, figure_path = _write_picture_png_asset(
            study=study,
            source_file=source_file,
            object_id=asset_object_id,
            pic=pic,
        )
        page_no, bbox = _page_bbox(pic)
        figtab_rows.append(
            {
                "study_folder": study,
                "source_file": source_file,
                "file_type": source_file,
                "type": "figure",
                "object_ref": picture_ref,
                "object_id": object_id,
                "caption": caption,
                "footnote": "",
                "page_no": page_no or "",
                "bbox": json.dumps(bbox or {}, ensure_ascii=False),
                "standardized_file_name": figure_stem,
                "file_path": figure_path,
                "source": "docling_json_base64" if figure_path else "docling_json",
                "metadata_source": "docling_figure_like_table",
                "needs_review": "false" if figure_path else "true",
            }
        )
    return True

def _process_picture_metadata(
    *,
    study: str,
    source_file: str,
    doc: Dict[str, Any],
    stream: List[StreamObject],
    stream_index: int,
    so: StreamObject,
    figtab_rows: List[Dict[str, Any]],
    caption_resolutions: Dict[str, Dict[str, Any]],
    consumed_text_refs: set[str],
    consumed_text_norms: set[str],
) -> None:
    pic = so.raw
    deterministic_caption_raw = _caption_text_with_adjacent_identifier(doc, stream, stream_index)
    deterministic_caption = (
        deterministic_caption_raw
        if _is_valid_caption_text(deterministic_caption_raw, "figure")
        else ""
    )
    deterministic_object_id = _object_id_from_caption(
        deterministic_caption,
        fallback=f"Picture {so.object_ref.split('/')[-1]}",
    )
    object_id, caption, _, metadata_source = _apply_caption_resolution(
        object_ref=so.object_ref,
        object_type="picture",
        deterministic_object_id=deterministic_object_id,
        deterministic_caption=deterministic_caption,
        deterministic_footnote="",
        resolutions=caption_resolutions,
    )
    if not caption or not _is_real_figure_object_id(object_id):
        return
    consumed_text_refs.update(_caption_text_refs_with_adjacent_identifier(doc, stream, stream_index))
    _remember_consumed_text(consumed_text_norms, caption)
    page_no, bbox = _page_bbox(pic)
    asset_object_id = _picture_asset_id_for_ref(object_id, so.object_ref, caption_resolutions)
    figure_stem, figure_path = _write_picture_png_asset(
        study=study,
        source_file=source_file,
        object_id=asset_object_id,
        pic=pic,
    )
    figtab_rows.append(
        {
            "study_folder": study,
            "source_file": source_file,
            "file_type": source_file,
            "type": "figure",
            "object_ref": so.object_ref,
            "object_id": object_id,
            "caption": caption,
            "footnote": "",
            "page_no": page_no or "",
            "bbox": json.dumps(bbox or {}, ensure_ascii=False),
            "standardized_file_name": figure_stem,
            "file_path": figure_path,
            "source": "docling_json_base64" if figure_path else "docling_json",
            "metadata_source": metadata_source,
            "needs_review": "false" if figure_path else "true",
        }
    )


def _process_source(
    *,
    study: str,
    study_dir: Path,
    source_file: str,
    json_path: Path,
    chunk_records: List[Dict[str, Any]],
    object_records: List[Dict[str, Any]],
    suspected_rows: List[Dict[str, str]],
    figtab_rows: List[Dict[str, Any]],
    excel_abstract: str = "",
) -> None:
    doc = _load_json(json_path)
    _apply_approved_docling_text_repairs(
        doc,
        study=study,
        source_file=source_file,
    )
    stream = _flatten_body(doc)
    section_scan_lines = [
        f"## {so.text}" if so.object_type == "text" and so.label == "section_header"
        else (so.text if so.object_type == "text" else "")
        for so in stream
    ]
    pre_results_start, pre_results_methods = find_pre_methods_results_window(
        section_scan_lines
    )
    caption_resolutions = _resolve_structured_captions_with_cache(
        study=study,
        source_file=source_file,
        doc=doc,
        stream=stream,
    )
    consumed_text_refs: set[str] = set()
    consumed_text_norms: set[str] = set()
    for resolution in caption_resolutions.values():
        _remember_consumed_text(consumed_text_norms, str(resolution.get("caption_text") or ""))
        _remember_consumed_text(consumed_text_norms, str(resolution.get("footnote_text") or ""))
    for idx, obj in enumerate(stream):
        if obj.object_type not in {"table", "picture"}:
            continue
        if obj.object_type == "table" and _is_empty_layout_table(doc, obj.raw):
            continue
        consumed_text_refs.update(_caption_text_refs_with_adjacent_identifier(doc, stream, idx))
        _remember_consumed_text(consumed_text_norms, _caption_text_with_adjacent_identifier(doc, stream, idx))
        if obj.object_type == "table":
            direct_footnote, direct_foot_refs = _linked_text(doc, obj.raw, "footnotes")
            consumed_text_refs.update(direct_foot_refs)
            _remember_consumed_text(consumed_text_norms, direct_footnote)
    table_counter = 0
    table_segment_indices: Dict[str, int] = {}
    table_output_groups = _table_continuation_output_groups(
        doc,
        stream,
        study_dir=study_dir,
        source_file=source_file,
        caption_resolutions=caption_resolutions,
    )
    # The structured-caption pass can attach the preceding table's caption to
    # a fragment whose caption is embedded in its own header row. Resolve that
    # high-confidence Camelot evidence before header carryover and table
    # output naming use the stale resolution.
    caption_pdf_path = _source_pdf_path(study_dir, source_file, doc)
    embedded_caption_resolutions: Dict[str, Dict[str, Any]] = {}
    for group in table_output_groups:
        first_idx = group[0]
        if not _table_has_embedded_caption_marker(doc, stream[first_idx].raw):
            continue
        recovered_caption, _payload = _camelot_embedded_caption_recovery(
            pdf_path=caption_pdf_path,
            table_obj=stream[first_idx].raw,
        )
        if not recovered_caption:
            continue
        recovered_table_id = _object_id_from_caption(recovered_caption, fallback="")
        if not recovered_table_id:
            continue
        embedded_caption_resolutions[stream[first_idx].object_ref] = {
            "object_ref": stream[first_idx].object_ref,
            "object_refs": [stream[first_idx].object_ref],
            "object_type": "table",
            "status": "resolved",
            "object_id": recovered_table_id,
            "caption_text": recovered_caption,
            "footnote_text": None,
            "source": "docling_camelot_embedded_caption",
        }
    if embedded_caption_resolutions:
        caption_resolutions = dict(caption_resolutions)
        caption_resolutions.update(embedded_caption_resolutions)
    table_output_group_by_idx = {
        idx: tuple(group)
        for group in table_output_groups
        for idx in group
    }
    processed_table_output_groups: set[Tuple[int, ...]] = set()

    current_label: Optional[str] = None
    current_heading_path: List[str] = []
    active_methods_major: Optional[int] = None
    active_results_major: Optional[int] = None

    main_records: List[Dict[str, Any]] = []
    fallback_main_records: List[Dict[str, Any]] = []
    methods_or_results_seen = False
    main_body_stopped = False
    awaiting_late_retained_section = False
    supplementary_body_stopped = False

    for i, so in enumerate(stream):
        object_records.append(
            {
                "study_folder": study,
                "source_file": source_file,
                **{k: v for k, v in asdict(so).items() if k != "raw"},
            }
        )
        if source_file != "main_paper" and supplementary_body_stopped:
            continue

        # Tables and pictures are processed independently of retained prose sections.
        if so.object_type == "table":
            if _is_empty_layout_table(doc, so.raw):
                continue
            table_output_group = table_output_group_by_idx.get(i)
            if table_output_group and table_output_group in processed_table_output_groups:
                continue
            if table_output_group:
                processed_table_output_groups.add(table_output_group)
                inherited_table_id = ""
                inherited_caption = ""
                first_member_idx = table_output_group[0]
                first_member_so = stream[first_member_idx]
                first_table_id = _resolved_table_id_for_index(
                    doc,
                    stream,
                    first_member_idx,
                    caption_resolutions=caption_resolutions,
                ) or first_member_so.object_ref.replace("/", "_")
                first_caption = str((caption_resolutions.get(first_member_so.object_ref) or {}).get("caption_text") or "")
                if not first_caption:
                    first_caption = _caption_text_with_adjacent_identifier(doc, stream, first_member_idx)
                first_body = _table_body_to_markdown(doc, first_member_so.raw)
                first_physical_rows = _table_physical_rows_from_block(first_body)
                first_header_decision = _resolve_table_header_rows_for_block(
                    study=study,
                    source_file=source_file,
                    table_id=first_table_id,
                    table_obj=first_member_so.raw,
                    caption_text=first_caption,
                    physical_rows=first_physical_rows,
                    suffix="_carryover",
                )
                carried_header_rows = _table_header_rows_for_carryover(
                    first_member_so.raw,
                    header_count=int(first_header_decision.get("header_count") or 1),
                )
                for member_pos, member_idx in enumerate(table_output_group):
                    member_so = stream[member_idx]
                    if _is_empty_layout_table(doc, member_so.raw):
                        continue
                    header_carryover_blocked = False
                    member_header_rows = None
                    member_expected_header_rows = carried_header_rows if member_pos > 0 else None
                    if member_pos > 0 and carried_header_rows:
                        if _table_columns_compatible_with_header(member_so.raw, carried_header_rows):
                            if _table_has_own_nonempty_header_rows(member_so.raw, carried_header_rows):
                                member_expected_header_rows = None
                            else:
                                member_header_rows = carried_header_rows
                        else:
                            header_carryover_blocked = True
                    if _process_figure_like_table_metadata(
                        study=study,
                        source_file=source_file,
                        doc=doc,
                        so=member_so,
                        figtab_rows=figtab_rows,
                        consumed_text_refs=consumed_text_refs,
                        consumed_text_norms=consumed_text_norms,
                    ):
                        continue
                    table_counter += 1
                    member_caption_resolutions = caption_resolutions
                    if member_pos > 0 and inherited_table_id and inherited_caption:
                        member_caption_resolutions = dict(caption_resolutions)
                        member_caption_resolutions[member_so.object_ref] = {
                            "object_ref": member_so.object_ref,
                            "object_refs": [member_so.object_ref],
                            "object_type": "table",
                            "status": "resolved",
                            "object_id": inherited_table_id,
                            "caption_text": inherited_caption,
                            "source": "docling_table_continuation",
                        }
                    table_records = _process_table(
                        study=study,
                        study_dir=study_dir,
                        source_file=source_file,
                        doc=doc,
                        stream=stream,
                        stream_index=member_idx,
                        table_counter=table_counter,
                        consumed_text_refs=consumed_text_refs,
                        consumed_text_norms=consumed_text_norms,
                        suspected_rows=suspected_rows,
                        figtab_rows=figtab_rows,
                        caption_resolutions=member_caption_resolutions,
                        table_segment_indices=table_segment_indices,
                        carried_header_rows=member_header_rows,
                        expected_header_rows=member_expected_header_rows,
                        header_carryover_blocked=header_carryover_blocked,
                    )
                    chunk_records.extend(table_records)
                    if member_pos == 0:
                        inherited_record = next(
                            (
                                record
                                for record in table_records
                                if _is_valid_caption_text(str(record.get("caption_text") or ""), "table")
                            ),
                            None,
                        )
                        if inherited_record:
                            inherited_table_id = str(
                                inherited_record.get("logical_table_id")
                                or inherited_record.get("table_id")
                                or ""
                            )
                            inherited_caption = str(inherited_record.get("caption_text") or "")
                continue
            if _process_figure_like_table_metadata(
                study=study,
                source_file=source_file,
                doc=doc,
                so=so,
                figtab_rows=figtab_rows,
                consumed_text_refs=consumed_text_refs,
                consumed_text_norms=consumed_text_norms,
            ):
                continue
            table_counter += 1
            table_records = _process_table(
                study=study,
                study_dir=study_dir,
                source_file=source_file,
                doc=doc,
                stream=stream,
                stream_index=i,
                table_counter=table_counter,
                consumed_text_refs=consumed_text_refs,
                consumed_text_norms=consumed_text_norms,
                suspected_rows=suspected_rows,
                figtab_rows=figtab_rows,
                caption_resolutions=caption_resolutions,
                table_segment_indices=table_segment_indices,
            )
            chunk_records.extend(table_records)
            continue
        if so.object_type == "picture":
            _process_picture_metadata(
                study=study,
                source_file=source_file,
                doc=doc,
                stream=stream,
                stream_index=i,
                so=so,
                figtab_rows=figtab_rows,
                caption_resolutions=caption_resolutions,
                consumed_text_refs=consumed_text_refs,
                consumed_text_norms=consumed_text_norms,
            )
            continue
        if so.object_type != "text":
            continue
        if so.object_ref in consumed_text_refs:
            continue
        if _is_consumed_text_piece(so.text, consumed_text_norms):
            continue
        if so.content_layer == "furniture" or so.label in {"page_header", "page_footer"}:
            continue
        if so.label == "footnote":
            # Table footnotes are attached above; other footnotes usually do not help extraction.
            continue
        if not so.text:
            continue
        if _is_caption_like_text(so.text):
            continue
        if source_file != "main_paper" and _is_supplementary_backmatter_stop(so.text):
            supplementary_body_stopped = True
            continue
        if so.label == "formula":
            so.text = _clean_formula_text(so.text)
            if not so.text:
                continue

        if source_file != "main_paper":
            if so.label == "section_header" and is_table_heading_text(
                normalize_heading(so.text)
            ):
                continue
            rec = _emit_text_record(
                study=study,
                source_file=source_file,
                file_type="supplementary_material",
                obj=so,
            )
            if rec.get("text"):
                chunk_records.append(rec)
            continue

        # Main paper section handling.
        if main_body_stopped and not awaiting_late_retained_section:
            continue

        if so.label == "section_header":
            hdr = so.text
            norm = normalize_heading(hdr)
            major_num = leading_major_section_number(hdr)
            new_label = classify_heading(hdr)

            if _is_retained_section_stop(hdr):
                current_label = None
                current_heading_path = []
                active_methods_major = None
                active_results_major = None
                if _is_conclusions_section(hdr):
                    # Conclusions are excluded, but a later Experimental or
                    # Materials/Methods section must still be discoverable.
                    main_body_stopped = False
                    awaiting_late_retained_section = True
                else:
                    awaiting_late_retained_section = False
                    main_body_stopped = True
                continue

            if (
                new_label is None
                and pre_results_start is not None
                and pre_results_methods is not None
                and pre_results_start <= i < pre_results_methods
            ):
                new_label = RESULTS_LABEL

            if (
                new_label is None
                and current_label == RESULTS_LABEL
                and active_results_major is not None
                and major_num is not None
                and major_num > active_results_major
            ):
                current_label = None
                current_heading_path = []
                active_results_major = None
                main_body_stopped = True
                continue

            if (
                new_label is None
                and current_label == METHODS_LABEL
                and active_methods_major is not None
                and major_num is not None
                and major_num > active_methods_major
            ):
                new_label = RESULTS_LABEL

            if new_label:
                if new_label in {METHODS_LABEL, RESULTS_LABEL}:
                    awaiting_late_retained_section = False
                current_label = new_label
                if new_label == METHODS_LABEL:
                    active_methods_major = major_num
                    active_results_major = None
                elif new_label == RESULTS_LABEL:
                    active_methods_major = None
                    active_results_major = major_num
                else:
                    active_methods_major = None
                    active_results_major = None
                current_heading_path = [hdr]
                if new_label in {METHODS_LABEL, RESULTS_LABEL}:
                    methods_or_results_seen = True
                if new_label in TARGET_MAIN_LABELS:
                    rec = _emit_text_record(
                        study=study,
                        source_file=source_file,
                        file_type="main_paper",
                        obj=so,
                        section_label=new_label,
                        section_path=current_heading_path,
                        is_heading=True,
                    )
                    main_records.append(rec)
                fallback_main_records.append(
                    _emit_text_record(
                        study=study,
                        source_file=source_file,
                        file_type="supplementary_material",
                        obj=so,
                        original_file_type="main_paper",
                        is_heading=True,
                    )
                )
                continue

            if is_table_heading_text(norm):
                continue
            if current_label == SECTION_LABELS["abstract"]:
                current_label = None
                current_heading_path = []
                active_methods_major = None
                active_results_major = None
            elif current_label in {METHODS_LABEL, RESULTS_LABEL}:
                # Keep unlabeled subheadings inside retained methods/results sections.
                current_heading_path = current_heading_path[:1] + [hdr]
                rec = _emit_text_record(
                    study=study,
                    source_file=source_file,
                    file_type="main_paper",
                    obj=so,
                    section_label=current_label,
                    section_path=current_heading_path,
                    is_heading=True,
                )
                main_records.append(rec)
                fallback_main_records.append(
                    _emit_text_record(
                        study=study,
                        source_file=source_file,
                        file_type="supplementary_material",
                        obj=so,
                        original_file_type="main_paper",
                        is_heading=True,
                    )
                )
                continue
            else:
                fallback_main_records.append(
                    _emit_text_record(
                        study=study,
                        source_file=source_file,
                        file_type="supplementary_material",
                        obj=so,
                        original_file_type="main_paper",
                        is_heading=True,
                    )
                )
                continue

            if awaiting_late_retained_section:
                # Do not route conclusion prose or intervening back matter
                # into the cleaned/fallback records while looking for a later
                # retained section.
                continue

            if current_label in {METHODS_LABEL, RESULTS_LABEL} and _is_retained_section_stop(so.text):
                current_label = None
                current_heading_path = []
                active_methods_major = None
                active_results_major = None
                main_body_stopped = True
                continue

        fallback_main_records.append(
            _emit_text_record(
                study=study,
                source_file=source_file,
                file_type="supplementary_material",
                obj=so,
                original_file_type="main_paper",
            )
        )

        if current_label in TARGET_MAIN_LABELS:
            rec = _emit_text_record(
                study=study,
                source_file=source_file,
                file_type="main_paper",
                obj=so,
                section_label=current_label,
                section_path=current_heading_path,
            )
            main_records.append(rec)

    if source_file == "main_paper":
        excel_record = _excel_abstract_record(study, excel_abstract)
        paper_abstract_records = [
            r for r in main_records
            if r.get("section_label") == SECTION_LABELS["abstract"]
        ]
        paper_abstract_refs = {
            ref
            for record in paper_abstract_records
            for ref in record.get("object_refs", [])
        }
        chunk_records.append(excel_record)

        if methods_or_results_seen:
            chunk_records.extend(
                r
                for r in main_records
                if r.get("section_label") != SECTION_LABELS["abstract"]
                and (r.get("text") or r.get("markdown"))
            )
        else:
            # Communication-style fallback: use the curated Excel abstract and
            # route the non-abstract main-paper body as supplementary material.
            routed = [
                r
                for r in fallback_main_records
                if not set(r.get("object_refs", [])).intersection(paper_abstract_refs)
            ]
            chunk_records.extend(routed)
            print(
                f"⚠️ Communication-style routing → study={study}, source={source_file} | "
                "no methods/results section detected; routed non-abstract main-paper text as supplementary_material"
            )


# Output writers --------------------------------------------------------------
def _write_jsonl(path: Path, rows: Iterable[Dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", errors="replace") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def _write_figure_table_csv(rows: List[Dict[str, Any]]) -> None:
    FIGURE_TABLE_CSV.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "study_folder",
        "source_file",
        "file_type",
        "type",
        "object_ref",
        "object_id",
        "caption",
        "footnote",
        "page_no",
        "bbox",
        "standardized_file_name",
        "file_path",
        "source",
        "metadata_source",
        "needs_review",
    ]
    with FIGURE_TABLE_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in fieldnames})


def _read_existing_figure_table_rows() -> List[Dict[str, Any]]:
    if not FIGURE_TABLE_CSV.is_file():
        return []
    try:
        with FIGURE_TABLE_CSV.open("r", encoding="utf-8-sig", newline="") as f:
            return [dict(row) for row in csv.DictReader(f) if row]
    except Exception:
        return []


def _merge_figure_table_rows(
    existing: List[Dict[str, Any]],
    new_rows: List[Dict[str, Any]],
    processed_studies: Iterable[str],
) -> List[Dict[str, Any]]:
    touched = set(processed_studies)
    kept = [row for row in existing if str(row.get("study_folder") or "") not in touched]
    return kept + new_rows


def _table_issue_identity_key(row: Dict[str, Any]) -> Tuple[str, str, str]:
    return (
        str(row.get("study", "") or "").strip().casefold(),
        str(row.get("file_type", "") or "").strip().casefold(),
        str(row.get("table_name", "") or "").strip().casefold(),
    )


def _table_issue_priority(issue: Any) -> int:
    """Prefer the repair issue when a table has multiple detected issues."""
    normalized = _normalize_table_issue_decision(issue)
    return {
        "table_repair_needed": 2,
        "uncertain": 2,
        "caption_missing_or_invalid": 1,
    }.get(normalized, 0)


def _deduplicate_table_issue_rows(
    rows: Iterable[Dict[str, str]],
) -> List[Dict[str, str]]:
    """Keep one actionable issue row per table, including legacy duplicates."""
    selected: Dict[Tuple[str, str, str], Dict[str, str]] = {}
    for row in rows:
        normalized_row = {
            key: str(row.get(key, "") or "") for key in TABLE_ISSUE_FIELDNAMES
        }
        identity = _table_issue_identity_key(normalized_row)
        if not all(identity):
            continue
        previous = selected.get(identity)
        if previous is None or _table_issue_priority(normalized_row["issue"]) > _table_issue_priority(previous["issue"]):
            selected[identity] = normalized_row

    return sorted(
        selected.values(),
        key=lambda row: tuple(
            str(row.get(key, "") or "").casefold() for key in TABLE_ISSUE_FIELDNAMES
        ),
    )


def _normalize_table_issue_decision(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _read_table_issue_decision_rows() -> List[Dict[str, str]]:
    if not TABLE_ISSUE_DECISIONS_CSV.is_file():
        print(
            f"[DECISIONS] No table issue decisions file found at "
            f"{TABLE_ISSUE_DECISIONS_CSV}; detected issues will be retained."
        )
        return []

    try:
        with TABLE_ISSUE_DECISIONS_CSV.open(
            "r", encoding="utf-8-sig", newline=""
        ) as f:
            reader = csv.DictReader(f)
            columns = {
                str(column).strip().casefold(): column
                for column in (reader.fieldnames or [])
            }
            decision_column = columns.get("decision")
            if decision_column is None:
                print(
                    f"[DECISIONS] {TABLE_ISSUE_DECISIONS_CSV} has no Decision column; "
                    "detected issues will be retained."
                )
                return []

            fixed_file_column = columns.get(TABLE_ISSUE_FIXED_FILE_COLUMN)
            rows: List[Dict[str, str]] = []
            for row in reader:
                if not row:
                    continue
                normalized = {
                    key: str(row.get(columns.get(key, key), "") or "").strip()
                    for key in TABLE_ISSUE_FIELDNAMES
                }
                normalized["Decision"] = str(row.get(decision_column, "") or "").strip()
                normalized[TABLE_ISSUE_FIXED_FILE_COLUMN] = str(
                    row.get(fixed_file_column, "") if fixed_file_column else ""
                ).strip()
                rows.append(normalized)
            return rows
    except Exception as exc:
        print(
            f"[DECISIONS] Could not read {TABLE_ISSUE_DECISIONS_CSV}: {exc}; "
            "detected issues will be retained.",
            file=sys.stderr,
        )
        return []


def _table_issue_decisions_from_rows(
    rows: Iterable[Dict[str, str]],
) -> Dict[Tuple[str, str, str], Dict[str, str]]:
    decisions: Dict[Tuple[str, str, str], Dict[str, str]] = {}
    for row in rows:
        key = _table_issue_identity_key(row)
        decision = _normalize_table_issue_decision(row.get("Decision"))
        if all(key) and decision:
            decisions[key] = {
                "decision": decision,
                "fixed_file": str(
                    row.get(TABLE_ISSUE_FIXED_FILE_COLUMN, "") or ""
                ).strip(),
            }
    return decisions


def _new_table_issue_decision_rows(
    existing_rows: Iterable[Dict[str, str]],
    issue_rows: Iterable[Dict[str, str]],
) -> List[Dict[str, str]]:
    existing_identities = {
        _table_issue_identity_key(row)
        for row in existing_rows
        if all(_table_issue_identity_key(row))
    }
    new_rows: List[Dict[str, str]] = []
    for issue_row in _deduplicate_table_issue_rows(issue_rows):
        identity = _table_issue_identity_key(issue_row)
        if identity in existing_identities:
            continue
        new_rows.append(
            {
                **issue_row,
                "Decision": "",
                TABLE_ISSUE_FIXED_FILE_COLUMN: "",
            }
        )
        existing_identities.add(identity)
    return new_rows


def _append_new_table_issue_decisions(
    existing_rows: Iterable[Dict[str, str]],
    issue_rows: Iterable[Dict[str, str]],
) -> int:
    new_rows = _new_table_issue_decision_rows(existing_rows, issue_rows)
    if not new_rows:
        return 0

    TABLE_ISSUE_DECISIONS_CSV.parent.mkdir(parents=True, exist_ok=True)
    has_content = (
        TABLE_ISSUE_DECISIONS_CSV.is_file()
        and TABLE_ISSUE_DECISIONS_CSV.stat().st_size > 0
    )
    encoding = "utf-8" if has_content else "utf-8-sig"
    with TABLE_ISSUE_DECISIONS_CSV.open("a", encoding=encoding, newline="") as f:
        writer = csv.DictWriter(f, fieldnames=TABLE_ISSUE_DECISION_FIELDNAMES)
        if not has_content:
            writer.writeheader()
        for row in new_rows:
            writer.writerow(
                {
                    key: str(row.get(key, "") or "")
                    for key in TABLE_ISSUE_DECISION_FIELDNAMES
                }
            )
    print(
        f"[DECISIONS] Appended {len(new_rows)} new table issue row(s) to "
        f"{TABLE_ISSUE_DECISIONS_CSV}"
    )
    return len(new_rows)


def _table_issue_decision_value(decision_info: Any) -> str:
    if isinstance(decision_info, dict):
        return _normalize_table_issue_decision(decision_info.get("decision"))
    return _normalize_table_issue_decision(decision_info)


def _discover_studies() -> List[str]:
    return [paths.study_id for paths in config.iter_active_study_paths()]


def process_study(
    study: str,
    excel_abstract: str,
    table_issue_decisions: Dict[Tuple[str, str, str], Dict[str, str]],
) -> Dict[str, Any]:
    study_dir = INPUT_DIR / study
    out_dir = OUTPUT_DIR / study
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "artifacts").mkdir(parents=True, exist_ok=True)
    tables_dir = out_dir / "tables"
    tables_dir.mkdir(parents=True, exist_ok=True)
    for stale_table in _iter_table_markdown_files(tables_dir):
        if _is_preserved_manual_table_copy(stale_table):
            continue
        try:
            stale_table.unlink()
        except OSError:
            pass
    figures_dir = out_dir / "figures"
    figures_dir.mkdir(parents=True, exist_ok=True)
    for stale_figure in figures_dir.glob("*.png"):
        try:
            stale_figure.unlink()
        except OSError:
            pass

    chunk_records: List[Dict[str, Any]] = []
    object_records: List[Dict[str, Any]] = []
    suspected_rows: List[Dict[str, str]] = []
    figtab_rows: List[Dict[str, Any]] = []

    sources = []
    for source_file in ("main_paper", "supplementary_material"):
        path = study_dir / f"{source_file}.docling.json"
        if path.is_file():
            sources.append((source_file, path))

    excel_supplementary_path = study_dir / EXCEL_SUPPLEMENTARY_TABLE_NAME

    if not sources and not excel_supplementary_path.is_file():
        return {"study": study, "status": "skipped", "reason": "no Docling JSON or curated supplementary Excel source found"}

    for source_file, json_path in sources:
        print(f"[structured input] {study}/{json_path.name} -> JSON-first records")
        _process_source(
            study=study,
            study_dir=study_dir,
            source_file=source_file,
            json_path=json_path,
            chunk_records=chunk_records,
            object_records=object_records,
            suspected_rows=suspected_rows,
            figtab_rows=figtab_rows,
            excel_abstract=excel_abstract if source_file == "main_paper" else "",
        )

    if excel_supplementary_path.is_file():
        print(f"[excel input] {study}/{excel_supplementary_path.name} -> markdown table records")
        excel_table_count = _process_excel_supplementary_tables(
            study=study,
            excel_path=excel_supplementary_path,
            chunk_records=chunk_records,
            object_records=object_records,
            figtab_rows=figtab_rows,
        )
        if excel_table_count:
            sources.append((EXCEL_SUPPLEMENTARY_TABLE_NAME, excel_supplementary_path))

    manual_table_overrides = _apply_manual_table_copy_overrides(
        tables_dir=tables_dir,
        chunk_records=chunk_records,
        figtab_rows=figtab_rows,
        table_issue_decisions=table_issue_decisions,
    )
    object_exclusions = _apply_approved_docling_object_exclusions(
        study=study,
        tables_dir=tables_dir,
        chunk_records=chunk_records,
        object_records=object_records,
        figtab_rows=figtab_rows,
    )
    table_proxy_warnings = _enumeration_table_proxy_warnings(study, chunk_records)

    _write_jsonl(out_dir / "chunk_input.jsonl", chunk_records)
    _write_jsonl(out_dir / "document_objects.jsonl", object_records)
    _write_jsonl(out_dir / "table_object_records.jsonl", [r for r in chunk_records if r.get("file_type") == "table"])

    return {
        "study": study,
        "status": "processed",
        "sources": [s for s, _ in sources],
        "chunk_input_records": len(chunk_records),
        "document_object_records": len(object_records),
        "table_records": sum(1 for r in chunk_records if r.get("file_type") == "table"),
        "table_issue_rows": suspected_rows,
        "figtab_rows": figtab_rows,
        "manual_table_overrides": manual_table_overrides,
        "object_exclusions": object_exclusions,
        "table_proxy_warnings": table_proxy_warnings,
    }


def main(*, check_only: bool = False) -> int:
    selected_paths = list(config.iter_active_study_paths())
    if not selected_paths:
        print("[ERROR] ACTIVE_STUDIES is empty; select at least one registered study.")
        return 2

    ready_paths = []
    missing_inputs = False
    for paths in selected_paths:
        problems = config.json_preprocess_preflight(paths)
        if problems:
            missing_inputs = True
            print(f"[ERROR] {paths.study_id}: " + " | ".join(problems), file=sys.stderr)
            continue
        # Check the existing artifact only when retention is enabled.  Besides
        # avoiding unnecessary network-filesystem access during an intentional
        # overwrite, this prevents transient Box/SMB file-handle errors from
        # blocking regeneration of an output that will be replaced anyway.
        if not config.OVERWRITE_EXISTING_OUTPUTS and paths.chunk_input_file.is_file():
            print(f"[RETAIN] {paths.study_id}: existing {paths.chunk_input_file} will not be replaced")
            continue
        ready_paths.append(paths)
        if check_only:
            print(
                f"[CHECK OK] {paths.study_id} ({paths.group}): "
                f"would regenerate {paths.processed_study_dir}"
            )

    if check_only:
        return 2 if missing_inputs else 0
    if not ready_paths:
        if not missing_inputs:
            print("[DONE] No JSON-first artifacts need regeneration.")
        return 2 if missing_inputs else 0

    try:
        abstract_by_study = _load_within_scope_abstracts(WITHIN_SCOPE_XLSX)
    except Exception as exc:
        print(f"[ERROR] Cannot load authoritative abstracts: {exc}", file=sys.stderr)
        return 2

    existing_table_issue_decision_rows = _read_table_issue_decision_rows()
    table_issue_decisions = _table_issue_decisions_from_rows(
        existing_table_issue_decision_rows
    )
    print(
        f"[DECISIONS] Loaded {len(table_issue_decisions)} completed table issue decisions from "
        f"{TABLE_ISSUE_DECISIONS_CSV}"
    )
    existing_figtab_rows = _read_existing_figure_table_rows()
    all_figtab_rows: List[Dict[str, Any]] = []
    all_table_issues: List[Dict[str, str]] = []
    processed_studies: List[str] = []
    report: List[Dict[str, Any]] = []

    processed = 0
    for paths in ready_paths:
        study = paths.study_id
        study_no = _study_number(study)
        excel_abstract = abstract_by_study.get(study_no or -1, "").strip()
        if not excel_abstract:
            reason = "missing or empty authoritative abstract in within_scope_records.xlsx"
            print(f"[ERROR] {study}: {reason}", file=sys.stderr)
            report.append({"study": study, "status": "failed", "error": reason})
            continue
        try:
            result = process_study(
                study,
                excel_abstract,
                table_issue_decisions,
            )
            report.append({k: v for k, v in result.items() if k not in {"figtab_rows", "table_issue_rows"}})
            all_figtab_rows.extend(result.get("figtab_rows") or [])
            all_table_issues.extend(result.get("table_issue_rows") or [])
            if result.get("status") == "processed":
                processed += 1
                processed_studies.append(study)
        except Exception as exc:
            print(f"[ERROR] Failed processing {study}: {exc}", file=sys.stderr)
            report.append({"study": study, "status": "failed", "error": str(exc)})

    # Do not create or rewrite shared artifacts if every selected study failed.
    if not processed_studies:
        print("[DONE] No JSON-first artifacts were written.")
        return 1

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    merged_figtab_rows = _merge_figure_table_rows(
        existing_figtab_rows,
        all_figtab_rows,
        processed_studies,
    )
    _append_new_table_issue_decisions(
        existing_table_issue_decision_rows,
        all_table_issues,
    )
    _write_figure_table_csv(merged_figtab_rows)
    PROCESSING_REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"[REPORT] Wrote {PROCESSING_REPORT}")
    print(f"[REPORT] Wrote {FIGURE_TABLE_CSV}")
    print(f"[DONE] Completed JSON-first preprocessing for {processed} of {len(ready_paths)} selected study folder(s).")
    if summarize_llm_usage is not None:
        try:
            summarize_llm_usage()
        except Exception:
            pass
    return 2 if missing_inputs else 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="validate selected input paths without writing")
    raise SystemExit(main(check_only=parser.parse_args().check))
